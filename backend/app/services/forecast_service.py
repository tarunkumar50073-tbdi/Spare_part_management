"""SCM customer forecast ingestion + peak daily requirement calculation.

Today the SCM app isn't live yet, so the 6-month schedule arrives as an
Excel upload (`ingest_excel`). Once the SCM app is live it will push the
same shape of data directly via API — `ingest_rows` is the shared entry
point both paths call, so swapping the source later is just adding a route
that calls `ingest_rows` instead of parsing a file.

Expected Excel columns (case/spacing-insensitive): Finished Good No,
Finished Good Name, Date, Qty. One row per Finished Good per date, `Qty`
already expressed as that day's requirement (if the SCM export is monthly,
convert to a daily figure before upload)."""
from datetime import datetime, date
from openpyxl import load_workbook
from app.extensions import db
from app.repositories.finished_good_repository import FinishedGoodRepository
from app.repositories.forecast_repository import ForecastRepository
from app.repositories.bom_repository import BOMRepository
from app.utils.exceptions import BadRequestError

_HEADER_ALIASES = {
    "fg_no": {"finishedgoodno", "fgno", "finished_good_no", "partno", "customerpartno"},
    "fg_name": {"finishedgoodname", "fgname", "finished_good_name", "partname"},
    "date": {"date", "forecastdate", "schedulingdate"},
    "qty": {"qty", "quantity", "dailyqty", "dailyrequirementqty", "requirement"},
}


def _normalize_header(text):
    return str(text or "").strip().lower().replace(" ", "").replace("_", "").replace("-", "")


def _match_columns(header_row):
    normalized = {_normalize_header(cell): idx for idx, cell in enumerate(header_row)}
    columns = {}
    for field, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                columns[field] = normalized[alias]
                break
    missing = [f for f in ("fg_no", "date", "qty") if f not in columns]
    if missing:
        raise BadRequestError(f"Excel is missing required column(s): {', '.join(missing)}")
    return columns


def _coerce_date(value, row_no):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    raise BadRequestError(f"Row {row_no}: unrecognized date value {value!r}")


class ForecastService:
    def __init__(self):
        self.fg_repo = FinishedGoodRepository()
        self.repo = ForecastRepository()
        self.bom_repo = BOMRepository()

    def ingest_excel(self, file_stream):
        wb = load_workbook(file_stream, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise BadRequestError("Excel file is empty")
        columns = _match_columns(header_row)

        parsed = []
        for row_no, row in enumerate(rows_iter, start=2):
            if row is None or all(v is None for v in row):
                continue
            fg_no = row[columns["fg_no"]]
            if fg_no is None or str(fg_no).strip() == "":
                continue
            fg_name = row[columns["fg_name"]] if "fg_name" in columns else None
            forecast_date = _coerce_date(row[columns["date"]], row_no)
            qty = row[columns["qty"]]
            if qty is None:
                continue
            parsed.append({
                "fg_no": str(fg_no).strip(),
                "fg_name": str(fg_name).strip() if fg_name else str(fg_no).strip(),
                "date": forecast_date,
                "qty": float(qty),
            })

        if not parsed:
            raise BadRequestError("No usable forecast rows found in the Excel file")
        return self.ingest_rows(parsed, source="EXCEL")

    def ingest_rows(self, rows, source="EXCEL"):
        """rows: list of {fg_no, fg_name, date, qty}. Shared by Excel upload
        today and the future SCM API push."""
        count = 0
        for row in rows:
            fg = self.fg_repo.get_or_create(row["fg_no"], row.get("fg_name") or row["fg_no"])
            self.repo.upsert(fg.id, row["date"], row["qty"], source=source)
            count += 1
        db.session.commit()
        return count

    def get_peak_requirement(self, part_id):
        series = self.repo.daily_requirement_series(part_id)
        if not series:
            return None
        peak_date, peak_qty = max(series, key=lambda r: r.requirement)
        return {
            "peak_daily_requirement": float(peak_qty),
            "peak_date": peak_date.isoformat(),
            "horizon_days": len(series),
        }

    def get_peak_requirement_for_all_bom_parts(self):
        rows = []
        for bom in self.bom_repo.list_all():
            part = bom.child_part
            if any(r["part_id"] == part.id for r in rows):
                continue
            peak = self.get_peak_requirement(part.id)
            rows.append({
                "part_id": part.id,
                "part_no": part.part_no,
                "part_name": part.part_name,
                "peak_daily_requirement": peak["peak_daily_requirement"] if peak else None,
                "peak_date": peak["peak_date"] if peak else None,
                "horizon_days": peak["horizon_days"] if peak else 0,
            })
        return rows

    def get_forecast_summary(self):
        summary = []
        for row in self.repo.summary_by_finished_good():
            fg = self.fg_repo.get_by_id(row.finished_good_id)
            summary.append({
                "finished_good_id": row.finished_good_id,
                "fg_no": fg.fg_no if fg else None,
                "fg_name": fg.fg_name if fg else None,
                "from_date": row.from_date.isoformat(),
                "to_date": row.to_date.isoformat(),
                "rows": row.rows,
            })
        return summary
