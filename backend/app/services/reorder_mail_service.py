"""Reorder / low-stock mail.

E-mails each flagged department the spare parts it needs to reorder (at or below
the reorder level), to that department's person-wise recipients split into
To / Cc / Bcc. A department whose name matches a section GROUP (Tool Room =
Press Tool Room + General Tool Room) gets the union of that group's parts; one
matching a single section (Maintenance) gets that section's parts; any other
flagged department (e.g. Purchase / Store) gets the whole store.

The mail itself only says how many parts there are and of which kind; the list
goes as the attached Excel sheet, worst first, with the supplier each part last
came from and the quantity to order. No parts table in the body — the store's
instruction of 14 Oct 2026 ("mail me only excel file jaaye, list likhi huyi
nahi"). In the sheet, In Stock and Status are coloured by how urgent the part
is (REORDER_BANDS) and a legend under the table says what the colours mean.
The Reports page's Reorder / Critical List download is coloured by the same
code (reorder_paint, xlsx_table).

A durable MailSendLog row per department per day makes the daily scheduler and a
manual "Send now" idempotent (one mail per department per day); force=True (the
button) overrides that once-a-day guard. Nothing here ever raises — a build or
send failure is reported, never 500'd, and the password is never logged. A
department whose list cannot be built is reported as skipped (BUILD_FAILED) and
the departments after it still get theirs.
"""
import io
import logging
import re
import threading
import time
from collections import Counter
from copy import copy
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app.extensions import db
from app.models.mail_send_log import MailSendLog
from app.models.store import Section, Store
from app.repositories.department_repository import DepartmentRepository
from app.services.email_engine import EmailEngine
from app.services.inventory_control_service import InventoryControlService
from app.services.spare_dashboard_service import SpareDashboardService
from app.services.spare_stock_service import SpareStockService
from app.services.stock_mis_report import LEVEL_BANDS, MUTED, TINTS, _grouped
from app.utils import mailer
from app.utils.excel_io import _append_safe_row, build_export, format_table

logger = logging.getLogger(__name__)

STORE_CODE = "MAINTENANCE"
# Parts that need a purchase order — out of stock, under the minimum, or at/below
# the reorder point (status strings from SpareStockService). "Not Counted" is
# excluded: an uncounted part's need can't be known.
REORDER_STATUSES = ("Out of Stock", "Below Minimum", "Reorder")
_ORDER = {"Out of Stock": 0, "Below Minimum": 1, "Reorder": 2}

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_COLUMNS = ("Section", "Category", "Part No", "Description", "Supplier", "Location", "UOM", "In Stock",
                 "Min", "Reorder Level", "Max", "Order Qty", "Rate", "Order Value", "Status")
# Why a department got no mail when its own list could not be built.
BUILD_FAILED = "list_build_failed"

# The reorder sheet's colours, the store's (14 Oct 2026: "reorder wali mail me
# bhi color karen" — he chose "3 rang: gaadha red / red / orange"). In Stock
# and Status say how urgent a part is, by the same rule: nothing on the shelf
# in solid dark red, at or below its min in red, the rest of the list — above
# the min, at or below the reorder level — in orange. The two reds are those
# of the Daily Stock Report (Stock Only)'s In Stock cell (stock_mis_report
# LEVEL_BANDS, where AT the min counts as red too) and the orange is the app's
# Reorder status tint, so a part reads the same in every sheet he gets. As
# (fill, text); Order Qty is only bold — it is the figure Purchase acts on.
REORDER_BANDS = {"zero": LEVEL_BANDS["zero"], "min": LEVEL_BANDS["low"], "reorder": TINTS["Reorder"][:2]}
REORDER_LEGEND = ("In Stock / Status: dark red = no stock · red = at or below min · "
                  "orange = at or below reorder level")


def _n(v):
    try:
        f = float(v)
        # Two decimals, as the quantities and levels are kept — the figure the
        # attached sheet shows. ("g" keeps six digits: 12345.25 read 12345.2.)
        return str(int(f)) if f == int(f) else f"{f:.2f}".rstrip("0").rstrip(".")
    except (TypeError, ValueError):
        return "-"


# ---- what the parts mails share (email_rules builds its lists with these too) ----

def supplier_names():
    """{part id as text: supplier name} — the supplier of each part's latest
    gate entry, the name the Critical · Reorder · Excess screen shows. One read
    for the whole store: ask once per run, never per part. A part that never
    came through the gate is not in it."""
    return {part_id: facts["supplier"]
            for part_id, facts in SpareStockService().part_facts(STORE_CODE).items() if facts.get("supplier")}


def part_cells(row, sections, suppliers):
    """What an Excel sheet says about one part, keyed by column heading; a
    sheet takes the columns it lists and adds its own figures. Category is the
    part's type (sub_category), as on the dashboard's own reorder sheet.
    Anything not known — no supplier, no section — is None: an empty cell. The
    stock row is only read."""
    return {
        "Section": sections.get(row.get("section_id")), "Category": row.get("sub_category"),
        "Part No": row.get("part_no"), "Description": row.get("part_name"),
        "Specification": row.get("specification"), "Supplier": suppliers.get(str(row.get("part_id"))),
        "Location": row.get("location"), "UOM": row.get("uom"), "In Stock": row.get("qty_on_hand"),
        "Min": row.get("min_stock_level"), "Reorder Level": row.get("reorder_level"),
        "Max": row.get("max_stock_level"), "Status": row.get("status"), "Rate": row.get("unit_price"),
        "Stock Value": row.get("stock_value"),
    }


def xlsx_attachment(filename, columns, rows, paint=None, legend=None):
    """A mail attachment holding one formatted Excel table (frozen bold header,
    borders, sized columns, autofilter). `rows` are dicts keyed by column
    heading. build_export guards every text cell against formula injection.

    `paint(ws, row number, row)` colours one data row once the table is laid
    out, and `legend` is one line under the table saying what the colours
    mean. Without either the sheet is the one build_export makes — what every
    attachment was before the colours, and what the others still are."""
    return {"filename": filename, "mime": XLSX_MIME, "content": xlsx_table(columns, rows, paint, legend)}


def xlsx_table(columns, rows, paint=None, legend=None):
    """The bytes of xlsx_attachment's sheet. The Reports page builds its
    reorder and excess downloads with it too (report_service), so a list is
    coloured the same wherever it opens."""
    if paint is None and legend is None:
        return build_export([(column, column) for column in columns], rows).getvalue()
    return _coloured_table(columns, rows, paint, legend)


def _coloured_table(columns, rows, paint, legend):
    """build_export's table — the same header, formula guard, borders, widths,
    frozen header and filter, built by the same helpers — then each data row
    handed to `paint`, then the legend: after one empty row, in column A. The
    legend is written after format_table, so it has no border, is not inside
    the filter and does not widen column A (its text runs on over the empty
    cells beside it)."""
    wb = Workbook()
    ws = wb.active
    ws.append(list(columns))
    for row in rows:
        _append_safe_row(ws, [row.get(column, "") for column in columns])
    format_table(ws)
    if paint is not None:
        for number, row in enumerate(rows, start=2):
            paint(ws, number, row)
    if legend:
        note = ws.cell(row=len(rows) + 3, column=1, value=legend)
        note.font = _font(note, italic=True, color=MUTED)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _font(cell, **change):
    """The cell's own font (name, size) with `change` made to it."""
    font = copy(cell.font)
    for name, value in change.items():
        setattr(font, name, value)
    return font


def mark_cell(cell, fill=None, ink=None):
    """Make one table cell bold — on a solid `fill`, its text in `ink`, when
    given. Only the font and the fill change: the border, alignment and number
    format format_table gave the cell stay as they are."""
    cell.font = _font(cell, bold=True, **({"color": ink} if ink else {}))
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def reorder_band(qty, minimum):
    """Which of REORDER_BANDS a part on the reorder list is in, by its stock
    against its min: "zero" with nothing on the shelf, "min" at or below a min
    that is set (a min stored as 0 is no min), "reorder" for the rest — the
    list holds only parts at or below the reorder level. None, no colour, for
    a stock not known (a part never counted is never on the list)."""
    if qty is None:
        return None
    if qty <= 0:
        return "zero"
    if minimum and qty <= minimum:
        return "min"
    return "reorder"


def reorder_paint(columns):
    """The `paint` of a reorder list laid out in `columns`: In Stock and
    Status in the part's REORDER_BANDS colour, Order Qty bold. The cells are
    found by their heading, so one rule colours every sheet of the list,
    whatever its column order — this mail's attachment and the Reports page's
    Reorder / Critical List download, which the store expects in the same
    colours (15 Oct 2026, on the Reports download: "bhai kahan kiya???")."""
    column = {heading: number for number, heading in enumerate(columns, start=1)}

    def paint(ws, number, row):
        band = reorder_band(row.get("In Stock"), row.get("Min"))
        if band:
            for heading in ("In Stock", "Status"):
                mark_cell(ws.cell(row=number, column=column[heading]), *REORDER_BANDS[band])
        mark_cell(ws.cell(row=number, column=column["Order Qty"]))
    return paint


_paint_reorder_row = reorder_paint(SHEET_COLUMNS)


def sheet_name(title, dept_name, day):
    """The file name of one department's sheet: "<title> - <department>
    DD.MM.YYYY.xlsx". A plain hyphen, not a long dash: a name that is not
    plain ASCII goes out RFC 2231-encoded and can show garbled in Outlook."""
    return f"{title} - {_file_safe(dept_name)} {day:%d.%m.%Y}.xlsx"


def _order_qty(row):
    """What to order for a part on the reorder list: the gap up to its maximum,
    or — with no maximum planned — the gap up to its minimum; None when neither
    gives a quantity. The Critical · Reorder · Excess screen's own rule."""
    return InventoryControlService._order_qty(row)[0]


def _file_safe(name):
    """A typed name as part of a file name: what Windows refuses in one, and
    any line break, becomes a space."""
    return " ".join(re.sub(r'[\\/:*?"<>|\x00-\x1f]', " ", name or "").split())


def _seconds_until(hour):
    now = datetime.now()
    target = now.replace(hour=hour, minute=0, second=0, microsecond=0)
    if target <= now:
        target += timedelta(days=1)
    return max(1, int((target - now).total_seconds()))


class ReorderMailService:
    _scheduler_started = False

    def __init__(self):
        self.dept_repo = DepartmentRepository()
        self.stock = SpareStockService()

    @staticmethod
    def _scopes():
        """(sections, groups) for matching a department to the parts it covers:
        - sections: active section name (lower) -> id
        - groups:   section group name (lower) -> set of that group's section ids
        A department named like a GROUP (e.g. "Tool Room") mails the union of the
        group's sections (Press Tool Room + General Tool Room); one named like a
        single section mails that section; anything else (Purchase/Store) mails
        the whole store. Only the spare store's own sections count: a name is
        unique per store, and another store's section of the same name must
        not take a department's parts away."""
        sections, groups = {}, {}
        for s in (Section.query.join(Store, Store.id == Section.store_id)
                  .filter(Store.code == STORE_CODE, Section.is_active.is_(True)).all()):
            sections[s.name.strip().lower()] = s.id
            grp = (s.group_name or "").strip().lower()
            if grp:
                groups.setdefault(grp, set()).add(s.id)
        return sections, groups

    @staticmethod
    def _section_ids_for(dept_name, sections, groups):
        """Section ids a department's reorder mail covers, or None for the whole
        store. Group (umbrella) first, then an exact section, else whole store."""
        key = (dept_name or "").strip().lower()
        if key in groups:
            return set(groups[key])
        if key in sections:
            return {sections[key]}
        return None

    def _reorder_rows(self, section_ids):
        """Parts at/below the reorder level for the given section ids (None =
        whole store), worst first. One store-wide read, then filtered, so a
        group spanning several sections needs no extra query."""
        rows = [r for r in self.stock.list_stock(STORE_CODE, None)
                if r.get("status") in REORDER_STATUSES
                and (section_ids is None or r.get("section_id") in section_ids)]
        rows.sort(key=lambda r: (_ORDER.get(r["status"], 9),
                                 -((r.get("min_stock_level") or 0) - (r.get("qty_on_hand") or 0))))
        return rows

    @staticmethod
    def _recipients(dept):
        """Active recipients with an address, split into (to, cc, bcc) by role."""
        buckets = {"TO": [], "CC": [], "BCC": []}
        for r in dept.recipients:
            email = (r.email or "").strip()
            if r.is_active and email:
                buckets.get((r.role or "TO").upper(), buckets["TO"]).append(email)
        return buckets["TO"], buckets["CC"], buckets["BCC"]

    @staticmethod
    def _wording(dept_name, count, day):
        """This mail's subject, opening and closing lines as Email Templates has
        them. The built-in text there is word for word what this mail always
        said, so nothing changes until an admin saves another."""
        # Imported here: email_rules runs this service for its reorder rule.
        from app.services.email_rules import mail_wording
        return mail_wording("reorder", department=dept_name, count=count, date=day)

    @staticmethod
    def _body(rows, intro, footer):
        """(text, html) of one department's mail: the opening lines, one line
        of totals over every part and the closing lines — the parts themselves
        are in the attached sheet. The closing lines end the mail (the system
        note, unless an admin reworded them), so the shared signature is left
        off."""
        # Imported here: email_rules runs this service for its reorder rule.
        from app.services.email_rules import _layout
        statuses = Counter(r.get("status") for r in rows)
        line = (f"{_grouped(len(rows))} parts to reorder · {_grouped(statuses['Out of Stock'])} zero stock · "
                f"{_grouped(statuses['Below Minimum'])} below minimum · {_grouped(statuses['Reorder'])} at reorder level")
        return _layout(intro, [{"line": line}], footer, signed=False)

    @staticmethod
    def _attachment(dept_name, rows, day, suppliers):
        """The whole list as an Excel sheet, in the mail's order, with what to
        order of each part and what that costs. A rate of 0 is a missing rate,
        as on the Critical · Reorder · Excess screen: no order value. In Stock
        and Status are coloured by REORDER_BANDS, with the legend under the
        table."""
        sections = SpareDashboardService._section_names(STORE_CODE)
        lines = []
        for r in rows:
            qty, rate = _order_qty(r), r.get("unit_price")
            lines.append({**part_cells(r, sections, suppliers), "Order Qty": qty,
                          "Order Value": round(qty * rate, 2) if qty and rate else None})
        return xlsx_attachment(sheet_name("Reorder List", dept_name, day), SHEET_COLUMNS, lines,
                               paint=_paint_reorder_row, legend=REORDER_LEGEND)

    def send_now(self, force=True):
        """Send the reorder mail for every flagged department that has parts to
        reorder and at least one recipient. Returns a per-department summary:
        {status, total_sent, departments:[{department, status, reason?, parts?,
        to?, cc?, bcc?}]}."""
        result = {"departments": [], "total_sent": 0}
        if not mailer.smtp_is_configured():
            return {**result, "status": "skipped", "reason": "smtp_not_configured"}

        today = date.today()
        sections, groups = self._scopes()
        suppliers = None
        for dept in self.dept_repo.list_all(active_only=True):
            if not dept.notify_reorder:
                continue
            entry = {"department": dept.name}
            to, cc, bcc = self._recipients(dept)
            if not (to or cc or bcc):
                result["departments"].append({**entry, "status": "skipped", "reason": "no_recipient"})
                continue
            # dept name -> its group / section; neither (Purchase/Store) -> whole store
            rows = self._reorder_rows(self._section_ids_for(dept.name, sections, groups))
            if not rows:
                result["departments"].append({**entry, "status": "skipped", "reason": "nothing_to_reorder"})
                continue
            key = f"reorder_{dept.id}"
            if not force and MailSendLog.query.filter_by(report_key=key, report_date=today).first():
                result["departments"].append({**entry, "status": "skipped", "reason": "already_sent_today"})
                continue
            wording = self._wording(dept.name, len(rows), today)
            subject = wording["subject"]
            if suppliers is None:
                suppliers = supplier_names()        # read once, for every department of this run
            try:
                text, html = self._body(rows, wording["intro"], wording["footer"])
                sheet = self._attachment(dept.name, rows, today, suppliers)
            except Exception:  # noqa: BLE001 — one department's list must not cost the others theirs
                db.session.rollback()
                logger.exception("the reorder list of %s could not be built", dept.name)
                result["departments"].append({**entry, "status": "skipped", "reason": BUILD_FAILED})
                continue
            # Through the mail engine: same result dict as the mailer's, and the
            # mail is on record — a failed attempt is retried from the queue.
            sent = EmailEngine().submit(
                mail_type="reorder", subject=subject, body=text, html=html, attachments=[sheet],
                to=to, cc=cc, bcc=bcc,
                trigger="MANUAL" if force else "SCHEDULED", store_code=STORE_CODE,
                dedup_key=EmailEngine.report_dedup_key(key, today, force), mode="sync")
            if not sent.get("ok"):
                result["departments"].append({**entry, "status": "skipped", "reason": sent.get("reason", "send_failed")})
                continue
            try:
                db.session.add(MailSendLog(report_key=key, report_date=today, subject=subject[:255],
                                           to_addrs=", ".join(to)[:500], cc_addrs=", ".join(cc + bcc)[:500]))
                db.session.commit()
            except Exception:  # noqa: BLE001 — the mail already went; a log race is harmless
                db.session.rollback()
            result["departments"].append({**entry, "status": "sent", "parts": len(rows),
                                          "to": len(to), "cc": len(cc), "bcc": len(bcc)})
            result["total_sent"] += 1
        return {**result, "status": "ok"}

    def status(self):
        """SMTP + schedule state and, per flagged department, how many recipients
        (by role) and parts currently need reordering — so the screen can preview
        what a send would do — and the sections its lists cover: their names,
        or None for the whole store (a department named like no section or
        group), so that fallback can be seen."""
        # Imported here: email_rules runs this service for its reorder rule.
        from app.services.email_rules import rule_settings
        # The schedule is the notification rule's now (the .env flags are only
        # its default until an admin saves it).
        rule, _ = rule_settings("reorder_daily")
        sections, groups = self._scopes()
        names = SpareDashboardService._section_names(STORE_CODE)
        depts = []
        for dept in self.dept_repo.list_all(active_only=True):
            if not dept.notify_reorder:
                continue
            to, cc, bcc = self._recipients(dept)
            covered = self._section_ids_for(dept.name, sections, groups)
            depts.append({
                "department": dept.name, "to": len(to), "cc": len(cc), "bcc": len(bcc),
                "reorder_parts": len(self._reorder_rows(covered)),
                "sections": None if covered is None else sorted(names[section_id] for section_id in covered),
            })
        return {
            "smtp_configured": mailer.smtp_is_configured(),
            "enabled": rule["enabled"],
            "hour": rule["hour"],
            "departments": depts,
        }

    # ---- daily scheduler (clone of the window-time mailer) ----
    def start_scheduler(self, app):
        """Start the once-a-day reorder mail thread. No-op unless
        REORDER_MAIL_ENABLED is true AND SMTP is configured, and only once per
        process. The durable MailSendLog makes a re-entry safe."""
        if (not bool(app.config.get("REORDER_MAIL_ENABLED"))
                or not bool(app.config.get("SMTP_HOST"))
                or ReorderMailService._scheduler_started):
            return
        ReorderMailService._scheduler_started = True
        hour = int(app.config.get("REORDER_MAIL_HOUR", 8))

        def loop():
            while True:
                time.sleep(_seconds_until(hour))
                try:
                    with app.app_context():
                        self.send_now(force=False)
                except Exception:  # noqa: BLE001 — never let the thread die
                    logger.exception("reorder daily mail failed")

        threading.Thread(target=loop, name="reorder-mailer", daemon=True).start()
