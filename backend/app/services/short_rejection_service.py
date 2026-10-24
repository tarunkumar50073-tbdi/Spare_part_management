"""Short & Rejection — what a supplier sent short, and what was rejected, on
the ERP GRN lines of a store's parts, and how long each has been pending.

The store asked for it on 17 Oct 2026: "short ki mail banani padegi agar kisi
supplier se short receive hua to mail jaayegi" and "rejection ka bhi banana
padega taaki pta chale kis supplier ki rejection padi h kitne dino se and kis
lot k against me reject hui — ye erp api se mil jaayega".

Never mixed. On hearing the plan the store asked "rejection and short mix nhi
rhega na": a short and a rejection are settled differently, so the service
reads ONE kind per call (KINDS) — its own lines, supplier-wise table, totals
and file — and a GRN line that is both short and rejected is listed under each
kind with only that kind's figure. Rejection is punched in ERP in two types
("rejection punch bhi erp se hoti h 2 type se srv rejection and line
rejection"; on screen "rejection me drop down se srv rejection, line
rejection"), and those two are never mixed either:

    short  Short = Challan Qty - Received Qty, when above 0: the supplier's
           challan said more than the store received.
    srv    SRV Rejection = the Reject Qty on the GRN line, when above 0:
           quality rejected it at receipt.
    line   Line Rejection = ERP's termsdb.line_rejection_report, which the app
           neither mirrors nor can reach yet: erp_line_rejections() says so,
           and the screen says that instead of showing a list.

The source of Short and SRV Rejection is ERP's GRN register
(termsdb.srv_report), mirrored read-only into receiving_register by
ERPSyncService. It is read in ONE function, erp_grn_lines(): when the ERP API
arrives it replaces that read and nothing else — the screen, the files and the
mail take the same rows from it.

A GRN line belongs to a store through its part: Part No = the ERP item code,
trimmed and case-blind, as ErpGrnPostingService matches it (the mirror's own
store column is CHILD_PART on every row, so it says nothing). Money is at the
part's CURRENT rate (unit_price): a part with no rate (or a rate of 0) has no
value — blank, never 0 — and the totals say how many lines are not valued.

Days pending, and when a line is closed. The store, 17 Oct 2026: a rejection
is done "rejection invoice bna kar bhej denge tab" — once the rejection
invoice is made and sent. ERP will tell the app that per GRN line, and later
how a short was settled; it tells neither yet. Closure is read in ONE
function, erp_closures(kind, rows): today it answers None ("ERP gives none"),
so every rejection line is "Open — rejection invoice not yet made", every
short line "Open", and the days pending are the days since the GRN date on the
plant clock. When ERP gives the reference (for a rejection: the invoice no. /
date), that function returns it and nothing else changes: the line reads
"Closed — invoice <no> (<date>)", its days stop on that date, and it leaves
the "oldest" and the age bands, which are about what is still pending. Age
bands: 0-7, 8-15, 16-30, 30+ days.

What the data cannot show is said, not left blank: the mirror keeps only the
last ERP_SYNC_SINCE_DAYS of GRNs, so the oldest GRN held is given and a window
reaching before it is flagged (an open line on an older GRN is not listed); a
window reaching past the last ERP refresh is flagged too — nothing is known
for those days, which is not "nothing short" (_coverage); and the kind's lines whose item code is no store's Part No are counted, since
no store would list them.

The window is on the GRN date: a Period (PERIODS) cut here on the plant
clock, or a Custom From - To; by default the last 90 days to today, at most
MAX_DAYS. Filters: a supplier — its code exactly (a click on the supplier-wise
table), else any code or name containing the typed text — and a part, by the
typed text in its Part No or description (or by its id). Nothing here writes.
"""
import io
from datetime import date, timedelta
from xml.sax.saxutils import escape

from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
from sqlalchemy import func, or_

from app.extensions import db
from app.models.base import utcnow
from app.models.part import Part
from app.models.receiving_register import ReceivingRegister
from app.services.consumption_analytics_service import _local_offset, _parse_day, _utc_naive
from app.services.consumption_table_service import preset_window
from app.services.reorder_mail_service import XLSX_MIME, _file_safe, mark_cell
from app.services.stock_mis_report import BELOW, HEALTHY, MONEY, REORDER, TINTS, TOTAL
from app.utils import pdf_io
from app.utils.excel_io import _append_safe_row, format_table
from app.utils.exceptions import BadRequestError, NotFoundError

TITLE = "Short & Rejection"

# The kinds, never mixed: each is its own tab (Short; Rejection with its Type
# dropdown SRV / Line), its own figures and its own file, named by `label`.
# `qty` / `value` head that kind's columns; `open` is a line's status while ERP
# reports no closure, `how` the word before the reference once it does; `noun`
# names the kind inside a sentence, `none` opens "none of it" (`nothing` whole);
# `days`,
# `closed` and `no_closure` are the notes said beside the figures — plain text,
# the PDF's font has no rupee sign.
_REJECTION_DAYS = "Days pending count from the GRN date until ERP reports the rejection invoice."
_REJECTION_CLOSED = ("A rejection is closed once ERP reports its rejection invoice (made and sent); "
                     "its days pending stop at the invoice date.")
_REJECTION_NONE = "ERP reports no rejection invoice yet, so every line here is open."
_REJECTION_OPEN = "Open — rejection invoice not yet made"
KINDS = {
    "short": {
        "label": "Short", "noun": "short", "qty": "Short Qty", "value": "Short Value", "open": "Open", "how": "ref",
        "what": "Short = Challan Qty - Received Qty: the supplier's challan said more than the store received.",
        "days": "Days pending count from the GRN date until ERP reports how the short is settled.",
        "closed": "A short is closed once ERP reports how it was settled; its days pending stop on that date.",
        "no_closure": "ERP reports no settlement of a short yet, so every line here is open.",
        "nothing": "Nothing short in this window.", "none": "Nothing short",
    },
    "srv": {
        "label": "SRV Rejection", "noun": "SRV rejection", "qty": "Rejected Qty", "value": "Rejected Value",
        "open": _REJECTION_OPEN,
        "how": "invoice",
        "what": "SRV Rejection = the qty quality rejected on the ERP GRN line (its Reject Qty).",
        "days": _REJECTION_DAYS, "closed": _REJECTION_CLOSED, "no_closure": _REJECTION_NONE,
        "nothing": "No SRV rejection in this window.", "none": "No SRV rejection",
    },
    "line": {
        "label": "Line Rejection", "noun": "line rejection", "qty": "Rejected Qty", "value": "Rejected Value",
        "open": _REJECTION_OPEN,
        "how": "invoice",
        "what": "Line Rejection = the qty rejected on the line, punched in ERP (line_rejection_report).",
        "days": _REJECTION_DAYS, "closed": _REJECTION_CLOSED, "no_closure": _REJECTION_NONE,
        "nothing": "No line rejection in this window.", "none": "No line rejection",
    },
}
# The other notes, the same for every kind.
BASIS = {
    "value": "Valued at each part's current rate. A part with no rate has no value - blank, not 0.",
    "store": "A GRN line is this store's when its ERP item code is one of the store's Part Nos.",
    "window": "The period is on the GRN date.",
}

DEFAULT_DAYS = 90
# The screen's Period dropdown, the Issuance / Consumption reports' kind:
# (key, label). Cut here on the plant clock, not in the browser — last month
# and the financial years (1 April - 31 March) as those reports cut them
# (consumption_table_service.preset_window). "custom" takes From - To.
PERIODS = (("30d", "Last 30 days"), ("90d", "Last 90 days"), ("month", "This month"),
           ("last_month", "Last month"), ("this_fy", "This financial year"),
           ("last_fy", "Last financial year"), ("custom", "Custom"))
_PERIOD = dict(PERIODS)
DEFAULT_PERIOD = "90d"
# Five years, as the Consumption screen allows: the register is small, and an
# old open rejection must stay reachable.
MAX_DAYS = 1830
# The ERP data is "not current" once its last refresh is more than this many
# plant days old — the screen's banner turns amber.
STALE_AFTER_DAYS = 2
# Item codes named when ERP lines match no store's part: enough to act on.
UNMATCHED_ITEMS = 10

# Days pending, in bands: (key, lowest, highest) — None = open-ended. "30+"
# starts at 31: thirty days is still the 16-30 band.
AGE_BANDS = (("0-7", 0, 7), ("8-15", 8, 15), ("16-30", 16, 30), ("30+", 31, None))
# The app's status tints as (fill, text), so a band reads the same here as on
# screen: green, amber, orange, red (stock_mis_report.TINTS).
AGE_TINTS = {"0-7": TINTS[HEALTHY][:2], "8-15": TINTS[BELOW][:2], "16-30": TINTS[REORDER][:2],
             "30+": TINTS["Critical"][:2]}
AGE_LEGEND = "Age (days pending): green = 0-7 days · amber = 8-15 · orange = 16-30 · red = more than 30"

FORMATS = {"excel": ("xlsx", XLSX_MIME), "pdf": ("pdf", "application/pdf")}

# ERP's table of line rejections — named where the Line Rejection type says why
# it shows nothing.
LINE_REJECTION_TABLE = "line_rejection_report"
MONEY_COLUMNS = ("Rate", "Short Value", "Rejected Value")
TOTAL_LABEL = "Total"
SUPPLIER_SHEET, LINES_SHEET = "Supplier-wise", "Lines"

# What ERP's GRN register does not carry yet, by line field -> how it is named.
# A field blank on every line is listed under `not_given`, so the screen can say
# "not given by ERP yet" instead of showing an empty column as if it were data.
ERP_FIELDS = {"gate_entry_no": "Lot (gate entry no. and date)", "challan_dt": "Challan date", "po_no": "PO no."}

NOTE_FONT = Font(italic=True, color="475569")
NOTE_STYLE = ParagraphStyle("ShortRejectionNote", fontName="Helvetica", fontSize=8, leading=10,
                            textColor=colors.HexColor("#475569"))


def supplier_columns(kind):
    """The supplier-wise sheet's headings: that kind's qty and value only."""
    k = KINDS[kind]
    return ("Supplier Code", "Supplier", "Lines", k["qty"], k["value"], "Oldest (days)", "Latest GRN")


def line_columns(kind):
    """The lines sheet's headings, in the screen's order: that kind's qty and
    value only — a GRN line also short (or also rejected) shows nothing of it
    here."""
    k = KINDS[kind]
    return ("GRN Date", "GRN No", "Lot (Gate Entry)", "Challan", "PO No", "Supplier Code", "Supplier",
            "Part No", "Description", "Challan Qty", "Received Qty", "Accepted Qty", k["qty"], "Rate",
            k["value"], "Days Pending", "Age", "Status")


# ---- the sources: the ONE reads an ERP API will replace ----

def erp_grn_lines(first, last):
    """(rows, as_of): every ERP GRN line dated `first`..`last` with something
    short or rejected on it, and how far the ERP data reaches. Read today from
    receiving_register, the read-only mirror of ERP's srv_report; the ERP API
    replaces this function — the same keys — and nothing else.

    rows: dicts with id, grn_no, grn_date, ven_cd, ven_name, gate_entry_no,
    gate_entry_dt, challan_no, challan_dt, po_no, item_cd, item_desc,
    challan_qty, receipt_qty, accept_qty, reject_qty.
    as_of: {"synced_at": when the mirror last took rows from ERP (UTC, naive),
    "oldest_grn_date" / "newest_grn_date": the GRN dates the data reaches,
    "keeps_days": how many days back the source keeps (None = all of them),
    "rows": the whole mirror's row count}."""
    RR = ReceivingRegister
    q = (db.session.query(RR.id, RR.grn_no, RR.grn_date, RR.ven_cd, RR.ven_name, RR.gate_entry_no,
                          RR.gate_entry_dt, RR.challan_no, RR.challan_dt, RR.po_no, RR.item_cd, RR.item_desc,
                          RR.challan_qty, RR.receipt_qty, RR.accept_qty, RR.reject_qty)
         .filter(RR.grn_date >= first, RR.grn_date <= last, RR.item_cd.isnot(None),
                 or_(RR.reject_qty > 0, RR.challan_qty - RR.receipt_qty > 0)))
    rows = [dict(r._mapping) for r in q]
    # The sync replaces the whole snapshot, so the newest row's stamp is the
    # moment of the last refresh.
    synced, oldest, newest, count = db.session.query(
        func.max(RR.created_at), func.min(RR.grn_date), func.max(RR.grn_date), func.count(RR.id)).one()
    # Each sync takes only the last ERP_SYNC_SINCE_DAYS of GRNs and swaps the
    # whole snapshot, so a short or a rejection older than that leaves the
    # mirror though nothing closed it: the screen must say where the data
    # starts, or an early window reads as "nothing short or rejected".
    keeps = current_app.config.get("ERP_SYNC_SINCE_DAYS") or None
    return rows, {"synced_at": _utc_naive(synced), "oldest_grn_date": _as_date(oldest),
                  "newest_grn_date": _as_date(newest), "keeps_days": keeps, "rows": count or 0}


def erp_line_rejections(first, last):
    """Line rejections dated `first`..`last` — punched in ERP
    (termsdb.line_rejection_report), the second type of rejection besides the
    SRV rejection on the GRN line. None while the app cannot read them: there
    is no mirror of that table (erp_sync_service) and the ERP database is not
    reachable from the app yet. When it is, this returns the window's rows with
    erp_grn_lines' keys — the GRN the rejected lot came in on, reject_qty the
    qty rejected on the line — and the Line Rejection type lists them as SRV
    Rejection lists its lines; nothing else needs to change."""
    return None


def erp_closures(kind, rows):
    """The ONE read of whether a line of `kind` is closed. `rows` are the
    source's rows, so a reader can ask ERP about those GRNs only.

    For a rejection (srv, line): the rejection invoice ERP made and sent
    against the GRN line ("rejection invoice bna kar bhej denge tab" — the
    store, 17 Oct 2026). For a short: how ERP settled it (ERP will tell later).

    Returns {grn_line_key(grn_no, item_cd): {"ref_no", "ref_date", "how"}}
    — ref_no the invoice (or settlement) no., ref_date a date, an ISO text or
    None, `how` optional (default KINDS[kind]["how"]: "invoice" for a
    rejection) — or None while ERP gives the app no closures of that kind at
    all: today, for every kind, since neither the srv_report mirror nor any
    other ERP read carries them. When ERP (or its API) gives them, this
    function returns them and the lines close themselves."""
    return None


def grn_line_key(grn_no, item_cd):
    """A GRN line as ERP names it — its GRN no. and item code, trimmed and
    case-blind. Not the mirror's row id: the sync swaps the whole snapshot, so
    the id changes on every refresh."""
    return _key(grn_no), _key(item_cd)


# ---- small helpers ----

def plant_today():
    return (_utc_naive(utcnow()) + _local_offset()).date()


def _key(code):
    return (code or "").strip().upper()


def _q(value):
    return round(float(value), 2) if value is not None else None


def _num(value):
    """A quantity or amount as a cell: 2, not 2.0; a fraction to 2 places."""
    if value is None:
        return None
    value = round(float(value), 2)
    return int(value) if value.is_integer() else value


def _dmy(day):
    return day.strftime("%d.%m.%Y") if day else None


def _text(value):
    return (value or "").strip() or None


def _as_date(value):
    """A GRN date from an aggregate: a date, or on SQLite possibly its text."""
    if value is None or isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _kind(kind):
    key = (kind or "").strip()
    if key not in KINDS:
        raise BadRequestError(f"kind must be one of {', '.join(KINDS)}")
    return key


def _quantities(row):
    """(challan, received, short, rejected) of one ERP GRN line, to the
    register's two places; short and rejected are 0 when not above 0."""
    challan, received = _q(row["challan_qty"]), _q(row["receipt_qty"])
    short = round(challan - received, 2) if challan is not None and received is not None else 0.0
    rejected = _q(row["reject_qty"]) or 0.0
    return challan, received, max(short, 0.0), max(rejected, 0.0)


def _kind_qty(kind, row):
    """The one figure of `kind` on a row: its short qty, or its rejected qty."""
    _, _, short, rejected = _quantities(row)
    return short if kind == "short" else rejected


def age_band(days):
    """The AGE_BANDS key of a number of days pending."""
    for key, low, high in AGE_BANDS:
        if days >= low and (high is None or days <= high):
            return key
    return AGE_BANDS[0][0]   # a GRN dated after today (a clock difference) is not overdue


def _dates(first, last):
    return _dmy(first) if first == last else f"{_dmy(first)} to {_dmy(last)}"


def _period(period, date_from, date_to):
    """The PERIODS key asked for: the given one, else Custom when a date is
    given (a link made before the Period dropdown), else DEFAULT_PERIOD."""
    key = (period or "").strip()
    if not key:
        return "custom" if date_from or date_to else DEFAULT_PERIOD
    if key not in _PERIOD:
        raise BadRequestError(f"period must be one of {', '.join(_PERIOD)}")
    return key


def _window(period, date_from, date_to, today):
    """(first, last) GRN dates of the window, `today` being the plant's. A
    period other than Custom ignores From - To. Custom: To defaults to today
    and is cut there (no GRN is dated later); From defaults to DEFAULT_DAYS
    before To."""
    if period in ("30d", "90d"):
        first, last = today - timedelta(days=int(period[:-1]) - 1), today
    elif period == "month":
        first, last = today.replace(day=1), today
    elif period != "custom":
        first, last = preset_window(period, today)
    else:
        last = min(_parse_day(date_to, "to") or today, today)
        first = _parse_day(date_from, "from") or last - timedelta(days=DEFAULT_DAYS - 1)
    if first > last:
        raise BadRequestError("'from' must be on or before 'to' (and not after today)")
    span = (last - first).days + 1
    if span > MAX_DAYS:
        raise BadRequestError(f"The window is {span} days; the most is {MAX_DAYS} (5 years)")
    return first, last


def _closure(found, kind):
    """(how, reference no., date or None) of a line ERP reports closed, or None
    while it is open. A report with no reference no. closes nothing."""
    if not isinstance(found, dict):
        return None
    no = _text(str(found.get("ref_no") or ""))
    if not no:
        return None
    how = _text(str(found.get("how") or "")) or KINDS[kind]["how"]
    return how, no, _as_date(found.get("ref_date"))


class _Group:
    """One supplier's (or the whole window's) lines of one kind, summed. A line
    with no rate adds no money and is counted in `unrated`; a group whose lines
    all lack a rate has a value of None (unknown — blank, never 0), a group
    with no line has 0. The oldest is of the OPEN lines: a closed line is no
    longer pending, however long it took."""

    def __init__(self, code=None, name=None):
        self.code, self.name = code, name
        self.lines = self.unrated = self.open = self.rated = 0
        self.qty, self.money = 0.0, 0.0
        self.oldest, self.latest = None, None

    def add(self, line):
        self.lines += 1
        self.name = self.name or line["supplier_name"]
        self.qty += line["qty"]
        if line["value"] is None:
            self.unrated += 1
        else:
            self.money += line["value"]
            self.rated += 1
        if line["status"] == "open":
            self.open += 1
            self.oldest = max(self.oldest or 0, line["days_pending"])
        if line["grn_date"] and (self.latest is None or line["grn_date"] > self.latest):
            self.latest = line["grn_date"]

    @property
    def value(self):
        if self.lines and not self.rated:
            return None
        return round(self.money, 2)

    def as_dict(self):
        return {
            "supplier_code": self.code, "supplier_name": self.name, "lines": self.lines,
            "qty": round(self.qty, 2), "value": self.value,
            "open_lines": self.open, "closed_lines": self.lines - self.open,
            "oldest_days": self.oldest, "latest_grn_date": self.latest, "unrated_lines": self.unrated,
        }


def _coverage(kind, first, last, as_of):
    """How much of the window `first`..`last` the ERP GRN data covers, from
    data_as_of `as_of`: (not_held, not_synced, known), each {"from", "to"}
    (ISO days) or None.

    not_held: the part before the oldest GRN the register holds — the sync
    keeps only its last ERP_SYNC_SINCE_DAYS. not_synced: the part after the
    last day the data reaches (data_until) — on the live data, last synced
    28 Sep 2026, every period that runs into October. No line can show in
    either, so the screen and the files say nothing is known there rather
    than let it read as "nothing short" / "nothing rejected". known: what is
    left between — None when nothing is, or when no ERP data was ever read.
    Line Rejection is not read from the register: its window is taken whole."""
    whole = {"from": first.isoformat(), "to": last.isoformat()}
    if kind == "line":
        return None, None, whole
    held, until = _as_date(as_of["held_from"]), _as_date(as_of["data_until"])
    if until is None:
        return None, whole, None
    start, end, not_held, not_synced = first, last, None, None
    if held and first < held:
        start = held
        not_held = {"from": first.isoformat(), "to": min(held - timedelta(days=1), last).isoformat()}
    if last > until:
        end = until
        not_synced = {"from": max(first, until + timedelta(days=1)).isoformat(), "to": last.isoformat()}
    known = {"from": start.isoformat(), "to": end.isoformat()} if start <= end else None
    return not_held, not_synced, known


def _lot(row):
    """The lot a line came in with — its gate entry no. and date — as one text."""
    no, day = _text(row.get("gate_entry_no")), row.get("gate_entry_dt")
    if not no:
        return None
    return f"{no} dt {_dmy(day)}" if day else no


def _challan(row):
    no, day = _text(row.get("challan_no")), row.get("challan_dt")
    if not no:
        return _dmy(day)
    return f"{no} dt {_dmy(day)}" if day else no


def _iso(value):
    return value.isoformat() if value else None


class ShortRejectionService:

    def report(self, store_code, kind=None, period=None, date_from=None, date_to=None, supplier=None, part=None,
               part_id=None, today=None):
        """ONE kind (KINDS) on the store's ERP GRN lines in the window: {kind,
        label, available, window, filters, title, lines, summary, totals,
        data_as_of, not_given, unmatched, closures_known, basis, ...}. Lines and
        suppliers are worst first — open before closed, the longest pending,
        then the dearest. `available` is False while the kind's source cannot
        be read (Line Rejection today): then there is no list at all, which is
        not "nothing rejected"."""
        kind = _kind(kind)
        today = today or plant_today()
        period = _period(period, date_from, date_to)
        first, last = _window(period, date_from, date_to, today)
        picked = self._part(store_code, part_id) if part_id not in (None, "") else None
        rows, as_of = erp_grn_lines(first, last)
        if kind == "line":
            rows = erp_line_rejections(first, last)
        available = rows is not None
        rows = [row for row in rows or () if _kind_qty(kind, row)]
        closures = erp_closures(kind, rows)
        parts = self._store_parts(store_code)

        lines = []
        for row in rows:
            p = parts.get(_key(row["item_cd"]))
            if p is None or (picked is not None and p["part_id"] != picked.id):
                continue
            closure = closures.get(grn_line_key(row["grn_no"], row["item_cd"])) if closures else None
            line = self._line(kind, row, p, today, closure)
            if line is not None:
                lines.append(line)
        lines = self._by_part_text(self._by_supplier(lines, supplier), part)
        lines.sort(key=lambda ln: (ln["status"] != "open", -(ln["days_pending"] or 0), -(ln["value"] or 0),
                                   ln["grn_no"] or "", ln["part_no"] or ""))

        groups, total = {}, _Group()
        for line in lines:
            key = line["supplier_code"] or line["supplier_name"] or ""
            groups.setdefault(key, _Group(line["supplier_code"]))
            groups[key].add(line)
            total.add(line)
        summary = sorted((g.as_dict() for g in groups.values()),
                         key=lambda s: (-(s["oldest_days"] or 0), -(s["value"] or 0),
                                        (s["supplier_name"] or s["supplier_code"] or "").lower()))
        totals = total.as_dict()
        del totals["supplier_code"], totals["supplier_name"]
        totals["suppliers"] = len(summary)
        totals["age_bands"] = {key: 0 for key, _, _ in AGE_BANDS}
        for line in lines:
            if line["age_band"]:
                totals["age_bands"][line["age_band"]] += 1
        # The parts of the window the register does not cover — before its
        # oldest GRN, after its last refresh: no line can show there, so the
        # screen and the files say so rather than let them read as "nothing
        # short" / "nothing rejected" (_coverage).
        data_as_of = self._as_of(as_of, today)
        not_held, not_synced, known = _coverage(kind, first, last, data_as_of)
        k = KINDS[kind]

        return {
            "store_code": store_code,
            "kind": kind,
            "label": k["label"],
            "kinds": [{"key": key, "label": spec["label"]} for key, spec in KINDS.items()],
            "columns": {"qty": k["qty"], "value": k["value"]},
            # False while the kind's source cannot be read at all — the screen
            # then says why instead of "nothing in this window".
            "available": available,
            "erp_table": LINE_REJECTION_TABLE if kind == "line" else "srv_report",
            "window": {"from": first.isoformat(), "to": last.isoformat(), "days": (last - first).days + 1,
                       "period": period, "label": _PERIOD[period],
                       "default_days": DEFAULT_DAYS, "not_held": not_held, "not_synced": not_synced,
                       "known": known},
            "filters": {"supplier": _text(supplier),
                        "part": _text(part),
                        "part_id": picked.id if picked else None,
                        "part_no": picked.part_no if picked else None,
                        "part_name": picked.part_name if picked else None},
            "title": self._title(kind, first, last, supplier, picked.part_no if picked else _text(part)),
            "lines": lines,
            "summary": summary,
            "totals": totals,
            "data_as_of": data_as_of,
            "not_given": [label for field, label in ERP_FIELDS.items() if lines and not any(ln[field] for ln in lines)],
            "unmatched": (None if not available or _text(supplier) or _text(part) or picked is not None
                          else self._unmatched(rows)),
            # Whether ERP reports closures of this kind at all: until it does,
            # "open" means "not known to be closed", and the screen says so.
            "closures_known": closures is not None,
            "age_bands": [{"key": key, "from": low, "to": high} for key, low, high in AGE_BANDS],
            "periods": [{"key": key, "label": label} for key, label in PERIODS],
            "basis": {"what": k["what"], "days": k["days"], "closed": k["closed"], "no_closure": k["no_closure"],
                      **BASIS},
        }

    # ---- reads ----

    @staticmethod
    def _part(store_code, part_id):
        try:
            pid = int(part_id)
        except (TypeError, ValueError):
            raise BadRequestError("part_id must be a whole number") from None
        p = Part.query.filter(Part.id == pid, Part.store_code == store_code).first()
        if p is None:
            raise NotFoundError(f"Part {pid} is not in store '{store_code}'")
        return p

    @staticmethod
    def _store_parts(store_code):
        """{ITEM CODE: part facts} of every part of the store, active ones
        winning a clash: a rejection is the supplier's, whether or not the part
        is still in use."""
        rows = (db.session.query(Part.id, Part.part_no, Part.part_name, Part.unit_price, Part.is_active)
                .filter(Part.store_code == store_code)
                .order_by(Part.is_active.asc(), Part.id.asc()))
        out = {}
        for pid, part_no, name, rate, _active in rows:
            if _key(part_no):
                out[_key(part_no)] = {"part_id": pid, "part_no": part_no, "part_name": name,
                                      "rate": float(rate) if rate else None}
        return out

    @staticmethod
    def _unmatched(rows):
        """The window's lines of the kind (`rows`, already of that kind) whose
        item code is no store's Part No — no store lists them, so without this
        count they would vanish without a word (in Oct 2026 an 80-qty short of
        a Schneider element was one): {lines, item_count, items: the first
        UNMATCHED_ITEMS codes, the most lines first}. Only for an unfiltered
        window: a supplier or a part filter narrows the list to lines that have
        a part."""
        known = {_key(no) for (no,) in db.session.query(Part.part_no).filter(Part.part_no.isnot(None))}
        items = {}
        for row in rows:
            code = _key(row["item_cd"])
            if code and code in known:
                continue
            # Named by the matching key — trimmed, upper case — so one code
            # typed two ways in ERP is one item.
            item = items.setdefault(code, {"item_cd": code or None, "item_desc": _text(row["item_desc"]),
                                           "lines": 0})
            item["lines"] += 1
        ranked = sorted(items.values(), key=lambda i: (-i["lines"], i["item_cd"] or ""))
        return {"lines": sum(i["lines"] for i in ranked), "item_count": len(ranked),
                "items": ranked[:UNMATCHED_ITEMS]}

    @staticmethod
    def _line(kind, row, part, today, closure=None):
        """One GRN line as the `kind` tab shows it — that kind's qty and value
        only — or None when, rounded to the register's two places, nothing of
        that kind is on it. `closure` is what erp_closures reports for it:
        while there is none the line is open and its days run to today;
        closed, they stop at the reference date (unknown when ERP gives no
        date) and it has no age band."""
        challan, received, short, rejected = _quantities(row)
        qty = short if kind == "short" else rejected
        if not qty:
            return None
        rate = part["rate"]
        grn_date = row["grn_date"]
        closed = _closure(closure, kind)
        if closed is None:
            days = max((today - grn_date).days, 0) if grn_date else 0
            status, status_text, band = "open", KINDS[kind]["open"], age_band(days)
            ref_no = ref_date = None
        else:
            how, ref_no, ref_date = closed
            days = max((ref_date - grn_date).days, 0) if ref_date and grn_date else None
            status, band = "closed", None
            status_text = f"Closed — {how} {ref_no}" + (f" ({_dmy(ref_date)})" if ref_date else "")
        return {
            "id": row["id"],
            "grn_date": _iso(grn_date),
            "grn_no": _text(row["grn_no"]),
            "gate_entry_no": _text(row.get("gate_entry_no")),
            "gate_entry_dt": _iso(row.get("gate_entry_dt")),
            "lot": _lot(row),
            "challan_no": _text(row.get("challan_no")),
            "challan_dt": _iso(row.get("challan_dt")),
            "po_no": _text(row.get("po_no")),
            "supplier_code": _text(row["ven_cd"]),
            "supplier_name": _text(row["ven_name"]),
            "part_id": part["part_id"], "part_no": part["part_no"],
            "description": _text(row["item_desc"]) or part["part_name"],
            "challan_qty": challan, "received_qty": received, "accepted_qty": _q(row["accept_qty"]),
            "qty": qty,
            "rate": rate,
            "value": round(qty * rate, 2) if rate else None,
            "days_pending": days,
            "age_band": band,
            "status": status,
            "status_text": status_text,
            "closed_ref": ref_no,
            "closed_date": _iso(ref_date),
        }

    @staticmethod
    def _by_part_text(lines, text):
        """The lines whose Part No or description holds the typed text, case-
        blind; a blank box keeps every line."""
        text = (text or "").strip().lower()
        if not text:
            return lines
        return [ln for ln in lines
                if text in (ln["part_no"] or "").lower() or text in (ln["description"] or "").lower()]

    @staticmethod
    def _by_supplier(lines, supplier):
        """The lines of one supplier: its code exactly when one line has that
        code (a click on the supplier-wise table), else every line whose code
        or name contains the text (a typed search)."""
        text = (supplier or "").strip().lower()
        if not text:
            return lines
        exact = [ln for ln in lines if (ln["supplier_code"] or "").lower() == text]
        if exact:
            return exact
        return [ln for ln in lines
                if text in (ln["supplier_code"] or "").lower() or text in (ln["supplier_name"] or "").lower()]

    @staticmethod
    def _title(kind, first, last, supplier, part):
        """The title, also the file's name: the kind ("Short", "SRV
        Rejection"), then each filter that narrows it (the supplier, the part
        text) — a file never passes for more than it holds — then the dates:
        "SRV Rejection 19.07.2026 to 17.10.2026"."""
        picked = [_text(supplier), _text(part)]
        return " ".join([KINDS[kind]["label"], *(p for p in picked if p), _dates(first, last)])

    @staticmethod
    def _as_of(as_of, today):
        """How current the ERP data is: its last refresh as a plant day, and
        whether that is more than STALE_AFTER_DAYS old (or never happened)."""
        synced = as_of["synced_at"]
        day = (synced + _local_offset()).date() if synced else None
        age = max((today - day).days, 0) if day else None
        newest, oldest = as_of["newest_grn_date"], as_of["oldest_grn_date"]
        # The last day the data reaches: the day of the last refresh — a GRN
        # booked in ERP after it is not here — or a later GRN date, should a
        # clock difference give one. None: no ERP data was ever read.
        until = max((d for d in (day, newest) if d), default=None)
        return {
            "source": "ERP GRN register (srv_report), mirrored read-only",
            "synced_at": _iso(synced),
            "synced_day": _iso(day),
            "newest_grn_date": _iso(newest),
            "data_until": _iso(until),
            # The oldest GRN the data holds, and how far back the source keeps:
            # an open line on an older GRN is not listed (erp_grn_lines).
            "held_from": _iso(oldest),
            "keeps_days": as_of["keeps_days"],
            "rows": as_of["rows"],
            "age_days": age,
            "stale": age is None or age > STALE_AFTER_DAYS,
            "stale_after_days": STALE_AFTER_DAYS,
        }

    # ---- the files ----

    def export(self, store_code, fmt, kind=None, **filters):
        """(file, download name, mimetype): ONE kind of the window on screen as
        ONE Excel workbook (sheets "Supplier-wise" and "Lines") or a PDF, named
        by the title — "Short DD.MM.YYYY to DD.MM.YYYY", "SRV Rejection ..." —
        with the notes the screen shows beside it: a file passed on has no
        screen around it. A kind whose source cannot be read has no file."""
        if fmt not in FORMATS:
            raise BadRequestError(f"format must be one of {', '.join(FORMATS)}")
        out = self.report(store_code, kind, **filters)
        if not out["available"]:
            raise BadRequestError(f"{out['label']}s are not read from ERP yet - there is nothing to download")
        ext, mime = FORMATS[fmt]
        buf = io.BytesIO(workbook(out)) if fmt == "excel" else _pdf(out)
        return buf, f"{_file_safe(out['title'])}.{ext}", mime


def file_notes(out):
    """The lines a file carries — what its figures cannot say themselves."""
    as_of, basis = out["data_as_of"], out["basis"]
    notes = []
    if as_of["synced_day"]:
        notes.append(f"ERP data up to {_dmy(date.fromisoformat(as_of['synced_day']))}"
                     + (" - not current" if as_of["stale"] else ""))
    else:
        notes.append("No ERP GRN data has been read yet")
    if as_of["held_from"] and out["kind"] != "line":
        keeps = f" (the ERP sync keeps the last {as_of['keeps_days']} days)" if as_of["keeps_days"] else ""
        notes.append(f"ERP GRN lines are held from {_dmy(date.fromisoformat(as_of['held_from']))}{keeps}: "
                     f"an open {KINDS[out['kind']]['noun']} on an older GRN is not listed.")
    gap = out["window"]["not_held"]
    if gap:
        notes.append(f"No GRN line is held for {_dates(date.fromisoformat(gap['from']), date.fromisoformat(gap['to']))}"
                     " - nothing can be listed for those dates.")
    # After the last ERP refresh nothing is known: a file of an empty October
    # must not pass for "nothing short / rejected in October". With no ERP
    # data at all the first note has said so already.
    late = out["window"].get("not_synced")
    if late and as_of.get("data_until"):
        notes.append(f"ERP data ends {_dmy(date.fromisoformat(as_of['data_until']))}: nothing is known for "
                     f"{_dates(date.fromisoformat(late['from']), date.fromisoformat(late['to']))} - a GRN line "
                     "booked in ERP on those dates is not here until the ERP sync runs again.")
    notes.append(basis["days"])
    notes.append(basis["closed"] if out.get("closures_known") else basis["no_closure"])
    notes.append(basis["what"])
    unrated = out["totals"]["unrated_lines"]
    notes.append(basis["value"] + (f" {unrated} line{'s have' if unrated != 1 else ' has'} no rate."
                                   if unrated else ""))
    if out["not_given"]:
        notes.append("Not given by ERP yet: " + ", ".join(out["not_given"]) + ".")
    unmatched = out.get("unmatched")
    if unmatched and unmatched["lines"]:
        codes = ", ".join(i["item_cd"] or "(no code)" for i in unmatched["items"])
        more = f" and {unmatched['item_count'] - len(unmatched['items'])} more" \
            if unmatched["item_count"] > len(unmatched["items"]) else ""
        n = unmatched["lines"]
        notes.append(f"{n} {KINDS[out['kind']]['noun']} ERP GRN line{'s are' if n != 1 else ' is'} for item codes that "
                     f"are no store's Part No, so no store lists {'them' if n != 1 else 'it'}: {codes}{more}.")
    return notes


def nothing_text(out):
    """What a file with no line says in place of its table: "nothing" only for
    the dates the ERP data covers (window.known)."""
    k, w = KINDS[out["kind"]], out["window"]
    known = w.get("known")
    if not known:
        return "No GRN line is held for these dates - nothing can be listed."
    if (known["from"], known["to"]) == (w["from"], w["to"]):
        return k["nothing"]
    return f"{k['none']} on the GRN lines dated {_dates(date.fromisoformat(known['from']), date.fromisoformat(known['to']))}."


def supplier_rows(out):
    """The supplier-wise sheet's rows, keyed by column heading, the Total last."""
    cols = out["columns"]

    def row(s, code, name):
        return {"Supplier Code": code, "Supplier": name, "Lines": s["lines"],
                cols["qty"]: _num(s["qty"]), cols["value"]: _num(s["value"]),
                "Oldest (days)": s["oldest_days"],
                "Latest GRN": _dmy(date.fromisoformat(s["latest_grn_date"])) if s["latest_grn_date"] else None}
    rows = [row(s, s["supplier_code"], s["supplier_name"]) for s in out["summary"]]
    t = out["totals"]
    return rows + [row(t, TOTAL_LABEL, f"{t['suppliers']} supplier{'' if t['suppliers'] == 1 else 's'}")]


def line_rows(out):
    """The lines sheet's rows, keyed by column heading, in the screen's order."""
    cols = out["columns"]
    return [{
        "GRN Date": _dmy(date.fromisoformat(ln["grn_date"])) if ln["grn_date"] else None,
        "GRN No": ln["grn_no"], "Lot (Gate Entry)": ln["lot"],
        "Challan": _challan({"challan_no": ln["challan_no"],
                             "challan_dt": date.fromisoformat(ln["challan_dt"]) if ln["challan_dt"] else None}),
        "PO No": ln["po_no"], "Supplier Code": ln["supplier_code"], "Supplier": ln["supplier_name"],
        "Part No": ln["part_no"], "Description": ln["description"],
        "Challan Qty": _num(ln["challan_qty"]), "Received Qty": _num(ln["received_qty"]),
        "Accepted Qty": _num(ln["accepted_qty"]), cols["qty"]: _num(ln["qty"]), "Rate": _num(ln["rate"]),
        cols["value"]: _num(ln["value"]),
        "Days Pending": ln["days_pending"], "Age": ln["age_band"], "Status": ln["status_text"],
    } for ln in out["lines"]]


def _sheet(ws, columns, rows, notes, total_last=False, paint=None, legend=None):
    """One table laid out as excel_io.build_export lays every export out —
    header in row 1, frozen and filterable, bordered, sized, each text through
    the formula guard — money in Indian grouping, the Total row bold, then the
    legend and the notes in grey italics after a blank row (written after the
    columns are sized, so a long note never widens column A)."""
    ws.append(list(columns))
    for r in rows:
        _append_safe_row(ws, [r.get(c) for c in columns])
    format_table(ws)
    at = {c: i for i, c in enumerate(columns, start=1)}
    for number in range(2, len(rows) + 2):
        for c in MONEY_COLUMNS:
            if c in at:
                ws.cell(row=number, column=at[c]).number_format = MONEY
        if paint is not None:
            paint(ws, number, rows[number - 2], at)
    if total_last and rows:
        for cell in ws[len(rows) + 1]:
            mark_cell(cell, TOTAL)
    line = len(rows) + 3
    for text in ([legend] if legend else []) + list(notes):
        ws.cell(row=line, column=1, value=text).font = NOTE_FONT
        line += 1


def _paint_age(ws, number, row, at):
    tint = AGE_TINTS.get(row.get("Age"))
    if tint:
        mark_cell(ws.cell(row=number, column=at["Age"]), *tint)


def workbook(out):
    """The bytes of the one workbook of ONE kind that the screen downloads:
    "Supplier-wise" (with its Total) and "Lines" (the Age cell coloured)."""
    wb = Workbook()
    notes = file_notes(out)
    ws = wb.active
    ws.title = SUPPLIER_SHEET
    _sheet(ws, supplier_columns(out["kind"]), supplier_rows(out), [out["title"], *notes], total_last=True)
    _sheet(wb.create_sheet(LINES_SHEET), line_columns(out["kind"]), line_rows(out), [out["title"], *notes],
           paint=_paint_age, legend=AGE_LEGEND)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def attachment(out):
    """The workbook as a mail attachment, named as the screen's download."""
    return {"filename": f"{_file_safe(out['title'])}.xlsx", "mime": XLSX_MIME, "content": workbook(out)}


def _pdf_table(columns, rows, room, bold_last=False):
    data = [list(columns)] + [["" if r.get(c) is None else str(r.get(c)) for c in columns] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(pdf_io.STYLE)
    if table.wrap(room, 10_000)[0] > room:
        table = pdf_io._fitted_table(data, room)
    if bold_last and rows:
        table.setStyle(TableStyle([("FONTNAME", (0, len(rows)), (-1, len(rows)), "Helvetica-Bold")]))
    return table


def _pdf(out):
    """The two tables as pdf_io draws every report — same look, on narrow
    margins, fitted when wider still — under the title and the notes. Text is
    escaped: a "&" in a supplier's name is a character, not markup."""
    buf = io.BytesIO()
    # Eighteen columns never fit the inch margins: narrow ones from the start,
    # so the two tables share one page width.
    doc = pdf_io._document(buf, pdf_io.FIT_MARGIN)
    room = pdf_io._room(doc)
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(out["title"]), styles["Heading2"])]
    story += [Paragraph(escape(note), NOTE_STYLE) for note in file_notes(out)]
    story += [Spacer(1, 8), Paragraph(SUPPLIER_SHEET, styles["Heading3"]),
              _pdf_table(supplier_columns(out["kind"]), supplier_rows(out), room, bold_last=True),
              Spacer(1, 12), Paragraph(LINES_SHEET, styles["Heading3"])]
    rows = line_rows(out)
    story.append(_pdf_table(line_columns(out["kind"]), rows, room) if rows
                 else Paragraph(escape(nothing_text(out)), NOTE_STYLE))
    doc.build(story)
    buf.seek(0)
    return buf
