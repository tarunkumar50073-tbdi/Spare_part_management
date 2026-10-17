"""The Issuance Report and the Consumption Report: what left the store and what
each department really used, cut the way the store asks, on screen and as Excel
/ PDF — the one place for both ("wahi mil jaaye app me, wahan select kare kis
part ki kab se kab tak consumption ya month wise ya year wise consumption mil
jaaye, alag se report me aane ki jarurat nhi, wahi se excel / pdf download ho
jaaye" — the store, 14 Oct 2026). The six consumption downloads the Reports
page had moved here. Issuance joined on 17 Oct 2026 as downloads only, never a
mail ("issuance and consumption ka bna de but mail nhi karenge bas excel/pdf
download ka option"), and the same day the two became two reports on two
screens ("issuance and consumption ki report alag alag chahiye").

Each report (ConsumptionTableService.REPORTS) names its views, their columns,
whose department it counts an issue under and its Full pack. The reading, the
builders, the notes and the files are shared, so the two add up the same way:
the Issuance Report's Qty Issued and Value over a range are the Consumption
Report's.

ISSUANCE REPORT — what left the store, when, to whom:
  * every ISSUE ledger row in the range, cancelled ones aside, GROSS: what came
    back later is not netted off here (what was really used is the
    Consumption Report's);
  * Department = the department written on the issue — who took it out, often
    the store's own people —, "Not recorded" when none was written;
  * figures: Issues, Parts, Qty Issued, Value.
  Views: list (every issue line, oldest first, a subtotal after each day,
  month or year — picked from the range's length, auto_group() — and
  cancelled issues listed, never counted), day, week, month, year, fy, part,
  department (taken by), person (issued to: the person, code and mobile).

CONSUMPTION REPORT — what each department really used:
  * an ISSUE ledger row in the range, cancelled ones aside;
  * counted for the department that owns the part's section (consumers()),
    never the one typed on the issue ("consumption maintenance and tool room
    karega store nhi", "store to material issue karta h" — the store, 14 Oct
    2026);
  * what came back against an issue counts against the issue — its day, week,
    month, year, part, machine — whenever it came back (_returned_qty_map), so
    Net Qty is what was really used;
  * figures: Issues, Parts, Qty Issued, Value, Qty Returned, Net Qty, Net Value.
  Views: day, week, month, year, fy, part, department (each consuming
  department, one with nothing consumed too unless the table is narrowed to a
  part, a section or a type), machine (the line and machine recorded at
  installation, named as the By machine chart names them, place_of), category
  (sub_category).

The period views: Date ("10.03.2026", only days with issues), Week (Monday to
Sunday, cut at the range's ends), Month ("Oct 2026"), Year (the calendar year)
and Financial Year ("FY 2026-27": 1 April to 31 March), one row per period
(the report's Department split is Department-wise's). Money is at each part's
CURRENT rate: the ledger keeps no rate per issue. Every view ends with a Total
row over every counted issue in it. The Full pack (pack()) is one workbook of the range: a Summary and the
report's sheets, all from one read, so every sheet's Total is the same.

What the figures cannot say is said beside them — on screen (`window`,
`unrated_issues`) and as note lines in the Excel / PDF:
  * days before the app's ledger (ledger_since) were never recorded: a range
    reaching back past it is counted in part, not a measured zero there;
  * an issue whose part has no rate adds nothing to Value / Net Value, so the
    money covers rated issues only, and how many are not is counted;
  * a cancelled issue on the Issue list is listed, not counted.

The filters are the screen's — part, section, the report's department, part
type — and the range is the screen's too (ConsumptionAnalyticsService._window:
the last 30 days by default, the end cut at today, at most MAX_DAYS = 5 years
so a year-wise table has room), counted in plant-clock days (window()). The
screen's Yesterday / Last month / financial-year presets are cut here, on the
plant clock (preset_window()), and handed to it with the filter lists.

Callers from before the split name no report: they get the Consumption Report,
except for the Issue list, which was issuance all along (_report()) — and
whose Department filter meant the consuming department then, so it still does
for them (Filters.by_consumer).
"""
import io
from collections import namedtuple
from datetime import date, datetime, time, timedelta
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
from sqlalchemy import func, or_

from app.extensions import db
from app.models.base import utcnow
from app.models.department import Department
from app.models.part import Part
from app.models.spare_stock import SpareStockLedger
from app.models.store import Store
from app.services.consumption_analytics_service import (
    DIE, PLACE_BASIS, ConsumptionAnalyticsService, _coverage, _utc_naive, consumer_clause, consumer_of,
    consumers, place_of, type_clause,
)
from app.services.reorder_mail_service import _file_safe
from app.services.report_service import FORMATS
from app.services.spare_report_service import (
    SpareReportService, _cells, _in, _issue_status, _line, _money, _stamp, plant_today, window,
)
from app.services.spare_stock_service import SpareStockService
from app.utils import pdf_io
from app.utils.excel_io import _append_safe_row, format_table
from app.utils.exceptions import BadRequestError

# One issue as the views read it: its ledger row (ISSUE_COLUMNS), its part, the
# department the report counts it under (the Consumption Report: the one that
# consumes it; the Issuance Report: the one written on it), the part's section,
# its plant day, qty issued, qty returned (always 0 on the Issuance Report,
# which does not net returns), and whether it was cancelled (listed on the
# Issue list only).
Use = namedtuple("Use", "issue part dept section day qty back cancelled")
# What the views read of an ISSUE row. Columns, not ORM rows: five years of
# issues (40,000 in the timing run) are read in half a second instead of three.
# The second line is what the Issue list and Person-wise show of each issue, in
# the Issue Register's words.
_L = SpareStockLedger
ISSUE_COLUMNS = (_L.id, _L.part_id, _L.qty_delta, _L.created_at, _L.line_name, _L.line_no, _L.machine,
                 _L.die_no, _L.confirmed_at.isnot(None).label("installed"),
                 _L.status, _L.actual_qty, _L.reference, _L.department, _L.issued_to, _L.issued_to_mobile,
                 _L.issued_to_code, _L.employee_id)
# The screen's filters: the Part and Section rows picked, the report's
# department and the part type typed (None = all). `by_consumer`: the
# department is the CONSUMING one though the report is Issuance — an Issue list
# asked for as before the split (17 Oct 2026), when the one screen's Department
# meant who consumed it in every view, the list too.
Filters = namedtuple("Filters", "part section department type by_consumer", defaults=(False,))
View = namedtuple("View", "key label columns build")
# A view built: its rows (the Total last), the Total's _Tally, and on the Issue
# list the rows that are a group's subtotal and the cancelled issues.
Built = namedtuple("Built", "rows total subtotals cancelled", defaults=((), ()))
# Everything one read of the range gives every view: the report, the counted
# issues, every issue listed (cancelled ones too), the report's departments in
# master order and the sort key they give, the range, whether its end was cut
# at today, the filters, the Issue list's grouping and the app ledger's first
# day.
Ctx = namedtuple("Ctx", "report uses listed depts order win clamped f group since")
# An Issue list grouping: period(plant day) -> (sort key, label).
Group = namedtuple("Group", "key label period")
# A report: its key, its name (the head of every title and file name), its
# views in the View dropdown's order, its Full pack's sheets, the "how it is
# counted" lines it carries (BASIS keys) and what its Department is called on
# the Full pack's Summary.
Report = namedtuple("Report", "key title views pack basis department")

# The figures: the Issuance Report's are gross; the Consumption Report's net
# what came back.
ISSUED = ("Issues", "Parts", "Qty Issued", "Value")
CONSUMED = ISSUED + ("Qty Returned", "Net Qty", "Net Value")

# The period views' own cells, both reports: one row per period, its figures
# beside it — the department split is Department-wise's (_by_period).
DATE, WEEK, MONTH, YEAR, FY = ("Date", "Day"), ("Week",), ("Month",), ("Year",), ("Financial Year",)

# The Consumption Report's columns. Department = the consuming department.
BY_DEPARTMENT_COLUMNS = ("Department", "Section") + CONSUMED
BY_DATE_COLUMNS = DATE + CONSUMED
BY_WEEK_COLUMNS = WEEK + CONSUMED
BY_MONTH_COLUMNS = MONTH + CONSUMED
BY_YEAR_COLUMNS = YEAR + CONSUMED
BY_FY_COLUMNS = FY + CONSUMED
BY_PART_COLUMNS = ("Part No", "Description", "Section", "Department", "UOM", "Issues", "Qty Issued",
                   "Qty Returned", "Net Qty", "Rate", "Value", "Net Value", "First Issued",
                   "Last Issued")
BY_MACHINE_COLUMNS = ("Department", "Line", "Machine") + CONSUMED
BY_CATEGORY_COLUMNS = ("Department", "Category") + CONSUMED

# The Issuance Report's. Department = the one written on the issue (who took
# it); a part is taken by many, so Part-wise has none.
ISSUED_BY_DATE_COLUMNS = DATE + ISSUED
ISSUED_BY_WEEK_COLUMNS = WEEK + ISSUED
ISSUED_BY_MONTH_COLUMNS = MONTH + ISSUED
ISSUED_BY_YEAR_COLUMNS = YEAR + ISSUED
ISSUED_BY_FY_COLUMNS = FY + ISSUED
ISSUED_BY_PART_COLUMNS = ("Part No", "Description", "Section", "UOM", "Issues", "Qty Issued", "Rate", "Value",
                          "First Issued", "Last Issued")
ISSUED_BY_DEPARTMENT_COLUMNS = ("Department",) + ISSUED
BY_PERSON_COLUMNS = ("Issued To", "Code", "Mobile", "Department") + ISSUED
# The Issue list, in the Issue Register's words: Department = the one written
# on the issue, Issued To = the person and the mobile.
LIST_COLUMNS = ("Date", "Time", "Issue No", "Reference", "Part No", "Description", "Section", "Department",
                "Issued To", "Qty Issued", "Rate", "Value", "Line / Machine / Die", "Status")

UNCATEGORISED = "Uncategorised"
# An issue with no department, or no person, written on it.
NOT_WRITTEN = "Not recorded"
# The last row of every view, named in its first column.
TOTAL = "Total"
# The first cell of an Issue list group's subtotal row: "Subtotal 10.03.2026".
SUBTOTAL = "Subtotal {}"

# The downloads the screen offers: format -> (file extension, mimetype).
TABLE_FORMATS = {fmt: FORMATS[fmt] for fmt in ("excel", "pdf")}

# The Issue list's subtotals follow the range's length (auto_group): a day
# each up to two months, a month each up to two years, a year each past that.
LIST_BY_DAY_DAYS = 62
LIST_BY_MONTH_DAYS = 731

# "How it is counted": each report carries its own lines (Report.basis).
BASIS = {
    "issuance": "Every issue in the range as it left the store, cancelled ones aside. Returns are not "
                "netted off: what came back, and so what was really used, is the Consumption Report's. "
                f"Department = the department written on the issue (who took it out); '{NOT_WRITTEN}' = "
                "none written.",
    "consumption": "Issues in the range, cancelled ones aside, each counted for the department that "
                   "owns the part's section, never the department written on the issue (the store only "
                   "hands parts out).",
    "returns": "What came back against an issue counts against that issue (its day, week, month, year, "
               "part and machine) whenever it came back. Net Qty = Qty Issued - Qty Returned: what was "
               "really consumed.",
    "value": "Qty x each part's current rate (valued at today's rate). An issue whose part has no rate "
             "adds nothing, so Value covers rated issues only (how many are not is said under the table); "
             "a row whose issues all lack a rate shows no value, not 0.",
    "period": "Plant-clock days. Weeks run Monday to Sunday, cut at the range's ends; Year-wise is the "
              "calendar year (January to December) and Financial year-wise the financial year (1 April to "
              "31 March, 'FY 2026-27'). Days before the app's ledger started were never recorded, so they "
              "are not a measured zero.",
    "list": "Issue list: every issue in the range, oldest first, with a subtotal after each day (a range "
            f"up to {LIST_BY_DAY_DAYS} days), month (up to 2 years) or year (longer). A cancelled issue is "
            "listed, marked Cancelled, with its quantity but no value, and is left out of every subtotal "
            "and total.",
    "person": "Person-wise: the person written on the issue, with the employee code and mobile kept at "
              f"issue time; one employee is one row. '{NOT_WRITTEN}' = no person written.",
    "place": PLACE_BASIS,
}

# The notes under the downloaded table: grey italics in the Excel, small grey
# lines under the title in the PDF.
NOTE_FONT = Font(italic=True, color="475569")
NOTE_STYLE = ParagraphStyle("ConsumptionNote", fontName="Helvetica", fontSize=8, leading=10,
                            textColor=colors.HexColor("#475569"))
# A group's subtotal and the Total, bold on a light tint; a cancelled issue grey.
SUBTOTAL_FILL = PatternFill("solid", fgColor="F1F5F9")
TOTAL_FILL = PatternFill("solid", fgColor="E2E8F0")
BOLD_FONT = Font(bold=True)
MUTED_FONT = Font(italic=True, color="94A3B8")
PDF_SUBTOTAL = colors.HexColor("#f1f5f9")
PDF_TOTAL = colors.HexColor("#e2e8f0")
PDF_MUTED = colors.HexColor("#94a3b8")

# The screen's ranges cut on the plant clock (the store, 17 Oct 2026: daily =
# yesterday, monthly = last month, yearly = the financial year, April - March).
PRESETS = ("yesterday", "last_month", "this_fy", "last_fy")
# Each report's Full pack sheets, in order: (sheet name, view); Year-wise is
# decided by the range (_year_view). A sheet name cannot hold "/", so the Line
# / Machine view's sheet says "&".
ISSUANCE_PACK = (("Date-wise", "day"), ("Month-wise", "month"), ("Year-wise", None), ("Part-wise", "part"),
                 ("Department-wise", "department"), ("Person-wise", "person"), ("Issue list", "list"))
CONSUMPTION_PACK = (("Date-wise", "day"), ("Month-wise", "month"), ("Year-wise", None), ("Part-wise", "part"),
                    ("Department-wise", "department"), ("Line & Machine-wise", "machine"),
                    ("Category-wise", "category"))
PACK = "Full pack"


def fy_start(day):
    """1 April of the financial year (1 April - 31 March) a day falls in."""
    return date(day.year if day.month >= 4 else day.year - 1, 4, 1)


def fy_label(day):
    """The financial year a day falls in, as the store writes it: FY 2026-27."""
    first = fy_start(day).year
    return f"FY {first}-{(first + 1) % 100:02d}"


def preset_window(key, today):
    """(first, last) plant day of one of the screen's presets, `today` being the
    plant's: yesterday; the whole previous calendar month; this financial year
    to today; the whole previous financial year (1 April - 31 March)."""
    if key == "yesterday":
        day = today - timedelta(days=1)
        return day, day
    if key == "last_month":
        last = today.replace(day=1) - timedelta(days=1)
        return last.replace(day=1), last
    if key == "this_fy":
        return fy_start(today), today
    if key == "last_fy":
        start = fy_start(today)
        return start.replace(year=start.year - 1), start - timedelta(days=1)
    raise BadRequestError(f"range must be one of {', '.join(PRESETS)}")


def presets(today=None):
    """{preset: {"from", "to"}} for the screen's filter bar, on the plant clock."""
    today = today or plant_today()
    out = {}
    for key in PRESETS:
        first, last = preset_window(key, today)
        out[key] = {"from": first.isoformat(), "to": last.isoformat()}
    return out


# The Issue list's groupings.
GROUPS = (
    Group("day", "Day", lambda d: (d, f"{d:%d.%m.%Y}")),
    Group("month", "Month", lambda d: ((d.year, d.month), f"{d:%b %Y}")),
    Group("year", "Year", lambda d: (d.year, f"{d.year}")),
    Group("fy", "Financial year", lambda d: (fy_start(d), fy_label(d))),
)


def auto_group(win):
    """The Issue list's subtotals for a range, in place of a toggle (the store,
    17 Oct 2026: "thoda simple ho"): by day up to LIST_BY_DAY_DAYS days, by
    month up to two years, by calendar year past that."""
    days = (win.last - win.first).days + 1
    return "day" if days <= LIST_BY_DAY_DAYS else "month" if days <= LIST_BY_MONTH_DAYS else "year"


def _year_view(win):
    """The Full pack's Year-wise view: financial years when the range holds a 1
    April (it reaches into a new financial year, or starts one), else calendar
    years."""
    april = date(win.first.year, 4, 1)
    if april < win.first:
        april = april.replace(year=april.year + 1)
    return "fy" if april <= win.last else "year"


def _dept_order(names, last=()):
    """Sort key for a department name: `names` in master order, then any other
    department (a section no department owns, parts with no section, a name
    no longer in the master) by name, then the `last` ones ("Not recorded")."""
    rank = {name.lower(): i for i, name in enumerate(names)}
    end = {name.lower() for name in last}
    return lambda name: (len(rank) + 1 if name.lower() in end else rank.get(name.lower(), len(rank)),
                         name.lower())


def _master_departments():
    """The active departments' names in master order."""
    return [d.name.strip() for d in Department.query.filter_by(is_active=True)
            .order_by(Department.sort_order, Department.name)]


def _written(i):
    """The department written on an issue — who took it out — or NOT_WRITTEN."""
    return (i.department or "").strip() or NOT_WRITTEN


def taker_clause(department):
    """SQL for "the issue was written to `department`" (case and spaces aside);
    NOT_WRITTEN also matches an issue with none."""
    name = department.strip().lower()
    match = func.lower(func.trim(SpareStockLedger.department)) == name
    if name == NOT_WRITTEN.lower():
        match = or_(match, SpareStockLedger.department.is_(None), func.trim(SpareStockLedger.department) == "")
    return match


def issue_departments(store_code):
    """The Issuance Report's Department choices: every department written on
    this store's issues — in master order, then any other name (an old or a
    switched-off one) by name, then NOT_WRITTEN when an issue carries none."""
    L = SpareStockLedger
    written = (db.session.query(func.trim(L.department)).join(Part, Part.id == L.part_id)
               .filter(Part.store_code == store_code, L.txn_type == "ISSUE").distinct())
    names, blank = {}, False
    for (name,) in written:
        if name:
            names.setdefault(name.lower(), name)
        else:
            blank = True
    master = _master_departments()
    spelt = {n.lower(): n for n in master}           # the master's spelling of a name typed otherwise
    out = sorted((spelt.get(key, name) for key, name in names.items()), key=_dept_order(master))
    return out + ([NOT_WRITTEN] if blank else [])


def _window(date_from, date_to):
    """(the range as a plant-clock Window, whether its end was cut at today):
    the screen's rules on the reports' plant clock."""
    first, last, clamped = ConsumptionAnalyticsService._window(date_from, date_to, plant_today())
    return window(first, last), clamped


def _parts(ids):
    """{id: Part} of these parts, a chunk at a time (SQLite's bound-parameter limit)."""
    ids, out = sorted(ids), {}
    for start in range(0, len(ids), 500):
        out.update((p.id, p) for p in Part.query.filter(Part.id.in_(ids[start:start + 500])))
    return out


def _dates(win):
    if win.first == win.last:
        return f"{win.first:%d.%m.%Y}"
    return f"{win.first:%d.%m.%Y} to {win.last:%d.%m.%Y}"


def _offset(win):
    """The plant clock the range was cut on (its UTC offset)."""
    return datetime.combine(win.first, time.min) - win.lo


def _plural(n, word):
    # Comma-grouped: the Indian and the western grouping agree below 1,00,000.
    return f"{n:,} {word}{'' if n == 1 else 's'}"


def _count(issues, cancelled):
    """An Issue list subtotal's count, in its Issue No cell: "3 issues", and
    "+ 1 cancelled" when the group lists one (not counted)."""
    return _plural(issues, "issue") + (f" + {cancelled} cancelled" if cancelled else "")


def _clean(v):
    return (v or "").strip() or None


def _issued_to(i):
    """Who the issue was handed to: the person and the mobile — the Issue
    Register's Issued To / Mobile."""
    return " · ".join(v for v in (_clean(i.issued_to), _clean(i.issued_to_mobile)) if v) or None


def _place(i):
    """Where the issue went, all the fitter recorded — line (_line), machine,
    die — in one cell; with none of them, what the Line / Machine view says
    instead (place_of): installed with no place given, or not recorded yet."""
    die = (i.die_no or "").strip()
    recorded = [v for v in (_line(i), (i.machine or "").strip(), DIE.format(die) if die else "") if v]
    return " / ".join(recorded) or place_of(None, None, None, None, i.installed)[1]


def _notes(out):
    """The lines a downloaded table carries under it — what its figures cannot
    say: the days before the app's ledger, never recorded, the issues left out
    of the money for want of a rate, and the cancelled issues the Issue list
    shows but does not count."""
    since, unrated = out["window"]["ledger_since"], out["unrated_issues"]
    gone = len(out.get("cancelled_rows") or ())
    notes = []
    if since is None:
        notes.append("The app's ledger has no rows yet; nothing in this range was recorded")
    elif out["window"]["from"] < since:
        notes.append(f"The app's ledger starts {datetime.fromisoformat(since):%d.%m.%Y}; "
                     "days before it were not recorded")
    if unrated:
        notes.append(f"Value covers rated issues only; {unrated} issue{'s have' if unrated != 1 else ' has'} "
                     "no rate")
    if gone:
        notes.append(f"{_plural(gone, 'cancelled issue')} listed in grey, left out of every subtotal and "
                     "the Total")
    return notes


def _marks(out):
    """(the rows drawn bold — each subtotal, then the Total —, the rows drawn
    grey: the cancelled issues), as indexes into out["rows"]."""
    return [*out.get("subtotal_rows", ()), len(out["rows"]) - 1], list(out.get("cancelled_rows", ()))


def _sheet(ws, columns, rows, notes, bold=(), muted=()):
    """The table on `ws` as excel_io.build_export lays every export out —
    header in row 1, frozen and filterable, each text through the formula
    guard — with the notes under the Total after a blank row: outside the
    filtered range, and written after the columns are sized, so a long note
    never widens column A. The `bold` rows (subtotals, the Total last) are bold
    on a tint, the `muted` ones grey."""
    ws.append(list(columns))
    for r in rows:
        _append_safe_row(ws, [r[c] for c in columns])
    format_table(ws)
    last = len(rows) - 1
    for i in bold:
        for cell in ws[i + 2]:
            cell.font, cell.fill = BOLD_FONT, TOTAL_FILL if i == last else SUBTOTAL_FILL
    for i in muted:
        for cell in ws[i + 2]:
            cell.font = MUTED_FONT
    for at, note in enumerate(notes, start=ws.max_row + 2):
        ws.cell(row=at, column=1, value=note).font = NOTE_FONT


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _excel(columns, rows, notes, bold=(), muted=()):
    """The table as one formatted sheet (_sheet)."""
    wb = Workbook()
    _sheet(wb.active, columns, rows, notes, bold, muted)
    return _save(wb)


def _pdf(title, columns, rows, notes, bold=(), muted=()):
    """The table as pdf_io.build_pdf_table draws every report — same look, on
    narrow margins when wide, fitted when wider still — with the notes under
    the title, ahead of the figures they qualify: the shared builder draws a
    title and a table only. The `bold` rows (subtotals, the Total last) are
    bold on a tint and measured so, the `muted` ones grey. Text is escaped: a
    "&" in a part no or a section name is a character, not markup."""
    buf = io.BytesIO()
    data = [list(columns)] + [["" if r[c] is None else str(r[c]) for c in columns] for r in rows]
    strong = [i + 1 for i in bold]                    # table rows: the header is row 0
    last = len(rows)
    look = TableStyle(
        [("FONTNAME", (0, r), (-1, r), pdf_io.BOLD) for r in strong]
        + [("BACKGROUND", (0, r), (-1, r), PDF_TOTAL if r == last else PDF_SUBTOTAL) for r in strong]
        + [("TEXTCOLOR", (0, i + 1), (-1, i + 1), PDF_MUTED) for i in muted])
    doc = pdf_io._document(buf, inch)
    table = Table(data, repeatRows=1)
    table.setStyle(pdf_io.STYLE)
    table.setStyle(look)
    width = table.wrap(pdf_io._room(doc), doc.height)[0]
    if width > pdf_io._room(doc):
        doc = pdf_io._document(buf, pdf_io.FIT_MARGIN)
        if width > pdf_io._room(doc):
            table = pdf_io._fitted_table(data, pdf_io._room(doc), bold=strong)
            table.setStyle(look)
    story = [Paragraph(escape(title), getSampleStyleSheet()["Heading2"])]
    story += [Paragraph(escape(note), NOTE_STYLE) for note in notes]
    doc.build(story + [Spacer(1, 8), table])
    buf.seek(0)
    return buf


class _Tally:
    """The figures of one group of counted issues. Money is at each part's
    CURRENT rate (the ledger keeps no rate per issue). An issue with no rate
    adds no money — counted in `unrated`, so the screen and the files can say
    the money covers rated issues only — and a group whose issues all lack a
    rate has Value / Net Value unknown (None), never 0; a group with no issue at
    all has 0. What came back against an issue is counted against the issue —
    its day, month, part, machine — whenever it came back, so Net Qty is what
    it really used (on the Issuance Report nothing is read as back: it is
    gross)."""

    def __init__(self, cells):
        self.cells, self.issues, self.parts = cells, 0, set()
        self.qty = self.back = 0.0
        self.value = self.net_value = None
        self.unrated = 0
        self.first = self.last = None

    @property
    def net(self):
        return round(self.qty - self.back, 2)

    def add(self, use):
        self.issues += 1
        self.parts.add(use.part.id)
        # Every sum kept to the 2 places its cell shows, so a view's order
        # compares what the reader sees: 0.1 + 0.2 ties with 0.3 (and falls
        # back to the name) instead of sorting ahead as 0.30000000000000004.
        self.qty = round(self.qty + use.qty, 2)
        self.back = round(self.back + use.back, 2)
        value = _money(use.qty, use.part.unit_price)
        if value is None:
            self.unrated += 1
        else:
            self.value = round((self.value or 0.0) + value, 2)
            # Net Qty x the same rate: 0 for an issue that all came back, not unknown.
            net_value = round((use.qty - use.back) * float(use.part.unit_price), 2)
            self.net_value = round((self.net_value or 0.0) + net_value, 2)
        self.first = min(self.first or use.day, use.day)
        self.last = max(self.last or use.day, use.day)

    def row(self, columns):
        """The group as a table row: its own cells, then its figures."""
        figures = {
            "Issues": self.issues, "Parts": len(self.parts), "Qty Issued": self.qty,
            "Value": self.value if self.issues else 0.0, "Qty Returned": self.back,
            "Net Qty": self.net, "Net Value": self.net_value if self.issues else 0.0,
            "First Issued": f"{self.first:%d.%m.%Y}" if self.first else None,
            "Last Issued": f"{self.last:%d.%m.%Y}" if self.last else None,
        }
        merged = {**figures, **self.cells}
        return {column: merged.get(column) for column in columns}


def _with_total(tallies, total, columns):
    """The groups' rows, then the Total row named in the first column, and the
    Total's _Tally, which the table reads its unrated count from."""
    total.cells = {columns[0]: TOTAL}
    return Built([t.row(columns) for t in tallies] + [total.row(columns)], total)


def _by_money(t):
    """Where the money went: Value first (Net Value on the Consumption Report),
    unrated groups after every rated one by quantity."""
    return t.net_value is None, -(t.net_value or 0), -t.net


class ConsumptionTableService:

    def table(self, store_code, view, date_from=None, date_to=None, part_id=None, section_id=None,
              department=None, part_type=None, group=None, report=None):
        """One view of one report over the range, rows keyed by column
        heading, the Total last. `report`: issuance or consumption (none: as
        before the split, _report()).

        `window` also says how much of the range the app's ledger covers: its
        first day (ledger_since, the charts' own figure) and the coverage —
        "none" (the range ends before it, or there is no ledger yet), "partial"
        (it starts inside the range) or "full". Days before it were never
        recorded, so an empty one there is not a measured zero.
        `unrated_issues`: the issues whose part has no rate, left out of Value
        and Net Value. On the Issue list (`group`: day, month, year or fy;
        none = from the range's length, auto_group()), `subtotal_rows` and
        `cancelled_rows` index the rows that are a group's subtotal and the
        cancelled issues it lists but does not count."""
        rep = self._report(report, view)
        chosen = self._view(rep, view)
        # No report named: a caller from before the split, whose Department
        # was the consuming one on the Issue list too — so it stays that.
        by_consumer = not (report or "").strip()
        c = self._read(store_code, rep, date_from, date_to, part_id, section_id, department, part_type, group,
                       by_consumer)
        return self._out(chosen, c, chosen.build(self, c, chosen.columns))

    def export(self, store_code, fmt, view, **filters):
        """(file, download name, mimetype) of the view on screen as Excel or
        PDF, named and titled by the report, the view, the filters and the
        dates, with the notes the screen shows beside it (_notes) — a file
        passed on has no screen around it to say them. Subtotals and the Total
        are bold, cancelled issues grey."""
        if fmt not in TABLE_FORMATS:
            raise BadRequestError(f"format must be one of {', '.join(TABLE_FORMATS)}")
        out = self.table(store_code, view, **filters)
        ext, mime = TABLE_FORMATS[fmt]
        notes, (bold, muted) = _notes(out), _marks(out)
        if fmt == "excel":
            buf = _excel(out["columns"], out["rows"], notes, bold, muted)
        else:
            buf = _pdf(out["title"], out["columns"], out["rows"], notes, bold, muted)
        return buf, f"{_file_safe(out['title'])}.{ext}", mime

    def pack(self, store_code, date_from=None, date_to=None, part_id=None, section_id=None, department=None,
             part_type=None, report=None):
        """(file, download name, mimetype) of a report's Full pack: ONE
        workbook of the range and filters — a Summary, then the report's sheets
        (Report.pack; Year-wise in financial years when the range holds a 1
        April, _year_view; the Issue list subtotalled as on screen) — every
        sheet from the same read, so their Totals agree. Named like "Issuance
        Report Full pack 01.04.2026 to 31.03.2027.xlsx", with the filters
        before the dates as the other files are."""
        rep = self._report(report)
        c = self._read(store_code, rep, date_from, date_to, part_id, section_id, department, part_type, None)
        year = _year_view(c.win)
        sheets = []
        for name, key in rep.pack:
            chosen = self._view(rep, key or year)
            sheets.append((name, self._out(chosen, c, chosen.build(self, c, chosen.columns))))
        wb = Workbook()
        self._summary(wb.active, store_code, c, sheets, year)
        for name, out in sheets:
            _sheet(wb.create_sheet(name), out["columns"], out["rows"], _notes(out), *_marks(out))
        title = self._named(f"{rep.title} {PACK}", c.f, c.win)
        return _save(wb), f"{_file_safe(title)}.xlsx", TABLE_FORMATS["excel"][1]

    # ---- what the screen picked ----

    @classmethod
    def _report(cls, key, view=None):
        """The report asked for. Callers from before the split (17 Oct 2026)
        name none: the Issue list was issuance all along, every other view
        consumption."""
        key = (key or "").strip()
        if not key:
            return cls.REPORTS["issuance" if (view or "").strip() == "list" else "consumption"]
        if key not in cls.REPORTS:
            raise BadRequestError(f"report must be one of {', '.join(cls.REPORTS)}")
        return cls.REPORTS[key]

    @staticmethod
    def _view(rep, key):
        chosen = next((v for v in rep.views if v.key == key), None)
        if chosen is None:
            raise BadRequestError(f"view must be one of {', '.join(v.key for v in rep.views)}")
        return chosen

    @staticmethod
    def _group(key, win):
        """The Issue list's grouping: the one asked for (links from before 17
        Oct 2026 carry it), else the range's own (auto_group)."""
        key = (key or "").strip() or auto_group(win)
        chosen = next((g for g in GROUPS if g.key == key), None)
        if chosen is None:
            raise BadRequestError(f"group must be one of {', '.join(g.key for g in GROUPS)}")
        return chosen

    @staticmethod
    def _filters(store_code, rep, part_id, section_id, department, part_type, by_consumer=False):
        # A part or section of another store is a 404, as on the charts.
        department = (department or "").strip() or None
        return Filters(
            ConsumptionAnalyticsService._part(store_code, part_id) if part_id is not None else None,
            ConsumptionAnalyticsService._section(store_code, section_id) if section_id is not None else None,
            department,
            (part_type or "").strip() or None,
            # Only the Issuance Report counts a department other than the consumer.
            bool(by_consumer and department and rep.key == "issuance"),
        )

    @staticmethod
    def _named(head, f, win):
        """A title, also a file's name: its head, then every filter that narrows
        it — so a file never passes for more than it holds — and the dates. An
        Issuance Report narrowed to a consuming department says so: its
        Department column is still who took the parts out."""
        dept = f"consumed by {f.department}" if f.by_consumer else f.department
        picked = [f.part.part_no if f.part else None, f.section.name if f.section else None, dept, f.type]
        return " ".join([head, *(p for p in picked if p), _dates(win)])

    def _title(self, view, c):
        """E.g. Consumption Report Month-wise CNTR3200 01.04.2026 to 14.10.2026;
        Issuance Report Issue list 16.10.2026."""
        return self._named(f"{c.report.title} {view.label}", c.f, c.win)

    def _out(self, chosen, c, built):
        """A built view as the screen and the files read it."""
        f, win, rep = c.f, c.win, c.report
        return {
            "report": rep.key,
            "report_title": rep.title,
            "view": chosen.key,
            "label": chosen.label,
            "title": self._title(chosen, c),
            "columns": list(chosen.columns),
            "rows": [_cells(r) for r in built.rows],
            "window": {"from": win.first.isoformat(), "to": win.last.isoformat(),
                       "days": (win.last - win.first).days + 1, "clamped_to_today": c.clamped,
                       "ledger_since": c.since.isoformat() if c.since else None,
                       "coverage": _coverage(win.first, win.last, c.since)},
            "unrated_issues": built.total.unrated,
            # The issues counted (the Issue list's cancelled ones aside).
            "issues": built.total.issues,
            "filters": {
                "part_id": f.part.id if f.part else None,
                "part_no": f.part.part_no if f.part else None,
                "part_name": f.part.part_name if f.part else None,
                "section_id": f.section.id if f.section else None,
                "section_name": f.section.name if f.section else None,
                "department": f.department,
                "type": f.type,
            },
            "basis": {k: BASIS[k] for k in rep.basis},
            "views": [{"key": v.key, "label": v.label} for v in rep.views],
            "group": c.group.key if chosen.key == "list" else None,
            "groups": [{"key": g.key, "label": g.label} for g in GROUPS],
            "subtotal_rows": list(built.subtotals),
            "cancelled_rows": list(built.cancelled),
            # The Full pack of the same report, range and filters, named as it downloads.
            "pack_title": self._named(f"{rep.title} {PACK}", f, win),
        }

    def _summary(self, ws, store_code, c, sheets, year):
        """The Full pack's first sheet: the report, the range, the filters and
        the totals — the Total row every other sheet ends on —, which years
        Year-wise is cut in, when it was prepared; then the notes and how it is
        counted."""
        f, win, rep = c.f, c.win, c.report
        first = sheets[0][1]
        total = first["rows"][-1]
        # The Issue list's notes say the cancelled issues it lists, too.
        listed = dict(sheets).get("Issue list", first)
        store = Store.query.filter_by(code=store_code).first()
        days = (win.last - win.first).days + 1
        part = f"{f.part.part_no} - {f.part.part_name}" if f.part and f.part.part_name else (
            f.part.part_no if f.part else "All parts")
        issuance = rep.key == "issuance"
        rows = [
            ("Store", f"{store.name} ({store_code})" if store else store_code),
            ("Report", rep.title),
            ("From", f"{win.first:%d.%m.%Y}"), ("To", f"{win.last:%d.%m.%Y}"), ("Days", days),
            ("Part", part),
            ("Section", f.section.name if f.section else "All sections"),
            (rep.department, f.department or "All departments"),
            ("Part type", f.type or "All part types"),
            ("Issues", len(c.uses)),
            ("Cancelled issues (listed, not counted)" if issuance else "Cancelled issues (not counted)",
             len(c.listed) - len(c.uses)),
            ("Parts", len({u.part.id for u in c.uses})),
            ("Qty Issued", total["Qty Issued"]),
            ("Value (today's rate)", total["Value"]),
        ]
        if not issuance:
            rows += [("Qty Returned", total["Qty Returned"]), ("Net Qty", total["Net Qty"]),
                     ("Net Value (today's rate)", total["Net Value"])]
        rows.append(("Year-wise sheet", "Financial years, 1 April to 31 March" if year == "fy"
                     else "Calendar years, January to December"))
        if issuance:
            rows.append(("Issue list subtotals", f"By {c.group.label.lower()}"))
        rows.append(("Prepared", _stamp(utcnow())))
        notes = _notes(listed) + ["How it is counted:"] + [BASIS[k] for k in rep.basis]
        _sheet(ws, ("Item", "Value"), [{"Item": k, "Value": v} for k, v in rows], notes)
        ws.title = "Summary"

    # ---- one read, every view ----

    def _read(self, store_code, rep, date_from, date_to, part_id, section_id, department, part_type, group,
              by_consumer=False):
        """The range read once for every view of a report (Ctx): its issues
        under the filters — cancelled ones listed, only the others counted —,
        the report's departments, the ledger's first day."""
        win, clamped = _window(date_from, date_to)
        grouping = self._group(group, win)
        f = self._filters(store_code, rep, part_id, section_id, department, part_type, by_consumer)
        listed, depts = self._issues(store_code, rep, win, f)
        since = ConsumptionAnalyticsService._ledger_since(store_code, _offset(win))
        order = _dept_order(depts, (NOT_WRITTEN,) if rep.key == "issuance" else ())
        return Ctx(rep, [u for u in listed if not u.cancelled], listed, depts, order, win, clamped, f, grouping,
                   since)

    @staticmethod
    def _issues(store_code, rep, win, f):
        """(the range's issues under the filters, oldest first, each marked
        cancelled or not; the report's departments in master order). The
        Consumption Report counts each issue for the department that owns its
        part's section, and what came back is every return raised against the
        issue (_returned_qty_map), whenever; the Issuance Report for the
        department written on it, gross. A cancelled issue was neither issued
        nor consumed: the Issue list shows it, no figure counts it."""
        L = SpareStockLedger
        issuance = rep.key == "issuance"
        where = [Part.store_code == store_code, L.txn_type == "ISSUE", *_in(L.created_at, win)]
        if f.part is not None:
            where.append(L.part_id == f.part.id)
        if f.section is not None:
            where.append(Part.section_id == f.section.id)
        if issuance:
            owner, depts = None, _master_departments()
            if f.by_consumer:                         # an old caller's consuming department
                where.append(consumer_clause(consumers(store_code)[0], f.department))
            elif f.department:
                where.append(taker_clause(f.department))
        else:
            owner, depts = consumers(store_code)
            if f.department:
                where.append(consumer_clause(owner, f.department))
        if f.type:
            where.append(type_clause(f.type))
        rows = (db.session.query(*ISSUE_COLUMNS).join(Part, Part.id == L.part_id)
                .filter(*where).order_by(L.created_at, L.id).all())
        parts = _parts({r.part_id for r in rows})
        sections = SpareReportService._sections(store_code)
        # Issuance is gross: what came back is the Consumption Report's, so it is not read.
        returned = {} if issuance else SpareStockService._returned_qty_map([r.id for r in rows])
        # The plant clock the range was cut on, read once rather than per issue.
        offset = _offset(win)
        uses = []
        for r in rows:
            part = parts[r.part_id]
            uses.append(Use(r, part, _written(r) if issuance else consumer_of(owner, part.section_id),
                            sections.get(part.section_id), (_utc_naive(r.created_at) + offset).date(),
                            abs(float(r.qty_delta)), returned.get(r.id, 0.0), r.status == "CANCELLED"))
        return uses, depts

    @staticmethod
    def _groups(c, key_fn):
        """The range's counted issues tallied in groups, key_fn(use) -> (group
        key, the group's own cells). Returns ({key: _Tally}, the Total's _Tally
        — of every counted issue, so its Parts are the distinct parts overall).
        Every view is this with its own grouping and order, so all of them add
        up the same way."""
        groups, total = {}, _Tally({})
        for use in c.uses:
            key, cells = key_fn(use)
            if key not in groups:
                groups[key] = _Tally(cells)
            groups[key].add(use)
            total.add(use)
        return groups, total

    def _by_period(self, c, columns, period):
        """A period view: period(plant day) -> (the period's sort key, its
        cells). ONE row per period, oldest first, only periods with issues
        listed. Not split by department: "how much on 07.10" or "in May" is
        read off one row, not added up from three (the store, 17 Oct 2026:
        "thoda simple ho jisme dropdown me date wise, month wise, year wise").
        Who took it or used it is Department-wise's, or the Department filter's."""
        groups, total = self._groups(c, lambda u: period(u.day))
        return _with_total([groups[k] for k in sorted(groups)], total, columns)

    def _by_date(self, c, columns):
        # The plant day of the issue: an issue at 01:30 is that morning's, not
        # the previous UTC day's.
        return self._by_period(c, columns, lambda d: (d, {"Date": f"{d:%d.%m.%Y}", "Day": f"{d:%a}"}))

    def _by_week(self, c, columns):
        def week(d):
            # Monday to Sunday; the range's first and last week only show the
            # days the range holds, so the label never claims days not counted.
            monday = d - timedelta(days=d.weekday())
            first, last = max(monday, c.win.first), min(monday + timedelta(days=6), c.win.last)
            return monday, {"Week": f"{first:%d.%m.%Y} - {last:%d.%m.%Y}"}

        return self._by_period(c, columns, week)

    def _by_month(self, c, columns):
        return self._by_period(c, columns, lambda d: (d.replace(day=1), {"Month": f"{d:%b %Y}"}))

    def _by_year(self, c, columns):
        # The calendar year, January to December; the financial year is its own view.
        return self._by_period(c, columns, lambda d: (d.year, {"Year": f"{d.year}"}))

    def _by_fy(self, c, columns):
        # The financial year, 1 April to 31 March, as the store writes it: FY 2026-27.
        return self._by_period(c, columns, lambda d: (fy_start(d), {"Financial Year": fy_label(d)}))

    def _by_part(self, c, columns):
        groups, total = self._groups(c, lambda u: (u.part.id, {
            "Part No": u.part.part_no, "Description": u.part.part_name, "Section": u.section,
            "Department": u.dept, "UOM": u.part.uom,
            "Rate": u.part.unit_price or None}))  # a rate of 0 is a missing rate
        ordered = sorted(groups.values(), key=lambda t: (*_by_money(t), t.cells["Part No"] or ""))
        return _with_total(ordered, total, columns)

    def _by_department(self, c, columns):
        """The Consumption Report's Department-wise: per consuming department
        and section."""
        f = c.f
        groups, total = self._groups(c, lambda u: (
            (u.dept.lower(), (u.section or "").lower()), {"Department": u.dept, "Section": u.section}))
        # Each consuming department is listed even with nothing consumed — the
        # store sees who used nothing — unless a part, section or type narrows
        # the table: then a zero says nothing, and only who used it is listed.
        if not (f.part or f.section or f.type):
            seen = {dept for dept, _ in groups}
            for name in c.depts:
                if name.lower() not in seen and (not f.department or name.lower() == f.department.lower()):
                    groups[(name.lower(), "")] = _Tally({"Department": name, "Section": None})
        # The consuming departments in master order, then any other consumer
        # (a section no department owns, parts with no section) by name.
        ordered = sorted(groups.items(), key=lambda kv: (c.order(kv[1].cells["Department"]), kv[0][1]))
        return _with_total([t for _, t in ordered], total, columns)

    def _by_taker(self, c, columns):
        """The Issuance Report's Department-wise: per department written on the
        issue (who took it out), in master order, "Not recorded" last."""
        groups, total = self._groups(c, lambda u: (u.dept.lower(), {"Department": u.dept}))
        ordered = sorted(groups.values(), key=lambda t: c.order(t.cells["Department"]))
        return _with_total(ordered, total, columns)

    def _by_person(self, c, columns):
        """Who the issues were handed to: one row per employee (one person
        picked from the Employee master is one row, whatever was typed), else
        per name and mobile as written, case and spaces aside. The latest
        name, code and mobile are shown, and every department the person took
        parts out for. Most money first; "Not recorded" (no person written)
        last."""
        groups, total = {}, _Tally({})
        for u in c.uses:
            i = u.issue
            name, mobile, code = _clean(i.issued_to), _clean(i.issued_to_mobile), _clean(i.issued_to_code)
            key = ("employee", i.employee_id) if i.employee_id else ("written", (name or "").lower(), mobile)
            t = groups.get(key)
            if t is None:
                t = groups[key] = _Tally({"Issued To": NOT_WRITTEN, "Code": None, "Mobile": None,
                                          "Department": []})
            # Issues are read oldest first, so the last word on the person stands.
            for column, v in (("Issued To", name), ("Code", code), ("Mobile", mobile)):
                if v:
                    t.cells[column] = v
            if u.dept not in t.cells["Department"]:
                t.cells["Department"].append(u.dept)
            t.add(u)
            total.add(u)
        for t in groups.values():
            t.cells["Department"] = ", ".join(t.cells["Department"])
        ordered = sorted(groups.values(), key=lambda t: (
            t.cells["Issued To"] == NOT_WRITTEN, *_by_money(t), t.cells["Issued To"].lower()))
        return _with_total(ordered, total, columns)

    def _by_machine(self, c, columns):
        def where(u):
            # The line / machine the fitter recorded at installation (the
            # latest update); with neither, its die or why there is none —
            # named as the By machine chart names it (place_of).
            i = u.issue
            line, machine = place_of(i.line_name, i.line_no, i.machine, i.die_no, i.installed)
            return ((u.dept.lower(), (line or "").lower(), (machine or "").lower()),
                    {"Department": u.dept, "Line": line, "Machine": machine})

        groups, total = self._groups(c, where)
        ordered = sorted(groups.items(), key=lambda kv: (c.order(kv[1].cells["Department"]), -kv[1].net,
                                                         kv[0][1:]))
        return _with_total([t for _, t in ordered], total, columns)

    def _by_category(self, c, columns):
        def kind(u):
            # The part's type (sub_category), as the Category column of Current Stock.
            category = (u.part.sub_category or "").strip() or UNCATEGORISED
            return (u.dept.lower(), category.lower()), {"Department": u.dept, "Category": category}

        groups, total = self._groups(c, kind)
        ordered = sorted(groups.items(), key=lambda kv: (c.order(kv[1].cells["Department"]), -kv[1].net,
                                                         kv[0][1]))
        return _with_total([t for _, t in ordered], total, columns)

    def _issue_list(self, c, columns):
        """Every issue in the range, oldest first, in the Issue Register's
        words; a subtotal row after each day / month / year (c.group), the
        Total last. Each issue's figures are a _Tally of that one issue, so its
        Value follows the same rules as every view's. A cancelled issue is
        listed with its quantity and nothing else, and no subtotal or total
        counts it."""
        rows, subtotals, cancelled = [], [], []
        total, struck = _Tally({}), 0
        open_ = None                                  # [key, label, _Tally, cancelled count]

        def close():
            if open_ is not None:
                _, label, tally, gone = open_
                tally.cells = {"Date": SUBTOTAL.format(label), "Issue No": _count(tally.issues, gone)}
                subtotals.append(len(rows))
                rows.append(tally.row(columns))

        for u in c.listed:
            key, label = c.group.period(u.day)
            if open_ is None or key != open_[0]:
                close()
                open_ = [key, label, _Tally({}), 0]
            if u.cancelled:
                open_[3] += 1
                struck += 1
                cancelled.append(len(rows))
            else:
                open_[2].add(u)
                total.add(u)
            rows.append(self._issue_row(u, columns))
        close()
        total.cells = {"Date": TOTAL, "Issue No": _count(total.issues, struck)}
        rows.append(total.row(columns))
        return Built(rows, total, subtotals, cancelled)

    @staticmethod
    def _issue_row(u, columns):
        i, part = u.issue, u.part
        on, at = _stamp(i.created_at).split(" ")
        cells = {
            "Date": on, "Time": at, "Issue No": i.id, "Reference": i.reference, "Part No": part.part_no,
            "Description": part.part_name, "Section": u.section, "Department": _clean(i.department),
            "Issued To": _issued_to(i), "Rate": part.unit_price or None,   # a rate of 0 is a missing rate
            "Line / Machine / Die": None if u.cancelled else _place(i), "Status": _issue_status(i),
        }
        if u.cancelled:
            return {column: cells.get(column) for column in columns} | {"Qty Issued": u.qty}
        one = _Tally(cells)
        one.add(u)
        return one.row(columns)

    # The two reports (the store, 17 Oct 2026: "issuance and consumption ki
    # report alag alag chahiye"), each with its View dropdown in its order.
    REPORTS = {
        "issuance": Report("issuance", "Issuance Report", (
            View("list", "Issue list", LIST_COLUMNS, _issue_list),
            View("day", "Date-wise", ISSUED_BY_DATE_COLUMNS, _by_date),
            View("week", "Week-wise", ISSUED_BY_WEEK_COLUMNS, _by_week),
            View("month", "Month-wise", ISSUED_BY_MONTH_COLUMNS, _by_month),
            View("year", "Year-wise", ISSUED_BY_YEAR_COLUMNS, _by_year),
            View("fy", "Financial year-wise", ISSUED_BY_FY_COLUMNS, _by_fy),
            View("part", "Part-wise", ISSUED_BY_PART_COLUMNS, _by_part),
            View("department", "Department-wise", ISSUED_BY_DEPARTMENT_COLUMNS, _by_taker),
            View("person", "Person-wise", BY_PERSON_COLUMNS, _by_person),
        ), ISSUANCE_PACK, ("issuance", "value", "period", "list", "person"), "Department (taken by)"),
        "consumption": Report("consumption", "Consumption Report", (
            View("day", "Date-wise", BY_DATE_COLUMNS, _by_date),
            View("week", "Week-wise", BY_WEEK_COLUMNS, _by_week),
            View("month", "Month-wise", BY_MONTH_COLUMNS, _by_month),
            View("year", "Year-wise", BY_YEAR_COLUMNS, _by_year),
            View("fy", "Financial year-wise", BY_FY_COLUMNS, _by_fy),
            View("part", "Part-wise", BY_PART_COLUMNS, _by_part),
            View("department", "Department-wise", BY_DEPARTMENT_COLUMNS, _by_department),
            View("machine", "Line / Machine-wise", BY_MACHINE_COLUMNS, _by_machine),
            View("category", "Category-wise", BY_CATEGORY_COLUMNS, _by_category),
        ), CONSUMPTION_PACK, ("consumption", "returns", "value", "period", "place"), "Consuming department"),
    }
