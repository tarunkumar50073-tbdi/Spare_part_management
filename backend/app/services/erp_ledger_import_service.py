"""ERP Stock-Ledger Import (Option A) — one-time Excel bridge until the ERP
API is live, and it feeds the SAME ledger the API will push to.

The ERP "Stock Ledger Valuation" export has a few metadata rows above the
real header, then one row per stock transaction. For the BOP-vs-in-house
Source Breakdown we only care about the two *receipt* sources:

    Doc Type "SR" (supplier receipt, e.g. ELITE ENGINEERS)  -> source_type BUY  (BOP)
    Doc Type "PP" (production receipt, PRODUCTION)          -> source_type MAKE (in-house)

Everything else in the file (opening balance OP, coating-out issues W to SAR
Coating, waste LM, minor RS/SA) is skipped here — that's the fuller ledger of
Option B. Rows are keyed by "<Doc No>#<Tran Seq>" (`ext_ref`) so re-importing
the same file, or a later ERP API push of the same transaction, is de-duped
instead of double-counted. Stock-on-hand is NOT adjusted: these are historical
receipts already reflected in ERP's live stock (which arrives via
ERPStockSyncService), so this import only backfills the movement history the
Source Breakdown reads from."""
from datetime import datetime, timezone

from openpyxl import load_workbook

from app.extensions import db
from app.repositories.part_repository import PartRepository
from app.repositories.store_stock_repository import StockLedgerRepository
from app.utils.exceptions import BadRequestError

# ERP Doc Type -> our received-stock source bucket.
DOC_TYPE_SOURCE = {"SR": "BUY", "PP": "MAKE"}

# Header cells we need, matched case-/space-insensitively.
_NEEDED = {
    "part_no": "itemcd",
    "date": "trandt",
    "doc_no": "docno",
    "doc_type": "doctype",
    "recv_qty": "recvqty",
    "tran_seq": "transeq",
}


def _norm(text):
    return "".join(ch for ch in str(text or "").strip().lower() if ch.isalnum())


def _parse_date(value):
    """ERP dates look like '01-Jun-2026'; fall back to a real datetime if
    openpyxl already parsed the cell as one."""
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc)
    for fmt in ("%d-%b-%Y", "%d-%b-%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).replace(tzinfo=timezone.utc)
        except (ValueError, TypeError):
            continue
    raise ValueError(f"Unreadable date: {value!r}")


class ERPLedgerImportService:
    def __init__(self):
        self.part_repo = PartRepository()
        self.ledger_repo = StockLedgerRepository()

    def import_excel(self, file_stream):
        wb = load_workbook(file_stream, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        header_idx, col = self._locate_header(rows)
        if header_idx is None:
            raise BadRequestError(
                "Could not find the ledger header row (needs columns Item Cd, Doc Type, Recv Qty, Tran Seq)."
            )

        seen_ext = set(self.ledger_repo.existing_ext_refs())
        part_cache = {}

        total = created = 0
        buy_qty = make_qty = 0.0
        skipped_duplicate = skipped_other = 0
        errors = []

        for excel_row_no, raw in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            if raw is None or all(v is None for v in raw):
                continue
            doc_type = str(raw[col["doc_type"]] or "").strip().upper()
            source_type = DOC_TYPE_SOURCE.get(doc_type)
            if not source_type:
                skipped_other += 1
                continue

            total += 1
            try:
                qty = float(raw[col["recv_qty"]] or 0)
                if qty <= 0:
                    skipped_other += 1
                    total -= 1
                    continue

                part_no = str(raw[col["part_no"]] or "").strip()
                if not part_no:
                    raise ValueError("Missing Item Cd")

                doc_no = str(raw[col["doc_no"]] or "").strip()
                tran_seq = str(raw[col["tran_seq"]] or "").strip()
                ext_ref = f"{doc_no}#{tran_seq}"
                if ext_ref in seen_ext:
                    skipped_duplicate += 1
                    total -= 1
                    continue

                if part_no not in part_cache:
                    part_cache[part_no] = self.part_repo.get_by_part_no(part_no)
                part = part_cache[part_no]
                if not part:
                    raise ValueError(f"Part '{part_no}' not in Part Master — add it there first")

                created_at = _parse_date(raw[col["date"]])
                self.ledger_repo.record(
                    part.id, "PART_STORE", qty,
                    reason="ERP_LEDGER_IMPORT", source_type=source_type,
                    ext_ref=ext_ref, created_at=created_at,
                )

                seen_ext.add(ext_ref)
                created += 1
                if source_type == "BUY":
                    buy_qty += qty
                else:
                    make_qty += qty
            except ValueError as e:
                errors.append({"row": excel_row_no, "reason": str(e)})

        db.session.commit()
        return {
            "total_rows": total,
            "created": created,
            "error_count": len(errors),
            "errors": errors,
            "buy_qty": buy_qty,
            "make_qty": make_qty,
            "skipped_duplicate": skipped_duplicate,
            "skipped_other": skipped_other,
        }

    @staticmethod
    def _locate_header(rows):
        """Scan the first ~20 rows for the real header (metadata rows precede
        it). Returns (row_index, {field: col_index})."""
        for idx, row in enumerate(rows[:20]):
            if row is None:
                continue
            norm_to_col = {_norm(cell): i for i, cell in enumerate(row) if cell is not None}
            col = {}
            for field, needle in _NEEDED.items():
                if needle in norm_to_col:
                    col[field] = norm_to_col[needle]
            if len(col) == len(_NEEDED):
                return idx, col
        return None, None
