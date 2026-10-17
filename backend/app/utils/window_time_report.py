"""Builds the plant's 'VEHICLE WINDOW TIME' monthly sheet — the exact format
the store fills by hand and mails every day — but auto-filled from the app's
own data and colour-coded to the app's legend:

    ON TIME  -> green      LATE -> amber      NOT COME -> red

Layout mirrors SAMPLE OF FORMAT.xlsx:
  * five fixed left columns: S.No | Arrival Plan (SCM) | Departure Plan (Store)
    | Supplier Name | Frequency
  * then one three-column block per calendar day: Actual Arrival | Actual
    Departure | Remarks, the date across the top of the block
  * a legend + tolerance note + signature line at the bottom

Only the plan columns and the actuals differ from the paper sheet; the store
person no longer types a single cell — the arrival/departure/remark come from
the same window-vs-gate-entry status the Live Dashboard already shows.
"""
import io
from calendar import monthrange, month_name
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils.excel_io import _excel_safe

FIXED_COLS = 5  # S.No, Arrival Plan, Departure Plan, Supplier Name, Frequency
_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
HEADER_FILL = PatternFill("solid", fgColor="FFC000")   # amber, matches the sample title
FREQ_FILL = PatternFill("solid", fgColor="FFF2CC")     # pale amber, like the sample's FREQUENCY
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# status -> (cell fill, font colour, remark word). Keyed on the app's own status
# strings; "Pending Arrival" is deliberately absent — a window still open (or a
# future day) is left blank, exactly as the paper sheet is until the truck comes.
STATUS_STYLE = {
    "On-Time": ("C6EFCE", "006100", "ON TIME"),
    "Late": ("FFEB9C", "9C6500", "LATE"),
    "Not Arrived": ("FFC7CE", "9C0006", "NOT COME"),
}
LEGEND = [("ON TIME", "C6EFCE", "006100"), ("LATE", "FFEB9C", "9C6500"), ("NOT COME", "FFC7CE", "9C0006")]


def _fmt_actual(iso):
    """An ISO datetime string ('2026-08-06T08:35:00') -> 'HH:MM'. Blank for
    None or anything unparseable."""
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso).strftime("%H:%M")
    except ValueError:
        return ""


def _day_lookup(status_rows):
    """A day's status list -> a lookup keyed both by (start, end, vendor_code)
    and (start, end, supplier_name), so a master row matches whether or not the
    supplier is milkrun-grouped (grouped rows carry the group's name)."""
    lut = {}
    for st in status_rows:
        lut[(st["window_start"], st["window_end"], st["vendor_code"])] = st
        lut.setdefault((st["window_start"], st["window_end"], st["supplier_name"]), st)
    return lut


def build_window_time_report(year, month, master_rows, status_by_day):
    """master_rows: [{vendor_code, supplier_name, window_start (time),
    window_end (time), frequency}] — the standing schedule, one row per slot,
    already in shift-day order. status_by_day: {day_int -> [status dict]} from
    the window-vs-arrival service. Returns a BytesIO .xlsx."""
    ndays = monthrange(year, month)[1]
    total_cols = FIXED_COLS + ndays * 3
    lut_by_day = {day: _day_lookup(rows) for day, rows in status_by_day.items()}

    wb = Workbook()
    ws = wb.active
    ws.title = "WINDOW TIME"

    # --- title ---
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title = ws.cell(1, 1, f"VEHICLE WINDOW TIME       Month :- {month_name[month].upper()} {year}")
    title.font = Font(bold=True, size=14)
    title.fill = HEADER_FILL
    title.alignment = Alignment(horizontal="left", vertical="center")

    # --- fixed headers (rows 2-3, merged vertically) ---
    fixed = ["S.No.", "Supplier Arrival Time Plan (SCM)",
             "Supplier Departure Time Plan (Store)", "Supplier Name", "FREQUENCY"]
    for col, label in enumerate(fixed, start=1):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        cell = ws.cell(2, col, label)
        cell.font = Font(bold=True, size=10)
        cell.alignment = CENTER
        cell.fill = HEADER_FILL
        cell.border = BORDER

    # --- per-day headers ---
    for d in range(1, ndays + 1):
        base = FIXED_COLS + (d - 1) * 3 + 1
        ws.merge_cells(start_row=2, start_column=base, end_row=2, end_column=base + 2)
        dcell = ws.cell(2, base, date(year, month, d).strftime("%d-%b (%a)"))
        dcell.font = Font(bold=True, size=10)
        dcell.alignment = CENTER
        dcell.fill = HEADER_FILL
        for off, sub in enumerate(("Arrival", "Departure", "Remarks")):
            sub_cell = ws.cell(3, base + off, sub)
            sub_cell.font = Font(bold=True, size=9)
            sub_cell.alignment = CENTER
            sub_cell.fill = HEADER_FILL
            sub_cell.border = BORDER

    # --- data rows ---
    for i, mr in enumerate(master_rows):
        r = 4 + i
        ws.cell(r, 1, i + 1).alignment = CENTER
        ws.cell(r, 2, mr["window_start"].strftime("%H:%M")).alignment = CENTER
        ws.cell(r, 3, mr["window_end"].strftime("%H:%M")).alignment = CENTER
        name_cell = ws.cell(r, 4, _excel_safe(mr["supplier_name"] or mr["vendor_code"]))
        name_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        freq_cell = ws.cell(r, 5, mr["frequency"])
        freq_cell.fill = FREQ_FILL
        freq_cell.alignment = CENTER
        for c in range(1, FIXED_COLS + 1):
            ws.cell(r, c).border = BORDER

        start_iso, end_iso = mr["window_start"].isoformat(), mr["window_end"].isoformat()
        for d in range(1, ndays + 1):
            base = FIXED_COLS + (d - 1) * 3 + 1
            lut = lut_by_day.get(d, {})
            st = lut.get((start_iso, end_iso, mr["vendor_code"])) or \
                lut.get((start_iso, end_iso, mr["supplier_name"]))
            arrival, departure, remark = ws.cell(r, base), ws.cell(r, base + 1), ws.cell(r, base + 2)
            for cell in (arrival, departure, remark):
                cell.border = BORDER
                cell.alignment = CENTER
            if not st:
                continue
            arrival.value = _fmt_actual(st.get("actual_arrival"))
            departure.value = _fmt_actual(st.get("actual_out"))
            style = STATUS_STYLE.get(st["status"])
            if style:
                fill_hex, font_hex, word = style
                remark.value = word
                fill = PatternFill("solid", fgColor=fill_hex)
                for cell in (arrival, departure, remark):
                    cell.fill = fill
                remark.font = Font(bold=True, size=9, color=font_hex)

    # --- legend + tolerance + signature ---
    legend_row = 4 + len(master_rows) + 2
    lg = ws.cell(legend_row, 1, "Legend")
    lg.font = Font(bold=True, size=11)
    tol = ws.cell(legend_row, 2, "TOLERANCE TIME :  PLUS / MINUS 10 MINUTES")
    tol.font = Font(italic=True, size=9)
    for j, (word, fill_hex, font_hex) in enumerate(LEGEND):
        cell = ws.cell(legend_row + 1 + j, 1, word)
        cell.fill = PatternFill("solid", fgColor=fill_hex)
        cell.font = Font(bold=True, size=10, color=font_hex)
        cell.alignment = CENTER
        cell.border = BORDER
    sign_row = legend_row + 1 + len(LEGEND) + 1
    ws.cell(sign_row, 1, "Prepared By :  DX_Store (auto-generated)").font = Font(italic=True, size=9)

    # --- widths, freeze, header heights ---
    for col, width in (("A", 6), ("B", 15), ("C", 15), ("D", 30), ("E", 14)):
        ws.column_dimensions[col].width = width
    for d in range(1, ndays + 1):
        base = FIXED_COLS + (d - 1) * 3 + 1
        for off in range(3):
            ws.column_dimensions[get_column_letter(base + off)].width = 9
    ws.freeze_panes = "F4"  # keep the plan block + headers visible while scrolling days
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
