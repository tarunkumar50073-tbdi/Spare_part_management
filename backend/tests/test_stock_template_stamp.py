"""Store Stock > Download Template / Upload Excel: an OLD copy of the template
must not put back stock figures that have moved since (store, 16 Oct 2026,
"Chhote data fix").

The template is pre-filled with today's stock and stamped (hidden "_meta"
sheet) with the moment it was made. On upload, a row whose part had a stock
movement after that moment AND whose figure differs from today's stock is not
applied — it is reported, the rest of the file applies. A file without the
stamp imports exactly as before.

Each test seeds its own store (the session DB is shared), and the seed's own
OPENING rows are moved an hour back, so "made before / after the template"
does not hinge on the same second.
"""
import io
from datetime import timedelta

import pytest
from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.part import Part
from app.models.spare_stock import SpareStock, SpareStockLedger
from app.models.store import Section, Store
from app.models.user import User
from app.services.spare_stock_service import (
    TEMPLATE_META_SHEET, TEMPLATE_STAMP_KEY, SpareStockService, _template_stamp,
)
from app.utils.excel_io import read_rows

_seq = {"n": 0}


def _next():
    _seq["n"] += 1
    return _seq["n"]


@pytest.fixture()
def store():
    n = _next()
    s = Store(code=f"TPL{n}", name=f"Template Test Store {n}", sort_order=960)
    db.session.add(s)
    db.session.commit()
    sec = Section(store_id=s.id, name=f"Template Section {n}")
    db.session.add(sec)
    db.session.commit()
    yield s
    s.is_active = False
    db.session.commit()


@pytest.fixture()
def parts(store):
    """A 10, B 5, C 7 on hand — counted an hour ago."""
    out = {}
    for tag, qty in (("A", 10), ("B", 5), ("C", 7)):
        p = Part(part_no=f"TPL-{_next()}-{tag}", part_name=f"Template part {tag}", uom="NO",
                 part_type="OTHER", store_code=store.code, location="R1-S1-B1")
        db.session.add(p)
        db.session.commit()
        SpareStockService().set_quantity(p.id, float(qty), txn_type="OPENING")
        out[tag] = p
    for p in out.values():
        for row in SpareStockLedger.query.filter_by(part_id=p.id).all():
            row.created_at = row.created_at - timedelta(hours=1)
        stock = SpareStock.query.filter_by(part_id=p.id).one()
        stock.last_counted_at = stock.last_counted_at - timedelta(hours=1)
    db.session.commit()
    return out


def _template(code):
    return SpareStockService().build_template(code).getvalue()


def _edited(template_bytes, changes):
    """The template with some Qty On Hand cells retyped ({part_no: qty})."""
    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=2):
        if row[0].value in changes:
            row[2].value = changes[row[0].value]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(data, code):
    return SpareStockService().bulk_import(io.BytesIO(data), store_code=code, user="keeper")


def _qty(part):
    return float(SpareStock.query.filter_by(part_id=part.id).one().qty_on_hand)


def _counted_at(part):
    return SpareStock.query.filter_by(part_id=part.id).one().last_counted_at


def _issue(part, qty):
    SpareStockService().issue(part.id, qty, "Maintenance", issued_to="Ravi")


# ---- the stamp itself ----

def test_template_is_stamped_on_a_hidden_sheet_that_is_not_data(store, parts):
    data = _template(store.code)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] != TEMPLATE_META_SHEET          # the data sheet opens first
    meta = wb[TEMPLATE_META_SHEET]
    assert meta.sheet_state == "hidden"
    assert meta["A1"].value == TEMPLATE_STAMP_KEY and _template_stamp(data) is not None

    # The importer reads the data sheet only: one row per part, no stamp row,
    # and no "which sheet?" question.
    rows = list(read_rows(io.BytesIO(data), {"part_no": {"partno"}, "qty_on_hand": {"qtyonhand"}},
                          ["part_no", "qty_on_hand"]))
    assert sorted(r["part_no"] for _, r in rows) == sorted(p.part_no for p in parts.values())

    report = _upload(data, store.code)
    assert (report["total_rows"], report["error_count"], report["stale_count"]) == (3, 0, 0)
    assert report["unchanged"] == 3 and report["template_made_at"]


def test_a_store_less_template_carries_no_stamp():
    data = SpareStockService().build_template(None).getvalue()
    assert TEMPLATE_META_SHEET not in load_workbook(io.BytesIO(data)).sheetnames
    assert _template_stamp(data) is None


# ---- applying ----

def test_a_fresh_template_applies(store, parts):
    data = _edited(_template(store.code), {parts["A"].part_no: 12, parts["B"].part_no: 4})
    report = _upload(data, store.code)
    assert (report["updated"], report["unchanged"], report["stale_count"], report["error_count"]) == (2, 1, 0, 0)
    assert (_qty(parts["A"]), _qty(parts["B"]), _qty(parts["C"])) == (12, 4, 7)


def test_an_old_template_skips_only_the_parts_that_moved(store, parts):
    old = _template(store.code)                 # A 10, B 5, C 7
    _issue(parts["A"], 2)                       # A is 8 now — after the template was made
    counted_before = _counted_at(parts["A"])

    # Uploaded later: A still says the old 10, B was recounted to 6, C unchanged.
    report = _upload(_edited(old, {parts["B"].part_no: 6}), store.code)

    assert _qty(parts["A"]) == 8                # the old figure did NOT come back
    assert _counted_at(parts["A"]) == counted_before   # and A is not marked counted
    assert (_qty(parts["B"]), _qty(parts["C"])) == (6, 7)
    assert (report["updated"], report["unchanged"], report["stale_count"], report["error_count"]) == (1, 1, 1, 1)

    stale = report["stale_rows"][0]
    assert (stale["part_no"], stale["file_qty"], stale["current_qty"]) == (parts["A"].part_no, 10, 8)
    assert stale["last_moved_at"]
    err = report["errors"][0]
    assert err["row"] == 2 + sorted(p.part_no for p in parts.values()).index(parts["A"].part_no)
    assert err["reason"].startswith("Skipped: stock moved since this template was made (")
    assert "download a fresh template" in err["reason"] and report["template_made_at"] in err["reason"]
    assert "file 10, stock now 8" in err["reason"]


def test_an_old_template_row_equal_to_todays_stock_applies(store, parts):
    old = _template(store.code)
    _issue(parts["A"], 2)                       # 8 now
    counted_before = _counted_at(parts["A"])
    report = _upload(_edited(old, {parts["A"].part_no: 8}), store.code)   # counted 8: right
    assert (report["stale_count"], report["error_count"], report["unchanged"]) == (0, 0, 3)
    assert _qty(parts["A"]) == 8
    assert _counted_at(parts["A"]) > counted_before    # the count is recorded as before


def test_a_file_without_the_stamp_imports_as_before(store, parts):
    _issue(parts["A"], 2)                       # 8 now
    wb = Workbook()
    ws = wb.active
    ws.append(["Part No", "Qty On Hand"])
    ws.append([parts["A"].part_no, 10])
    buf = io.BytesIO()
    wb.save(buf)
    report = _upload(buf.getvalue(), store.code)
    assert (report["updated"], report["stale_count"], report["error_count"]) == (1, 0, 0)
    assert report["template_made_at"] is None and report["stale_rows"] == []
    assert _qty(parts["A"]) == 10               # the store's own sheet is trusted, as today


def test_an_unreadable_stamp_counts_as_no_stamp(store, parts):
    old = _template(store.code)
    _issue(parts["A"], 2)
    wb = load_workbook(io.BytesIO(old))
    wb[TEMPLATE_META_SHEET]["B1"].value = "not a time"
    buf = io.BytesIO()
    wb.save(buf)
    assert _template_stamp(buf.getvalue()) is None
    report = _upload(buf.getvalue(), store.code)
    assert report["stale_count"] == 0 and _qty(parts["A"]) == 10


# ---- the same template uploaded again (review, 16 Oct 2026) ----
# Loading a physical count: upload, spot a typo, fix it in the SAME file,
# upload again. The first upload's own OPENING rows are not "stock moved".

def _restamped(template_bytes, shift):
    """The template with its stamp moved by `shift` — an older copy, made
    clearly before another template (stamps are compared to the second)."""
    from app.services.spare_stock_service import _template_stamp as stamp_of
    wb = load_workbook(io.BytesIO(template_bytes))
    wb[TEMPLATE_META_SHEET]["B1"].value = (stamp_of(template_bytes) + shift).isoformat()
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_re_uploading_the_same_template_with_a_correction_applies(store, parts):
    template = _template(store.code)
    first = _upload(_edited(template, {parts["A"].part_no: 11, parts["B"].part_no: 4}), store.code)
    assert (first["updated"], first["stale_count"]) == (2, 0)
    # The upload's rows carry the template's tag (and read as its document).
    opening = (SpareStockLedger.query.filter_by(part_id=parts["A"].id, txn_type="OPENING")
               .order_by(SpareStockLedger.id.desc()).first())
    assert opening.reference and opening.reference.startswith("Stock template ")

    # A typed 11 should have been 13: fixed in the same file, uploaded again.
    second = _upload(_edited(template, {parts["A"].part_no: 13, parts["B"].part_no: 4}), store.code)
    assert (second["updated"], second["unchanged"], second["stale_count"], second["error_count"]) == (1, 2, 0, 0)
    assert (_qty(parts["A"]), _qty(parts["B"]), _qty(parts["C"])) == (13, 4, 7)


def test_re_upload_still_skips_a_part_issued_between_the_uploads(store, parts):
    template = _template(store.code)
    _upload(_edited(template, {parts["A"].part_no: 11, parts["B"].part_no: 4}), store.code)
    _issue(parts["A"], 1)                       # a real movement: A is 10 now
    second = _upload(_edited(template, {parts["A"].part_no: 13, parts["B"].part_no: 6}), store.code)
    assert (second["stale_count"], second["error_count"]) == (1, 1)
    assert second["stale_rows"][0]["part_no"] == parts["A"].part_no
    assert second["errors"][0]["reason"].startswith("Skipped: stock moved since this template was made (")
    assert _qty(parts["A"]) == 10               # the issue stands
    assert _qty(parts["B"]) == 6                # B only had the template's own upload: applies


def test_another_templates_upload_still_stops_an_older_copy(store, parts):
    older = _restamped(_template(store.code), -timedelta(minutes=1))
    newer = _template(store.code)
    assert _upload(_edited(newer, {parts["A"].part_no: 12}), store.code)["updated"] == 1
    # The older copy still says A 10: the newer template's upload moved A since.
    report = _upload(older, store.code)
    assert report["stale_count"] == 1 and report["stale_rows"][0]["part_no"] == parts["A"].part_no
    assert _qty(parts["A"]) == 12


# ---- through the screen's own routes ----

def test_download_and_upload_routes_carry_the_guard(app, store, parts):
    _seq["n"] += 1
    admin = User(username=f"tpladmin{_seq['n']}", role="ADMIN", is_active=True)
    admin.set_password("pw")
    db.session.add(admin)
    db.session.commit()
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["user_id"] = admin.id

    resp = client.get(f"/api/v1/spare-stock/template?store_code={store.code}")
    assert resp.status_code == 200
    old = resp.data
    _issue(parts["A"], 2)
    resp = client.post("/api/v1/spare-stock/import",
                       data={"file": (io.BytesIO(old), "spare_opening_stock_template.xlsx"),
                             "store_code": store.code},
                       content_type="multipart/form-data")
    assert resp.status_code == 201, resp.get_json()
    data = resp.get_json()["data"]
    assert data["stale_count"] == 1 and data["stale_rows"][0]["part_no"] == parts["A"].part_no
    assert _qty(parts["A"]) == 8
