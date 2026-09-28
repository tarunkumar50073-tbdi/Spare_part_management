"""Line / Machine / Die master: the install screen picks from these lists, adds a
missing value on the spot (scan_part edit), and the master is editable by users
with the Equipment right. It is seeded from the values already on the ledger.

The test DB is shared across the session (no per-test rollback), so each test
uses a unique token in its names and asserts on the rows it created rather than
on global counts."""
import io

import pytest
from openpyxl import Workbook
from sqlalchemy.exc import IntegrityError

from app.extensions import db
from app.models.equipment import Equipment, normalise_code, normalise_name
from app.models.part import Part
from app.models.spare_flow import SpareTransaction
from app.models.user import User, UserStoreAccess
from app.services.equipment_service import EquipmentService
from app.utils.exceptions import BadRequestError, ConflictError, NotFoundError

_seq = {"n": 0}


def _tok():
    _seq["n"] += 1
    return f"T{_seq['n']:04d}"


def _client(app, perms):
    _seq["n"] += 1
    u = User(username=f"equser{_seq['n']}", role="STORE_USER", is_active=True)
    u.set_password("pw")
    u.set_permissions(perms)
    db.session.add(u)
    db.session.flush()
    db.session.add(UserStoreAccess(user_id=u.id, store_code="MAINTENANCE"))
    db.session.commit()
    c = app.test_client()
    with c.session_transaction() as s:
        s["user_id"] = u.id
    return c


@pytest.mark.parametrize("raw, want", [
    ("  Seat   Slider ", "Seat Slider"), ("Press-07", "Press-07"), ("", None), (None, None),
])
def test_normalise_name(raw, want):
    assert normalise_name(raw) == want


@pytest.mark.parametrize("raw, want", [
    (" l-03 ", "L-03"), ("press 07", "PRESS 07"), ("", None), (None, None),
])
def test_normalise_code(raw, want):
    assert normalise_code(raw) == want


def test_create_list_and_kind_filter():
    svc = EquipmentService()
    tok = _tok()
    svc.create({"kind": "line", "name": f"  Seat Slider {tok} ", "code": "l-03"}, user="t")
    svc.create({"kind": "MACHINE", "name": f"Press {tok}"})
    svc.create({"kind": "DIE", "name": f"Die {tok}"})

    line = [e for e in svc.list(kind="LINE", q=tok)]
    assert len(line) == 1
    assert line[0]["name"] == f"Seat Slider {tok}" and line[0]["code"] == "L-03"
    assert line[0]["kind"] == "LINE" and line[0]["is_active"] is True
    assert {e["name"] for e in svc.list(kind="machine", q=tok)} == {f"Press {tok}"}
    # q matches across kinds; kind filter narrows it.
    assert len(svc.list(q=tok)) == 3
    assert len(svc.list(kind="DIE", q=tok)) == 1


def test_duplicate_within_kind_refused_across_kinds_ok():
    svc = EquipmentService()
    tok = _tok()
    name = f"Cell {tok}"
    svc.create({"kind": "MACHINE", "name": name})
    # Same name, same kind (any case) -> conflict.
    with pytest.raises(ConflictError):
        svc.create({"kind": "MACHINE", "name": name.lower()})
    # Same name, different kind -> fine (a LINE and a MACHINE can share a name).
    svc.create({"kind": "LINE", "name": name})
    assert {e["kind"] for e in svc.list(q=name)} == {"LINE", "MACHINE"}


def test_bad_kind_and_missing_name_refused():
    svc = EquipmentService()
    with pytest.raises(BadRequestError):
        svc.create({"kind": "ROBOT", "name": "X"})
    with pytest.raises(BadRequestError):
        svc.create({"kind": "LINE", "name": "   "})
    with pytest.raises(NotFoundError):
        svc.update(999999, {"name": "Nope"})


def test_update_and_deactivate_keep_the_row():
    svc = EquipmentService()
    tok = _tok()
    eq = svc.create({"kind": "MACHINE", "name": f"Press {tok}"})
    upd = svc.update(eq["id"], {"name": f"Press {tok}A", "code": "m-09", "is_active": False}, user="admin")
    assert upd["name"] == f"Press {tok}A" and upd["code"] == "M-09"
    assert upd["is_active"] is False and upd["updated_by"] == "admin"
    # Deactivated, not deleted — still there for history.
    assert Equipment.query.filter_by(id=eq["id"]).count() == 1
    assert svc.list(kind="MACHINE", active=True, q=tok) == []
    assert [e["id"] for e in svc.list(kind="MACHINE", active=False, q=tok)] == [eq["id"]]


def test_excel_import_upserts_by_kind_and_name():
    svc = EquipmentService()
    tok = _tok()
    svc.create({"kind": "MACHINE", "name": f"Old Press {tok}", "code": "OLD"})
    wb = Workbook()
    ws = wb.active
    ws.append(["Kind (LINE/MACHINE/DIE)", "Name", "Code"])
    ws.append(["machine", f"Old Press {tok}", "NEW"])   # updates the code
    ws.append(["LINE", f"Weld Line {tok}", "L-11"])     # new
    ws.append(["DIE", f"Die {tok}", None])              # new
    ws.append(["ROBOT", f"Nope {tok}", None])           # bad kind
    ws.append(["MACHINE", "   ", None])                 # blank name
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    report = svc.import_file(buf, user="admin")
    assert report["total_rows"] == 5 and report["error_count"] == 2
    machines = svc.list(kind="MACHINE", q=tok)
    assert len(machines) == 1 and machines[0]["code"] == "NEW"   # upserted, not duplicated
    assert {e["name"] for e in svc.list(kind="LINE", q=tok)} == {f"Weld Line {tok}"}
    assert {e["name"] for e in svc.list(kind="DIE", q=tok)} == {f"Die {tok}"}


def test_seed_from_ledger_fills_and_is_idempotent():
    svc = EquipmentService()
    tok = _tok()
    ln1, ln2 = f"Seat Slider {tok}", f"Recliner {tok}"
    mc1, mc2 = f"Press {tok}A", f"Press {tok}B"
    die1, die2 = f"Die {tok}A", f"Die {tok}B"
    p = Part(part_no=f"SEED-{tok}", part_name="Seed Part", uom="PCS", part_type="OTHER")
    db.session.add(p)
    db.session.commit()
    for ln, lno, mc, die in [
        (ln1, "L-03", mc1, die1),
        (ln1, "L-03", mc1, None),      # same line + machine -> deduped
        (ln2, "L-05", mc2, die2),
        (None, None, mc1, None),       # machine only, no line/die
        ("-", "-", ".", "-"),          # junk placeholders -> skipped
    ]:
        db.session.add(SpareTransaction(
            part_id=p.id, store_code="MAINTENANCE", txn_type="INSTALLATION",
            line_name=ln, line_no=lno, machine=mc, die_no=die))
    db.session.commit()

    svc.seed_from_ledger(user="seed")
    lines = {e["name"]: e["code"] for e in svc.list(kind="LINE", q=tok)}
    assert lines == {ln1: "L-03", ln2: "L-05"}          # LINE carries its no.
    assert {e["name"] for e in svc.list(kind="MACHINE", q=tok)} == {mc1, mc2}
    assert {e["name"] for e in svc.list(kind="DIE", q=tok)} == {die1, die2}
    # Junk placeholders ("-", ".") never entered the master.
    assert all(e["name"] not in ("-", ".") for e in svc.list())
    # Every distinct value now has a master row, so a second run adds nothing.
    assert svc.seed_from_ledger(user="seed") == {"LINE": 0, "MACHINE": 0, "DIE": 0}


def test_case_insensitive_unique_index_backstops_races():
    # A racing insert that skipped the service pre-check must still be rejected by
    # the (kind, lower(name)) unique index, so two counters can't duplicate a row.
    tok = _tok()
    db.session.add(Equipment(kind="MACHINE", name=f"Press {tok}"))
    db.session.commit()
    db.session.add(Equipment(kind="MACHINE", name=f"press {tok}"))
    with pytest.raises(IntegrityError):
        db.session.commit()
    db.session.rollback()
    # A different kind with the same name is still allowed (index is per kind).
    db.session.add(Equipment(kind="LINE", name=f"Press {tok}"))
    db.session.commit()


def test_counter_can_add_but_only_the_master_right_edits(app):
    tok = _tok()
    counter = _client(app, {"scan_part": {"view": True, "edit": True}})
    r = counter.get("/api/v1/equipment?kind=MACHINE")
    assert r.status_code == 200
    r = counter.post("/api/v1/equipment", json={"kind": "MACHINE", "name": f"Quick {tok}"})
    assert r.status_code == 201
    eq_id = r.get_json()["data"]["id"]
    assert counter.put(f"/api/v1/equipment/{eq_id}", json={"name": "X"}).status_code == 403
    assert counter.post("/api/v1/equipment/import").status_code == 403
    assert counter.post("/api/v1/equipment/seed").status_code == 403

    master = _client(app, {"equipment": {"view": True, "edit": True}})
    assert master.put(f"/api/v1/equipment/{eq_id}", json={"name": f"Quick {tok} Fixed"}).status_code == 200
    r = master.get(f"/api/v1/equipment?q=Quick+{tok}")
    assert any(e["name"] == f"Quick {tok} Fixed" for e in r.get_json()["data"])

    nobody = _client(app, {"reports": {"view": True, "edit": False}})
    assert nobody.get("/api/v1/equipment").status_code == 403
