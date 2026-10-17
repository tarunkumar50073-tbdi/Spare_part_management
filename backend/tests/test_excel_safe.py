"""excel_io._excel_safe — what every Excel export passes its text through."""
import io

from openpyxl import load_workbook

from app.utils.excel_io import _excel_safe, build_export


def test_text_that_looks_like_a_formula_stays_text():
    assert _excel_safe("=SUM(A1:A9)") == "'=SUM(A1:A9)" and _excel_safe("-10 MM") == "'-10 MM"
    assert _excel_safe("Bearing 6204") == "Bearing 6204" and _excel_safe(12.5) == 12.5 and _excel_safe(None) is None


def test_a_control_character_in_a_name_does_not_cost_the_sheet():
    """A vertical tab pasted in with an ERP name is refused by the xlsx writer;
    it used to stop the whole export (and the mail carrying it)."""
    assert _excel_safe("Valve\x0bBody") == "ValveBody"
    assert _excel_safe("\x01=HYPERLINK(1)") == "'=HYPERLINK(1)"      # still guarded once the character is gone
    assert _excel_safe("line one\nline two\ttabbed") == "line one\nline two\ttabbed"   # legal white space stays
    book = load_workbook(build_export([("Part", "name")], [{"name": "Valve\x0bBody"}, {"name": "Seal\x1f"}]))
    assert [row[0] for row in book.active.iter_rows(values_only=True)] == ["Part", "ValveBody", "Seal"]


def test_an_export_shows_formula_like_text_as_typed_but_never_as_a_formula():
    """The workbook keeps the guard with Excel's quote prefix instead of a
    visible apostrophe; text the data itself starts with an apostrophe keeps it."""
    rows = [{"name": "=SUM(A1:A9)"}, {"name": "- MODEL-FT-650"}, {"name": "'as typed"}, {"name": "Bearing"}, {"name": 7}]
    sheet = load_workbook(build_export([("Part", "name")], rows)).active
    cells = [sheet.cell(row, 1) for row in range(2, 7)]
    assert [c.value for c in cells] == ["=SUM(A1:A9)", "- MODEL-FT-650", "'as typed", "Bearing", 7]
    assert [(c.data_type, bool(c.quotePrefix)) for c in cells[:3]] == [("s", True), ("s", True), ("s", False)]
    assert cells[4].data_type == "n"
