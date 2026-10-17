"""Notification Engine: the rules and their schedule, the templates, the
reminders / escalation and the manual alert.

No test here can send a real mail: an autouse fixture replaces
`mailer.send_email` with a stub that fails the test if a case forgot to install
its own.

The test database is shared by the whole pytest session, so every assertion is
on rows this file created (by id, by a unique department / address / tracker
key) and never on a table total. Rule runs are reserved per plant day, so each
case takes a Monday of its own far in the future (`_day`). The same fixture
starts every case from the built-in rules, templates and reminders, pins the
plant clock's offset, and puts the app settings back afterwards.
"""
import copy
import io
import uuid
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from sqlalchemy.exc import OperationalError

from app.extensions import db
from app.models.app_setting import AppSetting
from app.models.department import Department
from app.models.department_recipient import DepartmentRecipient
from app.models.email_engine import EmailMessage, NotificationState
from app.models.gate_entry import GateEntry
from app.models.mail_send_log import MailSendLog
from app.models.part import Part
from app.models.receiving_register import ReceivingRegister
from app.models.store import Section, Store
from app.models.user import User
from app.repositories.department_repository import DepartmentRepository
from app.repositories.report_recipient_repository import ReportRecipientRepository
from app.services import email_rules as rules
from app.services import short_rejection_service
from app.services.email_engine import EmailEngine
from app.services.reorder_mail_service import SHEET_COLUMNS, ReorderMailService, xlsx_attachment
from app.services.settings_service import SettingsService
from app.services.spare_management_service import SpareManagementService
from app.services.spare_stock_service import SpareStockService
from app.services.supplier_window_schedule_service import SupplierWindowScheduleService
from app.services.window_time_mail_service import WindowTimeMailService
from app.utils import mailer, rate_limit
from app.utils.excel_io import build_export
from app.utils.exceptions import AppException, BadRequestError, NotFoundError

_ABSENT = object()
_BASE = "/api/v1/mail-engine"
_SETTINGS = ("mail_engine", "smtp", "notification_rules", "mail_templates", "reminder_rules")
# The plant clock in these tests, whatever zone the machine running them is in.
_OFFSET = timedelta(hours=5, minutes=30)
# How every mail ends: a dashed rule, then the note — as typed text, and as the
# same two lines read in an HTML body.
_NOTE = "Note: This is a system-generated mail."
_SIGNATURE = "-" * 47 + "\n" + _NOTE
_SIGNATURE_HTML = "-" * 47 + "<br>" + _NOTE
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class Sender:
    """A stand-in for mailer.send_email that records every call."""

    def __init__(self):
        self.calls = []

    def __call__(self, **kw):
        self.calls.append(kw)
        return {"ok": True, "to": kw["to"], "cc": kw["cc"], "bcc": kw["bcc"]}


def _unpatched_send(**kw):
    raise AssertionError("this test reached mailer.send_email without stubbing it")


@pytest.fixture(autouse=True)
def isolated(app, monkeypatch):
    saved = {key: AppSetting.get_value(key, _ABSENT) for key in _SETTINGS}
    # Whatever .env says: the two old daily mails start switched off here.
    monkeypatch.setitem(app.config, "REORDER_MAIL_ENABLED", False)
    monkeypatch.setitem(app.config, "WINDOW_TIME_MAIL_ENABLED", False)
    AppSetting.query.filter(AppSetting.key.in_(_SETTINGS[2:])).delete(synchronize_session=False)
    # Mails other cases left waiting would be sent through this case's stub.
    EmailMessage.query.filter(EmailMessage.status.in_(["QUEUED", "SENDING"])).update(
        {"status": "CANCELLED", "next_attempt_at": None}, synchronize_session=False)
    db.session.commit()
    monkeypatch.setattr(mailer, "send_email", _unpatched_send)
    monkeypatch.setattr(mailer, "smtp_is_configured", lambda: True)
    monkeypatch.setattr(rules, "_local_offset", lambda: _OFFSET)
    for attr in ("_worker_started", "_thread", "_last_tick_at", "_last_error",
                 "_consecutive_failures", "_last_purge_day"):
        monkeypatch.setattr(EmailEngine, attr, getattr(EmailEngine, attr))
    try:
        yield
    finally:
        db.session.rollback()
        NotificationState.query.filter(
            NotificationState.key.like("tracker:MAINTENANCE:%") | (NotificationState.key == rules.RUN_MARKER)
            | NotificationState.key.like("rule:%")
        ).delete(synchronize_session=False)
        db.session.commit()
        for key, value in saved.items():
            if value is _ABSENT:
                AppSetting.query.filter_by(key=key).delete()
                db.session.commit()
            else:
                AppSetting.set_value(key, value)


@pytest.fixture()
def engine():
    return EmailEngine()


@pytest.fixture()
def send(monkeypatch):
    sender = Sender()
    monkeypatch.setattr(mailer, "send_email", sender)
    return sender


@pytest.fixture()
def client(app):
    return app.test_client()


_seq = {"n": 0, "week": 0}


def _day():
    """A plant-local Monday no other case uses."""
    _seq["week"] += 1
    day = date(2031, 1, 6) + timedelta(weeks=_seq["week"])
    assert day.weekday() == 0
    return day


def _at(day, hour, minute=0):
    """That plant-local day and time as the naive UTC the functions take as `now`."""
    return datetime.combine(day, time(hour, minute)) - _OFFSET


def _user(role="ADMIN"):
    _seq["n"] += 1
    user = User(username=f"ruleuser{_seq['n']}_{uuid.uuid4().hex[:6]}", role=role, is_active=True)
    user.set_password("pw")
    db.session.add(user)
    db.session.commit()
    return user


def _as(client, user):
    with client.session_transaction() as sess:
        sess["user_id"] = user.id
    return client


def _mail(tag):
    return f"{tag}.{uuid.uuid4().hex[:8]}@example.com"


def _row(message_id):
    db.session.expire_all()
    return db.session.get(EmailMessage, message_id)


def _max_id():
    return db.session.query(db.func.max(EmailMessage.id)).scalar() or 0


@pytest.fixture()
def make_dept():
    """Create a department with people — (name, address, role[, active]) — and
    remove it, its people, its reorder log and its reorder and excess mails
    afterwards. SQLite hands a removed department's id to the next one created,
    and the keys of those two mails carry that id: left behind, they would pass
    for the mails of some later case's department. `name` gives the department
    the name of a section or of a section group."""
    made = []

    def make(people, notify=False, active=True, name=None):
        dept = Department(name=name or f"Rule Dept {uuid.uuid4().hex[:8]}", notify_reorder=notify, is_active=active)
        db.session.add(dept)
        db.session.flush()
        for name, email, role, *rest in people:
            db.session.add(DepartmentRecipient(department_id=dept.id, name=name, email=email, role=role,
                                               is_active=rest[0] if rest else True))
        db.session.commit()
        made.append(dept.id)
        return dept

    try:
        yield make
    finally:
        db.session.rollback()
        for dept_id in made:
            for mine in (f"reorder\\_{dept_id}:%", f"excess\\_list:{dept_id}:%"):
                EmailMessage.query.filter(EmailMessage.dedup_key.like(mine, escape="\\")).delete(
                    synchronize_session=False)
            MailSendLog.query.filter_by(report_key=f"reorder_{dept_id}").delete()
            DepartmentRecipient.query.filter_by(department_id=dept_id).delete()
            Department.query.filter_by(id=dept_id).delete()
        db.session.commit()


@pytest.fixture()
def flagged(monkeypatch):
    """Make the given departments the only ones a per-department mail finds."""
    def install(*depts):
        monkeypatch.setattr(DepartmentRepository, "list_all", lambda self, active_only=True: list(depts))
    return install


@pytest.fixture()
def make_section():
    """Create a section of the spare store — under a section group when one is
    named — and remove it afterwards. `name` and `store` make it the namesake,
    in another store, of a section that already exists."""
    made = []

    def make(group=None, name=None, store="MAINTENANCE"):
        store = Store.query.filter_by(code=store).one()
        section = Section(store_id=store.id, name=name or f"Rule Section {uuid.uuid4().hex[:8]}", group_name=group)
        db.session.add(section)
        db.session.commit()
        made.append(section.id)
        return section

    try:
        yield make
    finally:
        db.session.rollback()
        if made:
            Section.query.filter(Section.id.in_(made)).delete(synchronize_session=False)
        db.session.commit()


def _part(no, name, status, qty, minimum, maximum=None, location="R1-S1-B1", section_id=7, **more):
    """A stock row with every field the stock list gives a part (the Daily
    Stock Report reads them all); what a test does not name is unknown."""
    return {"part_id": no, "part_no": no, "part_name": name, "status": status, "qty_on_hand": qty,
            "min_stock_level": minimum, "reorder_level": None, "max_stock_level": maximum,
            "location": location, "section_id": section_id, "sub_category": None, "specification": None,
            "uom": None, "unit_price": None, "stock_value": None, **more}


def _over(no, qty=30, maximum=10, rate=50.0, **more):
    """A stock row above its maximum: 20 over at ₹ 50 unless told otherwise."""
    return _part(no, f"Part {no}", "Above Maximum", qty, 2, maximum, unit_price=rate,
                 stock_value=None if rate is None else qty * rate, **more)


@pytest.fixture()
def gate_part():
    """Create a spare part of the store and — when a supplier is named — its
    gate entry from that supplier: the real rows the supplier column is read
    from. Returns the part's id; the rows are removed afterwards."""
    parts, entries = [], []

    def make(supplier=None):
        part = Part(part_no=f"RULE-{uuid.uuid4().hex[:10].upper()}", part_name="Rule test part", uom="NOS",
                    part_type="OTHER", store_code="MAINTENANCE")
        db.session.add(part)
        db.session.flush()
        parts.append(part.id)
        if supplier:
            # The gate keeps the item code as it was typed: matched whatever its case.
            entry = GateEntry(store_code="MAINTENANCE", item_code=part.part_no.lower(), ven_cd="V-RULE",
                              ven_name=supplier, gate_entry_dt=date(2026, 10, 1))
            db.session.add(entry)
            db.session.flush()
            entries.append(entry.id)
        db.session.commit()
        return part.id

    try:
        yield make
    finally:
        db.session.rollback()
        if entries:
            GateEntry.query.filter(GateEntry.id.in_(entries)).delete(synchronize_session=False)
        if parts:
            Part.query.filter(Part.id.in_(parts)).delete(synchronize_session=False)
        db.session.commit()


@pytest.fixture()
def facts_read(monkeypatch):
    """Record every read of the per-part facts the supplier comes from. A run
    reads them once, however many parts and departments it has."""
    asked, read = [], SpareStockService.part_facts

    def counted(self, store_code):
        asked.append(store_code)
        return read(self, store_code)

    monkeypatch.setattr(SpareStockService, "part_facts", counted)
    return asked


def _sheet(attachment):
    """(worksheet, its table as lists — the header and the data rows) of an
    .xlsx attachment — a stored EmailAttachment or the dict handed to the
    mailer. The table is what its filter covers. The reorder and excess sheets
    have one empty row and then a legend line in column A under it, which are
    not part of the table (their tests pin them); nothing else may be there."""
    content = attachment["content"] if isinstance(attachment, dict) else attachment.content
    sheet = load_workbook(io.BytesIO(content)).active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    end = range_boundaries(sheet.auto_filter.ref)[3]
    if len(rows) != end:
        blank, legend = rows[end:]
        assert set(blank) == {None} and isinstance(legend[0], str) and set(legend[1:]) == {None}
    return sheet, rows[:end]


def _layout_of(cell):
    """What a cell holds and how it is laid out — everything but its colours
    and boldness: the value, number format, formula-guard quote, font face and
    size, alignment and border."""
    side = cell.border
    return (cell.value, cell.number_format, cell.quotePrefix, cell.font.name, cell.font.sz,
            cell.alignment.horizontal, cell.alignment.vertical, cell.alignment.wrap_text,
            side.left.style, side.right.style, side.top.style, side.bottom.style)


def _colour_of(cell):
    """(fill, text colour, bold) of a cell: None for no fill and for the
    default text colour."""
    fill = cell.fill.fgColor.rgb if cell.fill.fill_type == "solid" else None
    color = cell.font.color
    return fill, color.rgb if color is not None and color.type == "rgb" else None, bool(cell.font.b)


def _plain_but_for(sheet, table, painted):
    """The coloured sheet is the plain build_export sheet of the same rows in
    everything but the colours of its `painted` columns: the same values,
    formula guard, fonts, borders and alignment in every cell, the same
    header, column widths, frozen header and filter."""
    plain = load_workbook(build_export([(heading, heading) for heading in table[0]],
                                       [dict(zip(table[0], line)) for line in table[1:]])).active
    assert plain.max_row == len(table)
    for number in range(1, len(table) + 1):
        for column, heading in enumerate(table[0], start=1):
            mine, theirs = sheet.cell(number, column), plain.cell(number, column)
            assert _layout_of(mine) == _layout_of(theirs), (number, heading)
            if number == 1 or heading not in painted:
                assert _colour_of(mine) == _colour_of(theirs), (number, heading)
    for column in range(1, len(table[0]) + 1):
        letter = get_column_letter(column)
        assert sheet.column_dimensions[letter].width == plain.column_dimensions[letter].width
    assert (sheet.freeze_panes, sheet.auto_filter.ref) == (plain.freeze_panes, plain.auto_filter.ref)


def _legend_of(sheet, data_rows):
    """The legend cell under a coloured sheet's table: after one empty row,
    in column A, and the last thing on the sheet."""
    below = data_rows + 3
    assert sheet.max_row == below
    assert [cell.value for cell in sheet[below - 1]] == [None] * sheet.max_column
    assert [cell.value for cell in sheet[below]][1:] == [None] * (sheet.max_column - 1)
    legend = sheet.cell(below, 1)
    # Not part of the table: no border, no fill, not bold — a quiet note in italics.
    assert legend.border.left.style is None and legend.border.bottom.style is None
    assert legend.fill.fill_type is None and not legend.font.b and legend.font.i
    return legend.value


def _no_list_in(html, text, part_nos):
    """A parts mail — the reorder list, the excess list, the stock report —
    carries its list only as the attached Excel (the store's instruction of
    14 Oct 2026): its body has no table, names no part and stays a few lines
    long, however long the list is."""
    assert "<table" not in html and "<tr" not in html and " | " not in text
    for part_no in part_nos:
        assert part_no not in html and part_no not in text
    assert len(html.encode("utf-8")) < 2 * 1024


@pytest.fixture()
def stock(monkeypatch):
    """Make the store's stock list a given set of rows (section 7 = "Maintenance")."""
    def install(rows):
        monkeypatch.setattr(SpareStockService, "list_stock",
                            lambda self, store_code, section_id=None, section_ids=None: list(rows))
        monkeypatch.setattr(rules.SpareDashboardService, "_section_names",
                            staticmethod(lambda store_code: {7: "Maintenance"}))
        return rows
    return install


def _item(key, label, status="OPEN", count=3, target=None, owner="Purchase", value=None,
          value_label=None, action="Raise PO", remarks=""):
    return {"key": key, "label": label, "status": status, "count": count, "value": value,
            "value_label": value_label, "action": action, "owner": owner, "target_date": target,
            "remarks": remarks}


@pytest.fixture()
def tracker(monkeypatch):
    """Make the Action Tracker a given list of lines."""
    def install(*items):
        monkeypatch.setattr(SpareManagementService, "management",
                            lambda self, store_code, **kw: {"tracker": {"items": list(items), "summary": {}}})
    return install


def _line_key():
    return f"rt{uuid.uuid4().hex[:10]}"


def _remind(key, owners=("owner@example.com",), escalate=("boss@example.com",), **settings):
    """Save reminder settings with one line's owners, bypassing the route's
    check that the key is a real tracker line (these cases use keys of their own)."""
    AppSetting.set_value("reminder_rules", {
        "enabled": True, **settings,
        "escalate_to": [{"email": e, "role": "TO"} for e in escalate],
        "owners": {key: {"department_ids": [], "emails": [{"email": e, "role": "TO"} for e in owners]}}})


def _line(view, key):
    return next(line for line in view["lines"] if line["key"] == key)


def _state(key):
    db.session.expire_all()
    return NotificationState.query.filter_by(key=f"tracker:MAINTENANCE:{key}").first()


def _reorder_subject(dept_name, today, count=2):
    return f"Spare Reorder Alert — {dept_name} — {count} part(s) [{today.strftime('%d.%m.%Y')}]"


def _reorder_text(dept_name):
    """The plain-text body of the reorder mail for _REORDER_ROWS with the
    built-in wording, written out — so a changed default or layout shows up
    here. The parts themselves are only in the attached sheet."""
    return "\n\n".join([
        "Dear Sir,",
        f"Please find attached the list of 2 spare part(s) for {dept_name} that are at or below the reorder "
        "level and need a purchase order.",
        "2 parts to reorder · 1 zero stock · 0 below minimum · 1 at reorder level",
        _SIGNATURE,
    ])


_REORDER_ROWS = [
    {"part_no": "P-1", "part_name": "Bearing 6204", "status": "Out of Stock", "qty_on_hand": 0,
     "min_stock_level": 2, "reorder_level": 3, "max_stock_level": 9, "location": "R1-S1-B1"},
    {"part_no": "P-2-WITH-A-LONG-NUMBER", "part_name": "V-belt", "status": "Reorder", "qty_on_hand": 2.5,
     "min_stock_level": 1, "reorder_level": 3, "max_stock_level": None, "location": None},
]


@pytest.fixture()
def reorder_for(monkeypatch):
    """Make the reorder mail see only the given department, with fixed parts."""
    def install(dept, rows=_REORDER_ROWS):
        monkeypatch.setattr(DepartmentRepository, "list_all", lambda self, active_only=True: [dept])
        monkeypatch.setattr(ReorderMailService, "_reorder_rows", lambda self, section_ids: list(rows))
    return install


@pytest.fixture()
def window_sheet(monkeypatch):
    """Give the window-time mail one recipient and a stand-in sheet. Yields the
    list of (year, month) sheets asked for. The send-log rows the case makes
    are removed (test_window_time_mail counts that table)."""
    built = []

    def workbook(self, year, month):
        built.append((year, month))
        return io.BytesIO(b"sheet"), "application/octet-stream", "xlsx"

    monkeypatch.setattr(ReportRecipientRepository, "recipients_for",
                        lambda self, report_key: (["sheet@example.com"], []))
    monkeypatch.setattr(SupplierWindowScheduleService, "window_time_workbook", workbook)
    logged = [row.id for row in MailSendLog.query.filter_by(report_key="window_time").all()]
    try:
        yield built
    finally:
        db.session.rollback()
        MailSendLog.query.filter(MailSendLog.report_key == "window_time",
                                 MailSendLog.id.notin_(logged)).delete(synchronize_session=False)
        db.session.commit()


# ---- the catalogue and its defaults ----

def test_catalogue_lists_six_rules_all_off_by_default(app, stock, tracker, make_dept, flagged):
    stock([_over("X-1"), _part("C-1", "Seal", "Out of Stock", 0, 2)])
    tracker(_item("critical", "Critical parts"), _item("excess", "Excess above max", status="CLEAR", count=0))
    flagged(make_dept([("Asha", _mail("to"), "TO")], notify=True))      # named like no section: the whole store
    listed = rules.list_rules()
    assert [r["key"] for r in listed] == ["reorder_daily", "excess_list", "stock_report", "tracker_digest",
                                          "window_time_daily", "short_rejection"]
    assert [r["audience"] for r in listed] == ["departments_flagged", "departments_flagged", "custom", "custom",
                                               "report_recipients", "custom"]
    assert [r["mail_type"] for r in listed] == ["reorder", "excess_list", "stock_report", "tracker_digest",
                                                "window_time", "short_rejection"]
    for rule in listed:
        assert set(rule) == {"key", "title", "description", "mail_type", "audience", "enabled", "hour", "minute",
                             "days", "department_ids", "emails", "recipients", "content", "last_run",
                             "next_run_at", "is_default"}
        assert rule["enabled"] is False and rule["is_default"] is True and rule["next_run_at"] is None
        assert set(rule["recipients"]) == {"to", "cc", "bcc"} and set(rule["content"]) == {"count", "label"}
    by_key = {r["key"]: r for r in listed}
    for key in ("excess_list", "stock_report"):
        assert (by_key[key]["hour"], by_key[key]["minute"], by_key[key]["days"]) == (8, 0, [0, 1, 2, 3, 4, 5])
        assert (by_key[key]["department_ids"], by_key[key]["emails"]) == ([], [])
    for key in ("stock_report", "tracker_digest", "short_rejection"):   # a free audience starts with no one in it
        assert by_key[key]["recipients"] == {"to": [], "cc": [], "bcc": []}
    assert (by_key["excess_list"]["title"], by_key["stock_report"]["title"]) == (
        "Excess parts list", "Daily stock report")
    # The parts mails carry their list as the attached Excel only, and say so.
    assert by_key["excess_list"]["description"] == (
        "Mails each department ticked for the reorder mail the list of its parts above their maximum level as "
        "an Excel sheet, the dearest excess first.")
    assert by_key["stock_report"]["description"] == (
        "The Daily Stock Report as an attached Excel workbook — summary, every part, critical & reorder, excess, "
        "by category — with one line of totals in the mail.")
    assert by_key["excess_list"]["content"] == {"count": 1, "label": "parts above maximum"}
    assert by_key["stock_report"]["content"] == {"count": 2, "label": "parts in the report"}
    assert by_key["tracker_digest"]["content"]["count"] == 1               # one OPEN line
    assert by_key["window_time_daily"]["content"]["count"] is None        # a sheet has no such number
    assert isinstance(by_key["short_rejection"]["content"]["count"], int)
    # The two old daily mails ran every day, at the hour .env gave them.
    assert by_key["reorder_daily"]["days"] == list(range(7))
    assert by_key["reorder_daily"]["hour"] == app.config["REORDER_MAIL_HOUR"]
    assert by_key["window_time_daily"]["hour"] == app.config["WINDOW_TIME_MAIL_HOUR"]
    assert rules.rules_summary() == {"enabled": 0, "total": 6, "next": None}


def test_the_critical_digest_rule_is_gone_and_what_was_saved_for_it_is_ignored(engine, stock, tracker, client):
    """The reorder mail already lists the zero-stock and below-minimum parts,
    so the separate critical mail was dropped. Its rows in the settings stay
    where an admin left them and decide nothing."""
    stock([_part("C-1", "Seal", "Out of Stock", 0, 2, 9)])
    tracker(_item("excess", "Excess above max", status="CLEAR", count=0))
    AppSetting.set_value("notification_rules", {"critical_digest": {
        "enabled": True, "hour": 0, "minute": 0, "days": list(range(7)), "department_ids": [],
        "emails": [{"email": "plant@example.com", "role": "TO"}]}})
    AppSetting.set_value("mail_templates", {"templates": {"critical_digest": {"subject": "Critical {count}"}}})
    day, before = _day(), _max_id()
    assert "critical_digest" not in [r["key"] for r in rules.list_rules(_at(day, 9))]
    assert rules.rules_summary(_at(day, 9)) == {"enabled": 0, "total": 6, "next": None}
    assert rules.run_due_rules(engine, _at(day, 9)) == []
    assert "critical_digest" not in [t["mail_type"] for t in rules.list_templates()["templates"]]
    for gone in (lambda: rules.rule_settings("critical_digest"), lambda: rules.describe_rule("critical_digest"),
                 lambda: rules.rule_is_due("critical_digest", _at(day, 9)),
                 lambda: rules.save_rule("critical_digest", {"enabled": False}),
                 lambda: rules.run_rule(engine, "critical_digest"), lambda: rules.preview_rule("critical_digest"),
                 lambda: rules.template("critical_digest"),
                 lambda: rules.save_template("critical_digest", {"subject": "x"})):
        with pytest.raises(NotFoundError):
            gone()

    c = _as(client, _user())
    assert "critical_digest" not in [r["key"] for r in c.get(_BASE + "/rules").get_json()["data"]]
    assert c.put(_BASE + "/rules/critical_digest", json={"enabled": True}).status_code == 404
    assert c.post(_BASE + "/rules/critical_digest/run").status_code == 404
    assert c.get(_BASE + "/rules/critical_digest/preview").status_code == 404
    assert c.put(_BASE + "/templates/critical_digest", json={"subject": "x"}).status_code == 404
    assert _max_id() == before                                # nothing was built, queued or recorded


def test_env_flags_are_the_default_only_until_the_rule_is_saved(app, monkeypatch):
    monkeypatch.setitem(app.config, "REORDER_MAIL_ENABLED", True)
    monkeypatch.setitem(app.config, "REORDER_MAIL_HOUR", 6)
    settings, is_default = rules.rule_settings("reorder_daily")
    assert is_default and settings["enabled"] is True and settings["hour"] == 6
    rules.save_rule("reorder_daily", {"minute": 30})
    monkeypatch.setitem(app.config, "REORDER_MAIL_ENABLED", False)
    monkeypatch.setitem(app.config, "REORDER_MAIL_HOUR", 11)
    settings, is_default = rules.rule_settings("reorder_daily")
    assert not is_default
    assert (settings["enabled"], settings["hour"], settings["minute"]) == (True, 6, 30)   # what the admin saw
    # The rules built here never read .env.
    assert rules.rule_settings("excess_list")[0]["enabled"] is False


# ---- the schedule ----

def test_rule_is_due_only_when_enabled_on_its_day_and_at_or_after_its_time(engine, stock):
    stock([_over("C-1")])
    day = _day()
    assert rules.rule_is_due("stock_report", _at(day, 12)) is False           # switched off
    rules.save_rule("stock_report", {"enabled": True, "hour": 8, "minute": 15,
                                     "emails": [{"email": "plant@example.com", "role": "TO"}]})
    assert rules.rule_is_due("stock_report", _at(day, 8, 14)) is False        # before its time
    assert rules.rule_is_due("stock_report", _at(day, 8, 15)) is True
    assert rules.rule_is_due("stock_report", _at(day + timedelta(days=6), 12)) is False   # Sunday is not ticked
    assert rules.describe_rule("stock_report", _at(day, 7))["next_run_at"] == _at(day, 8, 15).isoformat()
    assert rules.run_due_rules(engine, _at(day, 8, 14)) == []

    ran = rules.run_due_rules(engine, _at(day, 8, 15))
    assert [(r["key"], r["status"]) for r in ran] == [("stock_report", "queued")]
    message = ran[0]["messages"][0]
    row = _row(message["id"])
    assert row.dedup_key == f"stock_report:{day.isoformat()}" and row.status == "QUEUED"
    assert (row.trigger, row.triggered_by, row.rule_key, row.store_code) == (
        "RULE", "scheduler", "stock_report", "MAINTENANCE")
    assert row.to_list == ["plant@example.com"] and row.body_html

    # Already ran today: not due, and a second pass makes nothing.
    assert rules.rule_is_due("stock_report", _at(day, 9)) is False
    before = _max_id()
    assert rules.run_due_rules(engine, _at(day, 9)) == []
    assert _max_id() == before
    described = rules.describe_rule("stock_report", _at(day, 9))
    assert described["next_run_at"] == _at(day + timedelta(days=1), 8, 15).isoformat()
    assert described["last_run"]["status"] == "QUEUED" and described["last_run"]["detail"] == row.subject

    # The server was down at 08:15 on Tuesday: the rule still runs later that day, once.
    tuesday = day + timedelta(days=1)
    assert rules.rule_is_due("stock_report", _at(tuesday, 15, 30)) is True
    caught_up = rules.run_due_rules(engine, _at(tuesday, 15, 30))
    assert _row(caught_up[0]["messages"][0]["id"]).dedup_key == f"stock_report:{tuesday.isoformat()}"
    assert rules.rule_is_due("stock_report", _at(tuesday, 23, 59)) is False


def test_next_run_skips_days_that_are_not_ticked_and_the_summary_names_the_first(engine, stock, tracker):
    stock([_over("C-1")])
    tracker(_item("critical", "Critical parts"))
    day = _day()
    to = [{"email": "plant@example.com", "role": "TO"}]
    rules.save_rule("stock_report", {"enabled": True, "hour": 8, "days": [0], "emails": to})
    rules.save_rule("tracker_digest", {"enabled": True, "hour": 7, "minute": 30, "emails": to})
    assert rules.rules_summary(_at(day, 6)) == {
        "enabled": 2, "total": 6,
        "next": {"key": "tracker_digest", "title": "Action Tracker summary", "at": _at(day, 7, 30).isoformat()}}
    rules.run_due_rules(engine, _at(day, 8))
    # Mondays only: once today's run is done the next one is a week away.
    assert rules.describe_rule("stock_report", _at(day, 8, 1))["next_run_at"] == \
        _at(day + timedelta(days=7), 8).isoformat()
    assert rules.rules_summary(_at(day, 8, 1))["next"]["at"] == _at(day + timedelta(days=1), 7, 30).isoformat()
    # Switched on with no day at all, a rule never runs.
    rules.save_rule("stock_report", {"days": []})
    assert rules.describe_rule("stock_report", _at(day, 9))["next_run_at"] is None
    assert rules.rule_is_due("stock_report", _at(day + timedelta(days=7), 9)) is False


def test_one_failing_rule_does_not_stop_the_others(engine, stock, tracker, monkeypatch):
    tracker(_item("critical", "Critical parts"))
    day = _day()
    to = [{"email": "plant@example.com", "role": "TO"}]
    rules.save_rule("stock_report", {"enabled": True, "hour": 8, "emails": to})
    rules.save_rule("tracker_digest", {"enabled": True, "hour": 8, "emails": to})

    def broken(self, store_code, section_id=None, section_ids=None):
        raise RuntimeError("stock list is down")
    monkeypatch.setattr(SpareStockService, "list_stock", broken)
    with pytest.raises(RuntimeError, match="stock_report: RuntimeError: stock list is down"):
        rules.run_due_rules(engine, _at(day, 8))
    assert EmailMessage.query.filter_by(dedup_key=f"tracker_digest:{day.isoformat()}").one().status == "QUEUED"
    # Nothing was reserved for the rule that failed, so the next tick tries it again.
    assert EmailMessage.query.filter_by(dedup_key=f"stock_report:{day.isoformat()}").first() is None
    stock([_over("C-1")])
    assert [r["key"] for r in rules.run_due_rules(engine, _at(day, 8, 1))] == ["stock_report"]


def test_rule_settings_are_validated_on_save():
    for bad in ({"hour": 24}, {"hour": -1}, {"hour": "8"}, {"hour": True}, {"minute": 60}, {"days": [7]},
                {"days": "0,1"}, {"days": [1.5]}, {"enabled": "yes"}, {"nope": 1}, {}, None, [],
                {"emails": [{"email": "not-an-address", "role": "TO"}]},
                {"emails": [{"email": "a@example.com\nBcc: x@example.com"}]},
                {"emails": [{"email": "a@example.com", "role": "FROM"}]},
                {"emails": ["a@example.com"]}, {"emails": "a@example.com"},
                {"department_ids": [987654321]}, {"department_ids": ["1"]}):
        with pytest.raises(BadRequestError):
            rules.save_rule("stock_report", bad)
    assert rules.rule_settings("stock_report")[1] is True                    # nothing half-applied
    # The audience of the two old mails is kept on its own master screen, and
    # the excess list goes to the departments the reorder mail goes to.
    for fixed in ("reorder_daily", "excess_list", "window_time_daily"):
        with pytest.raises(BadRequestError, match="fixed"):
            rules.save_rule(fixed, {"emails": [{"email": "a@example.com", "role": "TO"}]})
        with pytest.raises(BadRequestError, match="fixed"):
            rules.save_rule(fixed, {"department_ids": []})
        assert rules.rule_settings(fixed)[1] is True
    with pytest.raises(NotFoundError):
        rules.save_rule("no_such_rule", {"enabled": True})
    with pytest.raises(NotFoundError):
        rules.preview_rule("no_such_rule")

    saved = rules.save_rule("stock_report", {"days": [5, 0, 5], "emails": [
        {"email": " One@Example.com ", "role": "cc"}, {"email": "one@example.com", "role": "TO"},
        {"email": "two@example.com"}]}, user="asha")
    assert saved["days"] == [0, 5] and saved["is_default"] is False
    assert saved["emails"] == [{"email": "One@Example.com", "role": "CC"}, {"email": "two@example.com", "role": "TO"}]
    assert AppSetting.query.filter_by(key="notification_rules").one().updated_by == "asha"
    # A row someone edited by hand cannot schedule a rule at hour 99.
    AppSetting.set_value("notification_rules", {"stock_report": {
        "enabled": True, "hour": 99, "minute": "x", "days": [0, 9, "a"], "emails": [{"email": "bad"}, "junk"],
        "department_ids": "all"}})
    settings, _ = rules.rule_settings("stock_report")
    assert settings == {"enabled": True, "hour": 8, "minute": 0, "days": [0], "department_ids": [], "emails": []}


def test_an_audience_saved_for_the_excess_list_before_it_went_department_wise_is_ignored(make_dept, flagged):
    """The excess list used to go to people chosen on the rule. What was saved
    then stays in the row and decides nothing: the mail goes to the departments
    ticked for the reorder mail."""
    chosen = make_dept([("Old", _mail("old"), "TO")])
    ticked_to = _mail("ticked")
    flagged(chosen, make_dept([("Asha", ticked_to, "TO")], notify=True))
    AppSetting.set_value("notification_rules", {"excess_list": {
        "enabled": True, "hour": 6, "minute": 30, "days": [0, 2], "department_ids": [chosen.id],
        "emails": [{"email": "plant@example.com", "role": "TO"}]}})
    settings, is_default = rules.rule_settings("excess_list")
    assert settings == {"enabled": True, "hour": 6, "minute": 30, "days": [0, 2], "department_ids": [], "emails": []}
    assert is_default is False
    described = rules.describe_rule("excess_list")
    assert described["audience"] == "departments_flagged"
    assert described["recipients"] == {"to": [ticked_to], "cc": [], "bcc": []}
    # The next save of the rule drops the old audience from the row.
    rules.save_rule("excess_list", {"minute": 45})
    assert AppSetting.get_value("notification_rules")["excess_list"] == {
        "enabled": True, "hour": 6, "minute": 45, "days": [0, 2], "department_ids": [], "emails": []}


# ---- recipients of a free audience ----

def test_custom_audience_takes_department_people_in_their_roles_plus_extra_addresses(make_dept):
    a, b, c, gone, extra = (_mail(t) for t in ("to", "cc", "bcc", "left", "extra"))
    dept = make_dept([("Asha", a, "TO"), ("Ravi", b, "CC"), ("Hidden", c, "BCC"),
                      ("Left", gone, "TO", False), ("No Address", None, "TO")])
    inactive = make_dept([("Old", _mail("old"), "TO")], active=False)
    saved = rules.save_rule("stock_report", {
        "department_ids": [dept.id, inactive.id],
        # An address the department already has in To stays in To, once.
        "emails": [{"email": extra, "role": "CC"}, {"email": a.upper(), "role": "BCC"}]})
    people = saved["recipients"]
    assert people["to"][0] == a and people["cc"] == [b, extra] and people["bcc"] == [c]
    assert gone not in people["to"] and len(people["to"]) == 2      # a ticked department counts, active or not
    assert saved["department_ids"] == [dept.id, inactive.id]

    # A department that was removed since adds no one and breaks nothing.
    AppSetting.set_value("notification_rules", {"stock_report": {
        "department_ids": [987654321, dept.id], "emails": []}})
    assert rules.describe_rule("stock_report")["recipients"] == {"to": [a], "cc": [b], "bcc": [c]}


def test_with_no_to_address_the_first_cc_is_promoted(make_dept):
    first, second, hidden = _mail("cc1"), _mail("cc2"), _mail("bcc")
    dept = make_dept([("One", first, "CC"), ("Two", second, "CC"), ("Hidden", hidden, "BCC")])
    people = rules.save_rule("tracker_digest", {"department_ids": [dept.id]})["recipients"]
    assert people == {"to": [first], "cc": [second], "bcc": [hidden]}
    # A Bcc alone is never moved into view.
    only_bcc = make_dept([("Hidden", hidden, "BCC")])
    people = rules.save_rule("tracker_digest", {"department_ids": [only_bcc.id]})["recipients"]
    assert people == {"to": [], "cc": [], "bcc": [hidden]}


# ---- the excess parts list ----

def _excess_subject(dept_name, day, count):
    return f"Excess Parts — {dept_name} — {count} part(s) above maximum [{day:%d.%m.%Y}]"


def _excess_file(dept_name, day):
    return f"Excess Parts - {dept_name} {day:%d.%m.%Y}.xlsx"


def test_excess_list_puts_the_dearest_excess_first_and_attaches_every_part(
        engine, stock, send, gate_part, facts_read, make_dept, flagged):
    supplied = gate_part("Bharat Seals <Pvt> & Co")
    rows = stock([
        _over("E-SMALL", qty=12, rate=5.0),                                         # 2 over: ₹ 10
        _over("E-NORATE-3", qty=13, rate=None),                                     # 3 over, no rate
        _over("E-BIG", part_id=supplied, sub_category="Seal kit", uom="NOS"),       # 20 over: ₹ 1,000
        _over("E-NORATE-9", qty=19, rate=None, location=None, section_id=None),     # 9 over, no rate
        _over("E-MID", qty=14.5, rate=100.0),                                       # 4.5 over: ₹ 450
        _part("E-NOMAX", "No maximum planned", "Above Maximum", 5, 2, None),        # nothing to be above
        _part("H-OK", "Healthy part", "Healthy", 9, 2, 20),
        _part("Z-MIN", "Seal", "Out of Stock", 0, 2, 9),
    ])
    untouched = copy.deepcopy(rows)
    to, cc, bcc = _mail("to"), _mail("cc"), _mail("bcc")
    # A department named like no section: its list is the whole store's.
    dept = make_dept([("Asha", to, "TO"), ("Ravi", cc, "CC"), ("Hidden", bcc, "BCC")], notify=True)
    flagged(dept)
    day = _day()
    result = rules.run_rule(engine, "excess_list", _at(day, 10))
    assert result["status"] == "queued" and result["reason"] is None and len(result["messages"]) == 1
    message = result["messages"][0]
    assert message["mail_type"] == "excess_list" and message["rule_key"] == "excess_list"
    assert message["status"] == "QUEUED" and message["trigger"] == "MANUAL"
    assert (message["to"], message["cc"], message["bcc"]) == ([to], [cc], [bcc])
    assert message["subject"] == _excess_subject(dept.name, day, 5)
    assert [a["filename"] for a in message["attachments"]] == [_excess_file(dept.name, day)]
    row = _row(message["id"])
    # The day's mail had not gone yet, so the button's mail is it.
    assert row.dedup_key == f"excess_list:{dept.id}:{day.isoformat()}"
    html, text = row.body_html, row.body_text
    # The money is that of the rated parts — 1,000 + 450 + 10 — and the line says how many it left out.
    totals = "5 parts above maximum · ₹ 1,460 tied up · 2 without a rate, not valued"
    assert f"<strong>{totals}</strong>" in html
    assert text == "\n\n".join([
        f"Dear Sir,\n\nPlease find attached the list of 5 spare part(s) for {dept.name} that are above their "
        "maximum level.",
        totals,
        "Please stop ordering these parts and review their maximum level.",
        _SIGNATURE,
    ])
    _no_list_in(html, text, [r["part_no"] for r in rows])
    # Nor a supplier's name: that is in the sheet.
    assert "Bharat Seals" not in html and "<Pvt>" not in html
    assert "http" not in html and "<style" not in html and "<link" not in html     # nothing external
    assert html.endswith(f"{_SIGNATURE_HTML}</p></body></html>") and html.count(_SIGNATURE_HTML) == 1
    assert rows == untouched                                   # the stock list's own rows were only read
    assert facts_read == ["MAINTENANCE"]                       # the suppliers: one read for the whole run

    stored = row.attachments[0]
    assert (stored.filename, stored.mime) == (_excess_file(dept.name, day), _XLSX)
    assert stored.filename.isascii()                           # no long dash: Outlook shows the name as it is
    sheet, table = _sheet(stored)
    assert table[0] == ["Section", "Category", "Part No", "Description", "Supplier", "Location", "UOM",
                        "In Stock", "Max", "Excess Qty", "Rate", "Excess Value"]
    # Dearest excess first; the parts with no rate after them, largest quantity first. Only the parts
    # above a planned maximum: no E-NOMAX, H-OK or Z-MIN.
    assert [line[2] for line in table[1:]] == ["E-BIG", "E-MID", "E-SMALL", "E-NORATE-9", "E-NORATE-3"]
    # A supplier's name is data: it is kept as it was typed.
    assert table[1] == ["Maintenance", "Seal kit", "E-BIG", "Part E-BIG", "Bharat Seals <Pvt> & Co", "R1-S1-B1",
                        "NOS", 30, 10, 20, 50, 1000]
    assert table[2] == ["Maintenance", None, "E-MID", "Part E-MID", None, "R1-S1-B1", None, 14.5, 10, 4.5, 100, 450]
    # What is not known is an empty cell: no section, no supplier, no location, no rate, no value.
    assert table[4] == [None, None, "E-NORATE-9", "Part E-NORATE-9", None, None, None, 19, 10, 9, None, None]
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:L6" and sheet["A1"].font.bold

    engine.process_due()
    assert _row(message["id"]).status == "SENT"
    sent = send.calls[-1]
    assert sent["html"] == html and sent["body"] == text and sent["subject"] == message["subject"]
    assert [(a["filename"], a["mime"]) for a in sent["attachments"]] == [(stored.filename, _XLSX)]
    assert sent["attachments"][0]["content"] == stored.content

    # With a rate on every part the totals line stops at the money.
    stock([_over("E-BIG"), _over("E-SMALL", qty=12, rate=5.0)])
    rated = rules.preview_rule("excess_list", _at(day, 10))
    assert "2 parts above maximum · ₹ 1,010 tied up" in rated["text"].split("\n\n")
    assert "without a rate" not in rated["text"] and "without a rate" not in rated["html"]


def test_a_long_excess_list_is_one_line_in_the_mail_and_every_part_in_the_sheet(engine, stock, make_dept, flagged):
    parts = [f"X-{n:03d}" for n in range(130)]
    stock([_over(no, qty=210 - n, rate=1.0) for n, no in enumerate(parts)])
    flagged(make_dept([("Asha", _mail("to"), "TO")], notify=True))
    preview = rules.preview_rule("excess_list")
    assert preview["count"] == 130
    assert "130 parts above maximum · ₹ 17,615 tied up" in preview["text"].split("\n\n")
    _no_list_in(preview["html"], preview["text"], parts)
    result = rules.run_rule(engine, "excess_list", _at(_day(), 10))
    row = _row(result["messages"][0]["id"])
    _no_list_in(row.body_html, row.body_text, parts)
    _, table = _sheet(row.attachments[0])
    assert [line[2] for line in table[1:]] == parts                # every part, the dearest excess first


def test_excess_sheet_has_in_stock_yellow_and_the_excess_value_coloured_by_the_money_it_ties_up(stock):
    """The store's choice (14 Oct 2026, "In Stock yellow + Excess Value ₹ ke
    hisaab se"): every In Stock cell in the yellow of his stock report's
    above-max stock; the Excess Value in solid dark amber from ₹ 1,00,000,
    yellow from ₹ 25,000 — AT a threshold counts — and no colour below that or
    with no rate. Nothing else is coloured, and the legend says what the
    colours mean, with the thresholds as they are set."""
    assert (rules.EXCESS_HIGH_RUPEES, rules.EXCESS_WATCH_RUPEES) == (100_000, 25_000)
    stock([
        _over("E-2L", qty=2010, rate=100.0),                # 2,000 over: ₹ 2,00,000
        _over("E-1L", qty=1010, rate=100.0),                # 1,000 over: ₹ 1,00,000 exactly
        _over("E-JUST-UNDER", qty=11, rate=99999.99),       # 1 over: ₹ 99,999.99
        _over("E-25K", qty=260, rate=100.0),                # 250 over: ₹ 25,000 exactly
        _over("E-24999", qty=11, rate=24999.0),             # 1 over: ₹ 24,999
        _over("E-NORATE", qty=12, rate=None),               # 2 over, no rate: no value
    ])
    found = rules._excess_rows()
    attachment = rules._excess_sheet("Tool Room", found, date(2026, 10, 14), {7: "Maintenance"}, {})
    assert (attachment["filename"], attachment["mime"]) == ("Excess Parts - Tool Room 14.10.2026.xlsx", _XLSX)
    sheet, table = _sheet(attachment)
    assert table[0] == list(rules.EXCESS_COLUMNS)
    assert [(line[2], line[11]) for line in table[1:]] == [
        ("E-2L", 200000), ("E-1L", 100000), ("E-JUST-UNDER", 99999.99), ("E-25K", 25000), ("E-24999", 24999),
        ("E-NORATE", None)]
    columns = {heading: number for number, heading in enumerate(table[0], start=1)}

    yellow, dark_amber, none = ("00FEF9C3", "00854D0E", True), ("00D97706", "00FFFFFF", True), (None, None, False)
    for number, value_colour in enumerate([dark_amber, dark_amber, yellow, yellow, none, none], start=2):
        assert _colour_of(sheet.cell(number, columns["In Stock"])) == yellow
        assert _colour_of(sheet.cell(number, columns["Excess Value"])) == value_colour
        for heading, column in columns.items():
            if heading not in ("In Stock", "Excess Value"):
                assert _colour_of(sheet.cell(number, column)) == none, (number, heading)
    _plain_but_for(sheet, table, ("In Stock", "Excess Value"))
    assert sheet.auto_filter.ref == "A1:L7" and sheet.freeze_panes == "A2"
    assert _legend_of(sheet, len(found)) == (
        "In Stock: yellow = above max · Excess Value: dark amber = ₹1,00,000 or more · yellow = ₹25,000 to ₹1,00,000")


def test_excess_list_goes_to_each_department_with_the_parts_of_its_own_sections(
        engine, send, make_dept, make_section, flagged, facts_read, monkeypatch):
    """Routed like the reorder mail. A department named like a section group
    gets the parts of every section of the group, one named like a section
    that section's, any other the whole store's — each mailed to its own
    people. A ticked department with no part above maximum, or with no one to
    mail, gets nothing, and neither does one that is not ticked."""
    group = f"Rule Group {uuid.uuid4().hex[:8]}"
    press, general, own, quiet = make_section(group), make_section(group), make_section(), make_section()
    rows = [
        _over("E-PRESS", section_id=press.id),                              # 20 over: ₹ 1,000
        _over("E-GENERAL", qty=14, rate=100.0, section_id=general.id),      # 4 over: ₹ 400
        _over("E-OWN", qty=12, rate=None, section_id=own.id),               # 2 over, no rate
        _over("E-NOSECTION", qty=11, rate=5.0, section_id=None),            # 1 over: ₹ 5
        _part("H-OK", "Healthy part", "Healthy", 9, 2, 20, section_id=quiet.id),
    ]
    untouched = copy.deepcopy(rows)
    read = {"stock": 0, "sections": 0}
    section_names = rules.SpareDashboardService._section_names

    def list_stock(self, store_code, section_id=None, section_ids=None):
        read["stock"] += 1
        return list(rows)

    def counted_names(store_code):
        read["sections"] += 1
        return section_names(store_code)
    monkeypatch.setattr(SpareStockService, "list_stock", list_stock)
    monkeypatch.setattr(rules.SpareDashboardService, "_section_names", staticmethod(counted_names))

    people = {}

    def dept(tag, name=None):
        people[tag] = [_mail(f"{tag}.{role}") for role in ("to", "cc", "bcc")]
        to, cc, bcc = people[tag]
        return make_dept([("Asha", to, "TO"), ("Ravi", cc, "CC"), ("Hidden", bcc, "BCC")], notify=True, name=name)
    tool = dept("tool", name=group.upper())                 # like the group, whatever its letter case
    maint = dept("maint", name=own.name)
    silent = dept("silent", name=quiet.name)                # its section has nothing above maximum
    nobody = make_dept([("No Address", None, "TO"), ("Left", _mail("left"), "TO", False)], notify=True)
    store = dept("store")
    unticked = make_dept([("Meena", _mail("unticked"), "TO")])
    flagged(tool, maint, silent, nobody, store, unticked)

    described = rules.describe_rule("excess_list")
    # "Store" is sent the whole store: every part above maximum is in some mail, and is counted once.
    assert described["content"] == {"count": 4, "label": "parts above maximum"}
    assert described["recipients"]["to"] == [people[tag][0] for tag in ("tool", "maint", "silent", "store")]
    # Without such a department the count is what the ticked ones are sent — not the part no section holds…
    flagged(tool, maint, silent, unticked)
    assert rules.describe_rule("excess_list")["content"]["count"] == 3
    flagged(maint, unticked)
    assert rules.describe_rule("excess_list")["content"]["count"] == 1
    # …and nothing while no department is ticked, as the preview says.
    flagged(unticked)
    assert rules.describe_rule("excess_list")["content"]["count"] == 0
    flagged(tool, maint, silent, nobody, store, unticked)

    day, before = _day(), _max_id()
    read.update(stock=0, sections=0)
    del facts_read[:]
    result = rules.run_rule(engine, "excess_list", _at(day, 10))
    assert result["status"] == "queued" and result["reason"] is None
    assert EmailMessage.query.filter(EmailMessage.id > before).count() == 3     # no row for the ones left out
    # One read of the stock list, of the section names and of the suppliers serves every department.
    assert read == {"stock": 1, "sections": 1} and facts_read == ["MAINTENANCE"] and rows == untouched

    expected = (
        (tool, "tool", ["E-PRESS", "E-GENERAL"], "2 parts above maximum · ₹ 1,400 tied up"),
        (maint, "maint", ["E-OWN"], "1 parts above maximum · ₹ 0 tied up · 1 without a rate, not valued"),
        (store, "store", ["E-PRESS", "E-GENERAL", "E-NOSECTION", "E-OWN"],
         "4 parts above maximum · ₹ 1,405 tied up · 1 without a rate, not valued"),
    )
    assert len(result["messages"]) == len(expected)
    for message, (department, tag, parts, totals) in zip(result["messages"], expected):
        to, cc, bcc = people[tag]
        assert (message["to"], message["cc"], message["bcc"]) == ([to], [cc], [bcc])
        assert (message["status"], message["trigger"], message["rule_key"]) == ("QUEUED", "MANUAL", "excess_list")
        assert message["subject"] == _excess_subject(department.name, day, len(parts))
        row = _row(message["id"])
        blocks = row.body_text.split("\n\n")
        assert blocks[1] == (f"Please find attached the list of {len(parts)} spare part(s) for {department.name} "
                             "that are above their maximum level.")
        assert blocks[2] == totals and f"<strong>{totals}</strong>" in row.body_html
        _no_list_in(row.body_html, row.body_text, [r["part_no"] for r in rows])
        stored = row.attachments[0]
        assert stored.filename == _excess_file(department.name, day) and stored.filename.isascii()
        _, table = _sheet(stored)
        assert [line[2] for line in table[1:]] == parts                   # its own, dearest first
    # The Section column tells the two sections of a group apart.
    _, tool_table = _sheet(_row(result["messages"][0]["id"]).attachments[0])
    assert [(line[0], line[2], line[9], line[11]) for line in tool_table[1:]] == [
        (press.name, "E-PRESS", 20, 1000), (general.name, "E-GENERAL", 4, 400)]

    engine.process_due()
    assert [(c["to"], c["cc"], c["bcc"]) for c in send.calls] == [
        ([to], [cc], [bcc]) for to, cc, bcc in (people[tag] for tag in ("tool", "maint", "store"))]
    assert [c["attachments"][0]["filename"] for c in send.calls] == [
        _excess_file(department.name, day) for department in (tool, maint, store)]


def test_a_namesake_section_in_another_store_does_not_take_a_departments_parts_away(
        stock, send, make_dept, make_section, flagged):
    """A section's name is unique within its store only. These lists are the
    spare store's, so only its own sections say which parts a department gets."""
    group = f"Rule Group {uuid.uuid4().hex[:8]}"
    own, grouped = make_section(), make_section(group)
    # Made later, so with the higher ids: the same name and the same group in the Child Parts Store.
    make_section(name=own.name, store="CHILD_PART")
    make_section(group, store="CHILD_PART")
    sections, groups = ReorderMailService._scopes()
    assert ReorderMailService._section_ids_for(own.name, sections, groups) == {own.id}
    assert ReorderMailService._section_ids_for(group, sections, groups) == {grouped.id}

    stock([_over("E-OWN", section_id=own.id), _part("R-OWN", "Seal", "Out of Stock", 0, 5, 10, section_id=own.id),
           _part("R-OTHER", "Belt", "Out of Stock", 0, 5, 10, section_id=grouped.id)])
    flagged(make_dept([("Asha", _mail("to"), "TO")], notify=True, name=own.name))
    excess = rules.preview_rule("excess_list")
    assert excess["count"] == 1 and "1 parts above maximum · ₹ 1,000 tied up" in excess["text"].split("\n\n")
    out = ReorderMailService().send_now(force=True)
    assert [(d["status"], d.get("parts")) for d in out["departments"]] == [("sent", 1)]
    _, table = _sheet(send.calls[0]["attachments"][0])
    assert [line[2] for line in table[1:]] == ["R-OWN"]


def test_a_department_whose_excess_list_cannot_be_built_does_not_block_the_others(
        engine, stock, send, make_dept, flagged, monkeypatch):
    stock([_over("E-1")])
    tos = [_mail(tag) for tag in ("first", "second", "third")]
    first, second, third = (make_dept([("Asha", to, "TO")], notify=True) for to in tos)
    flagged(first, second, third)
    rules.save_rule("excess_list", {"enabled": True, "hour": 8})
    day, before = _day(), _max_id()
    build, broken = rules.xlsx_attachment, [second.name]

    def sheet(filename, columns, lines, **colours):
        if any(name in filename for name in broken):
            raise ValueError("a cell openpyxl refuses")
        return build(filename, columns, lines, **colours)
    monkeypatch.setattr(rules, "xlsx_attachment", sheet)

    def mailed():
        made = EmailMessage.query.filter(EmailMessage.id > before).order_by(EmailMessage.id.asc()).all()
        return [(m.to_list, m.status) for m in made]

    failure = f"excess_list: RuntimeError: the excess list of {second.name} could not be built"
    with pytest.raises(RuntimeError, match=failure):
        rules.run_due_rules(engine, _at(day, 8))
    assert mailed() == [([tos[0]], "QUEUED"), ([tos[2]], "QUEUED")]     # the one after it has its list too
    # Not done for the day: every tick says so on the engine's health, and mails no one twice.
    assert rules.rule_is_due("excess_list", _at(day, 8, 1)) is True
    with pytest.raises(RuntimeError, match=failure):
        rules.run_due_rules(engine, _at(day, 8, 1))
    assert len(mailed()) == 2

    # Once its list can be built the department gets it, and the run is through.
    del broken[:]
    ran = rules.run_due_rules(engine, _at(day, 8, 2))
    assert [(r["key"], r["status"]) for r in ran] == [("excess_list", "queued")]
    assert [m["to"] for m in ran[0]["messages"]] == [[tos[1]]]
    assert rules.rule_is_due("excess_list", _at(day, 8, 3)) is False
    engine.process_due()
    assert sorted(c["to"][0] for c in send.calls) == sorted(tos)

    # From the button the others are mailed as well; with no list built at all the answer says why.
    broken.append(second.name)
    manual = rules.run_rule(engine, "excess_list", _at(day, 9))
    assert manual["status"] == "queued" and [m["to"] for m in manual["messages"]] == [[tos[0]], [tos[2]]]
    broken.extend([first.name, third.name])
    assert rules.run_rule(engine, "excess_list", _at(day, 9)) == {
        "status": "skipped", "reason": "list_build_failed", "messages": []}


def test_excess_run_that_mails_no_one_leaves_one_skipped_row_with_the_reason(
        engine, stock, make_dept, make_section, flagged):
    rules.save_rule("excess_list", {"enabled": True, "hour": 8})
    quiet = make_section()
    unticked = make_dept([("Meena", _mail("unticked"), "TO")])
    ticked = make_dept([("Asha", _mail("ticked"), "TO")], notify=True)
    silent = make_dept([("Ravi", _mail("silent"), "TO")], notify=True, name=quiet.name)
    nobody = make_dept([("No Address", None, "TO"), ("Left", _mail("left"), "TO", False)], notify=True)

    def scheduled_run_says(reason, words):
        day, before = _day(), _max_id()
        preview = rules.preview_rule("excess_list", _at(day, 7))
        assert (preview["text"], preview["html"], preview["count"]) == (
            f"No mail would go out right now: {words}.", None, 0)
        ran = rules.run_due_rules(engine, _at(day, 8))
        assert [(r["key"], r["status"], r["reason"]) for r in ran] == [("excess_list", "skipped", reason)]
        made = EmailMessage.query.filter(EmailMessage.id > before).all()
        assert len(made) == 1                                  # one row for the rule, not one per department
        row = made[0]
        assert (row.dedup_key, row.mail_type, row.rule_key, row.status, row.last_error, row.trigger) == (
            f"excess_list:{day.isoformat()}", "excess_list", "excess_list", "SKIPPED", reason, "RULE")
        assert row.to_list == [] and row.attachments == [] and row.body_text == f"Nothing was sent: {words}."
        # The day is taken: later ticks do not build the list again.
        assert rules.rule_is_due("excess_list", _at(day, 8, 1)) is False
        assert rules.run_due_rules(engine, _at(day, 8, 1)) == [] and _max_id() == row.id
        last = rules.describe_rule("excess_list", _at(day, 9))["last_run"]
        assert (last["status"], last["detail"]) == ("SKIPPED", f"Nothing sent — {words}")
        # From the button the same answer comes back, on record, without taking a day.
        manual = rules.run_rule(engine, "excess_list", _at(day, 9))
        assert (manual["status"], manual["reason"], len(manual["messages"])) == ("skipped", reason, 1)
        assert manual["messages"][0]["status"] == "SKIPPED" and _row(manual["messages"][0]["id"]).dedup_key is None

    # Nobody is ticked for the reorder mail — the stock list is not even looked at.
    stock([_over("E-1", section_id=None)])
    flagged(unticked)
    scheduled_run_says("no_department_flagged", "no department is ticked for the reorder mail")
    # The ticked department that has the part has no one to mail; the one with people has no such part.
    flagged(unticked, silent, nobody)
    scheduled_run_says("no_recipient", "no one to mail")
    # No part of the store is above its maximum.
    stock([_part("H-OK", "Healthy part", "Healthy", 9, 2, 20), _part("Z-MIN", "Seal", "Out of Stock", 0, 2, 9)])
    flagged(ticked, silent, nobody)
    scheduled_run_says("nothing_to_report", "nothing to report")


def test_scheduled_excess_run_mails_each_department_once_a_day_even_when_it_broke_half_way(
        engine, stock, send, make_dept, flagged, monkeypatch):
    stock([_over("E-1")])
    first_to, second_to = _mail("first"), _mail("second")
    first = make_dept([("Asha", first_to, "TO")], notify=True)
    second = make_dept([("Ravi", second_to, "TO")], notify=True)
    flagged(first, second)
    rules.save_rule("excess_list", {"enabled": True, "hour": 8})
    day = _day()
    first_key, second_key = (f"excess_list:{dept.id}:{day.isoformat()}" for dept in (first, second))
    submit, locked = EmailEngine.submit, []

    def locked_once(self, **kw):
        if kw["to"] == [second_to] and not locked:
            locked.append(kw["subject"])
            raise OperationalError("INSERT INTO email_messages", {}, Exception("database is locked"))
        return submit(self, **kw)
    monkeypatch.setattr(EmailEngine, "submit", locked_once)
    built, build = [], rules.xlsx_attachment

    def recorded(filename, columns, lines, **colours):
        built.append(filename)
        return build(filename, columns, lines, **colours)
    monkeypatch.setattr(rules, "xlsx_attachment", recorded)

    with pytest.raises(RuntimeError, match="excess_list: OperationalError"):
        rules.run_due_rules(engine, _at(day, 8))
    assert len(locked) == 1
    assert EmailMessage.query.filter_by(dedup_key=first_key).one().status == "QUEUED"
    assert EmailMessage.query.filter_by(dedup_key=second_key).first() is None
    # The run is not through: the rule is still due, and still shows today's time.
    assert rules.rule_is_due("excess_list", _at(day, 8, 1)) is True
    assert rules.describe_rule("excess_list", _at(day, 8, 1))["next_run_at"] == _at(day, 8).isoformat()

    # The next tick takes it up again: the second department only, and no second sheet for the first.
    del built[:]
    ran = rules.run_due_rules(engine, _at(day, 8, 1))
    assert [(r["key"], r["status"], r["reason"]) for r in ran] == [("excess_list", "queued", None)]
    assert [m["to"] for m in ran[0]["messages"]] == [[second_to]] and built == [_excess_file(second.name, day)]
    reserved = EmailMessage.query.filter_by(dedup_key=second_key).one()
    assert (reserved.trigger, reserved.triggered_by, reserved.rule_key) == ("RULE", "scheduler", "excess_list")
    assert EmailMessage.query.filter(EmailMessage.to_addrs.like(f"%{first_to}%")).count() == 1
    # Now it is done for the day.
    assert rules.rule_is_due("excess_list", _at(day, 8, 2)) is False
    before = _max_id()
    assert rules.run_due_rules(engine, _at(day, 8, 2)) == [] and _max_id() == before
    assert rules.describe_rule("excess_list", _at(day, 9))["next_run_at"] == \
        _at(day + timedelta(days=1), 8).isoformat()
    engine.process_due()
    assert [c["to"] for c in send.calls] == [[first_to], [second_to]]          # each department once

    # "Run now" mails every department again, whatever went out earlier today…
    manual = rules.run_rule(engine, "excess_list", _at(day, 9))
    assert manual["status"] == "queued" and [m["to"] for m in manual["messages"]] == [[first_to], [second_to]]
    assert all(m["trigger"] == "MANUAL" and _row(m["id"]).dedup_key is None for m in manual["messages"])
    # …and takes nothing from tomorrow's run.
    tomorrow = day + timedelta(days=1)
    assert rules.rule_is_due("excess_list", _at(tomorrow, 8)) is True
    ran = rules.run_due_rules(engine, _at(tomorrow, 8))
    assert [m["to"] for m in ran[0]["messages"]] == [[first_to], [second_to]]
    assert rules.rule_is_due("excess_list", _at(tomorrow, 8, 1)) is False


def test_run_now_before_the_hour_is_the_departments_excess_mail_of_the_day(engine, stock, make_dept, flagged, send):
    """Pressing "Run now" at 07:30 and the schedule at 08:00 must not mail a
    department twice: the button's mail takes the day when it is still free."""
    stock([_over("E-1")])
    to = _mail("to")
    dept = make_dept([("Asha", to, "TO")], notify=True)
    flagged(dept)
    rules.save_rule("excess_list", {"enabled": True, "hour": 8})
    day = _day()
    early = rules.run_rule(engine, "excess_list", _at(day, 7, 30))
    assert early["status"] == "queued" and [m["to"] for m in early["messages"]] == [[to]]
    assert _row(early["messages"][0]["id"]).dedup_key == f"excess_list:{dept.id}:{day.isoformat()}"
    # The scheduled run finds the department's mail on record and adds none.
    before = _max_id()
    ran = rules.run_due_rules(engine, _at(day, 8))
    assert [(r["key"], r["status"], r["reason"]) for r in ran] == [("excess_list", "skipped", "already_sent_today")]
    assert _max_id() == before and rules.rule_is_due("excess_list", _at(day, 8, 1)) is False
    # The button still mails again when it is pressed again — with no key of its own.
    again = rules.run_rule(engine, "excess_list", _at(day, 9))
    assert [m["to"] for m in again["messages"]] == [[to]] and _row(again["messages"][0]["id"]).dedup_key is None
    # Tomorrow's scheduled run is untouched.
    assert rules.rule_is_due("excess_list", _at(day + timedelta(days=1), 8)) is True


def test_two_ticks_racing_for_the_excess_run_mail_a_department_once(engine, stock, make_dept, flagged, monkeypatch):
    stock([_over("E-1")])
    to = _mail("to")
    dept = make_dept([("Asha", to, "TO")], notify=True)
    flagged(dept)
    rules.save_rule("excess_list", {"enabled": True, "hour": 8})
    day, before = _day(), _max_id()
    submit = EmailEngine.submit

    def beaten_to_it(self, **kw):
        # The other tick gets the department's mail in between this tick's look and its own submit.
        submit(self, **kw)
        return submit(self, **kw)
    monkeypatch.setattr(EmailEngine, "submit", beaten_to_it)
    ran = rules.run_due_rules(engine, _at(day, 8))
    assert ran == [{"key": "excess_list", "status": "skipped", "reason": "already_sent_today", "messages": []}]
    # The key is the reservation: one mail, and no "nothing was sent" row beside it.
    made = EmailMessage.query.filter(EmailMessage.id > before).all()
    assert [(m.dedup_key, m.status, m.to_list) for m in made] == [
        (f"excess_list:{dept.id}:{day.isoformat()}", "QUEUED", [to])]
    assert rules.rule_is_due("excess_list", _at(day, 8, 1)) is False


def test_scheduled_excess_run_that_cannot_read_the_stock_is_tried_again(
        engine, stock, make_dept, flagged, monkeypatch):
    to = _mail("to")
    dept = make_dept([("Asha", to, "TO")], notify=True)
    flagged(dept)
    rules.save_rule("excess_list", {"enabled": True, "hour": 8})
    day, before = _day(), _max_id()

    def broken(self, store_code, section_id=None, section_ids=None):
        raise RuntimeError("stock list is down")
    with monkeypatch.context() as patch:
        patch.setattr(SpareStockService, "list_stock", broken)
        assert rules.describe_rule("excess_list")["content"]["count"] is None      # not known, never 0
        with pytest.raises(RuntimeError, match="excess_list: RuntimeError: stock list is down"):
            rules.run_due_rules(engine, _at(day, 8))
    # Nothing was recorded or reserved, so the day stays free for the next tick.
    assert _max_id() == before and rules.rule_is_due("excess_list", _at(day, 8, 1)) is True
    stock([_over("E-1")])
    ran = rules.run_due_rules(engine, _at(day, 8, 1))
    assert [(r["key"], r["status"]) for r in ran] == [("excess_list", "queued")]
    assert EmailMessage.query.filter_by(dedup_key=f"excess_list:{dept.id}:{day.isoformat()}").one().to_list == [to]


def test_excess_preview_shows_the_first_of_several_mails_and_says_so(stock, make_dept, flagged, monkeypatch):
    stock([_over("E-1"), _over("E-2", qty=12, rate=None)])
    to, cc, bcc, other = _mail("to"), _mail("cc"), _mail("bcc"), _mail("other")
    first = make_dept([("Asha", to, "TO"), ("Ravi", cc, "CC"), ("Hidden", bcc, "BCC")], notify=True)
    nobody = make_dept([("No Address", None, "TO")], notify=True)
    second = make_dept([("Meena", other, "TO")], notify=True)
    unticked = make_dept([("Kiran", _mail("unticked"), "TO")])
    flagged(first, nobody, second, unticked)
    monkeypatch.setattr(rules, "xlsx_attachment", lambda *args: pytest.fail("a preview built an Excel sheet"))
    monkeypatch.setattr(rules, "daily_stock_report", lambda *args: pytest.fail("a preview built the stock report"))
    day, before = _day(), _max_id()
    preview = rules.preview_rule("excess_list", _at(day, 10))
    assert set(preview) == {"subject", "text", "html", "to", "cc", "bcc", "count"}
    assert preview["subject"] == _excess_subject(first.name, day, 2) and preview["count"] == 2
    assert (preview["to"], preview["cc"], preview["bcc"]) == ([to], [cc], [bcc])       # that department's own
    said = f"[One mail per department — 2 would go out now. This is the one for {first.name}.]"
    totals = "2 parts above maximum · ₹ 1,000 tied up · 1 without a rate, not valued"
    assert preview["text"] == "\n\n".join([
        said,
        f"Dear Sir,\n\nPlease find attached the list of 2 spare part(s) for {first.name} that are above their "
        "maximum level.",
        totals,
        "Please stop ordering these parts and review their maximum level.",
        _SIGNATURE,
    ])
    # The screen shows the HTML body: it says so too, above the mail.
    assert f">{said}</p>" in preview["html"]
    assert preview["html"].index(said) < preview["html"].index("Dear Sir,") < preview["html"].index(totals)
    _no_list_in(preview["html"], preview["text"], ["E-1", "E-2"])
    assert _max_id() == before

    # One mail alone needs no such line.
    flagged(nobody, second)
    alone = rules.preview_rule("excess_list", _at(day, 10))
    assert alone["subject"] == _excess_subject(second.name, day, 2) and alone["to"] == [other]
    assert alone["text"].startswith("Dear Sir,") and "One mail per department" not in alone["html"]
    # And with no mail to show, the preview says why.
    flagged(nobody)
    assert rules.preview_rule("excess_list")["text"] == "No mail would go out right now: no one to mail."
    flagged(unticked)
    assert rules.preview_rule("excess_list") == {
        "subject": "", "text": "No mail would go out right now: no department is ticked for the reorder mail.",
        "html": None, "to": [], "cc": [], "bcc": [], "count": 0}


def test_rule_with_content_but_no_one_to_mail_is_skipped_by_the_engine(engine, stock):
    stock([_over("C-1")])
    result = rules.run_rule(engine, "stock_report", _at(_day(), 9))
    assert (result["status"], result["reason"]) == ("skipped", "no_recipient")
    # No sheet is built for a mail that goes to no one.
    assert result["messages"][0]["status"] == "SKIPPED" and result["messages"][0]["attachments"] == []


# ---- the current stock report ----

def test_stock_report_is_one_line_of_totals_with_the_daily_stock_report_attached(
        engine, stock, send, gate_part, facts_read):
    supplied = gate_part("Kiran Bearings")
    rows = stock([
        _part("S-ZERO", "Seal", "Out of Stock", 0, 2, 9, unit_price=40.0, stock_value=0.0),
        _part("S-LOW", "Belt", "Below Minimum", 1, 5, 20, unit_price=100.0, stock_value=100.0),
        _part("S-ORD", "=SUM(A1:A9)", "Reorder", 3, 2, 9, unit_price=10.0, stock_value=30.0),
        _part("S-OVER", "Gear", "Above Maximum", 30, 2, 10, section_id=None, unit_price=50.0, stock_value=1500.0),
        _part("S-OK", "Bearing 6204", "Healthy", 9, 2, 20, part_id=supplied, unit_price=25.0, stock_value=225.0,
              uom="NOS", sub_category="Bearing", specification="6204 ZZ"),
        _part("S-NOCNT", "Never counted", "Not Counted", None, 3),
    ])
    untouched = copy.deepcopy(rows)
    rules.save_rule("stock_report", {"emails": [{"email": "plant@example.com", "role": "TO"}]})
    day, store = _day(), rules._store_name()

    # The preview is the same mail — without the sheet, and so without the supplier read.
    preview = rules.preview_rule("stock_report", _at(day, 10))
    assert preview["count"] == 6 and preview["to"] == ["plant@example.com"] and facts_read == []

    result = rules.run_rule(engine, "stock_report", _at(day, 10))
    assert result["status"] == "queued" and len(result["messages"]) == 1
    message = result["messages"][0]
    assert (message["mail_type"], message["rule_key"]) == ("stock_report", "stock_report")
    assert message["subject"] == f"Daily Stock Report — {store} [{day:%d.%m.%Y}]" == preview["subject"]
    row = _row(message["id"])
    html, text = row.body_html, row.body_text
    assert (html, text) == (preview["html"], preview["text"])
    assert text == "\n\n".join([
        f"Dear Sir,\n\nPlease find attached the daily stock report of {store}.",
        "6 parts · 5 counted · stock value ₹ 1,855",
        _SIGNATURE,                                           # no closing lines of its own by default
    ])
    assert "<strong>6 parts · 5 counted · stock value ₹ 1,855</strong>" in html
    # The totals by status and by section are the workbook's Summary now, not tables in the mail.
    _no_list_in(html, text, [r["part_no"] for r in rows])
    assert html.endswith(f"{_SIGNATURE_HTML}</p></body></html>") and html.count(_SIGNATURE_HTML) == 1
    assert rows == untouched and facts_read == ["MAINTENANCE"]

    # The attachment is the Daily Stock Report — the workbook the Reports page
    # downloads (its own tests are in test_stock_mis_report.py).
    stored = row.attachments[0]
    assert (stored.filename, stored.mime) == (f"Daily Stock Report {day:%d.%m.%Y}.xlsx", _XLSX)
    assert message["attachments"] == [{"filename": stored.filename, "size": len(stored.content)}]
    book = load_workbook(io.BytesIO(stored.content))
    assert book.sheetnames == ["Summary", "Stock Detail", "Critical & Reorder", "Excess", "Category Summary"]
    detail = [list(line) for line in book["Stock Detail"].iter_rows(values_only=True)]
    head = next(i for i, line in enumerate(detail) if "Part No" in line)
    col = {name: detail[head].index(name) for name in ("Part No", "Description", "Supplier", "Status")}
    by_no = {line[col["Part No"]]: line for line in detail[head + 1:] if line[col["Part No"]]}
    # Every part of the store, counted or not.
    assert sorted(by_no) == ["S-LOW", "S-NOCNT", "S-OK", "S-ORD", "S-OVER", "S-ZERO"]
    assert by_no["S-OK"][col["Supplier"]] == "Kiran Bearings" and by_no["S-OVER"][col["Supplier"]] is None
    assert by_no["S-NOCNT"][col["Status"]] == "Not Counted"
    # Text that looks like a formula is stored as text (Excel's quote prefix), never as a formula.
    formula_like = next(c for line in book["Stock Detail"].iter_rows() for c in line if c.value == "=SUM(A1:A9)")
    assert formula_like.data_type == "s" and formula_like.quotePrefix

    engine.process_due()
    sent = send.calls[-1]
    assert sent["subject"] == message["subject"] and sent["html"] == html
    assert sent["attachments"][0]["filename"] == stored.filename and sent["attachments"][0]["content"] == stored.content


def test_stock_report_of_an_empty_store_is_skipped_and_its_closing_lines_can_be_set(engine, stock):
    rules.save_rule("stock_report", {"emails": [{"email": "plant@example.com", "role": "TO"}]})
    stock([])
    empty = rules.preview_rule("stock_report")
    assert (empty["count"], empty["html"]) == (0, None)
    assert empty["text"] == "No mail would go out right now: nothing to report."
    skipped = rules.run_rule(engine, "stock_report", _at(_day(), 9))
    assert (skipped["status"], skipped["reason"]) == ("skipped", "nothing_to_report")

    stock([_part("S-OK", "Bearing", "Healthy", 9, 2, 20, unit_price=10.0, stock_value=90.0)])
    template = rules.template("stock_report")
    assert template["footer"] == "" and template["is_default"] is True
    totals = "1 parts · 1 counted · stock value ₹ 90"
    assert rules.preview_rule("stock_report")["text"].endswith(f"\n\n{totals}\n\n{_SIGNATURE}")
    saved = rules.save_template("stock_report", {"footer": "Stock as counted by the store."})
    assert saved["footer"] == "Stock as counted by the store." and saved["is_default"] is False
    assert rules.preview_rule("stock_report")["text"].endswith(
        f"\n\n{totals}\n\nStock as counted by the store.\n\n{_SIGNATURE}")
    assert rules.save_template("stock_report", {"footer": ""})["is_default"] is True


# ---- the action tracker digest ----

def test_tracker_digest_puts_open_lines_first(engine, tracker):
    day = _day()
    tracker(
        _item("no_rate", "Parts without a rate", status="CLEAR", count=0, owner="Store + Accounts"),
        _item("erp_gap", "ERP vs physical gap", status="UNKNOWN", count=None, owner="Store + ERP team"),
        _item("critical", "Critical parts", count=495, value=1234567.4, value_label="to restock to max",
              target=(day - timedelta(days=7)).isoformat(), owner="Purchase <urgent>"),
        _item("excess", "Excess above max", count=12, value=0.0, target=(day + timedelta(days=3)).isoformat(),
              action="Stop ordering · review max level", owner="Store + Planning"),
    )
    rules.save_rule("tracker_digest", {"emails": [{"email": "plant@example.com", "role": "TO"}]})
    result = rules.run_rule(engine, "tracker_digest", _at(day, 10))
    assert result["status"] == "queued"
    message = result["messages"][0]
    assert message["subject"] == f"Action Tracker — 2 open point(s) [{day:%d.%m.%Y}]"
    row = _row(message["id"])
    for body in (row.body_html, row.body_text):
        order = [body.index(label) for label in ("Critical parts", "Excess above max", "ERP vs physical gap",
                                                 "Parts without a rate")]
        assert order == sorted(order)                        # OPEN, then unknown, then CLEAR
        assert "2 open · 1 overdue · 1 clear" in body
        assert "₹ 12,34,567 to restock to max" in body and "Stop ordering · review max level" in body
        assert f"{day - timedelta(days=7):%d.%m.%Y}" in body and f"{day + timedelta(days=3):%d.%m.%Y}" in body
    for header in ("Problem", "Count", "₹ at stake", "Action needed", "Action by", "Target date", "Status"):
        assert f">{header}</th>" in row.body_html and header in row.body_text
    assert ">495</td>" in row.body_html and ">OVERDUE</td>" in row.body_html and ">CLEAR</td>" in row.body_html
    assert "Purchase &lt;urgent&gt;" in row.body_html and "<urgent>" not in row.body_html
    assert "Critical parts | 495 | ₹ 12,34,567 to restock to max | Raise PO | Purchase <urgent> | " in row.body_text


def test_tracker_digest_waits_for_an_unreadable_tracker_and_is_skipped_with_nothing_open(engine, tracker, monkeypatch):
    rules.save_rule("tracker_digest", {"enabled": True, "hour": 8,
                                       "emails": [{"email": "plant@example.com", "role": "TO"}]})
    day = _day()
    reserved = EmailMessage.query.filter_by(dedup_key=f"tracker_digest:{day.isoformat()}")

    def down(self, store_code, **kw):
        raise RuntimeError("ERP mirror is down")
    with monkeypatch.context() as patch:
        patch.setattr(SpareManagementService, "management", down)
        assert rules.describe_rule("tracker_digest")["content"]["count"] is None      # not known, never 0
        preview = rules.preview_rule("tracker_digest")
        assert preview["count"] is None and "could not be read" in preview["text"] and preview["html"] is None
        # A scheduled run does not give the day up over a failure that passes:
        # it is raised (so it shows on the engine's health) and nothing is reserved.
        with pytest.raises(RuntimeError, match="tracker_digest: RuntimeError: the Action Tracker could not be read"):
            rules.run_due_rules(engine, _at(day, 8))
        assert reserved.first() is None and rules.rule_is_due("tracker_digest", _at(day, 8, 1)) is True
        # From the button the reason goes on record instead, and takes no day.
        manual = rules.run_rule(engine, "tracker_digest", _at(day, 8, 1))
        assert (manual["status"], manual["reason"]) == ("skipped", "tracker_unavailable")
        assert _row(manual["messages"][0]["id"]).dedup_key is None
    # Readable again a tick later: the day's mail goes out after all.
    tracker(_item("critical", "Critical parts"))
    ran = rules.run_due_rules(engine, _at(day, 8, 2))
    assert [(r["key"], r["status"]) for r in ran] == [("tracker_digest", "queued")]
    assert reserved.one().status == "QUEUED"

    tracker(_item("excess", "Excess above max", status="CLEAR", count=0))
    manual = rules.run_rule(engine, "tracker_digest", _at(day, 9))
    assert (manual["status"], manual["reason"]) == ("skipped", "nothing_to_report")


# ---- the two rules that hand over to the old senders ----

def test_reorder_rule_that_sends_nothing_leaves_one_skipped_row_for_the_day(engine, monkeypatch):
    calls = []

    def send_now(self, force=True):
        calls.append(force)
        return {"status": "ok", "total_sent": 0, "departments": [
            {"department": "Maintenance", "status": "skipped", "reason": "nothing_to_reorder"},
            {"department": "Store", "status": "skipped", "reason": "no_recipient"}]}
    monkeypatch.setattr(ReorderMailService, "send_now", send_now)
    rules.save_rule("reorder_daily", {"enabled": True, "hour": 6})
    day = _day()
    ran = rules.run_due_rules(engine, _at(day, 6))
    assert calls == [False]                                   # the sender keeps its own once-a-day guard
    assert (ran[0]["key"], ran[0]["status"], ran[0]["reason"]) == ("reorder_daily", "skipped", "nothing_to_reorder")
    row = EmailMessage.query.filter_by(dedup_key=f"reorder_daily:{day.isoformat()}").one()
    assert (row.mail_type, row.status, row.last_error, row.trigger, row.rule_key) == (
        "reorder", "SKIPPED", "nothing_to_reorder", "SCHEDULED", "reorder_daily")
    assert rules.run_due_rules(engine, _at(day, 6, 1)) == [] and calls == [False]      # not asked again every tick
    assert rules.describe_rule("reorder_daily", _at(day, 7))["last_run"]["detail"] == \
        "Nothing sent — no part to reorder"


def test_reorder_rule_sends_through_its_own_service_once_a_day(engine, send, make_dept, reorder_for):
    to, cc = _mail("to"), _mail("cc")
    dept = make_dept([("Asha", to, "TO"), ("Ravi", cc, "CC")], notify=True)
    reorder_for(dept)
    rules.save_rule("reorder_daily", {"enabled": True, "hour": 6, "minute": 0})
    described = rules.describe_rule("reorder_daily")
    assert described["recipients"] == {"to": [to], "cc": [cc], "bcc": []}
    assert described["content"] == {"count": 2, "label": "parts to reorder"}

    day = _day()
    ran = rules.run_due_rules(engine, _at(day, 6))
    assert [r["key"] for r in ran] == ["reorder_daily"] and ran[0]["status"] == "sent" and ran[0]["reason"] is None
    message = next(m for m in ran[0]["messages"] if m["to"] == [to])
    assert (message["mail_type"], message["status"], message["trigger"]) == ("reorder", "SENT", "SCHEDULED")
    assert len([c for c in send.calls if c["to"] == [to]]) == 1
    # The sender reserves under its own key (per department), so the rule's own
    # run of the day is marked done once the sender is through.
    assert EmailMessage.query.filter_by(dedup_key=f"reorder_daily:{day.isoformat()}").first() is None
    assert rules.rule_is_due("reorder_daily", _at(day, 7)) is False
    assert rules.run_due_rules(engine, _at(day, 7)) == []
    assert len([c for c in send.calls if c["to"] == [to]]) == 1
    assert rules.rule_is_due("reorder_daily", _at(day + timedelta(days=1), 7)) is True

    # The button ignores the schedule and the once-a-day guard, like the sender's own "send now".
    forced = rules.run_rule(engine, "reorder_daily")
    assert forced["status"] == "sent" and forced["messages"][-1]["trigger"] == "MANUAL"
    assert len([c for c in send.calls if c["to"] == [to]]) == 2


def test_reorder_mail_is_one_line_of_totals_and_an_excel_of_every_part_with_supplier_and_order_qty(
        send, make_dept, reorder_for, gate_part, facts_read, monkeypatch):
    supplied, never_at_the_gate = gate_part("Kiran Bearings <Pune> & Sons"), gate_part()
    rows = [
        # A maximum is planned: order up to it — 9, at ₹ 12.50 each.
        {"part_id": supplied, "part_no": "RO-ZERO", "part_name": "Bearing <6204>", "status": "Out of Stock",
         "qty_on_hand": 0, "min_stock_level": 2, "reorder_level": 3, "max_stock_level": 9,
         "location": "R1-S1-B1", "section_id": 7, "sub_category": "Bearing", "uom": "NOS", "unit_price": 12.5},
        # No maximum: order up to the minimum — 2.5. No rate, so no order value.
        {"part_id": never_at_the_gate, "part_no": "RO-NOMAX", "part_name": "V-belt", "status": "Below Minimum",
         "qty_on_hand": 1.5, "min_stock_level": 4, "reorder_level": None, "max_stock_level": None,
         "location": None, "section_id": None, "unit_price": None},
        # No maximum and not under the minimum: no quantity can be worked out.
        {"part_no": "RO-ATMIN", "part_name": "Coupling", "status": "Reorder", "qty_on_hand": 5,
         "min_stock_level": 2, "reorder_level": 6, "max_stock_level": None, "location": "R2-S1-B4",
         "section_id": 7, "unit_price": 80.0},
    ]
    untouched = copy.deepcopy(rows)
    to, cc, bcc = _mail("to"), _mail("cc"), _mail("bcc")
    dept = make_dept([("Asha", to, "TO"), ("Ravi", cc, "CC"), ("Hidden", bcc, "BCC")], notify=True)
    reorder_for(dept, rows)
    monkeypatch.setattr(rules.SpareDashboardService, "_section_names",
                        staticmethod(lambda store_code: {7: "Maintenance"}))
    today = date.today()

    out = ReorderMailService().send_now(force=True)
    # The result is what it always was.
    assert out == {"status": "ok", "total_sent": 1, "departments": [
        {"department": dept.name, "status": "sent", "parts": 3, "to": 1, "cc": 1, "bcc": 1}]}
    assert len(send.calls) == 1 and facts_read == ["MAINTENANCE"]
    sent = send.calls[0]
    assert (sent["to"], sent["cc"], sent["bcc"]) == ([to], [cc], [bcc])
    assert sent["subject"] == _reorder_subject(dept.name, today, 3)
    html, text = sent["html"], sent["body"]
    totals = "3 parts to reorder · 1 zero stock · 1 below minimum · 1 at reorder level"
    assert f"<strong>{totals}</strong>" in html
    # The template's closing ends the mail, once — no second signature under it.
    assert text == "\n\n".join([
        f"Dear Sir,\n\nPlease find attached the list of 3 spare part(s) for {dept.name} that are at or below the "
        "reorder level and need a purchase order.",
        totals,
        _SIGNATURE,
    ])
    assert html.endswith(f"{_SIGNATURE_HTML}</p></body></html>") and html.count(_SIGNATURE_HTML) == 1
    _no_list_in(html, text, [r["part_no"] for r in rows])
    # Nor a description or a supplier: those are in the sheet.
    for said in ("Bearing", "Kiran", "V-belt", "Coupling"):
        assert said not in html and said not in text
    assert "http" not in html and "<style" not in html and "<link" not in html
    assert rows == untouched

    assert [(a["filename"], a["mime"]) for a in sent["attachments"]] == [
        (f"Reorder List - {dept.name} {today:%d.%m.%Y}.xlsx", _XLSX)]
    sheet, table = _sheet(sent["attachments"][0])
    # Worst first, as they were given. The supplier: the latest gate entry's vendor, kept as it was typed,
    # and empty for a part that never came through the gate.
    assert table == [
        ["Section", "Category", "Part No", "Description", "Supplier", "Location", "UOM", "In Stock", "Min",
         "Reorder Level", "Max", "Order Qty", "Rate", "Order Value", "Status"],
        ["Maintenance", "Bearing", "RO-ZERO", "Bearing <6204>", "Kiran Bearings <Pune> & Sons", "R1-S1-B1", "NOS",
         0, 2, 3, 9, 9, 12.5, 112.5, "Out of Stock"],
        [None, None, "RO-NOMAX", "V-belt", None, None, None, 1.5, 4, None, None, 2.5, None, None, "Below Minimum"],
        ["Maintenance", None, "RO-ATMIN", "Coupling", None, "R2-S1-B4", None, 5, 2, 6, None, None, 80, None,
         "Reorder"],
    ]
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:O4" and sheet["A1"].font.bold

    # On record through the engine, html and sheet included.
    row = EmailMessage.query.filter(EmailMessage.dedup_key.like(f"reorder\\_{dept.id}:%", escape="\\")).one()
    assert (row.mail_type, row.status, row.trigger, row.body_html, row.body_text) == (
        "reorder", "SENT", "MANUAL", html, text)
    assert [a.filename for a in row.attachments] == [sent["attachments"][0]["filename"]]
    # The once-a-day guard and its log row are untouched: an unforced run finds today's mail.
    assert MailSendLog.query.filter_by(report_key=f"reorder_{dept.id}", report_date=today).count() == 1
    again = ReorderMailService().send_now(force=False)
    assert again == {"status": "ok", "total_sent": 0, "departments": [
        {"department": dept.name, "status": "skipped", "reason": "already_sent_today"}]}
    assert len(send.calls) == 1 and facts_read == ["MAINTENANCE"]      # and nothing was read or built for it


def test_a_long_reorder_list_is_one_line_in_the_mail_and_every_part_in_the_sheet(send, make_dept, reorder_for):
    rows = [{"part_no": f"RO-{n:03d}", "part_name": f"Part {n}", "status": "Out of Stock" if n < 40 else "Reorder",
             "qty_on_hand": 0 if n < 40 else 3, "min_stock_level": 2, "reorder_level": 3, "max_stock_level": 10,
             "location": "R1-S1-B1"} for n in range(125)]
    to = _mail("to")
    dept = make_dept([("Asha", to, "TO")], notify=True)
    reorder_for(dept, rows)
    out = ReorderMailService().send_now(force=True)
    assert out["departments"] == [{"department": dept.name, "status": "sent", "parts": 125, "to": 1, "cc": 0,
                                   "bcc": 0}]
    sent = send.calls[-1]
    assert sent["subject"] == _reorder_subject(dept.name, date.today(), 125)
    totals = "125 parts to reorder · 40 zero stock · 0 below minimum · 85 at reorder level"
    assert totals in sent["body"].split("\n\n") and f"<strong>{totals}</strong>" in sent["html"]
    _no_list_in(sent["html"], sent["body"], [r["part_no"] for r in rows])
    _, table = _sheet(sent["attachments"][0])
    assert [line[2] for line in table[1:]] == [r["part_no"] for r in rows]      # every part, worst first
    assert {line[11] for line in table[1:41]} == {10} and {line[11] for line in table[41:]} == {7}


def test_a_fractional_quantity_keeps_its_two_decimals_in_the_sheet_and_in_a_mail_table(send, make_dept, reorder_for):
    """Quantities and levels are kept to two decimals. The reorder sheet holds
    all of a large fractional one, and so does a table a mail still carries
    (the manual alert's critical parts) — not six digits of it, and never an
    exponent."""
    rows = [{"part_no": "RO-FRAC", "part_name": "Hydraulic oil", "status": "Below Minimum",
             "qty_on_hand": 12345.25, "min_stock_level": 20000, "reorder_level": 20000.5,
             "max_stock_level": 1234567.5, "location": "R1-S1-B1"}]
    reorder_for(make_dept([("Asha", _mail("to"), "TO")], notify=True), rows)
    ReorderMailService().send_now(force=True)
    _, table = _sheet(send.calls[-1]["attachments"][0])
    assert table[1][7:12] == [12345.25, 20000, 20000.5, 1234567.5, 1222222.25]
    assert [rules._n(value) for value in (12345.25, 20000.5, 1222222.25)] == ["12345.25", "20000.5", "1222222.25"]
    # Whole numbers and short fractions read as they always did.
    assert [rules._n(value) for value in (0, 7.0, 2.5, 0.05, 10885, None)] == ["0", "7", "2.5", "0.05", "10885", "-"]


def test_reorder_sheet_colours_in_stock_and_status_by_how_urgent_the_part_is(monkeypatch):
    """The store's three colours (14 Oct 2026, "3 rang: gaadha red / red /
    orange"): In Stock and Status in solid dark red with nothing on the shelf,
    in red at or below the min — AT the min counts, as on his stock report —
    and in orange for the rest of the list. The colour goes by the stock
    against the min, not by the Status text: a part AT its min is "Reorder" in
    the stock list and red here. Order Qty is bold; nothing else is coloured,
    and the legend says what the colours mean."""
    monkeypatch.setattr(rules.SpareDashboardService, "_section_names",
                        staticmethod(lambda store_code: {7: "Maintenance"}))

    def part(no, status, qty, minimum, reorder, **more):
        return {"part_no": no, "part_name": f"Part {no}", "status": status, "qty_on_hand": qty,
                "min_stock_level": minimum, "reorder_level": reorder, "max_stock_level": 10,
                "location": "R1-S1-B1", "section_id": 7, "uom": "NOS", "unit_price": 5.0, **more}
    rows = [
        part("R-ZERO", "Out of Stock", 0, 2, 4),
        part("R-NEGATIVE", "Out of Stock", -1, 2, 4),           # booked below zero: nothing on the shelf either
        part("R-BELOW", "Below Minimum", 1, 4, 6),
        part("R-ATMIN", "Reorder", 4, 4, 6),
        part("R-ATREORDER", "Reorder", 6, 4, 6),
        # A min stored as 0 is no min. A text the formula guard keeps from Excel's formula bar.
        part("R-NOMIN", "Reorder", 3, 0, 5, part_name="-MODEL FT 650"),
    ]
    untouched = copy.deepcopy(rows)
    attachment = ReorderMailService._attachment("Maintenance", rows, date(2026, 10, 14), {})
    assert (attachment["filename"], attachment["mime"]) == ("Reorder List - Maintenance 14.10.2026.xlsx", _XLSX)
    assert rows == untouched
    sheet, table = _sheet(attachment)
    assert table[0] == list(SHEET_COLUMNS) and [line[2] for line in table[1:]] == [r["part_no"] for r in rows]
    assert [line[7] for line in table[1:]] == [0, -1, 1, 4, 6, 3]
    columns = {heading: number for number, heading in enumerate(table[0], start=1)}

    dark_red, red, orange = ("00DC2626", "00FFFFFF", True), ("00FFE4E6", "009F1239", True), ("00FFEDD5", "009A3412", True)
    for number, colour in enumerate([dark_red, dark_red, red, red, orange, orange], start=2):
        assert _colour_of(sheet.cell(number, columns["In Stock"])) == colour
        assert _colour_of(sheet.cell(number, columns["Status"])) == colour
        assert _colour_of(sheet.cell(number, columns["Order Qty"])) == (None, None, True)
        for heading, column in columns.items():
            if heading not in ("In Stock", "Status", "Order Qty"):
                assert _colour_of(sheet.cell(number, column)) == (None, None, False), (number, heading)
    # Everything else is the sheet the mail always had: header, widths, borders, alignment, guard, filter.
    _plain_but_for(sheet, table, ("In Stock", "Status", "Order Qty"))
    assert sheet["D7"].value == "-MODEL FT 650" and sheet["D7"].quotePrefix
    assert sheet.auto_filter.ref == "A1:O7" and sheet.freeze_panes == "A2"
    assert _legend_of(sheet, len(rows)) == (
        "In Stock / Status: dark red = no stock · red = at or below min · orange = at or below reorder level")


def test_without_colours_an_attachment_is_the_plain_build_export_sheet():
    """Only the reorder and excess sheets are coloured. Asked without `paint`
    and `legend`, xlsx_attachment hands over the sheet build_export makes, as
    it always did: the same cells and looks, no fill or bold in a data cell,
    nothing under the table."""
    columns = ("Part No", "Description", "In Stock", "Status")
    rows = [{"Part No": "P-1", "Description": '=HYPERLINK("http://x")', "In Stock": 0, "Status": "Out of Stock"},
            {"Part No": "P-2", "Description": "Seal", "In Stock": 4.5, "Status": "Reorder"},
            {"Part No": "P-3"}]
    attachment = xlsx_attachment("Plain - 14.10.2026.xlsx", columns, rows)
    assert (attachment["filename"], attachment["mime"]) == ("Plain - 14.10.2026.xlsx", _XLSX)
    sheet = load_workbook(io.BytesIO(attachment["content"])).active
    plain = load_workbook(build_export([(column, column) for column in columns], rows)).active
    assert (sheet.max_row, sheet.max_column) == (plain.max_row, plain.max_column) == (4, 4)
    for number in range(1, 5):
        for column in range(1, 5):
            mine, theirs = sheet.cell(number, column), plain.cell(number, column)
            assert (_layout_of(mine), _colour_of(mine), mine.font.i) == (
                _layout_of(theirs), _colour_of(theirs), theirs.font.i)
            if number > 1:
                assert _colour_of(mine) == (None, None, False)
    assert sheet["B2"].value == '=HYPERLINK("http://x")' and sheet["B2"].data_type == "s" and sheet["B2"].quotePrefix
    for letter in "ABCD":
        assert sheet.column_dimensions[letter].width == plain.column_dimensions[letter].width
    assert (sheet.freeze_panes, sheet.auto_filter.ref) == (plain.freeze_panes, plain.auto_filter.ref) == ("A2", "A1:D4")


def test_the_widest_table_a_mail_still_carries_stays_well_under_the_size_gmail_clips_a_mail_at():
    """Gmail cuts a mail off past about 102 KB, and a parts table says its cell
    style in every cell. The widest table a mail still carries — the manual
    alert's worst TOP_PARTS critical parts, with cells as long as the store's
    real ones — has to stay well under that, and still be a bordered table with
    a tinted bold header in a client that reads inline styles only."""
    rows = [[f"22040001{n:05d}", "BEARING BALL 6204 ZZ C3 SKF MAKE", "Below Minimum", "12", "20", "60", "R3-S2-B11",
             "Press Tool Room"] for n in range(rules.TOP_PARTS)]
    block = {"line": f"{rules.TOP_PARTS} critical parts", "headers": rules.CRITICAL_HEADERS, "rows": rows}
    _, html = rules._layout("Please act on these.", [block], "")
    assert len(html.encode("utf-8")) < 40 * 1024
    assert "<style" not in html and "class=" not in html
    table = html[html.index("<table"):html.index("</table>")]
    # The font and the colour are said once, on the table; every cell keeps its border.
    assert "font-family:Arial,Helvetica,sans-serif" in table[:table.index(">")]
    assert "color:#1f2937" in table[:table.index(">")]
    assert table.count(f'<td style="{rules._CELL}">') == rules.TOP_PARTS * len(rules.CRITICAL_HEADERS) == 400
    assert table.count(f'<th style="{rules._HEAD}">') == len(rules.CRITICAL_HEADERS)
    assert "border:1px solid" in rules._CELL
    for said in ("border:1px solid", "background:#f3f4f6", "font-weight:bold"):
        assert said in rules._HEAD


def test_reorder_run_reads_the_suppliers_once_for_all_its_departments(send, make_dept, facts_read, monkeypatch):
    first = make_dept([("Asha", _mail("first"), "TO")], notify=True)
    second = make_dept([("Ravi", _mail("second"), "TO")], notify=True)
    monkeypatch.setattr(DepartmentRepository, "list_all", lambda self, active_only=True: [first, second])
    monkeypatch.setattr(ReorderMailService, "_reorder_rows", lambda self, section_ids: list(_REORDER_ROWS))
    assert ReorderMailService().send_now(force=True)["total_sent"] == 2
    assert len(send.calls) == 2 and facts_read == ["MAINTENANCE"]
    # A department's name is typed by an admin; in the file name it cannot break out of it.
    assert ReorderMailService._attachment('Tool/Room: "A"\r\nX', [], date(2026, 10, 13), {})["filename"] == \
        "Reorder List - Tool Room A X 13.10.2026.xlsx"


def test_window_time_rule_reports_why_its_service_sent_nothing(engine, monkeypatch):
    asked = []

    def send_current_month(self, force=False, at=None):
        asked.append((force, at))
        return {"report_date": "2031-01-01", "to": [], "cc": [], "status": "skipped", "reason": "no_recipient_email"}
    monkeypatch.setattr(WindowTimeMailService, "send_current_month", send_current_month)
    result = rules.run_rule(engine, "window_time_daily")
    assert asked == [(True, None)]                            # the button: forced, and dated by the clock
    assert (result["status"], result["reason"]) == ("skipped", "no_recipient_email")
    row = _row(result["messages"][0]["id"])
    assert (row.mail_type, row.status, row.last_error, row.trigger, row.rule_key, row.dedup_key) == (
        "window_time", "SKIPPED", "no_recipient_email", "MANUAL", "window_time_daily", None)

    # On schedule: not forced, and dated by the rule's own time on the plant clock.
    rules.save_rule("window_time_daily", {"enabled": True, "hour": 7, "minute": 15})
    day = _day()
    ran = rules.run_due_rules(engine, _at(day, 9, 40))
    assert asked[-1] == (False, datetime.combine(day, time(7, 15)))
    assert (ran[0]["status"], ran[0]["reason"]) == ("skipped", "no_recipient_email")
    # No one to mail is final for the day: on record once, not asked again every tick.
    assert EmailMessage.query.filter_by(dedup_key=f"window_time_daily:{day.isoformat()}").one().status == "SKIPPED"
    assert rules.run_due_rules(engine, _at(day, 9, 41)) == [] and len(asked) == 2


def test_window_time_sheet_caught_up_late_is_the_one_the_rule_time_would_have_mailed(engine, send, window_sheet):
    """A 07:00 rule mails the sheet of the day just closed. Run late — the
    server was off at 07:00 — it is still that sheet: today's belongs to
    tomorrow morning's run, which would otherwise find it already sent."""
    rules.save_rule("window_time_daily", {"enabled": True, "hour": 7, "minute": 0})
    day = _day()

    def run(now):
        before = len(send.calls)
        ran = rules.run_due_rules(engine, now)
        return [(r["key"], r["status"], r["reason"]) for r in ran], [c["subject"] for c in send.calls[before:]]

    def sheet_of(reported):
        return [("window_time_daily", "sent", None)], [f"SUPPLIER WINDOW TIME ON DATED {reported:%d.%m.%Y}"]

    assert run(_at(day, 9)) == sheet_of(day - timedelta(days=1))               # caught up after the 08:30 shift start
    assert run(_at(day, 9, 1)) == ([], [])
    assert run(_at(day + timedelta(days=1), 7)) == sheet_of(day)               # on time the next two mornings
    assert run(_at(day + timedelta(days=2), 7)) == sheet_of(day + timedelta(days=1))
    assert window_sheet == [((day + timedelta(days=n)).year, (day + timedelta(days=n)).month) for n in (-1, 0, 1)]
    assert send.calls[-1]["to"] == ["sheet@example.com"] and send.calls[-1]["attachments"][0]["content"] == b"sheet"
    # A rule set after the shift start mails the day it runs in.
    rules.save_rule("window_time_daily", {"hour": 9, "minute": 30})
    assert run(_at(day + timedelta(days=3), 11)) == sheet_of(day + timedelta(days=3))


def test_window_time_sheet_that_cannot_be_built_is_tried_again_the_same_day(engine, send, window_sheet, monkeypatch):
    rules.save_rule("window_time_daily", {"enabled": True, "hour": 7})
    day = _day()
    reserved = EmailMessage.query.filter_by(dedup_key=f"window_time_daily:{day.isoformat()}")

    def locked(self, year, month):
        raise RuntimeError("database is locked")
    with monkeypatch.context() as patch:
        patch.setattr(SupplierWindowScheduleService, "window_time_workbook", locked)
        with pytest.raises(RuntimeError,
                           match="window_time_daily: RuntimeError: the window-time sheet could not be built"):
            rules.run_due_rules(engine, _at(day, 7))
        assert reserved.first() is None and rules.rule_is_due("window_time_daily", _at(day, 7, 1)) is True
        # From the button the failure goes on record as a run that sent nothing.
        manual = rules.run_rule(engine, "window_time_daily")
        assert (manual["status"], manual["reason"]) == ("skipped", "report_build_failed")
        assert _row(manual["messages"][0]["id"]).dedup_key is None
    assert send.calls == []
    ran = rules.run_due_rules(engine, _at(day, 7, 1))                          # the next tick
    assert [(r["key"], r["status"]) for r in ran] == [("window_time_daily", "sent")] and len(send.calls) == 1
    assert rules.run_due_rules(engine, _at(day, 7, 2)) == []


def test_report_rules_wait_for_a_mail_server_instead_of_giving_the_day_up(
        engine, send, window_sheet, make_dept, reorder_for, monkeypatch):
    to = _mail("to")
    reorder_for(make_dept([("Asha", to, "TO")], notify=True))
    rules.save_rule("reorder_daily", {"enabled": True, "hour": 6})
    rules.save_rule("window_time_daily", {"enabled": True, "hour": 7})
    day = _day()
    before = _max_id()
    with monkeypatch.context() as patch:
        patch.setattr(mailer, "smtp_is_configured", lambda: False)
        waiting = [{"key": key, "status": "skipped", "reason": "smtp_not_configured", "messages": []}
                   for key in ("reorder_daily", "window_time_daily")]
        assert rules.run_due_rules(engine, _at(day, 7)) == waiting
        assert rules.run_due_rules(engine, _at(day, 7, 1)) == waiting
        assert _max_id() == before                             # nothing on record, tick after tick…
        assert rules.rule_is_due("reorder_daily", _at(day, 8)) and rules.rule_is_due("window_time_daily", _at(day, 8))
        # The button still leaves the reason on record.
        manual = rules.run_rule(engine, "reorder_daily")
        assert (manual["status"], manual["reason"], manual["messages"][0]["status"]) == (
            "skipped", "smtp_not_configured", "SKIPPED")
    # …and with the server set up later that day, both mails go.
    ran = rules.run_due_rules(engine, _at(day, 11))
    assert [(r["key"], r["status"]) for r in ran] == [("reorder_daily", "sent"), ("window_time_daily", "sent")]
    assert sorted(c["to"][0] for c in send.calls) == sorted([to, "sheet@example.com"])
    assert rules.run_due_rules(engine, _at(day, 11, 1)) == []


def test_rules_built_here_wait_for_a_mail_server_instead_of_queueing_a_mail_a_day(
        engine, stock, tracker, send, make_dept, flagged, monkeypatch):
    """Queued on schedule with no mail server, a rule would leave one mail per
    day in the queue and all of them would go out together the day the server
    is entered — the same department handed three dated excess lists at once."""
    stock([_over("E-1")])
    tracker(_item("critical", "Critical parts"))
    to = _mail("to")
    flagged(make_dept([("Asha", to, "TO")], notify=True))
    keys = ("excess_list", "stock_report", "tracker_digest")
    every_day = {"enabled": True, "hour": 8, "days": list(range(7))}
    rules.save_rule("excess_list", every_day)
    for key in keys[1:]:
        rules.save_rule(key, {**every_day, "emails": [{"email": "plant@example.com", "role": "TO"}]})
    day, before = _day(), _max_id()
    third = day + timedelta(days=2)
    with monkeypatch.context() as patch:
        patch.setattr(mailer, "smtp_is_configured", lambda: False)
        waiting = [{"key": key, "status": "skipped", "reason": "smtp_not_configured", "messages": []} for key in keys]
        for n in range(3):
            assert rules.run_due_rules(engine, _at(day + timedelta(days=n), 8)) == waiting
            assert rules.run_due_rules(engine, _at(day + timedelta(days=n), 8, 1)) == waiting
        assert _max_id() == before                             # nothing on record, day after day…
        assert all(rules.rule_is_due(key, _at(third, 9)) for key in keys)
        # The button still queues: that mail was asked for, and waits for the server.
        manual = rules.run_rule(engine, "tracker_digest", _at(third, 9))
        assert manual["status"] == "queued" and _max_id() == manual["messages"][0]["id"]
    # …and with the server set up on the third day each rule mails once, for that day.
    ran = rules.run_due_rules(engine, _at(third, 11))
    assert [(r["key"], r["status"]) for r in ran] == [(key, "queued") for key in keys]
    engine.process_due()
    stamp = f"[{third:%d.%m.%Y}]"
    assert sorted(c["subject"].split(" — ")[0] for c in send.calls if stamp in c["subject"]) == [
        "Action Tracker", "Action Tracker", "Daily Stock Report", "Excess Parts"]
    assert len(send.calls) == 4 and [c["to"] for c in send.calls if "Excess Parts" in c["subject"]] == [[to]]
    assert rules.run_due_rules(engine, _at(third, 11, 1)) == []


def test_reorder_run_that_broke_half_way_is_finished_on_the_next_tick(engine, send, make_dept, monkeypatch):
    first_to, second_to = _mail("first"), _mail("second")
    first = make_dept([("Asha", first_to, "TO")], notify=True)
    second = make_dept([("Ravi", second_to, "TO")], notify=True)
    monkeypatch.setattr(DepartmentRepository, "list_all", lambda self, active_only=True: [first, second])
    monkeypatch.setattr(ReorderMailService, "_reorder_rows", lambda self, section_ids: list(_REORDER_ROWS))
    rules.save_rule("reorder_daily", {"enabled": True, "hour": 6})
    day = _day()
    submit, locked = EmailEngine.submit, []

    def locked_once(self, **kw):
        if kw["to"] == [second_to] and not locked:
            locked.append(kw["subject"])
            raise OperationalError("INSERT INTO email_messages", {}, Exception("database is locked"))
        return submit(self, **kw)
    monkeypatch.setattr(EmailEngine, "submit", locked_once)

    with pytest.raises(RuntimeError, match="reorder_daily: OperationalError"):
        rules.run_due_rules(engine, _at(day, 6))
    assert [c["to"] for c in send.calls] == [[first_to]] and len(locked) == 1
    # The first department's mail is a scheduled reorder mail of this plant day
    # (the sender stamps it with the real clock — put it on this case's day).
    # It must not pass for the whole run.
    EmailMessage.query.filter(EmailMessage.to_addrs.like(f"%{first_to}%")).update(
        {"created_at": _at(day, 6)}, synchronize_session=False)
    db.session.commit()
    assert rules.rule_is_due("reorder_daily", _at(day, 6, 1)) is True
    ran = rules.run_due_rules(engine, _at(day, 6, 1))
    assert [(r["key"], r["status"]) for r in ran] == [("reorder_daily", "sent")]
    assert [c["to"] for c in send.calls] == [[first_to], [second_to]]          # each department once
    assert rules.rule_is_due("reorder_daily", _at(day, 6, 2)) is False
    assert rules.run_due_rules(engine, _at(day, 6, 2)) == []


def test_a_department_whose_reorder_list_cannot_be_built_does_not_block_the_others(
        client, engine, send, make_dept, monkeypatch):
    tos = [_mail(tag) for tag in ("first", "second", "third")]
    first, second, third = (make_dept([("Asha", to, "TO")], notify=True) for to in tos)
    monkeypatch.setattr(DepartmentRepository, "list_all", lambda self, active_only=True: [first, second, third])
    monkeypatch.setattr(ReorderMailService, "_reorder_rows", lambda self, section_ids: list(_REORDER_ROWS))
    rules.save_rule("reorder_daily", {"enabled": True, "hour": 6})
    day = _day()
    build, broken = ReorderMailService._attachment, [second.name]

    def sheet(dept_name, rows, on, suppliers):
        if dept_name in broken:
            raise ValueError("a cell openpyxl refuses")
        return build(dept_name, rows, on, suppliers)
    monkeypatch.setattr(ReorderMailService, "_attachment", staticmethod(sheet))

    failure = f"reorder_daily: RuntimeError: the reorder list of {second.name} could not be built"
    with pytest.raises(RuntimeError, match=failure):
        rules.run_due_rules(engine, _at(day, 6))
    assert [c["to"] for c in send.calls] == [[tos[0]], [tos[2]]]            # the one after it has its list too
    # Not done for the day: every tick says so on the engine's health, and mails no one twice.
    assert rules.rule_is_due("reorder_daily", _at(day, 6, 1)) is True
    with pytest.raises(RuntimeError, match=failure):
        rules.run_due_rules(engine, _at(day, 6, 1))
    assert len(send.calls) == 2

    # Once its list can be built the department gets it, and the run is through.
    del broken[:]
    ran = rules.run_due_rules(engine, _at(day, 6, 2))
    assert [(r["key"], r["status"]) for r in ran] == [("reorder_daily", "sent")]
    assert [c["to"] for c in send.calls[2:]] == [[tos[1]]]
    assert rules.rule_is_due("reorder_daily", _at(day, 6, 3)) is False

    # "Send now" never raises: it mails the others and reports the one it could not build.
    broken.append(second.name)
    out = ReorderMailService().send_now(force=True)
    assert (out["status"], out["total_sent"]) == ("ok", 2)
    assert [(d["department"], d["status"], d.get("reason")) for d in out["departments"]] == [
        (first.name, "sent", None), (second.name, "skipped", "list_build_failed"), (third.name, "sent", None)]
    assert [c["to"] for c in send.calls[3:]] == [[tos[0]], [tos[2]]]
    resp = _as(client, _user()).post("/api/v1/departments/reorder-mail/send")
    assert resp.status_code == 200 and resp.get_json()["data"]["total_sent"] == 2
    # With no list built at all, the rule's own "Run now" puts the reason on record.
    broken.extend([first.name, third.name])
    manual = rules.run_rule(engine, "reorder_daily")
    assert (manual["status"], manual["reason"]) == ("skipped", "list_build_failed")
    assert _row(manual["messages"][0]["id"]).body_text == "Nothing was sent: a department's list could not be built."


def test_the_status_screens_show_the_rule_that_now_decides_the_daily_mails(app, monkeypatch):
    monkeypatch.setitem(app.config, "WINDOW_TIME_MAIL_HOUR", 7)
    monkeypatch.setitem(app.config, "REORDER_MAIL_HOUR", 8)
    # Until a rule is saved the .env flags are its default, and what is shown.
    email = SettingsService().system_status()["email"]
    assert (email["daily_mail_enabled"], email["daily_mail_hour"]) == (False, 7)
    rules.save_rule("window_time_daily", {"enabled": True, "hour": 9, "minute": 30})
    rules.save_rule("reorder_daily", {"enabled": True, "hour": 6})
    email = SettingsService().system_status()["email"]
    assert (email["daily_mail_enabled"], email["daily_mail_hour"]) == (True, 9)
    window = WindowTimeMailService().status()
    assert (window["enabled"], window["hour"]) == (True, 9)
    reorder = ReorderMailService().status()
    assert (reorder["enabled"], reorder["hour"]) == (True, 6)


def test_reorder_status_names_the_sections_each_ticked_department_is_sent(make_dept, make_section, monkeypatch):
    """A department named like no section or group is sent the whole store.
    The status says which it is, so a rename that turned a department's own
    list into the whole store's can be seen instead of guessed."""
    group = f"Rule Group {uuid.uuid4().hex[:8]}"
    press, general, own = make_section(group), make_section(group), make_section()
    tool = make_dept([("Asha", _mail("tool"), "TO")], notify=True, name=group.upper())
    maint = make_dept([("Ravi", _mail("maint"), "TO")], notify=True, name=own.name)
    store = make_dept([("Meena", _mail("store"), "TO")], notify=True)
    unticked = make_dept([("Kiran", _mail("unticked"), "TO")], name=press.name)
    monkeypatch.setattr(DepartmentRepository, "list_all",
                        lambda self, active_only=True: [tool, maint, store, unticked])
    monkeypatch.setattr(ReorderMailService, "_reorder_rows", lambda self, section_ids: [])

    def covered():
        return {d["department"]: d["sections"] for d in ReorderMailService().status()["departments"]}
    assert covered() == {tool.name: sorted([press.name, general.name]), maint.name: [own.name], store.name: None}
    for entry in ReorderMailService().status()["departments"]:
        assert set(entry) == {"department", "to", "cc", "bcc", "reorder_parts", "sections"}

    # The section is renamed: its department no longer matches it, and is sent everything.
    own.name = f"Renamed {uuid.uuid4().hex[:8]}"
    db.session.commit()
    assert covered()[maint.name] is None


def test_run_now_does_not_go_past_the_kill_switch(engine, stock, make_dept, flagged):
    stock([_over("C-1")])
    rules.save_rule("stock_report", {"emails": [{"email": "plant@example.com", "role": "TO"}]})
    flagged(make_dept([("Asha", _mail("to"), "TO")], notify=True))
    engine.save_settings({"enabled": False})
    before = _max_id()
    for key in ("stock_report", "excess_list"):
        assert rules.run_rule(engine, key) == {"status": "skipped", "reason": "mail_engine_disabled", "messages": []}
    assert rules.run_reminders(engine) == {"reminded": 0, "escalated": 0, "messages": []}
    assert _max_id() == before


# ---- preview ----

def test_preview_builds_the_mail_but_stores_and_sends_nothing(
        engine, stock, tracker, send, make_dept, reorder_for, monkeypatch):
    stock([_over("X-1"), _part("C-1", "Seal", "Out of Stock", 0, 2)])
    tracker(_item("critical", "Critical parts"))
    to = _mail("to")
    dept = make_dept([("Asha", to, "TO")], notify=True)
    reorder_for(dept)
    for key in ("stock_report", "tracker_digest"):
        rules.save_rule(key, {"emails": [{"email": "plant@example.com", "role": "TO"}]})
    # A preview shows the mail, not its attachment: no Excel sheet is built for it.
    monkeypatch.setattr(rules, "xlsx_attachment", lambda *args: pytest.fail("a preview built an Excel sheet"))
    monkeypatch.setattr(ReorderMailService, "_attachment",
                        lambda *args: pytest.fail("a preview built an Excel sheet"))
    before, states = _max_id(), NotificationState.query.count()
    day = _day()
    previews = {spec["key"]: rules.preview_rule(spec["key"], _at(day, 10)) for spec in rules.RULES}
    assert _max_id() == before and send.calls == [] and NotificationState.query.count() == states
    for preview in previews.values():
        assert set(preview) == {"subject", "text", "html", "to", "cc", "bcc", "count"}

    excess = previews["excess_list"]                          # the department's own mail, like the reorder list
    assert excess["count"] == 1 and excess["to"] == [to]
    assert excess["subject"] == _excess_subject(dept.name, day, 1)
    assert "1 parts above maximum · ₹ 1,000 tied up" in excess["text"].split("\n\n")
    _no_list_in(excess["html"], excess["text"], ["X-1", "C-1"])
    report = previews["stock_report"]
    assert report["count"] == 2 and report["to"] == ["plant@example.com"]
    assert report["subject"] == f"Daily Stock Report — {rules._store_name()} [{day:%d.%m.%Y}]"
    assert "2 parts · 2 counted · stock value ₹ 1,500" in report["html"]
    _no_list_in(report["html"], report["text"], ["X-1", "C-1"])
    assert previews["tracker_digest"]["count"] == 1 and "Critical parts" in previews["tracker_digest"]["html"]

    reorder = previews["reorder_daily"]
    assert (reorder["subject"], reorder["text"]) == (_reorder_subject(dept.name, day), _reorder_text(dept.name))
    assert reorder["to"] == [to] and reorder["count"] == 2
    _no_list_in(reorder["html"], reorder["text"], [r["part_no"] for r in _REORDER_ROWS])

    window = previews["window_time_daily"]
    assert window["subject"].startswith("SUPPLIER WINDOW TIME ON DATED ") and window["count"] is None
    assert window["text"] == f"Dear Sir,\n\nPlease find attached supplier window time sheet.\n\n{_SIGNATURE}"

    # With nothing to send the preview says so instead of showing an empty mail.
    stock([])
    for key in ("excess_list", "stock_report"):
        empty = rules.preview_rule(key)
        assert empty["count"] == 0 and empty["html"] is None and empty["subject"] == ""
        assert empty["text"] == "No mail would go out right now: nothing to report."


def test_reorder_preview_shows_the_first_of_several_mails_and_says_so(make_dept, monkeypatch):
    first = make_dept([("Asha", _mail("a"), "TO")], notify=True)
    nobody = make_dept([("No Address", None, "TO")], notify=True)
    second = make_dept([("Ravi", _mail("b"), "CC")], notify=True)
    unticked = make_dept([("Meena", _mail("c"), "TO")])
    monkeypatch.setattr(DepartmentRepository, "list_all",
                        lambda self, active_only=True: [first, nobody, second, unticked])
    monkeypatch.setattr(ReorderMailService, "_reorder_rows", lambda self, section_ids: list(_REORDER_ROWS))
    day = _day()
    preview = rules.preview_rule("reorder_daily", _at(day, 10))
    assert preview["subject"] == _reorder_subject(first.name, day) and preview["count"] == 2
    said = f"[One mail per department — 2 would go out now. This is the one for {first.name}.]"
    assert preview["text"] == f"{said}\n\n{_reorder_text(first.name)}"
    # The screen shows the HTML body: it says so too, above the mail.
    assert f">{said}</p>" in preview["html"]
    assert preview["html"].index(said) < preview["html"].index("Dear Sir,") < preview["html"].index(
        "2 parts to reorder")
    # No part to reorder anywhere: no mail, and the reason.
    monkeypatch.setattr(ReorderMailService, "_reorder_rows", lambda self, section_ids: [])
    assert rules.preview_rule("reorder_daily")["text"] == \
        "No mail would go out right now: no part to reorder."
    monkeypatch.setattr(DepartmentRepository, "list_all", lambda self, active_only=True: [unticked])
    assert rules.preview_rule("reorder_daily")["text"] == \
        "No mail would go out right now: no department is ticked for the reorder mail."


# ---- templates ----

def test_default_templates_word_the_reorder_and_window_time_mails(engine, send, make_dept, reorder_for):
    to = _mail("to")
    dept = make_dept([("Asha", to, "TO")], notify=True)
    reorder_for(dept)
    ReorderMailService().send_now(force=True)
    sent = next(c for c in send.calls if c["to"] == [to])
    # The subject is what the reorder mail always said; the opening points to the attached list.
    assert sent["subject"] == _reorder_subject(dept.name, date.today())
    assert sent["body"] == _reorder_text(dept.name)
    assert sent["html"].startswith("<!DOCTYPE html>") and "<table" not in sent["html"]   # an HTML mail, no table
    assert sent["html"].count(_SIGNATURE_HTML) == 1 and len(sent["attachments"]) == 1

    assert WindowTimeMailService._wording(date(2026, 10, 12)) == (
        "SUPPLIER WINDOW TIME ON DATED 12.10.2026",
        f"Dear Sir,\n\nPlease find attached supplier window time sheet.\n\n{_SIGNATURE}")

    listed = rules.list_templates()
    assert set(listed) == {"identity", "templates"}
    assert listed["identity"] == {"signature": _SIGNATURE, "reply_to": None}
    assert [t["mail_type"] for t in listed["templates"]] == [
        "reorder", "window_time", "excess_list", "stock_report", "tracker_digest", "tracker_reminder",
        "tracker_escalation", "short_rejection"]
    for tpl in listed["templates"]:
        assert set(tpl) == {"mail_type", "title", "subject", "intro", "footer", "defaults", "is_default",
                            "placeholders"}
        assert tpl["is_default"] is True and tpl["defaults"] == {k: tpl[k] for k in ("subject", "intro", "footer")}
        assert "{date}" in tpl["placeholders"] and "{store}" in tpl["placeholders"]
    # The two mails that go out one per department can name the department.
    assert [t["mail_type"] for t in listed["templates"] if "{department}" in t["placeholders"]] == [
        "reorder", "excess_list"]
    by_type = {t["mail_type"]: t for t in listed["templates"]}
    assert by_type["reorder"]["defaults"] == {
        "subject": "Spare Reorder Alert — {department} — {count} part(s) [{date}]",
        "intro": "Dear Sir,\n\nPlease find attached the list of {count} spare part(s) for {department} that are "
                 "at or below the reorder level and need a purchase order.",
        "footer": _SIGNATURE}
    assert by_type["excess_list"]["placeholders"] == ["{department}", "{count}", "{date}", "{store}"]
    assert by_type["excess_list"]["defaults"] == {
        "subject": "Excess Parts — {department} — {count} part(s) above maximum [{date}]",
        "intro": "Dear Sir,\n\nPlease find attached the list of {count} spare part(s) for {department} that are "
                 "above their maximum level.",
        "footer": "Please stop ordering these parts and review their maximum level."}
    assert by_type["stock_report"]["defaults"] == {
        "subject": "Daily Stock Report — {store} [{date}]",
        "intro": "Dear Sir,\n\nPlease find attached the daily stock report of {store}.", "footer": ""}


def test_a_saved_subject_intro_and_footer_are_used(engine, send, make_dept, reorder_for):
    to = _mail("to")
    dept = make_dept([("Asha", to, "TO")], notify=True)
    reorder_for(dept)
    saved = rules.save_template("reorder", {
        "subject": "  Reorder for {department}: {count} parts on {date}  ",
        "intro": "Hello,\r\n\r\nPlease order these {count} parts for {store}:\n",
        "footer": "Thanks,\nStore team"}, user="asha")
    assert saved["is_default"] is False
    assert saved["subject"] == "Reorder for {department}: {count} parts on {date}"
    assert saved["intro"] == "Hello,\n\nPlease order these {count} parts for {store}:"
    assert saved["defaults"]["subject"] == "Spare Reorder Alert — {department} — {count} part(s) [{date}]"
    assert AppSetting.query.filter_by(key="mail_templates").one().updated_by == "asha"

    ReorderMailService().send_now(force=True)
    sent = next(c for c in send.calls if c["to"] == [to])
    assert sent["subject"] == f"Reorder for {dept.name}: 2 parts on {date.today():%d.%m.%Y}"
    assert sent["body"].split("\n") == [
        "Hello,", "", f"Please order these 2 parts for {rules._store_name()}:", "",
        "2 parts to reorder · 1 zero stock · 0 below minimum · 1 at reorder level", "", "Thanks,", "Store team"]
    # The saved closing ends the HTML body too; the system note it replaced is not added back.
    assert sent["html"].endswith("Thanks,<br>Store team</p></body></html>") and _NOTE not in sent["html"]

    # A field saved empty — or the built-in text typed back in — goes back to the default.
    back = rules.save_template("reorder", {"subject": "", "footer": _SIGNATURE})
    assert back["subject"] == back["defaults"]["subject"] and back["footer"] == back["defaults"]["footer"]
    assert back["is_default"] is False and back["intro"].startswith("Hello,")       # the intro is still edited
    assert rules.save_template("reorder", {"intro": "   "})["is_default"] is True
    ReorderMailService().send_now(force=True)
    again = [c for c in send.calls if c["to"] == [to]][-1]
    assert (again["subject"], again["body"]) == (_reorder_subject(dept.name, date.today()),
                                                 _reorder_text(dept.name))


def test_placeholders_are_filled_once_and_unknown_ones_stay_as_typed():
    store = rules._store_name()
    rules.save_template("reorder", {
        "subject": "{department} / {count} / {date} / {store} / {nope} / {Date} / {problem} / {}",
        "intro": "{count}{count} for {department}", "footer": "{store}"})
    # A value is never read as a placeholder itself.
    words = rules.mail_wording("reorder", department="{count} & {store}", count=7, date=date(2026, 1, 5))
    assert words["subject"] == (f"{{count}} & {{store}} / 7 / 05.01.2026 / {store} / "
                                "{nope} / {Date} / {problem} / {}")
    assert words["intro"] == "77 for {count} & {store}" and words["footer"] == store
    # {department} belongs to the two mails that go out per department: elsewhere it is left as typed.
    for mail_type in ("excess_list", "stock_report"):
        rules.save_template(mail_type, {"subject": "{department} {count} {date}"})
    assert rules.mail_wording("excess_list", department="Tool Room", count=3, date=date(2026, 12, 31))["subject"] == \
        "Tool Room 3 31.12.2026"
    assert rules.mail_wording("stock_report", count=3, date=date(2026, 12, 31))["subject"] == \
        "{department} 3 31.12.2026"


def test_templates_and_identity_are_validated():
    for bad in ({"subject": "two\nlines"}, {"subject": "x" * 201}, {"intro": "x" * 2001}, {"footer": 5},
                {"subject": None}, {"body": "x"}, {}, None, "text"):
        with pytest.raises(BadRequestError):
            rules.save_template("reorder", bad)
    assert rules.template("reorder")["is_default"] is True
    for unknown in ("manual_alert", "smtp_test", "critical_digest", "nope"):
        with pytest.raises(NotFoundError):
            rules.save_template(unknown, {"subject": "x"})

    for bad in ({"reply_to": "not-an-address"}, {"reply_to": "a@example.com\nBcc: b@example.com"},
                {"signature": 5}, {"signature": "x" * 201}, {"name": "x"}, {}, None):
        with pytest.raises(BadRequestError):
            rules.save_identity(bad)
    assert rules.save_identity({"signature": " Plant Store Team ", "reply_to": " store@example.com "}, "asha") == {
        "signature": "Plant Store Team", "reply_to": "store@example.com"}
    assert rules.save_identity({"reply_to": None}) == {"signature": "Plant Store Team", "reply_to": None}
    assert rules.save_identity({"reply_to": ""})["reply_to"] is None
    assert rules.save_identity({"signature": ""})["signature"] == _SIGNATURE         # empty = the built-in one


def test_the_saved_reply_to_address_goes_on_every_mail_the_engine_sends(engine, send):
    alert = {"subject": "Note", "message": "Text", "emails": [{"email": "guest@example.com"}]}
    rules.send_alert(engine, alert)
    engine.process_due()
    assert "reply_to" not in send.calls[-1]                    # none saved: the mailer is not even told
    rules.save_identity({"reply_to": "store.inbox@example.com"})
    rules.send_alert(engine, alert)
    engine.process_due()
    assert send.calls[-1]["subject"] == "Note" and send.calls[-1]["reply_to"] == "store.inbox@example.com"
    # The mails that are sent inline carry it too.
    SettingsService().send_test_mail("probe@example.com", "asha")
    assert send.calls[-1]["to"] == ["probe@example.com"] and send.calls[-1]["reply_to"] == "store.inbox@example.com"
    # Taken off again, the next mail goes without it.
    rules.save_identity({"reply_to": None})
    rules.send_alert(engine, alert)
    engine.process_due()
    assert len(send.calls) == 4 and "reply_to" not in send.calls[-1]


def test_saved_wording_and_signature_reach_the_html_mail_escaped(engine, stock, make_dept, flagged):
    stock([_over("C-1")])
    rules.save_template("excess_list", {
        "subject": "Excess: {count} <now> for {department}",
        "intro": "Team,\n\n<i>Urgent</i> list of {store} for {department}:\nsecond line",
        "footer": "Reply to <store@example.com> & act"})
    rules.save_identity({"signature": "Store <Team>\nTBDI"})
    # A department's name is typed by an admin: it is data too.
    tag = uuid.uuid4().hex[:6]
    flagged(make_dept([("Asha", _mail("to"), "TO")], notify=True, name=f"Dies & <Tools> {tag}"))
    preview = rules.preview_rule("excess_list")
    assert preview["subject"] == f"Excess: 1 <now> for Dies & <Tools> {tag}"      # a header is not HTML
    html, text = preview["html"], preview["text"]
    assert "&lt;i&gt;Urgent&lt;/i&gt; list of" in html and "<i>" not in html
    assert f"for Dies &amp; &lt;Tools&gt; {tag}:<br>second line</p>" in html and "<Tools>" not in html
    assert "Reply to &lt;store@example.com&gt; &amp; act" in html
    assert "Store &lt;Team&gt;<br>TBDI</p></body></html>" in html         # the signature closes the mail
    totals = "1 parts above maximum · ₹ 1,000 tied up"
    assert html.index("Team,") < html.index(totals) < html.index("Reply to") < html.index("Store &lt;Team&gt;")
    assert text.startswith("Team,\n\n<i>Urgent</i> list of") and text.endswith("Store <Team>\nTBDI")
    assert text.index(totals) < text.index("Reply to <store@example.com> & act")


# ---- the manual alert ----

def test_send_alert_queues_one_mail_with_the_message_escaped(engine, stock, make_dept):
    stock([_part("C-9", "Gear <box>", "Below Minimum", 1, 4, 9)])
    to, cc = _mail("to"), _mail("cc")
    dept = make_dept([("Asha", to, "TO"), ("Ravi", cc, "CC")])
    hostile = '<script>alert("x")</script>\r\n\r\nLine & <b>bold</b>\nnext line'
    message = rules.send_alert(engine, {
        "subject": "  Stock check today  ", "message": hostile, "department_ids": [dept.id],
        "emails": [{"email": "guest@example.com", "role": "BCC"}], "include": "critical"})
    assert message["status"] == "QUEUED" and message["mail_type"] == "manual_alert"
    assert message["subject"] == "Stock check today" and message["trigger"] == "MANUAL"
    assert (message["to"], message["cc"], message["bcc"]) == ([to], [cc], ["guest@example.com"])
    assert message["rule_key"] is None and message["store_code"] == "MAINTENANCE"
    row = _row(message["id"])
    html, text = row.body_html, row.body_text
    assert "<script" not in html and "<b>" not in html
    assert "&lt;script&gt;alert(&quot;x&quot;)&lt;/script&gt;" in html
    assert "Line &amp; &lt;b&gt;bold&lt;/b&gt;<br>next line" in html
    assert text.startswith('<script>alert("x")</script>\n\nLine & <b>bold</b>\nnext line')
    # The list asked for sits under the message, the signature under both.
    assert "1 critical parts · 0 zero stock · 1 below minimum" in html and "Gear &lt;box&gt;" in html
    assert html.index("next line") < html.index("<table") < html.index(_SIGNATURE_HTML)
    assert "C-9 | Gear <box> | Below Minimum | 1 | 4 | 9 | R1-S1-B1 | Maintenance" in text
    assert text.endswith(_SIGNATURE)

    plain = rules.send_alert(engine, {"subject": "Note", "message": "Only this.",
                                      "emails": [{"email": "guest@example.com"}]})
    assert _row(plain["id"]).body_text == f"Only this.\n\n{_SIGNATURE}"
    assert "<table" not in _row(plain["id"]).body_html

    stock([])
    none_critical = rules.send_alert(engine, {"subject": "Note", "message": "See below.", "include": "critical",
                                              "emails": [{"email": "guest@example.com"}]})
    assert "No part is out of stock or below its minimum right now." in _row(none_critical["id"]).body_text


def test_alert_critical_list_puts_the_worst_parts_first_under_the_totals(engine, stock):
    rows = stock([
        _part("B-GAP1", "Belt", "Below Minimum", 4, 5, 20),
        _part("Z-MIN2", "Seal <b>kit</b> & <img src=x onerror=alert(1)>", "Out of Stock", 0, 2, None, None, None),
        _part("H-OK", "Healthy part", "Healthy", 9, 2, 20),
        _part("B-GAP8", "Chain", "Below Minimum", 2, 10, 30, "R2-S1-B4"),
        _part("Z-MIN9", "Bearing", "Out of Stock", 0, 9, 12.5),
        _part("N-CNT", "Not counted", "Not Counted", None, 3),
        _part("R-ORD", "At reorder", "Reorder", 3, 2, 9),
    ])
    untouched = copy.deepcopy(rows)
    alert = {"subject": "Critical parts", "message": "Please act on these.", "include": "critical",
             "emails": [{"email": "plant@example.com"}]}
    row = _row(rules.send_alert(engine, alert)["id"])
    html, text = row.body_html, row.body_text
    totals = "4 critical parts · 2 zero stock · 2 below minimum"
    assert totals in html and totals in text
    for body in (html, text):
        # No stock at all first (largest minimum first), then the largest gap under the minimum.
        order = [body.index(no) for no in ("Z-MIN9", "Z-MIN2", "B-GAP8", "B-GAP1")]
        assert order == sorted(order)
        for absent in ("H-OK", "N-CNT", "R-ORD"):
            assert absent not in body
    assert html.endswith(f"{_SIGNATURE_HTML}</p></body></html>")    # the shared signature closes the mail
    assert "<th" in html and html.count("<tr") == 5           # the header row and the four parts
    for header in ("Part No", "Description", "Status", "Stock", "Min", "Max", "Location", "Section"):
        assert f">{header}</th>" in html and header in text
    assert ">12.5</td>" in html and ">R2-S1-B4</td>" in html and ">Maintenance</td>" in html
    # A description is data: it is shown, never run.
    assert "<img" not in html and "<b>kit</b>" not in html
    assert "Seal &lt;b&gt;kit&lt;/b&gt; &amp; &lt;img src=x onerror=alert(1)&gt;" in html
    assert "Seal <b>kit</b> & <img src=x onerror=alert(1)>" in text
    assert "http" not in html and "<style" not in html and "<link" not in html     # nothing external
    assert text.endswith(_SIGNATURE) and row.attachments == []
    assert rows == untouched                                   # the stock list's own rows were only read

    # Fifty parts in the table, all of them in the totals.
    stock([_part(f"Q-{n:03d}", f"Part {n}", "Out of Stock", 0, 100 - n) for n in range(60)])
    row = _row(rules.send_alert(engine, alert)["id"])
    assert "60 critical parts · 60 zero stock · 0 below minimum" in row.body_text
    assert row.body_html.count("<tr") == 51
    assert "Q-049" in row.body_html and "Q-050" not in row.body_html
    assert "Only the 50 worst parts are listed. The Store Stock screen has all 60." in row.body_text


def test_send_alert_can_carry_the_tracker_table_and_says_when_it_cannot(engine, tracker, monkeypatch):
    tracker(_item("excess", "Excess above max", count=12))
    base = {"subject": "Tracker", "message": "Please look.", "emails": [{"email": "guest@example.com"}],
            "include": "tracker"}
    row = _row(rules.send_alert(engine, base)["id"])
    assert "Excess above max" in row.body_html and "1 open · 0 overdue · 0 clear" in row.body_text

    def down(self, store_code, **kw):
        raise RuntimeError("ERP mirror is down")
    monkeypatch.setattr(SpareManagementService, "management", down)
    before = _max_id()
    with pytest.raises(AppException) as failed:
        rules.send_alert(engine, base)
    assert failed.value.status_code == 503 and "could not be read" in failed.value.message
    assert _max_id() == before


def test_send_alert_is_validated(engine, make_dept):
    nobody = make_dept([("No Address", None, "TO")])
    ok = {"subject": "Note", "message": "Text", "emails": [{"email": "guest@example.com", "role": "TO"}]}
    before = _max_id()
    for bad in ({**ok, "subject": ""}, {**ok, "subject": "   "}, {**ok, "subject": "x" * 201},
                {**ok, "subject": "two\nlines"}, {**ok, "subject": None}, {**ok, "message": ""},
                {**ok, "message": "x" * 5001}, {**ok, "message": 5}, {**ok, "include": "everything"},
                {**ok, "emails": [{"email": "broken", "role": "TO"}]}, {**ok, "department_ids": [987654321]},
                {**ok, "extra": 1}, {"subject": "Note", "message": "Text"},
                {"subject": "Note", "message": "Text", "department_ids": [nobody.id]}, None, []):
        with pytest.raises(BadRequestError):
            rules.send_alert(engine, bad)
    assert _max_id() == before
    with pytest.raises(BadRequestError, match="no one to send to"):
        rules.send_alert(engine, {"subject": "Note", "message": "Text", "department_ids": [nobody.id]})


# ---- reminders and escalation ----

def test_reminder_settings_defaults_and_validation(tracker):
    tracker(_item("excess", "Excess above max", count=12, owner="Store + Planning"))
    view = rules.reminders_view()
    assert {k: v for k, v in view.items() if k != "lines"} == {
        "enabled": False, "hour": 9, "minute": 0, "remind_before_days": 2, "repeat_every_days": 1,
        "escalate_after_days": 3, "escalate_to": [], "owners": {}}
    assert set(view["lines"][0]) == {"key", "label", "owner", "status", "target_date", "overdue", "days_to_target",
                                     "to", "level", "last_notified_at", "next_action", "note"}
    assert view["lines"][0]["owner"] == "Store + Planning" and view["lines"][0]["note"] == "No target date set"

    for bad in ({"repeat_every_days": 0}, {"remind_before_days": -1}, {"escalate_after_days": 366},
                {"hour": 24}, {"minute": "5"}, {"enabled": 1}, {"nope": True}, {}, None,
                {"escalate_to": [{"email": "broken"}]}, {"escalate_to": "boss@example.com"},
                {"owners": {"not_a_tracker_line": {"department_ids": [], "emails": []}}},
                {"owners": {"excess": {"emails": [{"email": "broken"}]}}},
                {"owners": {"excess": {"department_ids": [987654321]}}}, {"owners": ["excess"]}):
        with pytest.raises(BadRequestError):
            rules.save_reminders(bad)
    assert rules.reminder_settings()["owners"] == {}

    saved = rules.save_reminders({"enabled": True, "hour": 7, "repeat_every_days": 2,
                                  "escalate_to": [{"email": "boss@example.com", "role": "to"}],
                                  "owners": {"excess": {"emails": [{"email": "owner@example.com"}]}}}, "asha")
    assert (saved["enabled"], saved["hour"], saved["repeat_every_days"]) == (True, 7, 2)
    assert saved["escalate_to"] == [{"email": "boss@example.com", "role": "TO"}]
    assert saved["owners"] == {
        "excess": {"department_ids": [], "emails": [{"email": "owner@example.com", "role": "TO"}]}}
    assert _line(saved, "excess")["to"] == ["owner@example.com"]
    # Another line's owners are added without undoing the first; an emptied line is cleared.
    merged = rules.save_reminders({"owners": {"critical": {"emails": [{"email": "po@example.com"}]}}})
    assert set(merged["owners"]) == {"excess", "critical"}
    cleared = rules.save_reminders({"owners": {"excess": {"department_ids": [], "emails": []}}})
    assert set(cleared["owners"]) == {"critical"} and cleared["hour"] == 7


def test_a_line_is_reminded_inside_the_window_and_at_most_once_per_interval(engine, tracker, make_dept):
    key, day = _line_key(), _day()
    owner, copied = _mail("owner"), _mail("copied")
    dept = make_dept([("Copied", copied, "CC")])
    AppSetting.set_value("reminder_rules", {
        "enabled": True, "remind_before_days": 2, "repeat_every_days": 3, "escalate_after_days": 30,
        "escalate_to": [{"email": "boss@example.com", "role": "TO"}],
        "owners": {key: {"department_ids": [dept.id], "emails": [{"email": owner, "role": "TO"}]}}})
    target = day + timedelta(days=5)
    tracker(_item(key, "Excess above max <stores>", count=12, value=5000.0, target=target.isoformat(),
                  owner="Store + Planning", remarks="Waiting for the planning review"))

    line = _line(rules.reminders_view(_at(day, 9)), key)        # five days to go: outside the window
    assert (line["next_action"], line["days_to_target"], line["overdue"]) == ("none", 5, False)
    assert line["note"] == "Due in 5 days — reminders start 2 days before the target date"
    assert line["to"] == [owner, copied] and line["level"] == 0 and line["last_notified_at"] is None
    assert rules.run_reminders(engine, _at(day, 9)) == {"reminded": 0, "escalated": 0, "messages": []}

    start = day + timedelta(days=3)                              # two days to go: the window opens
    line = _line(rules.reminders_view(_at(start, 9)), key)
    assert (line["next_action"], line["note"]) == ("remind", "Due in 2 days")
    result = rules.run_reminders(engine, _at(start, 9))
    assert (result["reminded"], result["escalated"]) == (1, 0)
    message = result["messages"][0]
    assert (message["mail_type"], message["status"], message["trigger"]) == ("tracker_reminder", "QUEUED", "MANUAL")
    assert message["to"] == [owner] and message["cc"] == [copied]
    assert message["subject"] == f"Reminder — Excess above max <stores> [{start:%d.%m.%Y}]"
    row = _row(message["id"])
    assert "Excess above max &lt;stores&gt;" in row.body_html and "<stores>" not in row.body_html
    for body in (row.body_html, row.body_text):
        assert "Due in 2 days" in body and f"{target:%d.%m.%Y}" in body and "₹ 5,000" in body
        assert "Remarks: Waiting for the planning review" in body and "Store + Planning" in body
    state = _state(key)
    assert (state.level, state.last_count) == (1, 12) and state.last_notified_at == _at(start, 9)

    # Mailed already: nothing more that day, nor the next two (repeat every 3 days).
    for later in (start, start + timedelta(days=1), start + timedelta(days=2)):
        line = _line(rules.reminders_view(_at(later, 10)), key)
        assert line["next_action"] == "none" and (line["level"], line["last_notified_at"]) == (
            1, _at(start, 9).isoformat())
        assert f"mailed on {start:%d.%m.%Y}, next mail on {start + timedelta(days=3):%d.%m.%Y}" in line["note"]
        assert rules.run_reminders(engine, _at(later, 10))["reminded"] == 0
    third = start + timedelta(days=3)                            # the target date has passed by a day
    line = _line(rules.reminders_view(_at(third, 9)), key)
    assert (line["next_action"], line["note"], line["overdue"]) == ("remind", "Overdue by 1 day", True)
    assert rules.run_reminders(engine, _at(third, 9))["reminded"] == 1
    assert _state(key).last_notified_at == _at(third, 9)


def test_an_overdue_line_is_escalated_with_the_owners_in_cc(engine, tracker):
    key, day = _line_key(), _day()
    _remind(key, owners=("owner@example.com", "second@example.com"), escalate=("boss@example.com",),
            remind_before_days=2, repeat_every_days=1, escalate_after_days=3)
    target = day - timedelta(days=2)
    tracker(_item(key, "Critical parts", count=40, target=target.isoformat()))
    line = _line(rules.reminders_view(_at(day, 9)), key)        # overdue by 2: still a reminder
    assert (line["next_action"], line["note"]) == ("remind", "Overdue by 2 days")
    assert rules.run_reminders(engine, _at(day, 9))["reminded"] == 1 and _state(key).level == 1

    next_day = day + timedelta(days=1)                           # overdue by 3: escalation
    line = _line(rules.reminders_view(_at(next_day, 9)), key)
    assert (line["next_action"], line["note"], line["level"]) == ("escalate", "Overdue by 3 days", 1)
    result = rules.run_reminders(engine, _at(next_day, 9))
    assert (result["reminded"], result["escalated"]) == (0, 1)
    message = result["messages"][0]
    assert message["mail_type"] == "tracker_escalation"
    assert message["to"] == ["boss@example.com"] and message["cc"] == ["owner@example.com", "second@example.com"]
    assert message["subject"] == f"Escalation — Critical parts is overdue [{next_day:%d.%m.%Y}]"
    assert "Overdue by 3 days" in _row(message["id"]).body_text and ">OVERDUE</td>" in _row(message["id"]).body_html
    assert _state(key).level == 2
    view = _line(rules.reminders_view(_at(next_day, 12)), key)
    assert (view["level"], view["next_action"]) == (2, "none")
    # Escalated again a day later — once a day at most.
    assert rules.run_reminders(engine, _at(next_day, 18))["escalated"] == 0
    assert rules.run_reminders(engine, _at(next_day + timedelta(days=1), 9))["escalated"] == 1


def test_without_an_escalation_address_the_owners_go_on_being_reminded(engine, tracker):
    key, day = _line_key(), _day()
    _remind(key, escalate=(), escalate_after_days=0)
    tracker(_item(key, "RGP overdue", target=(day - timedelta(days=9)).isoformat()))
    line = _line(rules.reminders_view(_at(day, 9)), key)
    assert line["next_action"] == "remind"
    assert line["note"] == "Overdue by 9 days — no escalation address is set, so the owners are reminded"
    result = rules.run_reminders(engine, _at(day, 9))
    assert (result["reminded"], result["escalated"]) == (1, 0)
    assert result["messages"][0]["to"] == ["owner@example.com"]


def test_lines_with_no_target_date_or_no_one_to_mail_do_nothing_and_say_why(engine, tracker):
    undated, unowned, unknown, day = _line_key(), _line_key(), _line_key(), _day()
    _remind(undated, escalate=(), escalate_after_days=3)
    overdue = (day - timedelta(days=1)).isoformat()
    tracker(_item(undated, "No date", target=None), _item(unowned, "No owner", target=overdue),
            _item(unknown, "Not readable", status="UNKNOWN", count=None, target=overdue))
    view = rules.reminders_view(_at(day, 9))
    assert (_line(view, undated)["next_action"], _line(view, undated)["note"]) == ("none", "No target date set")
    assert _line(view, undated)["days_to_target"] is None and _line(view, undated)["overdue"] is False
    assert (_line(view, unowned)["next_action"], _line(view, unowned)["note"]) == ("none", "No one to mail")
    assert _line(view, unowned)["to"] == [] and _line(view, unowned)["overdue"] is True
    assert (_line(view, unknown)["next_action"], _line(view, unknown)["note"]) == (
        "none", "Could not be read right now")
    before = _max_id()
    assert rules.run_reminders(engine, _at(day, 9)) == {"reminded": 0, "escalated": 0, "messages": []}
    assert _max_id() == before
    assert _state(undated).first_seen_at == _at(day, 9) and _state(undated).level == 0    # seen open, never mailed
    assert _state(unknown) is None

    # With only an escalation address, an overdue line is escalated even with no owner.
    _remind(undated, escalate=("boss@example.com",), escalate_after_days=1)
    result = rules.run_reminders(engine, _at(day, 10))
    assert result["escalated"] == 1 and result["messages"][0]["to"] == ["boss@example.com"]
    assert result["messages"][0]["cc"] == [] and _state(unowned).level == 2


def test_a_line_that_goes_clear_loses_its_state(engine, tracker):
    key, day = _line_key(), _day()
    _remind(key)
    target = (day + timedelta(days=1)).isoformat()
    tracker(_item(key, "Returns awaiting inspection", count=4, target=target))
    assert rules.run_reminders(engine, _at(day, 9))["reminded"] == 1 and _state(key).level == 1

    tracker(_item(key, "Returns awaiting inspection", status="CLEAR", count=0, target=target))
    line = _line(rules.reminders_view(_at(day, 12)), key)
    assert (line["level"], line["last_notified_at"], line["next_action"]) == (0, None, "none")
    assert line["note"] == "Clear — nothing to chase" and line["overdue"] is False
    assert rules.run_reminders(engine, _at(day, 12)) == {"reminded": 0, "escalated": 0, "messages": []}
    assert _state(key) is None
    # Open again the next day: the chase starts over.
    tracker(_item(key, "Returns awaiting inspection", count=2, target=target))
    again = rules.run_reminders(engine, _at(day + timedelta(days=1), 9))
    assert again["reminded"] == 1 and (_state(key).level, _state(key).last_count) == (1, 2)


def test_a_target_date_moved_out_starts_the_chase_over(engine, tracker):
    key, day = _line_key(), _day()
    _remind(key, escalate_after_days=1)
    tracker(_item(key, "ERP gap", target=(day - timedelta(days=5)).isoformat()))
    assert rules.run_reminders(engine, _at(day, 9))["escalated"] == 1 and _state(key).level == 2
    tracker(_item(key, "ERP gap", target=(day + timedelta(days=20)).isoformat()))
    assert rules.run_reminders(engine, _at(day + timedelta(days=1), 9))["messages"] == []
    assert _state(key).level == 0


def test_scheduled_reminder_run_happens_once_a_day_at_its_time(engine, tracker, monkeypatch):
    key, day = _line_key(), _day()
    _remind(key, hour=9, minute=30, repeat_every_days=1)
    tracker(_item(key, "Kiosk requests open", target=(day - timedelta(days=1)).isoformat()))
    assert rules.reminders_are_due(_at(day, 9, 29)) is False
    assert rules.run_due_reminders(engine, _at(day, 9, 29)) is None
    first = rules.run_due_reminders(engine, _at(day, 9, 30))
    assert first["reminded"] == 1 and first["messages"][0]["trigger"] == "RULE"
    assert first["messages"][0]["triggered_by"] == "scheduler"
    # Checked today: later ticks do not read the tracker again.
    asked = []
    with monkeypatch.context() as patch:
        patch.setattr(SpareManagementService, "management", lambda self, store_code, **kw: asked.append(1))
        assert rules.run_due_reminders(engine, _at(day, 16)) is None and asked == []
    tomorrow = day + timedelta(days=1)
    assert rules.reminders_are_due(_at(tomorrow, 9)) is False
    assert rules.run_due_reminders(engine, _at(tomorrow, 23))["reminded"] == 1     # caught up late in the day

    # Switched off, nothing is due whatever the time.
    AppSetting.set_value("reminder_rules", {"enabled": False, "hour": 0})
    assert rules.reminders_are_due(_at(tomorrow + timedelta(days=1), 12)) is False


def test_reminder_run_with_an_unreadable_tracker_fails_loudly_and_is_tried_again(engine, tracker, monkeypatch):
    key, day = _line_key(), _day()
    _remind(key, hour=9)

    def down(self, store_code, **kw):
        raise RuntimeError("ERP mirror is down")
    monkeypatch.setattr(SpareManagementService, "management", down)
    assert rules.reminders_view(_at(day, 9))["lines"] == []                # the settings still load
    with pytest.raises(AppException) as failed:
        rules.run_due_reminders(engine, _at(day, 9))
    assert failed.value.status_code == 503
    assert rules.reminders_are_due(_at(day, 9, 1)) is True                 # the day was not marked as done
    tracker(_item(key, "Critical parts", target=day.isoformat()))
    assert rules.run_due_reminders(engine, _at(day, 9, 1))["reminded"] == 1


# ---- the worker ----

def test_tick_runs_a_due_rule_and_sends_its_mail(engine, stock, send, make_dept, flagged):
    stock([_over("C-1")])
    address = _mail("tick")
    dept = make_dept([("Asha", address, "TO")], notify=True)
    flagged(dept)
    rules.save_rule("excess_list", {"enabled": True, "hour": 0, "minute": 0, "days": list(range(7))})
    summary = engine.tick()
    assert summary["errors"] == [] and summary["sent"] >= 1
    sent = next(c for c in send.calls if c["to"] == [address])
    assert "1 parts above maximum · ₹ 1,000 tied up" in sent["body"].split("\n\n")
    assert [a["filename"][:-15] for a in sent["attachments"]] == [f"Excess Parts - {dept.name} "]
    _, table = _sheet(sent["attachments"][0])
    assert [line[2] for line in table[1:]] == ["C-1"]
    row = EmailMessage.query.filter(EmailMessage.to_addrs.like(f"%{address}%")).one()
    assert (row.status, row.trigger, row.triggered_by) == ("SENT", "RULE", "scheduler")
    # The next tick finds the day taken.
    engine.tick()
    assert len([c for c in send.calls if c["to"] == [address]]) == 1
    status = engine.status()["rules"]
    assert status["enabled"] == 1 and status["total"] == 6 and status["next"]["key"] == "excess_list"


# ---- the routes ----

_ROUTES = (
    ("get", "/rules"), ("put", "/rules/excess_list"), ("post", "/rules/excess_list/run"),
    ("get", "/rules/excess_list/preview"), ("get", "/templates"), ("put", "/templates/reorder"),
    ("put", "/templates-identity"), ("post", "/send-alert"), ("get", "/reminders"), ("put", "/reminders"),
    ("post", "/reminders/run"),
)


def test_every_rule_route_is_admin_only(client):
    c = _as(client, _user(role="STORE_USER"))
    for method, path in _ROUTES:
        resp = c.get(_BASE + path) if method == "get" else getattr(c, method)(_BASE + path, json={})
        assert resp.status_code == 403, (method, path, resp.status_code)


def test_machine_key_can_read_but_cannot_save_run_or_send(app, stock, tracker):
    stock([_over("C-1")])
    tracker(_item("excess", "Excess above max"))
    machine = app.test_client()                                # no session cookie
    key = {"X-API-Key": app.config["API_KEY"]}
    before = _max_id()
    for method, path in _ROUTES:
        if method == "get":
            assert machine.get(_BASE + path, headers=key).status_code == 200, path
            assert machine.get(_BASE + path).status_code == 401, path          # and with no credential at all
        else:
            body = {"enabled": True, "subject": "x", "message": "y", "emails": [{"email": "a@example.com"}]}
            assert getattr(machine, method)(_BASE + path, json=body, headers=key).status_code == 403, (method, path)
    assert rules.rule_settings("excess_list")[1] is True and rules.template("reorder")["is_default"] is True
    assert rules.reminder_settings()["enabled"] is False and _max_id() == before


def test_rule_routes_save_validate_run_and_preview(client, stock, tracker, make_dept, flagged):
    stock([_over("C-1")])
    tracker(_item("excess", "Excess above max"))
    admin = _user()
    c = _as(client, admin)
    listed = c.get(_BASE + "/rules").get_json()
    assert listed["success"] is True and len(listed["data"]) == 6

    for bad in ({"hour": 24}, {"days": [0, 7]}, {"days": "mon"}, {"minute": -1},
                {"emails": [{"email": "not-an-address", "role": "TO"}]}, {"unknown": 1}):
        resp = c.put(_BASE + "/rules/stock_report", json=bad)
        assert resp.status_code == 400 and resp.get_json()["success"] is False, bad
    assert c.put(_BASE + "/rules/stock_report", data="not json", content_type="text/plain").status_code == 400
    # A rule whose audience is kept elsewhere cannot be re-addressed here.
    for fixed in ("window_time_daily", "reorder_daily", "excess_list"):
        assert c.put(_BASE + f"/rules/{fixed}", json={"emails": []}).status_code == 400
        assert c.put(_BASE + f"/rules/{fixed}", json={"department_ids": []}).status_code == 400
    assert c.put(_BASE + "/rules/nope", json={"enabled": True}).status_code == 404
    assert c.post(_BASE + "/rules/nope/run").status_code == 404
    assert c.get(_BASE + "/rules/nope/preview").status_code == 404

    address = _mail("route")
    resp = c.put(_BASE + "/rules/stock_report", json={"enabled": True, "hour": 6, "minute": 45, "days": [2, 0],
                                                      "emails": [{"email": address, "role": "TO"}]})
    assert resp.status_code == 200
    rule = resp.get_json()["data"]
    assert (rule["enabled"], rule["hour"], rule["minute"], rule["days"]) == (True, 6, 45, [0, 2])
    assert rule["recipients"]["to"] == [address] and rule["is_default"] is False and rule["next_run_at"]
    assert AppSetting.query.filter_by(key="notification_rules").one().updated_by == admin.username
    status = c.get(_BASE + "/status").get_json()["data"]["rules"]
    assert status["enabled"] == 1 and status["next"]["key"] == "stock_report"

    preview = c.get(_BASE + "/rules/stock_report/preview").get_json()["data"]
    assert preview["to"] == [address] and preview["count"] == 1 and "1 parts · 1 counted" in preview["html"]
    ran = c.post(_BASE + "/rules/stock_report/run").get_json()["data"]
    assert ran["status"] == "queued" and ran["reason"] is None
    row = _row(ran["messages"][0]["id"])
    assert (row.status, row.trigger, row.triggered_by, row.to_list) == ("QUEUED", "MANUAL", admin.username, [address])
    assert [a["filename"][-5:] for a in ran["messages"][0]["attachments"]] == [".xlsx"]

    # The excess list: its schedule is saved here, its mails go one to each ticked department.
    first_to, second_to = _mail("first"), _mail("second")
    first = make_dept([("Asha", first_to, "TO")], notify=True)
    flagged(first, make_dept([("Ravi", second_to, "TO")], notify=True))
    rule = c.put(_BASE + "/rules/excess_list", json={"hour": 7, "days": [1]}).get_json()["data"]
    assert (rule["audience"], rule["hour"], rule["days"]) == ("departments_flagged", 7, [1])
    assert rule["recipients"]["to"] == [first_to, second_to] and rule["content"]["count"] == 1
    preview = c.get(_BASE + "/rules/excess_list/preview").get_json()["data"]
    assert preview["to"] == [first_to] and preview["count"] == 1
    assert "<strong>1 parts above maximum · ₹ 1,000 tied up</strong>" in preview["html"]
    assert preview["text"].startswith(
        f"[One mail per department — 2 would go out now. This is the one for {first.name}.]\n\n")
    resp = c.post(_BASE + "/rules/excess_list/run").get_json()
    assert resp["message"] == "Mail queued" and resp["data"]["status"] == "queued"
    assert [m["to"] for m in resp["data"]["messages"]] == [[first_to], [second_to]]
    assert {(m["trigger"], m["triggered_by"]) for m in resp["data"]["messages"]} == {("MANUAL", admin.username)}


def test_template_and_reminder_routes(client, tracker):
    tracker(_item("excess", "Excess above max"))
    admin = _user()
    c = _as(client, admin)
    data = c.get(_BASE + "/templates").get_json()["data"]
    assert data["identity"]["reply_to"] is None and len(data["templates"]) == 8

    assert c.put(_BASE + "/templates/reorder", json={"subject": "a\nb"}).status_code == 400
    assert c.put(_BASE + "/templates/nope", json={"subject": "x"}).status_code == 404
    saved = c.put(_BASE + "/templates/excess_list", json={"footer": "Call the store."}).get_json()["data"]
    assert saved["footer"] == "Call the store." and saved["is_default"] is False
    assert c.put(_BASE + "/templates-identity", json={"reply_to": "broken"}).status_code == 400
    identity = c.put(_BASE + "/templates-identity", json={"reply_to": "store@example.com"}).get_json()["data"]
    assert identity == {"signature": _SIGNATURE, "reply_to": "store@example.com"}

    assert c.put(_BASE + "/reminders", json={"repeat_every_days": 0}).status_code == 400
    assert c.put(_BASE + "/reminders", json={"escalate_to": [{"email": "broken"}]}).status_code == 400
    saved = c.put(_BASE + "/reminders", json={"enabled": True, "escalate_after_days": 5}).get_json()["data"]
    assert saved["enabled"] is True and saved["escalate_after_days"] == 5 and len(saved["lines"]) == 1
    assert AppSetting.query.filter_by(key="reminder_rules").one().updated_by == admin.username
    assert c.get(_BASE + "/reminders").get_json()["data"]["lines"][0]["key"] == "excess"
    ran = c.post(_BASE + "/reminders/run").get_json()
    assert ran["data"] == {"reminded": 0, "escalated": 0, "messages": []}
    assert ran["message"] == "0 reminded, 0 escalated"


def test_send_alert_route_queues_validates_and_is_rate_limited(app, client):
    admin = _user()
    c = _as(client, admin)
    address = _mail("alert")
    body = {"subject": "Line stop", "message": "Press 4 is waiting for a seal kit.", "include": "none",
            "department_ids": [], "emails": [{"email": address, "role": "TO"}]}
    resp = c.post(_BASE + "/send-alert", json=body)
    assert resp.status_code == 200
    message = resp.get_json()["data"]
    assert (message["status"], message["mail_type"], message["to"]) == ("QUEUED", "manual_alert", [address])
    row = _row(message["id"])
    assert row.triggered_by == admin.username and "Press 4 is waiting for a seal kit." in row.body_html

    assert c.post(_BASE + "/send-alert", json={**body, "emails": []}).status_code == 400          # no recipient
    assert c.post(_BASE + "/send-alert", json={**body, "subject": ""}).status_code == 400
    assert c.post(_BASE + "/send-alert", json={**body, "message": "x" * 5001}).status_code == 400
    assert c.post(_BASE + "/send-alert", json={**body, "emails": [{"email": "broken"}]}).status_code == 400
    assert c.post(_BASE + "/send-alert", data="not json", content_type="text/plain").status_code == 400

    app.config["RATE_LIMIT_ENABLED"] = True
    rate_limit.reset()
    try:
        codes = [c.post(_BASE + "/send-alert", json=body).status_code for _ in range(7)]
    finally:
        app.config["RATE_LIMIT_ENABLED"] = False
        rate_limit.reset()
    assert codes == [200] * 5 + [429, 429]


def test_a_footer_that_already_ends_with_the_system_note_does_not_get_it_twice(engine, stock):
    """The store typed the system note into the stock report's footer on Email
    Templates: the mail must not end with it twice."""
    stock([_part("S-OK", "Bearing", "Healthy", 9, 2, 20, unit_price=10.0, stock_value=90.0)])
    rules.save_template("stock_report", {"footer": _SIGNATURE})
    preview = rules.preview_rule("stock_report")
    assert preview["text"].endswith(_SIGNATURE) and preview["text"].count(_NOTE) == 1
    assert preview["html"].count(_NOTE) == 1


# ---- the short & rejection list (the store's request of 17 Oct 2026) ----

@pytest.fixture()
def grn_lines():
    """grn_lines(day, *lines): one spare part and one ERP GRN mirror row per
    line, dated `days` before the plant day `day`; all removed afterwards. A
    line is (days, challan qty, received qty, rejected qty, rate, supplier
    code, supplier name). Returns the parts, in order."""
    parts, rows = [], []

    def make(day, *lines):
        made = []
        for days, challan, received, rejected, rate, code, name in lines:
            part = Part(part_no=f"SRM-{uuid.uuid4().hex[:10].upper()}", part_name="Short rule part", uom="NOS",
                        part_type="OTHER", store_code="MAINTENANCE", unit_price=rate)
            db.session.add(part)
            db.session.flush()
            parts.append(part.id)
            row = ReceivingRegister(store_code="CHILD_PART", grn_no=f"GRN-{uuid.uuid4().hex[:6]}",
                                    grn_date=day - timedelta(days=days), ven_cd=code, ven_name=name,
                                    challan_no="CH-9", item_cd=part.part_no, item_desc="Seal kit",
                                    challan_qty=challan, receipt_qty=received,
                                    accept_qty=received - (rejected or 0), reject_qty=rejected)
            db.session.add(row)
            db.session.flush()
            rows.append(row.id)
            made.append(part)
        db.session.commit()
        return made

    try:
        yield make
    finally:
        db.session.rollback()
        ReceivingRegister.query.filter(ReceivingRegister.id.in_(rows)).delete(synchronize_session=False)
        Part.query.filter(Part.id.in_(parts)).delete(synchronize_session=False)
        db.session.commit()


_STORE_DESK = [{"email": "store@example.com", "role": "TO"}]


def test_short_rejection_rule_ships_off_with_no_one_to_mail(engine, grn_lines, send):
    """Mail is off until the store has seen the list: the rule starts switched
    off with no audience, the scheduler never runs it, and the button mails no
    one — nothing is sent either way."""
    day = _day()
    grn_lines(day, (3, 10, 10, 2, 100.0, "VK-0016", "Keyence India"))
    rule = rules.describe_rule("short_rejection", _at(day, 9))
    assert (rule["enabled"], rule["is_default"], rule["audience"], rule["mail_type"]) == (
        False, True, "custom", "short_rejection")
    assert (rule["department_ids"], rule["emails"], rule["next_run_at"]) == ([], [], None)
    assert rule["recipients"] == {"to": [], "cc": [], "bcc": []}
    assert (rule["title"], rule["content"]["label"]) == ("Short & Rejection list", "GRN lines short or rejected")
    before = _max_id()
    assert not any(rules.rule_is_due("short_rejection", _at(day, hour, 59)) for hour in (0, 8, 12, 23))
    assert rules.run_due_rules(engine, _at(day, 23, 59)) == []
    assert _max_id() == before
    result = rules.run_rule(engine, "short_rejection", _at(day, 9))
    assert (result["status"], result["reason"]) == ("skipped", "no_recipient")
    assert result["messages"][0]["attachments"] == []                      # no workbook for a mail to no one
    assert send.calls == []


def test_short_rejection_mail_is_one_line_of_totals_with_the_workbook_attached(engine, grn_lines, send, monkeypatch):
    day = _day()
    keyence, short, unrated = grn_lines(
        day,
        (40, 2, 2, 2, 25000.0, "VK-0016", "Keyence India"),       # 2 rejected x 25,000
        (5, 10, 7, None, 100.0, "VS-0013", "Sar Coating"),         # 3 short x 100
        (12, 4, 4, 1, None, "VS-0013", "Sar Coating"))             # 1 rejected, no rate
    rules.save_rule("short_rejection", {"emails": _STORE_DESK})
    store = rules._store_name()
    template = rules.template("short_rejection")
    assert template["is_default"] is True and template["defaults"] == {
        "subject": "Short & Rejection — {count} GRN line(s) [{date}]",
        "intro": "Dear Sir,\n\nPlease find attached the GRN lines of {store} of the last 90 days that a supplier "
                 "sent short or quality rejected, supplier-wise and line by line with the days pending.",
        "footer": ""}

    # The preview is the same mail — without building the workbook.
    with monkeypatch.context() as patch:
        patch.setattr(short_rejection_service, "workbook", lambda out: pytest.fail("a preview built the workbook"))
        preview = rules.preview_rule("short_rejection", _at(day, 10))
    assert preview["count"] == 3 and preview["to"] == ["store@example.com"]

    result = rules.run_rule(engine, "short_rejection", _at(day, 10))
    assert result["status"] == "queued" and len(result["messages"]) == 1
    message = result["messages"][0]
    assert (message["mail_type"], message["rule_key"]) == ("short_rejection", "short_rejection")
    assert message["subject"] == f"Short & Rejection — 3 GRN line(s) [{day:%d.%m.%Y}]" == preview["subject"]
    row = _row(message["id"])
    html, text = row.body_html, row.body_text
    assert (html, text) == (preview["html"], preview["text"])
    synced = short_rejection_service.ShortRejectionService().report(
        "MAINTENANCE", today=day)["data_as_of"]["synced_day"]
    totals = ("3 GRN lines · 1 short (₹ 300) · 2 rejected (₹ 50,000) · 2 suppliers · oldest 40 days pending"
              " · 1 without a rate, not valued")
    assert text == "\n\n".join([
        f"Dear Sir,\n\nPlease find attached the GRN lines of {store} of the last 90 days that a supplier sent "
        "short or quality rejected, supplier-wise and line by line with the days pending.",
        totals,
        f"ERP data up to {date.fromisoformat(synced):%d.%m.%Y}. Days pending are counted from the GRN date until "
        "ERP tells the app the item is closed (a replacement or a debit note). ERP gives no closure yet, so every "
        "line here counts as open.",
        _SIGNATURE,                                       # the dashed line and the system note, once
    ])
    assert f"<strong>{totals}</strong>" in html
    _no_list_in(html, text, [keyence.part_no, short.part_no, unrated.part_no])
    assert html.endswith(f"{_SIGNATURE_HTML}</p></body></html>") and html.count(_SIGNATURE_HTML) == 1

    # The attachment is the screen's own workbook, for the last 90 days.
    stored = row.attachments[0]
    first = day - timedelta(days=89)
    assert (stored.filename, stored.mime) == (f"Short & Rejection {first:%d.%m.%Y} to {day:%d.%m.%Y}.xlsx", _XLSX)
    assert message["attachments"] == [{"filename": stored.filename, "size": len(stored.content)}]
    book = load_workbook(io.BytesIO(stored.content))
    assert book.sheetnames == ["Supplier-wise", "Lines"]
    suppliers = [line[:2] for line in book["Supplier-wise"].iter_rows(min_row=2, max_row=4, values_only=True)]
    assert suppliers == [("VK-0016", "Keyence India"), ("VS-0013", "Sar Coating"), ("Total", "2 suppliers")]
    head = [c.value for c in book["Lines"][1]]
    listed = [line[head.index("Part No")] for line in book["Lines"].iter_rows(min_row=2, max_row=4, values_only=True)]
    assert listed == [keyence.part_no, unrated.part_no, short.part_no]       # the longest pending first
    assert send.calls == []                                                  # queued, never sent by a test


def test_short_rejection_rule_is_skipped_when_nothing_is_short_or_rejected(engine, grn_lines, send):
    day = _day()
    grn_lines(day, (200, 5, 3, 1, 10.0, "VX-1", "Old Supplier"))             # before the last 90 days
    rules.save_rule("short_rejection", {"emails": _STORE_DESK})
    preview = rules.preview_rule("short_rejection", _at(day, 10))
    assert (preview["count"], preview["html"], preview["subject"]) == (0, None, "")
    assert preview["text"] == "No mail would go out right now: nothing short or rejected."
    result = rules.run_rule(engine, "short_rejection", _at(day, 10))
    assert (result["status"], result["reason"]) == ("skipped", "nothing_short_or_rejected")
    row = _row(result["messages"][0]["id"])
    assert (row.status, row.body_text, row.attachments) == (
        "SKIPPED", "Nothing was sent: nothing short or rejected.", [])
    # Scheduled: the day is taken by the skip, and the rule is not tried again.
    rules.save_rule("short_rejection", {"enabled": True, "hour": 8, "days": list(range(7))})
    ran = rules.run_due_rules(engine, _at(day, 8))
    assert [(r["key"], r["status"], r["reason"]) for r in ran] == [
        ("short_rejection", "skipped", "nothing_short_or_rejected")]
    assert rules.rule_is_due("short_rejection", _at(day, 9)) is False
    assert send.calls == []


def test_a_short_rejection_list_that_cannot_be_read_is_tried_again(engine, send, monkeypatch):
    day = _day()
    rules.save_rule("short_rejection", {"enabled": True, "hour": 8, "days": list(range(7)), "emails": _STORE_DESK})

    def broken(self, *args, **kwargs):
        raise RuntimeError("the register is down")

    with monkeypatch.context() as patch:
        patch.setattr(short_rejection_service.ShortRejectionService, "report", broken)
        with pytest.raises(RuntimeError, match="short_rejection: RuntimeError: the Short & Rejection list could not"):
            rules.run_due_rules(engine, _at(day, 8))
        assert EmailMessage.query.filter_by(dedup_key=f"short_rejection:{day.isoformat()}").first() is None
        assert rules.rule_is_due("short_rejection", _at(day, 8, 1)) is True
        assert rules.preview_rule("short_rejection", _at(day, 9))["text"] == \
            "No mail would go out right now: the Short & Rejection list could not be read."
    # Readable again: the next tick runs it (nothing in this window, so a skip) and the day is done.
    ran = rules.run_due_rules(engine, _at(day, 8, 2))
    assert [(r["key"], r["reason"]) for r in ran] == [("short_rejection", "nothing_short_or_rejected")]
    assert rules.rule_is_due("short_rejection", _at(day, 9)) is False and send.calls == []
