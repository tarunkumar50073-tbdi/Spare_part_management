"""Short & Rejection (/api/v1/short-rejection, ShortRejectionService) — the
store's request of 17 Oct 2026: which supplier sent short or had a rejection,
since how many days, and against which lot. ONE kind per report, never mixed
("rejection and short mix nhi rhega na"): Short, SRV Rejection, and Line
Rejection ("rejection me drop down se srv rejection, line rejection"), which
ERP holds but the app cannot read yet.

The session DB is shared, so every case reads a plant day of its own far in the
future (`_day`, passed as `today`) with GRN lines dated around it, and removes
its rows afterwards; the route cases look only at their own lines.
"""
import io
import uuid
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.part import Part
from app.models.receiving_register import ReceivingRegister
from app.models.user import User, UserStoreAccess
from app.services import short_rejection_service as sr
from app.services.short_rejection_service import AGE_BANDS, ShortRejectionService, age_band
from app.utils.exceptions import BadRequestError, NotFoundError

STORE = "MAINTENANCE"
_seq = {"n": 0, "week": 0}


def _next():
    _seq["n"] += 1
    return _seq["n"]


def _day():
    """A plant day no other case uses, years after any real GRN."""
    _seq["week"] += 1
    return date(2033, 1, 3) + timedelta(weeks=26 * _seq["week"])


@pytest.fixture()
def make():
    """make.part(tag, rate=, store=) and make.grn(part_no, grn_date, ...): real
    parts and ERP GRN mirror rows, removed after the case."""
    parts, rows = [], []

    class Make:
        @staticmethod
        def part(tag="P", rate=100.0, store=STORE, active=True):
            p = Part(part_no=f"SR-{_next()}-{tag}-{uuid.uuid4().hex[:4]}".upper(), part_name=f"Short {tag}",
                     uom="NOS", part_type="OTHER", store_code=store, unit_price=rate, is_active=active)
            db.session.add(p)
            db.session.commit()
            parts.append(p.id)
            return p

        @staticmethod
        def grn(item_cd, grn_date, challan=10, received=10, accepted=None, rejected=None, supplier=("VS-1", "Sharma Steels"),
                grn_no=None, gate=None, desc="ERP text"):
            r = ReceivingRegister(
                store_code="CHILD_PART", grn_no=grn_no or f"GRN-SR-{_next()}", grn_date=grn_date,
                ven_cd=supplier[0], ven_name=supplier[1], gate_entry_no=gate[0] if gate else None,
                gate_entry_dt=gate[1] if gate else None, challan_no="CH-1", challan_dt=None, po_no=None,
                item_cd=item_cd, item_desc=desc, challan_qty=challan, receipt_qty=received,
                accept_qty=accepted, reject_qty=rejected)
            db.session.add(r)
            db.session.commit()
            rows.append(r.id)
            return r

    try:
        yield Make
    finally:
        db.session.rollback()
        if rows:
            ReceivingRegister.query.filter(ReceivingRegister.id.in_(rows)).delete(synchronize_session=False)
        if parts:
            Part.query.filter(Part.id.in_(parts)).delete(synchronize_session=False)
        db.session.commit()


def _report(today, kind="srv", **kw):
    return ShortRejectionService().report(STORE, kind, today=today, **kw)


def _mine(lines, *parts):
    ids = {p.id for p in parts}
    return [ln for ln in lines if ln["part_id"] in ids]


# ---- the figures, one kind at a time ----

def test_short_is_challan_minus_received_and_srv_rejection_is_the_reject_qty(make):
    today = _day()
    a, b, c, d = make.part("A", rate=100), make.part("B", rate=40), make.part("C"), make.part("D")
    make.grn(a.part_no, today - timedelta(days=3), challan=10, received=7, accepted=7)        # 3 short
    make.grn(b.part_no, today - timedelta(days=4), challan=5, received=5, accepted=3, rejected=2)  # 2 rejected
    make.grn(c.part_no, today - timedelta(days=5), challan=8, received=6.5, accepted=5.25, rejected=1.25)  # both
    make.grn(d.part_no, today - timedelta(days=6), challan=4, received=6, accepted=6)         # more than the challan
    make.grn(d.part_no, today - timedelta(days=6), challan=4, received=4, accepted=4, rejected=0)  # nothing wrong

    short = _report(today, "short")
    by_part = {ln["part_id"]: ln for ln in short["lines"]}
    assert set(by_part) == {a.id, c.id}                 # an excess receipt is not short; a clean line is not listed
    la, lc = by_part[a.id], by_part[c.id]
    assert (la["qty"], la["value"], la["rate"]) == (3.0, 300.0, 100.0)
    assert (lc["qty"], lc["value"]) == (1.5, 150.0)
    assert (lc["challan_qty"], lc["received_qty"], lc["accepted_qty"]) == (8.0, 6.5, 5.25)
    assert lc["description"] == "ERP text" and lc["supplier_code"] == "VS-1" and lc["supplier_name"] == "Sharma Steels"
    assert la["days_pending"] == 3 and la["age_band"] == "0-7"
    t = short["totals"]
    assert (t["lines"], t["qty"], t["value"], t["suppliers"]) == (2, 4.5, 450.0, 1)

    srv = _report(today, "srv")
    by_part = {ln["part_id"]: ln for ln in srv["lines"]}
    assert set(by_part) == {b.id, c.id}
    assert (by_part[b.id]["qty"], by_part[b.id]["value"]) == (2.0, 80.0)
    assert (by_part[c.id]["qty"], by_part[c.id]["value"]) == (1.25, 125.0)
    t = srv["totals"]
    assert (t["lines"], t["qty"], t["value"], t["suppliers"]) == (2, 3.25, 205.0, 1)
    assert (short["kind"], short["label"], srv["kind"], srv["label"]) == ("short", "Short", "srv", "SRV Rejection")


def test_short_and_rejection_are_never_mixed(make):
    """A GRN line both short and rejected is in each list with only that list's
    figure; nothing of the other kind is on its line, its supplier row, its
    totals or its file."""
    today = _day()
    p = make.part("MIX", rate=10)
    make.grn(p.part_no, today - timedelta(days=4), challan=10, received=7, accepted=5, rejected=2)
    short, srv = _report(today, "short"), _report(today, "srv")
    (s_line,), (r_line,) = _mine(short["lines"], p), _mine(srv["lines"], p)
    assert (s_line["qty"], s_line["value"]) == (3.0, 30.0)
    assert (r_line["qty"], r_line["value"]) == (2.0, 20.0)
    for out in (short, srv):
        line, = _mine(out["lines"], p)
        assert not {"short_qty", "rejected_qty", "short_value", "rejected_value", "kind"} & set(line)
        assert not {"short_lines", "rejected_lines", "short_qty", "rejected_qty"} & set(out["totals"])
        assert not {"short_lines", "rejected_lines"} & set(out["summary"][0])
    assert (short["columns"], srv["columns"]) == ({"qty": "Short Qty", "value": "Short Value"},
                                                  {"qty": "Rejected Qty", "value": "Rejected Value"})
    short_book = load_workbook(io.BytesIO(sr.workbook(short)))
    srv_book = load_workbook(io.BytesIO(sr.workbook(srv)))
    for book, mine, other in ((short_book, "Short", "Rejected"), (srv_book, "Rejected", "Short")):
        for sheet in ("Supplier-wise", "Lines"):
            head = [c.value for c in book[sheet][1]]
            assert f"{mine} Qty" in head and f"{mine} Value" in head
            assert not any(str(h).startswith(other) for h in head), (sheet, head)
    assert sr.supplier_columns("short") == ("Supplier Code", "Supplier", "Lines", "Short Qty", "Short Value",
                                            "Oldest (days)", "Latest GRN")
    assert sr.line_columns("srv")[12:15] == ("Rejected Qty", "Rate", "Rejected Value")


def test_a_kind_must_be_named(make):
    for bad in (None, "", "both", "rejection"):
        with pytest.raises(BadRequestError):
            ShortRejectionService().report(STORE, bad, today=_day())
    assert list(sr.KINDS) == ["short", "srv", "line"]
    assert [k["label"] for k in sr.KINDS.values()] == ["Short", "SRV Rejection", "Line Rejection"]


def test_a_line_belongs_to_the_store_through_its_part_no_trimmed_and_case_blind(make):
    today = _day()
    spare = make.part("SP")
    child = make.part("CH", store="CHILD_PART")
    make.grn(f"  {spare.part_no.lower()} ", today - timedelta(days=1), rejected=1, accepted=9)
    make.grn(f"{spare.part_no.lower()}\t", today - timedelta(days=1), challan=5, received=4, accepted=4)
    make.grn(child.part_no, today - timedelta(days=1), rejected=1, accepted=9)
    make.grn("NOT-A-PART-SR", today - timedelta(days=1), rejected=1, accepted=9)

    out = _report(today)
    assert [ln["part_id"] for ln in out["lines"]] == [spare.id]
    assert out["lines"][0]["part_no"] == spare.part_no
    assert [ln["part_id"] for ln in _report(today, "short")["lines"]] == [spare.id]
    child_lines = ShortRejectionService().report("CHILD_PART", "srv", today=today)["lines"]
    assert [ln["part_id"] for ln in child_lines] == [child.id]
    # The line no store has a part for is not lost without a word: counted and named.
    assert out["unmatched"] == {"lines": 1, "item_count": 1,
                                "items": [{"item_cd": "NOT-A-PART-SR", "item_desc": "ERP text", "lines": 1}]}
    assert _report(today, "short")["unmatched"] == {"lines": 0, "item_count": 0, "items": []}


def test_lines_of_items_in_no_part_master_are_counted_per_kind_and_named_in_the_files(make):
    today = _day()
    p = make.part("UM")
    make.grn(p.part_no, today - timedelta(days=1), rejected=1, accepted=9)
    make.grn("CNEI1427-SR", today - timedelta(days=2), challan=80, received=0, accepted=0, desc="Schneider NO Element")
    make.grn("cnei1427-sr ", today - timedelta(days=3), challan=5, received=4, accepted=3, rejected=1)   # both
    make.grn("CNMM6021-SR", today - timedelta(days=4), challan=6, received=5, accepted=5)
    make.grn("CLEAN-SR", today - timedelta(days=4), challan=6, received=6, accepted=6, rejected=0)   # nothing wrong
    short = _report(today, "short")["unmatched"]
    assert (short["lines"], short["item_count"]) == (3, 2)
    assert [(i["item_cd"], i["lines"]) for i in short["items"]] == [("CNEI1427-SR", 2), ("CNMM6021-SR", 1)]
    out = _report(today, "srv")
    assert (out["unmatched"]["lines"], out["unmatched"]["item_count"]) == (1, 1)
    assert [ln["part_id"] for ln in out["lines"]] == [p.id]                     # never listed as a store's line
    notes = sr.file_notes(_report(today, "short"))
    assert any(n.startswith("3 short ERP GRN lines are for item codes that are no store's Part No")
               and "CNEI1427-SR, CNMM6021-SR." in n for n in notes)
    # A supplier or a part filter narrows the list to lines with a part: no count.
    assert _report(today, supplier="VS-1")["unmatched"] is None
    assert _report(today, part_id=p.id)["unmatched"] is None
    # More codes than are named: the rest are counted.
    for i in range(sr.UNMATCHED_ITEMS + 2):
        make.grn(f"ZZ-SR-{i:02d}", today - timedelta(days=5), rejected=1, accepted=9)
    many = _report(today, "srv")
    assert many["unmatched"]["item_count"] == 13 and len(many["unmatched"]["items"]) == sr.UNMATCHED_ITEMS
    assert any("and 3 more." in n for n in sr.file_notes(many))


def test_a_rejection_of_a_part_no_longer_active_still_shows(make):
    today = _day()
    old = make.part("OLD", active=False)
    make.grn(old.part_no, today - timedelta(days=2), rejected=1, accepted=0)
    assert [ln["part_id"] for ln in _report(today)["lines"]] == [old.id]


def test_the_window_is_on_the_grn_date_and_defaults_to_the_last_90_days(make):
    today = _day()
    p = make.part("W")
    for days in (0, 89, 90, 200):
        make.grn(p.part_no, today - timedelta(days=days), grn_no=f"W{days}-{_next()}", rejected=1, accepted=0)

    default = _report(today)
    # A GRN dated today: the data reaches today, so all of the window is known.
    assert default["window"] == {"from": (today - timedelta(days=89)).isoformat(), "to": today.isoformat(),
                                 "days": 90, "period": "90d", "label": "Last 90 days", "default_days": 90,
                                 "not_held": None, "not_synced": None,
                                 "known": {"from": (today - timedelta(days=89)).isoformat(), "to": today.isoformat()}}
    assert sorted(ln["days_pending"] for ln in default["lines"]) == [0, 89]
    # From - To are both inside the window.
    picked = _report(today, date_from=(today - timedelta(days=200)).isoformat(),
                     date_to=(today - timedelta(days=90)).isoformat())
    assert sorted(ln["days_pending"] for ln in picked["lines"]) == [90, 200]
    # A To after today is cut at today; no GRN is dated later.
    later = _report(today, date_to=(today + timedelta(days=30)).isoformat())
    assert later["window"]["to"] == today.isoformat()
    for bad in ({"date_from": today.isoformat(), "date_to": (today - timedelta(days=1)).isoformat()},
                {"date_from": "17-10-2026"}, {"date_to": "yesterday"},
                {"date_from": (today - timedelta(days=sr.MAX_DAYS)).isoformat()}):
        with pytest.raises(BadRequestError):
            _report(today, **bad)
    assert len(_report(today, date_from=(today - timedelta(days=sr.MAX_DAYS - 1)).isoformat())["lines"]) == 4


def test_age_bands_at_their_edges():
    assert [age_band(d) for d in (0, 7, 8, 15, 16, 30, 31, 400)] == [
        "0-7", "0-7", "8-15", "8-15", "16-30", "16-30", "30+", "30+"]
    assert [key for key, _, _ in AGE_BANDS] == ["0-7", "8-15", "16-30", "30+"]


@pytest.mark.parametrize("kind", ["short", "srv"])
def test_days_pending_are_days_since_the_grn_date_and_banded(make, kind):
    today = _day()
    p = make.part("AGE")
    for days in (7, 8, 15, 16, 30, 31):
        if kind == "short":
            make.grn(p.part_no, today - timedelta(days=days), grn_no=f"AGE{days}-{_next()}", challan=5, received=4)
        else:
            make.grn(p.part_no, today - timedelta(days=days), grn_no=f"AGE{days}-{_next()}", rejected=1, accepted=0)
    out = _report(today, kind)
    assert [(ln["days_pending"], ln["age_band"]) for ln in out["lines"]] == [
        (31, "30+"), (30, "16-30"), (16, "16-30"), (15, "8-15"), (8, "8-15"), (7, "0-7")]   # oldest first
    assert out["totals"]["age_bands"] == {"0-7": 1, "8-15": 2, "16-30": 2, "30+": 1}
    assert out["totals"]["oldest_days"] == 31


@pytest.mark.parametrize("kind", ["short", "srv"])
def test_supplier_summary_adds_up_to_the_lines_and_puts_the_worst_first(make, kind):
    today = _day()
    p, q, r = make.part("S1", rate=10), make.part("S2", rate=1000), make.part("S3", rate=50)
    old = ("VA-1", "Alpha Tools")
    dear = ("VB-2", "Beta Bearings")

    def line(part, days, qty, supplier):
        if kind == "short":
            make.grn(part.part_no, today - timedelta(days=days), challan=10, received=10 - qty, supplier=supplier)
        else:
            make.grn(part.part_no, today - timedelta(days=days), rejected=qty, accepted=10 - qty, supplier=supplier)
        # The other kind on the same supplier, in the same window: never counted here.
        other = {"challan": 10, "received": 10, "rejected": 4, "accepted": 6} if kind == "short" else \
            {"challan": 10, "received": 6}
        make.grn(part.part_no, today - timedelta(days=days + 1), supplier=supplier, **other)

    line(p, 40, 2, old)       # ₹ 20
    line(p, 5, 3, old)        # ₹ 30
    line(q, 10, 2, dear)      # ₹ 2,000
    line(r, 10, 1, dear)      # ₹ 50

    out = _report(today, kind)
    summary = out["summary"]
    assert [s["supplier_code"] for s in summary] == ["VA-1", "VB-2"]            # the longest pending first
    alpha, beta = summary
    assert (alpha["lines"], alpha["qty"], alpha["value"]) == (2, 5.0, 50.0)
    assert (alpha["oldest_days"], alpha["latest_grn_date"]) == (40, (today - timedelta(days=5)).isoformat())
    assert (beta["supplier_name"], beta["lines"], beta["qty"], beta["value"], beta["oldest_days"]) == (
        "Beta Bearings", 2, 3.0, 2050.0, 10)
    t = out["totals"]
    for key in ("lines", "qty", "value", "unrated_lines", "open_lines"):
        assert t[key] == round(sum(s[key] for s in summary), 2), key
    lines = out["lines"]
    assert len(lines) == t["lines"] == 4
    assert t["qty"] == sum(ln["qty"] for ln in lines) and t["value"] == sum(ln["value"] for ln in lines)
    assert t["suppliers"] == 2 and t["oldest_days"] == 40
    assert t["latest_grn_date"] == (today - timedelta(days=5)).isoformat()
    # Same days pending: the dearer line first.
    assert [ln["part_id"] for ln in lines] == [p.id, q.id, r.id, p.id]


def test_a_part_with_no_rate_has_no_value_blank_never_zero(make):
    today = _day()
    none, zero, rated = make.part("N", rate=None), make.part("Z", rate=0), make.part("R", rate=10)
    only_unrated = ("VU-1", "Unrated Co")
    make.grn(none.part_no, today - timedelta(days=1), rejected=2, accepted=0, supplier=only_unrated)
    make.grn(zero.part_no, today - timedelta(days=1), rejected=3, accepted=0, supplier=only_unrated)
    make.grn(rated.part_no, today - timedelta(days=1), rejected=1, accepted=0, supplier=("VR-1", "Rated Co"))
    make.grn(none.part_no, today - timedelta(days=1), rejected=1, accepted=0, supplier=("VR-1", "Rated Co"))

    out = _report(today)
    unrated = [ln for ln in out["lines"] if ln["part_id"] in (none.id, zero.id)]
    assert len(unrated) == 3
    assert all(ln["rate"] is None and ln["value"] is None for ln in unrated)
    by = {s["supplier_code"]: s for s in out["summary"]}
    # Every line unrated: the value is not known — None, not 0.
    assert (by["VU-1"]["value"], by["VU-1"]["unrated_lines"], by["VU-1"]["qty"]) == (None, 2, 5.0)
    # Some rated: the rated ones' money, and the others counted.
    assert (by["VR-1"]["value"], by["VR-1"]["unrated_lines"]) == (10.0, 1)
    assert out["totals"]["unrated_lines"] == 3 and out["totals"]["value"] == 10.0
    # Nothing of the kind at all: 0, which is known.
    assert _report(today, "short")["totals"]["value"] == 0.0
    # The files leave the cell blank and say how many lines have no rate.
    book = load_workbook(io.BytesIO(sr.workbook(out)))
    head = [c.value for c in book["Lines"][1]]
    values = [row[head.index("Rejected Value")] for row in book["Lines"].iter_rows(min_row=2, values_only=True)
              if row[head.index("Part No")] == none.part_no]
    assert values == [None, None]
    sup = {row[0]: row for row in book["Supplier-wise"].iter_rows(min_row=2, values_only=True) if row[0]}
    assert sup["VU-1"][4] is None
    notes = [row[0] for row in book["Lines"].iter_rows(values_only=True) if row[0]]
    assert any("3 lines have no rate" in str(n) for n in notes)


def test_filters_supplier_exactly_or_by_text_and_part(make):
    today = _day()
    p, q = make.part("F1"), make.part("F2")
    make.grn(p.part_no, today - timedelta(days=2), rejected=1, accepted=0, supplier=("VK-0016", "Keyence India"))
    make.grn(q.part_no, today - timedelta(days=2), rejected=1, accepted=0, supplier=("VK-00161", "Kirloskar"))

    codes = lambda out: sorted({ln["supplier_code"] for ln in out["lines"]})  # noqa: E731
    assert codes(_report(today, supplier="vk-0016")) == ["VK-0016"]            # a code: that supplier only
    assert codes(_report(today, supplier="VK-001")) == ["VK-0016", "VK-00161"]   # text: in a code
    assert codes(_report(today, supplier="kirlo")) == ["VK-00161"]             # text: in a name
    assert _report(today, supplier="nobody")["lines"] == []
    clicked = _report(today, supplier="VK-0016")
    assert [s["supplier_code"] for s in clicked["summary"]] == ["VK-0016"]    # the click-through: its lines only
    assert clicked["totals"]["lines"] == 1
    by_part = _report(today, part_id=q.id)
    assert [ln["part_id"] for ln in by_part["lines"]] == [q.id]
    assert by_part["filters"] == {"supplier": None, "part": None, "part_id": q.id, "part_no": q.part_no,
                                  "part_name": q.part_name}
    assert by_part["title"].startswith(f"SRV Rejection {q.part_no} ")
    child = make.part("FC", store="CHILD_PART")
    with pytest.raises(NotFoundError):
        _report(today, part_id=child.id)
    with pytest.raises(BadRequestError):
        _report(today, part_id="abc")


def test_lot_erp_age_and_what_erp_does_not_give(make):
    today = _day()
    p, q = make.part("L1"), make.part("L2")
    make.grn(p.part_no, today - timedelta(days=3), rejected=1, accepted=0, gate=("GE-77", today - timedelta(days=4)))
    make.grn(q.part_no, today - timedelta(days=3), rejected=1, accepted=0)
    out = _report(today)
    lots = {ln["part_id"]: ln["lot"] for ln in out["lines"]}
    assert lots == {p.id: f"GE-77 dt {(today - timedelta(days=4)):%d.%m.%Y}", q.id: None}
    gate = {ln["part_id"]: (ln["gate_entry_no"], ln["gate_entry_dt"]) for ln in out["lines"]}
    assert gate[p.id] == ("GE-77", (today - timedelta(days=4)).isoformat())
    # One line has a lot, so the lot is not "not given"; no line has a PO no. or a challan date.
    assert out["not_given"] == ["Challan date", "PO no."]
    as_of = out["data_as_of"]
    assert as_of["synced_at"] and as_of["rows"] >= 2
    # The mirror was written today and this case's plant day is years later: not current.
    assert as_of["stale"] is True and as_of["age_days"] > sr.STALE_AFTER_DAYS
    assert set(out["basis"]) >= {"what", "days", "closed", "no_closure", "value"}
    assert out["basis"]["days"] == "Days pending count from the GRN date until ERP reports the rejection invoice."
    assert _report(today, "short")["basis"]["days"] == \
        "Days pending count from the GRN date until ERP reports how the short is settled."


def test_the_oldest_grn_held_is_given_and_a_window_before_it_is_flagged(make, monkeypatch):
    """The ERP sync keeps only the last ERP_SYNC_SINCE_DAYS of GRNs: an open
    line on an older GRN leaves the data though nothing closed it, so a window
    that reaches before the oldest GRN held must say so, not show "nothing"."""
    today = _day()
    p = make.part("HELD")
    make.grn(p.part_no, today - timedelta(days=89), rejected=1, accepted=0)   # on the window's first day
    make.grn(p.part_no, today - timedelta(days=10), rejected=1, accepted=0)
    out = _report(today)
    as_of = out["data_as_of"]
    # The register is shared by the cases: its oldest GRN is never later than this case's.
    assert as_of["held_from"] <= (today - timedelta(days=89)).isoformat()
    assert as_of["keeps_days"] == (sr.current_app.config.get("ERP_SYNC_SINCE_DAYS") or None)
    assert out["window"]["not_held"] is None                                   # the window is all held

    # The source as it stands on the live register: held from 31 May 2026, kept 120 days.
    held = today - timedelta(days=60)
    real = sr.erp_grn_lines

    def source(first, last):
        rows, as_of = real(first, last)
        return rows, {**as_of, "oldest_grn_date": held, "keeps_days": 120}

    monkeypatch.setattr(sr, "erp_grn_lines", source)
    out = _report(today)                                                        # the last 90 days
    assert out["data_as_of"]["held_from"] == held.isoformat() and out["data_as_of"]["keeps_days"] == 120
    assert out["window"]["not_held"] == {"from": (today - timedelta(days=89)).isoformat(),
                                         "to": (held - timedelta(days=1)).isoformat()}
    notes = sr.file_notes(out)
    assert f"ERP GRN lines are held from {held:%d.%m.%Y} (the ERP sync keeps the last 120 days): " \
           "an open SRV rejection on an older GRN is not listed." in notes
    assert f"No GRN line is held for {today - timedelta(days=89):%d.%m.%Y} to {held - timedelta(days=1):%d.%m.%Y}" \
           " - nothing can be listed for those dates." in notes
    # A window wholly before it: all of it is not held.
    early = _report(today, date_from=(today - timedelta(days=200)).isoformat(),
                    date_to=(today - timedelta(days=100)).isoformat())
    assert early["window"]["not_held"] == {"from": (today - timedelta(days=200)).isoformat(),
                                           "to": (today - timedelta(days=100)).isoformat()}
    # The workbook carries the same lines under its tables.
    book = load_workbook(io.BytesIO(sr.workbook(out)))
    written = [r[0] for r in book["Lines"].iter_rows(values_only=True) if r[0]]
    assert any(str(n).startswith("No GRN line is held for") for n in written)
    # Line Rejection is not read from the GRN register: no such flag there.
    assert _report(today, "line")["window"]["not_held"] is None


def _as_of_like_live(monkeypatch, synced, newest, held):
    """The register's facts as on the live data (synced 28 Sep 2026, newest GRN
    27 Sep, held from 31 May): `synced` the plant day of the last refresh (05:30
    IST = 00:00 UTC), None for no ERP data ever read."""
    real = sr.erp_grn_lines

    def source(first, last):
        rows, as_of = real(first, last)
        return rows, {**as_of, "synced_at": datetime.combine(synced, time(0, 0)) if synced else None,
                      "newest_grn_date": newest, "oldest_grn_date": held, "keeps_days": 120}

    monkeypatch.setattr(sr, "erp_grn_lines", source)


@pytest.mark.parametrize("kind", ["short", "srv"])
def test_a_window_past_the_last_erp_refresh_is_flagged_not_called_empty(make, monkeypatch, kind):
    """The data ends at the last ERP refresh: a period reaching past it (on the
    live data, every period running into October) must say nothing is known
    for those days — not "nothing short" / "no SRV rejection"."""
    today = _day()
    synced, newest, held = today - timedelta(days=19), today - timedelta(days=20), today - timedelta(days=139)
    _as_of_like_live(monkeypatch, synced, newest, held)
    noun = sr.KINDS[kind]["none"]
    iso = date.isoformat

    out = _report(today, kind)                                                  # the last 90 days
    assert out["data_as_of"]["data_until"] == iso(synced) and out["data_as_of"]["stale"] is True
    w = out["window"]
    assert (w["not_held"], w["not_synced"], w["known"]) == (
        None, {"from": iso(synced + timedelta(days=1)), "to": iso(today)},
        {"from": iso(today - timedelta(days=89)), "to": iso(synced)})
    notes = sr.file_notes(out)
    gone = f"{synced + timedelta(days=1):%d.%m.%Y} to {today:%d.%m.%Y}"
    assert f"ERP data ends {synced:%d.%m.%Y}: nothing is known for {gone} - a GRN line booked in ERP on those " \
           "dates is not here until the ERP sync runs again." in notes
    # Nothing in the known part: "none" only for the dates the data covers.
    assert out["lines"] == [] and sr.nothing_text(out) == \
        f"{noun} on the GRN lines dated {today - timedelta(days=89):%d.%m.%Y} to {synced:%d.%m.%Y}."

    # A window wholly after the data: nothing of it is known.
    late = _report(today, kind, date_from=iso(today - timedelta(days=10)))
    assert (late["window"]["not_synced"], late["window"]["known"]) == (
        {"from": iso(today - timedelta(days=10)), "to": iso(today)}, None)
    assert sr.nothing_text(late) == "No GRN line is held for these dates - nothing can be listed."
    buf, name, _ = ShortRejectionService().export(STORE, "pdf", kind=kind, today=today,
                                                  date_from=iso(today - timedelta(days=10)))
    assert buf.getvalue().startswith(b"%PDF")
    book = load_workbook(io.BytesIO(sr.workbook(late)))
    written = [str(r[0]) for r in book["Lines"].iter_rows(values_only=True) if r[0]]
    assert any(n.startswith(f"ERP data ends {synced:%d.%m.%Y}: nothing is known for ") for n in written)

    # Before the oldest GRN held AND after the last refresh: both said, the middle known.
    wide = _report(today, kind, date_from=iso(today - timedelta(days=200)))
    assert wide["window"]["not_held"] == {"from": iso(today - timedelta(days=200)),
                                          "to": iso(held - timedelta(days=1))}
    assert wide["window"]["known"] == {"from": iso(held), "to": iso(synced)}
    # A window wholly inside the data: nothing flagged, the whole of it known.
    inside = _report(today, kind, date_from=iso(today - timedelta(days=60)), date_to=iso(today - timedelta(days=30)))
    assert (inside["window"]["not_held"], inside["window"]["not_synced"]) == (None, None)
    assert inside["window"]["known"] == {"from": iso(today - timedelta(days=60)), "to": iso(today - timedelta(days=30))}
    assert sr.nothing_text(inside) == sr.KINDS[kind]["nothing"]
    assert not any(n.startswith("ERP data ends") for n in sr.file_notes(inside))
    # Line Rejection is not read from the register: not flagged by it.
    line = _report(today, "line")["window"]
    assert (line["not_held"], line["not_synced"], line["known"]) == (
        None, None, {"from": iso(today - timedelta(days=89)), "to": iso(today)})


def test_the_data_reaches_the_later_of_the_refresh_day_and_the_newest_grn(make, monkeypatch):
    today = _day()
    # Refreshed today: the window is all known, even with no GRN dated today.
    _as_of_like_live(monkeypatch, today, today - timedelta(days=3), today - timedelta(days=100))
    out = _report(today)
    assert out["data_as_of"]["data_until"] == today.isoformat() and out["window"]["not_synced"] is None
    # A GRN dated after the refresh day (a clock difference): the data reaches it.
    _as_of_like_live(monkeypatch, today - timedelta(days=5), today - timedelta(days=2),
                     today - timedelta(days=100))
    out = _report(today)
    assert out["data_as_of"]["data_until"] == (today - timedelta(days=2)).isoformat()
    assert out["window"]["not_synced"] == {"from": (today - timedelta(days=1)).isoformat(), "to": today.isoformat()}
    # No ERP data ever read: nothing known at all, and the files say that once.
    _as_of_like_live(monkeypatch, None, None, None)
    out = _report(today)
    assert out["data_as_of"]["data_until"] is None
    assert (out["window"]["not_held"], out["window"]["known"]) == (None, None)
    assert out["window"]["not_synced"] == {"from": out["window"]["from"], "to": today.isoformat()}
    notes = sr.file_notes(out)
    assert notes[0] == "No ERP GRN data has been read yet" and not any(n.startswith("ERP data ends") for n in notes)
    assert sr.nothing_text(out) == "No GRN line is held for these dates - nothing can be listed."


@pytest.mark.parametrize("kind", ["short", "srv"])
def test_an_empty_window_is_an_empty_list_with_zero_totals(make, kind):
    out = _report(_day(), kind)
    assert out["available"] is True
    assert out["lines"] == [] and out["summary"] == [] and out["not_given"] == []
    t = out["totals"]
    assert (t["lines"], t["suppliers"], t["qty"], t["value"], t["oldest_days"]) == (0, 0, 0.0, 0.0, None)


# ---- the Period dropdown ----

def test_periods_are_cut_on_the_plant_day():
    """The screen's Period dropdown, cut on the server's plant day — last month
    and the financial years (1 April - 31 March) as the Issuance / Consumption
    reports cut them."""
    def window(period, today, date_from=None, date_to=None):
        first, last = sr._window(period, date_from, date_to, today)
        return first.isoformat(), last.isoformat()

    today = date(2026, 10, 17)
    assert window("30d", today) == ("2026-09-18", "2026-10-17")
    assert window("90d", today) == ("2026-07-20", "2026-10-17")
    assert window("month", today) == ("2026-10-01", "2026-10-17")
    assert window("last_month", today) == ("2026-09-01", "2026-09-30")
    assert window("this_fy", today) == ("2026-04-01", "2026-10-17")
    assert window("last_fy", today) == ("2025-04-01", "2026-03-31")
    # Across a year end and before April: last month is December, the FY began last April.
    january = date(2027, 1, 10)
    assert window("last_month", january) == ("2026-12-01", "2026-12-31")
    assert window("this_fy", january) == ("2026-04-01", "2027-01-10")
    assert window("last_fy", january) == ("2025-04-01", "2026-03-31")
    # A period other than Custom ignores From - To; Custom takes them.
    assert window("30d", today, "2026-01-01", "2026-01-31") == ("2026-09-18", "2026-10-17")
    assert window("custom", today, "2026-01-01", "2026-01-31") == ("2026-01-01", "2026-01-31")
    # No period: Custom when a date is given (an older link), else the last 90 days.
    assert sr._period(None, None, None) == "90d" == sr.DEFAULT_PERIOD
    assert sr._period("", "2026-01-01", None) == "custom"
    assert [key for key, _ in sr.PERIODS] == ["30d", "90d", "month", "last_month", "this_fy", "last_fy", "custom"]
    with pytest.raises(BadRequestError):
        sr._period("week", None, None)


def test_a_period_picks_the_window_of_the_report(make):
    today = _day()
    p = make.part("PER")
    make.grn(p.part_no, today - timedelta(days=20), grn_no=f"P20-{_next()}", rejected=1, accepted=0)
    make.grn(p.part_no, today - timedelta(days=60), grn_no=f"P60-{_next()}", rejected=1, accepted=0)
    last30 = _report(today, period="30d")
    assert (last30["window"]["period"], last30["window"]["label"], last30["window"]["days"]) == (
        "30d", "Last 30 days", 30)
    assert [ln["days_pending"] for ln in last30["lines"]] == [20]
    assert sorted(ln["days_pending"] for ln in _report(today, period="90d")["lines"]) == [20, 60]
    assert last30["periods"][0] == {"key": "30d", "label": "Last 30 days"}
    first = today - timedelta(days=29)
    assert last30["title"] == f"SRV Rejection {first:%d.%m.%Y} to {today:%d.%m.%Y}"
    assert _report(today, "short", period="30d")["title"] == f"Short {first:%d.%m.%Y} to {today:%d.%m.%Y}"


def test_the_part_box_finds_the_typed_text_in_the_part_no_or_description(make):
    today = _day()
    p, q = make.part("PT1"), make.part("PT2")
    make.grn(p.part_no, today - timedelta(days=2), rejected=1, accepted=0, desc="Proximity sensor")
    make.grn(q.part_no, today - timedelta(days=2), rejected=1, accepted=0, desc="Seal kit")
    ids = lambda out: sorted(ln["part_id"] for ln in out["lines"])  # noqa: E731
    assert ids(_report(today, part=f"  {p.part_no.lower()} ")) == [p.id]        # the Part No, case- and space-blind
    assert ids(_report(today, part="SEAL")) == [q.id]                           # the description
    assert ids(_report(today, part="  ")) == sorted([p.id, q.id])               # a blank box keeps every line
    out = _report(today, part="seal")
    assert out["filters"]["part"] == "seal" and out["title"].startswith("SRV Rejection seal ")
    assert out["unmatched"] is None                                             # narrowed: no count of the others


# ---- open / closed: the ONE closure read ----

def test_every_line_is_open_until_erp_reports_its_closure(make):
    today = _day()
    p = make.part("OPEN")
    make.grn(p.part_no, today - timedelta(days=12), challan=5, received=4, accepted=3, rejected=1)
    srv, short = _report(today, "srv"), _report(today, "short")
    (line,) = srv["lines"]
    assert (line["status"], line["status_text"], line["closed_ref"], line["closed_date"]) == (
        "open", "Open — rejection invoice not yet made", None, None)
    assert (line["days_pending"], line["age_band"]) == (12, "8-15")
    (sline,) = short["lines"]
    assert (sline["status"], sline["status_text"]) == ("open", "Open")
    assert srv["closures_known"] is False and short["closures_known"] is False   # ERP gives none yet
    assert all(sr.erp_closures(kind, []) is None for kind in sr.KINDS)
    notes = sr.file_notes(srv)
    assert srv["basis"]["no_closure"] in notes and srv["basis"]["closed"] not in notes
    assert (srv["totals"]["open_lines"], srv["totals"]["closed_lines"]) == (1, 0)


def test_a_rejection_invoice_from_erp_closes_the_line_and_stops_its_days(make, monkeypatch):
    """The closure hook: when ERP reports "rejection invoice no. / date" per
    GRN line (keyed by GRN no. and item code), the line is closed, its days
    stop at the invoice date, and it leaves the oldest and the age bands."""
    today = _day()
    p, q, r = make.part("CL1", rate=10), make.part("CL2", rate=10), make.part("CL3", rate=10)
    vendor = ("VC-1", "Closing Co")
    closed = make.grn(p.part_no, today - timedelta(days=40), rejected=2, accepted=8, supplier=vendor, grn_no="grn-cl-1 ")
    still_open = make.grn(q.part_no, today - timedelta(days=5), rejected=1, accepted=9, supplier=vendor)
    no_date = make.grn(r.part_no, today - timedelta(days=50), rejected=1, accepted=9, supplier=vendor)
    blank = make.grn(q.part_no, today - timedelta(days=9), rejected=1, accepted=9, supplier=vendor)
    asked = []

    def closures(kind, rows):
        asked.append((kind, len(rows)))
        if kind != "srv":
            return None
        return {
            sr.grn_line_key("GRN-CL-1", p.part_no.lower()): {"ref_no": "RI/26/0042",
                                                            "ref_date": (today - timedelta(days=10)).isoformat()},
            sr.grn_line_key(no_date.grn_no, r.part_no): {"ref_no": "RI/26/0050", "ref_date": None},
            sr.grn_line_key(blank.grn_no, q.part_no): {"ref_no": "  ", "ref_date": today},   # no no.: still open
        }

    monkeypatch.setattr(sr, "erp_closures", closures)
    out = _report(today, "srv")
    assert asked == [("srv", 4)] and out["closures_known"] is True
    by_id = {ln["id"]: ln for ln in out["lines"]}
    c = by_id[closed.id]
    assert (c["status"], c["status_text"]) == ("closed", f"Closed — invoice RI/26/0042 ({today - timedelta(days=10):%d.%m.%Y})")
    assert (c["days_pending"], c["age_band"], c["closed_ref"]) == (30, None, "RI/26/0042")   # GRN to invoice, not to today
    nd = by_id[no_date.id]
    assert (nd["status"], nd["status_text"], nd["days_pending"], nd["age_band"]) == (
        "closed", "Closed — invoice RI/26/0050", None, None)
    assert by_id[still_open.id]["status"] == by_id[blank.id]["status"] == "open"
    # Open lines first, the longest pending first; the closed ones after.
    assert [ln["id"] for ln in out["lines"]] == [blank.id, still_open.id, closed.id, no_date.id]
    t = out["totals"]
    assert (t["open_lines"], t["closed_lines"], t["oldest_days"]) == (2, 2, 9)
    assert t["age_bands"] == {"0-7": 1, "8-15": 1, "16-30": 0, "30+": 0}
    (vc,) = out["summary"]
    assert (vc["open_lines"], vc["closed_lines"], vc["oldest_days"], vc["lines"]) == (2, 2, 9, 4)
    # The files carry the status; a closed line's age is not coloured.
    book = load_workbook(io.BytesIO(sr.workbook(out)))
    head = [cell.value for cell in book["Lines"][1]]
    rows = list(book["Lines"].iter_rows(min_row=2, max_row=5))
    status = [row[head.index("Status")].value for row in rows]
    assert status == ["Open — rejection invoice not yet made"] * 2 + [
        f"Closed — invoice RI/26/0042 ({today - timedelta(days=10):%d.%m.%Y})", "Closed — invoice RI/26/0050"]
    assert rows[2][head.index("Days Pending")].value == 30 and rows[2][head.index("Age")].fill.fill_type is None
    notes = sr.file_notes(out)
    assert out["basis"]["closed"] in notes and out["basis"]["no_closure"] not in notes
    # The same hook settles a short, in ERP's own words; the SRV closure never touches it.
    monkeypatch.setattr(sr, "erp_closures", lambda kind, rows: {
        sr.grn_line_key("GRN-SH-1", r.part_no): {"ref_no": "DN/26/7", "ref_date": today,
                                                                       "how": "debit note"}} if kind == "short" else None)
    make.grn(r.part_no, today - timedelta(days=6), challan=5, received=3, grn_no="GRN-SH-1", supplier=vendor)
    (settled,) = _report(today, "short")["lines"]
    assert (settled["status"], settled["status_text"], settled["days_pending"]) == (
        "closed", f"Closed — debit note DN/26/7 ({today:%d.%m.%Y})", 6)
    assert all(ln["status"] == "open" for ln in _report(today, "srv")["lines"])


# ---- Line Rejection: in ERP, not readable yet ----

def test_line_rejections_are_said_to_be_not_read_yet(make):
    """Line rejections are punched in ERP (line_rejection_report), which the
    app neither mirrors nor reaches: said, never filled with made-up rows —
    and an SRV rejection is never shown under Line Rejection."""
    today = _day()
    p = make.part("LR")
    make.grn(p.part_no, today - timedelta(days=3), rejected=2, accepted=8)
    out = _report(today, "line")
    assert (out["kind"], out["label"], out["available"], out["erp_table"]) == (
        "line", "Line Rejection", False, "line_rejection_report")
    assert (out["lines"], out["summary"], out["unmatched"], out["totals"]["lines"]) == ([], [], None, 0)
    assert out["data_as_of"]["synced_at"]                                       # how old the ERP data is, still said
    assert sr.erp_line_rejections(date(2026, 1, 1), date(2026, 1, 31)) is None
    for fmt in ("excel", "pdf"):
        with pytest.raises(BadRequestError, match="Line Rejections are not read from ERP yet"):
            ShortRejectionService().export(STORE, fmt, kind="line", today=today)


def test_once_erp_gives_line_rejections_they_list_on_their_own(make, monkeypatch):
    """The hook: when erp_line_rejections returns rows, the Line Rejection
    type lists them — and only them, never the GRN's SRV rejections."""
    today = _day()
    p, q = make.part("LR1", rate=20), make.part("LR2", rate=20)
    make.grn(q.part_no, today - timedelta(days=3), rejected=2, accepted=8)        # an SRV rejection
    line_row = {"id": -1, "grn_no": "GRN-L-1", "grn_date": today - timedelta(days=18), "ven_cd": "VL-1",
                "ven_name": "Line Co", "gate_entry_no": None, "gate_entry_dt": None, "challan_no": None,
                "challan_dt": None, "po_no": None, "item_cd": p.part_no, "item_desc": "Bush", "challan_qty": 10,
                "receipt_qty": 10, "accept_qty": 10, "reject_qty": 3}
    monkeypatch.setattr(sr, "erp_line_rejections", lambda first, last: [line_row])
    out = _report(today, "line")
    assert out["available"] is True
    assert [(ln["part_id"], ln["qty"], ln["value"], ln["days_pending"], ln["age_band"]) for ln in out["lines"]] == [
        (p.id, 3.0, 60.0, 18, "16-30")]
    assert out["lines"][0]["status_text"] == "Open — rejection invoice not yet made"
    assert out["title"].startswith("Line Rejection ")
    assert q.id not in {ln["part_id"] for ln in out["lines"]}
    assert p.id not in {ln["part_id"] for ln in _report(today, "srv")["lines"]}


# ---- the mail rule: on hold ----

def test_the_mail_rule_stays_off_with_no_one_to_mail():
    """The store put every mail on hold on 17 Oct 2026 ("ye mail wale kaam baad
    me kar dena"): the Short & Rejection rule ships switched off, with no one in
    its audience and no .env flag that could switch it on at start-up, so no
    scheduled or catch-up run ever finds it due."""
    from app.services import email_rules as rules
    spec = rules._RULE["short_rejection"]
    assert spec["audience"] == "custom" and "env" not in spec
    assert rules._rule_defaults(spec) == {"enabled": False, "hour": 8, "minute": 0, "days": [0, 1, 2, 3, 4, 5],
                                          "department_ids": [], "emails": []}
    settings, is_default = rules.rule_settings("short_rejection")
    assert is_default is True and settings["enabled"] is False
    assert (settings["department_ids"], settings["emails"]) == ([], [])
    day = date(2033, 1, 3)
    assert not any(rules.rule_is_due("short_rejection", datetime.combine(day, time(hour, 59)))
                   for hour in range(24))
    # Its mail, were it ever switched on, takes the kinds one by one — never mixed.
    assert rules.SHORT_REJECTION_KINDS == ("short", "srv")


# ---- the files ----

def test_excel_is_one_workbook_of_one_kind_with_a_supplier_wise_and_a_lines_sheet(make):
    today = _day()
    p = make.part("X", rate=12.5)
    make.grn(p.part_no, today - timedelta(days=20), challan=10, received=6, rejected=1, accepted=5,
             supplier=("VX-9", "=HYPERLINK(\"x\")"))
    first = today - timedelta(days=89)
    for kind, label, qty, value in (("short", "Short", 4, 50), ("srv", "SRV Rejection", 1, 12.5)):
        buf, name, mime = ShortRejectionService().export(STORE, "excel", kind=kind, today=today)
        assert name == f"{label} {first:%d.%m.%Y} to {today:%d.%m.%Y}.xlsx"
        assert mime == sr.XLSX_MIME
        book = load_workbook(io.BytesIO(buf.getvalue()))
        assert book.sheetnames == ["Supplier-wise", "Lines"]
        sup = [list(r) for r in book["Supplier-wise"].iter_rows(values_only=True)]
        assert tuple(sup[0]) == sr.supplier_columns(kind)
        assert sup[1][:7] == ["VX-9", "=HYPERLINK(\"x\")", 1, qty, value, 20, f"{today - timedelta(days=20):%d.%m.%Y}"]
        assert sup[2][:5] == ["Total", "1 supplier", 1, qty, value]               # the Total row last
        assert book["Supplier-wise"].cell(row=2, column=2).data_type == "s"        # a formula-looking name stays text
        assert book["Supplier-wise"]["A3"].font.bold
        lines = [list(r) for r in book["Lines"].iter_rows(values_only=True)]
        assert tuple(lines[0]) == sr.line_columns(kind)
        row = dict(zip(sr.line_columns(kind), lines[1]))
        cols = sr.KINDS[kind]
        assert (row["Part No"], row[cols["qty"]], row[cols["value"]], row["Rate"]) == (p.part_no, qty, value, 12.5)
        assert (row["Challan Qty"], row["Received Qty"], row["Accepted Qty"]) == (10, 6, 5)
        assert (row["Days Pending"], row["Age"], row["GRN Date"]) == (20, "16-30", f"{today - timedelta(days=20):%d.%m.%Y}")
        age = book["Lines"].cell(row=2, column=sr.line_columns(kind).index("Age") + 1)
        assert age.fill.fgColor.rgb.endswith(sr.AGE_TINTS["16-30"][0])
        assert row["Status"] == cols["open"]
        notes = [r[0] for r in lines if r[0]]
        assert cols["days"] in notes and cols["no_closure"] in notes                 # the basis travels with the file
        assert cols["what"] in notes


def test_pdf_and_a_bad_format(make):
    today = _day()
    p = make.part("PDF")
    make.grn(p.part_no, today - timedelta(days=2), rejected=1, accepted=0, supplier=("VP-1", "Smith & Sons"))
    buf, name, mime = ShortRejectionService().export(STORE, "pdf", kind="srv", today=today)
    assert buf.getvalue().startswith(b"%PDF") and name.startswith("SRV Rejection ") and name.endswith(".pdf")
    assert mime == "application/pdf"
    empty, name, _ = ShortRejectionService().export(STORE, "pdf", kind="short", today=_day())   # nothing in the window
    assert empty.getvalue().startswith(b"%PDF") and name.startswith("Short ")
    with pytest.raises(BadRequestError):
        ShortRejectionService().export(STORE, "csv", kind="srv", today=today)


# ---- the routes ----

@pytest.fixture()
def client(app):
    return app.test_client()


def _login(client, perms=None, stores=(STORE,)):
    u = User(username=f"sruser{_next()}_{uuid.uuid4().hex[:6]}", role="STORE_USER", is_active=True)
    u.set_password("pw")
    if perms:
        u.set_permissions(perms)
    db.session.add(u)
    db.session.flush()
    for code in stores:
        db.session.add(UserStoreAccess(user_id=u.id, store_code=code))
    db.session.commit()
    with client.session_transaction() as s:
        s["user_id"] = u.id
    return client


URL = "/api/v1/short-rejection"


def test_routes_need_the_receiving_right_and_the_store(client):
    c = _login(client)
    assert c.get(f"{URL}?store_code={STORE}&kind=srv").status_code == 403
    assert c.get(f"{URL}/export?store_code={STORE}&kind=short&format=excel").status_code == 403
    c = _login(client, perms={"short_rejection": {"view": True}})
    assert c.get(f"{URL}?store_code=CHILD_PART&kind=srv").status_code == 403
    assert c.get(f"{URL}/export?store_code=CHILD_PART&kind=srv&format=excel").status_code == 403
    assert c.get(f"{URL}?kind=srv").status_code == 400
    assert client.application.test_client().get(f"{URL}?store_code={STORE}&kind=srv").status_code == 401


def test_routes_read_and_export_one_kind(client, make):
    today = sr.plant_today()
    p = make.part("RT", rate=10)
    make.grn(p.part_no, today - timedelta(days=9), challan=5, received=3, rejected=2, accepted=1,
             supplier=("VR-77", "Route Rubber"))
    c = _login(client, perms={"short_rejection": {"view": True}})
    srv = c.get(f"{URL}?store_code={STORE}&kind=srv").get_json()["data"]
    assert {"kind", "label", "available", "summary", "lines", "totals", "data_as_of", "basis", "window", "filters",
            "title"} <= set(srv)
    assert [(ln["qty"], ln["value"], ln["days_pending"], ln["age_band"]) for ln in _mine(srv["lines"], p)] == [
        (2.0, 20.0, 9, "8-15")]
    short = c.get(f"{URL}?store_code={STORE}&kind=short").get_json()["data"]
    assert [(ln["qty"], ln["value"], ln["status_text"]) for ln in _mine(short["lines"], p)] == [(2.0, 20.0, "Open")]
    narrowed = c.get(f"{URL}?store_code={STORE}&kind=srv&supplier=VR-77&part_id={p.id}").get_json()["data"]
    assert [ln["part_id"] for ln in narrowed["lines"]] == [p.id] and narrowed["summary"][0]["supplier_code"] == "VR-77"
    assert c.get(f"{URL}?store_code={STORE}").status_code == 400                   # no kind: never a mixed list
    assert c.get(f"{URL}?store_code={STORE}&kind=both").status_code == 400
    assert c.get(f"{URL}?store_code={STORE}&kind=srv&from=2026-13-01").status_code == 400
    assert c.get(f"{URL}?store_code={STORE}&kind=srv&part_id=999999999").status_code == 404

    xl = c.get(f"{URL}/export?store_code={STORE}&kind=srv&format=excel&supplier=VR-77")
    assert xl.status_code == 200 and xl.mimetype == sr.XLSX_MIME
    assert "SRV Rejection VR-77 " in xl.headers["Content-Disposition"]
    assert load_workbook(io.BytesIO(xl.data)).sheetnames == ["Supplier-wise", "Lines"]
    sxl = c.get(f"{URL}/export?store_code={STORE}&kind=short&format=excel&supplier=VR-77")
    assert sxl.status_code == 200 and "Short VR-77 " in sxl.headers["Content-Disposition"]
    assert "Short Qty" in [cell.value for cell in load_workbook(io.BytesIO(sxl.data))["Lines"][1]]
    pdf = c.get(f"{URL}/export?store_code={STORE}&kind=short&format=pdf")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")
    assert c.get(f"{URL}/export?store_code={STORE}&kind=srv&format=csv").status_code == 400
    assert c.get(f"{URL}/export?store_code={STORE}&format=excel").status_code == 400


def test_routes_line_rejection_says_why_and_has_no_file(client):
    c = _login(client, perms={"short_rejection": {"view": True}})
    body = c.get(f"{URL}?store_code={STORE}&kind=line").get_json()["data"]
    assert (body["available"], body["erp_table"], body["lines"]) == (False, "line_rejection_report", [])
    res = c.get(f"{URL}/export?store_code={STORE}&kind=line&format=excel")
    assert res.status_code == 400 and "not read from ERP yet" in res.get_json()["message"]


def test_routes_take_the_period_and_the_part_text(client, make):
    today = sr.plant_today()
    p = make.part("RP", rate=10)
    make.grn(p.part_no, today - timedelta(days=3), rejected=1, accepted=0, supplier=("VR-88", "Period Plastics"))
    c = _login(client, perms={"short_rejection": {"view": True}})
    body = c.get(f"{URL}?store_code={STORE}&kind=srv&period=30d&part={p.part_no.lower()}").get_json()["data"]
    assert (body["window"]["period"], body["window"]["days"]) == ("30d", 30)
    assert [ln["part_id"] for ln in body["lines"]] == [p.id]
    assert body["lines"][0]["status_text"] == "Open — rejection invoice not yet made"
    assert body["closures_known"] is False
    # A period other than Custom ignores From - To; an unknown one is refused.
    ignored = c.get(f"{URL}?store_code={STORE}&kind=srv&period=30d&from=2020-01-01&to=2020-01-31").get_json()["data"]
    assert ignored["window"]["to"] == today.isoformat()
    assert c.get(f"{URL}?store_code={STORE}&kind=srv&period=week").status_code == 400
    last_fy = c.get(f"{URL}/export?store_code={STORE}&kind=srv&format=excel&period=last_fy")
    first, last = sr.preset_window("last_fy", today)
    assert last_fy.status_code == 200
    assert f"SRV Rejection {first:%d.%m.%Y} to {last:%d.%m.%Y}.xlsx" in last_fy.headers["Content-Disposition"]
    pdf = c.get(f"{URL}/export?store_code={STORE}&kind=short&format=pdf&period=custom&from={today.isoformat()}&part=zz-none")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")
