"""Urgent Challan Register (the store's spec, 21 Oct 2026): urgent material that
came on a delivery challan before its PO / invoice. A register that never moves
stock — the ERP GRN after quality stays the one way stock goes up — with
duplicate challans, repeated submits, over-invoicing and a GRN line linked
twice all refused, partial invoicing reconciled, and every change kept.

Everything here is CR-<tag>-<n>: this file's own store, parts, suppliers and
users, switched off afterwards."""
import io
import re
import uuid
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.challan_register import ChallanHistory, ChallanLink, ChallanRegister
from app.models.part import Part
from app.models.receiving_register import ReceivingRegister
from app.models.spare_stock import SpareStock, SpareStockLedger
from app.models.store import Section, Store
from app.models.supplier import Supplier
from app.models.user import PERMISSION_MODULES, User, UserStoreAccess
from app.services import spare_report_service
from app.services.challan_register_service import ChallanRegisterService
from app.services.spare_stock_service import SpareStockService
from app.utils.exceptions import BadRequestError, ConflictError, NotFoundError

_TAG = uuid.uuid4().hex[:6].upper()
_seq = {"n": 0}
URL = "/api/v1/challans"
TODAY = date(2026, 10, 21)


def _n():
    _seq["n"] += 1
    return _seq["n"]


@pytest.fixture(autouse=True)
def plant_today(monkeypatch):
    monkeypatch.setattr(spare_report_service, "plant_today", lambda: TODAY)


@pytest.fixture()
def make():
    made = {"stores": [], "users": []}

    def store():
        code = f"CR{_TAG}{_n()}"
        s = Store(code=code, name=f"Challan Store {code}")
        db.session.add(s)
        db.session.flush()
        db.session.add(Section(store_id=s.id, name=f"Sec {code}"))
        db.session.commit()
        made["stores"].append(code)
        return code

    def part(store_code, qty=5, uom="NOS"):
        n = _n()
        p = Part(part_no=f"CR-{_TAG}-{n}", part_name=f"Challan Part {n}", uom=uom, part_type="OTHER",
                 store_code=store_code)
        db.session.add(p)
        db.session.commit()
        if qty is not None:
            SpareStockService().set_quantity(p.id, qty, txn_type="OPENING")
        return p

    def supplier():
        n = _n()
        s = Supplier(vendor_code=f"CRV{_TAG}{n}"[:20], supplier_name=f"Challan Supplier {n}", status="Active")
        db.session.add(s)
        db.session.commit()
        return s

    def user(stores=(), perms=None, role="STORE_USER"):
        u = User(username=f"cr{_TAG.lower()}{_n()}", role=role, is_active=True)
        u.set_password("pw")
        if perms:
            u.set_permissions(perms)
        db.session.add(u)
        db.session.flush()
        for code in stores:
            db.session.add(UserStoreAccess(user_id=u.id, store_code=code))
        db.session.commit()
        made["users"].append(u.id)
        return u

    yield type("Make", (), {"store": staticmethod(store), "part": staticmethod(part),
                            "supplier": staticmethod(supplier), "user": staticmethod(user)})
    db.session.rollback()
    Store.query.filter(Store.code.in_(made["stores"])).update({"is_active": False}, synchronize_session=False)
    User.query.filter(User.id.in_(made["users"])).update({"is_active": False}, synchronize_session=False)
    db.session.commit()


@pytest.fixture()
def client(app):
    return app.test_client()


def _as(client, user):
    with client.session_transaction() as sess:
        sess["user_id"] = user.id
    return client


svc = ChallanRegisterService()
# The PO an invoice is made against — every invoice in these tests carries it.
PO = {"po_status": "Approved", "po_reference": "PO-1"}


def _entry(part, **over):
    data = {"part_id": part.id, "challan_no": f"DC-{_n()}", "challan_date": "2026-10-19",
            "receipt_date": "2026-10-20", "supplier_source": "NAME_ONLY", "supplier_name": "Urgent Traders",
            "challan_qty": 5, "po_status": "Pending Approval", "urgency_reason": "Breakdown"}
    data.update(over)
    return data


def _stock(part):
    """(qty on hand, ledger rows) of a part — what the register must never change."""
    row = SpareStock.query.filter_by(part_id=part.id).first()
    return (float(row.qty_on_hand) if row and row.qty_on_hand is not None else None,
            SpareStockLedger.query.filter_by(part_id=part.id).count())


def _grn_line(part, grn_no, ven_cd="VEN-1", accept=None, reject=None, challan_no=None, grn_date=TODAY):
    rr = ReceivingRegister(store_code="CHILD_PART", grn_no=grn_no, grn_date=grn_date, ven_cd=ven_cd,
                           ven_name="ERP Vendor", item_cd=part.part_no, item_desc=part.part_name,
                           challan_no=challan_no, receipt_qty=5, accept_qty=accept, reject_qty=reject)
    db.session.add(rr)
    db.session.commit()
    return rr


# ---- registering -------------------------------------------------------------------------------

def test_registering_a_challan_writes_the_register_only(make):
    store = make.store()
    part = make.part(store, qty=7)
    before = _stock(part)
    out = svc.create(store, _entry(part), user="store.user")
    assert re.fullmatch(r"CHL-20261021-\d{3}", out["challan_id"]) and out["status"] == "Pending Invoice"
    assert (out["invoice_status"], out["quality_status"], out["part_description"], out["unit_of_measure"]) == (
        "Pending", "Pending", part.part_name, "NOS")
    assert [h["action"] for h in out["history"]] == ["CREATED"] and out["history"][0]["by"] == "store.user"
    assert _stock(part) == before                                            # no stock, no ledger row
    seq = int(out["challan_id"][-3:])
    assert svc.create(store, _entry(part))["challan_id"] == f"CHL-20261021-{seq + 1:03d}"


@pytest.mark.parametrize("over, message", [
    ({"challan_no": " "}, "Challan No. is required"),
    ({"challan_qty": 0}, "more than 0"),
    ({"challan_qty": "abc"}, "must be a number"),
    ({"challan_qty": 1.234}, "at most 2 decimals"),
    ({"challan_date": "2026-10-25"}, "cannot be in the future"),
    ({"receipt_date": "2026-10-18"}, "before the challan date"),
    ({"po_status": "Approved"}, "PO reference is required"),
    ({"po_status": "Maybe"}, "PO status must be one of"),
    ({"urgency_reason": "Other"}, "Remarks when the urgency is Other"),
    ({"supplier_source": None}, "Select the supplier"),
    ({"supplier_source": "SUPPLIER_MASTER"}, "supplier master"),
    ({"supplier_source": "NAME_ONLY", "ven_cd": "X1"}, "A vendor code comes only from ERP"),
    ({"supplier_source": "ERP_GRN", "ven_cd": "NOPE"}, "is not an ERP vendor"),
    ({"part_id": None}, "Pick the part"),
])
def test_entry_is_validated(make, over, message):
    store = make.store()
    part = make.part(store)
    with pytest.raises(BadRequestError, match=message):
        svc.create(store, _entry(part, **over))


def test_supplier_from_erp_and_master_never_creates_a_master_record(make):
    store = make.store()
    part = make.part(store)
    _grn_line(part, "OLD-GRN", ven_cd="VEN-77", accept=1, grn_date=date(2026, 9, 1))
    masters = Supplier.query.count()
    erp = svc.create(store, _entry(part, supplier_source="ERP_GRN", ven_cd=" ven-77 ", supplier_name=None))
    assert (erp["ven_cd"], erp["supplier_name"]) == ("VEN-77", "ERP Vendor")
    sup = make.supplier()
    master = svc.create(store, _entry(part, supplier_source="SUPPLIER_MASTER", supplier_id=sup.id,
                                      supplier_name=None))
    assert (master["supplier_id"], master["supplier_name"]) == (sup.id, sup.supplier_name)
    assert Supplier.query.count() == masters + 1                             # only the one this test made
    with pytest.raises(BadRequestError, match="Part .* is not in this store"):
        svc.create(make.store(), _entry(part))


def test_the_same_challan_line_is_refused_until_cancelled(make):
    store = make.store()
    part, other = make.part(store), make.part(store)
    first = svc.create(store, _entry(part, challan_no="DC 501"))
    with pytest.raises(ConflictError, match=f"already registered as {first['challan_id']}"):
        svc.create(store, _entry(part, challan_no="dc501"))                 # spacing / case aside
    svc.create(store, _entry(other, challan_no="DC 501"))                   # another part on the same challan
    svc.create(store, _entry(part, challan_no="DC 502"))                    # the next lot: its own challan
    svc.create(store, _entry(part, challan_no="DC 501", supplier_name="Another Trader"))
    svc.cancel(first["id"], {"reason": "Entered twice"})
    again = svc.create(store, _entry(part, challan_no="DC 501"))
    assert again["challan_id"] != first["challan_id"]
    assert svc.detail(first["id"])["status"] == "Cancelled"                  # kept for the audit


def test_a_repeated_submit_returns_the_first_record(make):
    store = make.store()
    part = make.part(store)
    one = svc.create(store, _entry(part, request_id="form-1"))
    two = svc.create(store, _entry(part, request_id="form-1", challan_no="ANOTHER"))
    assert two["id"] == one["id"] and two["replayed"] is True
    assert ChallanRegister.query.filter_by(store_code=store).count() == 1


# ---- the invoice ---------------------------------------------------------------------------------

def test_partial_invoices_reconcile_and_over_invoicing_is_refused(make):
    store = make.store()
    part = make.part(store, qty=3)
    before = _stock(part)
    row = svc.create(store, _entry(part, challan_qty=5))
    one = svc.link_invoice(row["id"], {"invoice_no": "INV-1", "invoice_date": "2026-10-21", **PO, "invoice_qty": 3})
    assert (one["status"], one["invoice_status"], one["invoiced_qty"], one["remaining_to_invoice"]) == (
        "Pending Invoice", "Pending", 3, 2)
    with pytest.raises(BadRequestError, match="Over-invoicing: .* has 2 NOS left to invoice, not 4"):
        svc.link_invoice(row["id"], {"invoice_no": "INV-2", "invoice_date": "2026-10-21", **PO, "invoice_qty": 4})
    with pytest.raises(ConflictError, match="Invoice INV-1 is already linked"):
        svc.link_invoice(row["id"], {"invoice_no": " inv-1", "invoice_date": "2026-10-21", **PO, "invoice_qty": 1})
    two = svc.link_invoice(row["id"], {"invoice_no": "INV-2", "invoice_date": "2026-10-21", **PO})   # the rest
    assert (two["status"], two["invoice_status"], two["invoiced_qty"], two["invoice_no"]) == (
        "Invoice Received", "Received", 5, "INV-1, INV-2")
    with pytest.raises(BadRequestError, match="already fully invoiced"):
        svc.link_invoice(row["id"], {"invoice_no": "INV-3", "invoice_date": "2026-10-21", **PO})
    assert _stock(part) == before


def test_one_invoice_may_cover_several_challans_and_repeats_are_idempotent(make):
    store = make.store()
    a, b = make.part(store), make.part(store)
    ca = svc.create(store, _entry(a, challan_no="DC-A"))
    cb = svc.create(store, _entry(b, challan_no="DC-B"))
    svc.link_invoice(ca["id"], {"invoice_no": "INV-9", "invoice_date": "2026-10-21", **PO})
    out = svc.link_invoice(cb["id"], {"invoice_no": "INV-9", "invoice_date": "2026-10-21", **PO, "request_id": "r-1"})
    assert out["invoice_also_on"] == [ca["challan_id"]]
    again = svc.link_invoice(cb["id"], {"invoice_no": "INV-9", "invoice_date": "2026-10-21", **PO, "request_id": "r-1"})
    assert again["replayed"] is True and ChallanLink.query.filter_by(challan_ref=cb["id"]).count() == 1


def test_the_po_is_only_ever_what_the_person_sets(make):
    store = make.store()
    part = make.part(store)
    row = svc.create(store, _entry(part))
    assert row["po_status"] == "Pending Approval"                           # never approved by itself
    # A supplier invoices only against a PO: the invoice brings its PO no.
    with pytest.raises(BadRequestError, match="An invoice comes only against a PO"):
        svc.link_invoice(row["id"], {"invoice_no": "INV-P", "invoice_date": "2026-10-21"})
    with pytest.raises(BadRequestError, match="PO reference is required"):
        svc.link_invoice(row["id"], {"invoice_no": "INV-P", "invoice_date": "2026-10-21", "po_status": "Approved"})
    assert ChallanLink.query.filter_by(challan_ref=row["id"]).count() == 0
    out = svc.link_invoice(row["id"], {"invoice_no": "INV-P", "invoice_date": "2026-10-21",
                                       "po_status": "Approved", "po_reference": "PO-77"})
    assert (out["po_status"], out["po_reference"]) == ("Approved", "PO-77")
    assert [h["action"] for h in out["history"]][-2:] == ["PO_UPDATED", "INVOICE_LINKED"]
    with pytest.raises(BadRequestError, match="An invoice comes only against a PO"):
        svc.update(row["id"], {"po_status": "Pending Approval"})            # never back once invoiced
    out = svc.update(row["id"], {"po_reference": "PO-78"})
    assert out["po_reference"] == "PO-78" and out["history"][-1]["action"] == "UPDATED"
    with pytest.raises(BadRequestError, match="cannot change once an invoice is linked"):
        svc.update(row["id"], {"challan_qty": 9})
    # A line that needs no PO (Not Applicable) still takes its invoice.
    free = svc.create(store, _entry(part, challan_no="FOC-1", po_status="Not Applicable"))
    assert svc.link_invoice(free["id"], {"invoice_no": "INV-F", "invoice_date": "2026-10-21"})["po_status"] ==         "Not Applicable"


# ---- the GRN -------------------------------------------------------------------------------------

def test_grn_completed_only_once_erp_quality_accepts_and_stock_is_never_posted(make):
    store = make.store()
    part = make.part(store, qty=2)
    before = _stock(part)
    row = svc.create(store, _entry(part, challan_no="DC-77"))
    inv = svc.link_invoice(row["id"], {"invoice_no": "INV-7", "invoice_date": "2026-10-21", **PO})
    link_id = inv["links"][0]["id"]
    rr = _grn_line(part, "GRN-7", challan_no="DC-77")
    cands = svc.grn_candidates(row["id"])
    assert cands[0]["id"] == rr.id and cands[0]["matches_challan_no"] is True
    out = svc.link_grn(row["id"], link_id, {"receiving_register_id": rr.id})
    assert (out["status"], out["quality_status"], out["grn_reference"]) == ("GRN Pending", "Pending", "GRN-7")
    rr.accept_qty = 5
    db.session.commit()
    assert svc.refresh_from_erp(store)["changed"] == 1
    done = svc.detail(row["id"])
    assert (done["status"], done["quality_status"]) == ("GRN Completed", "OK") and done["completed_at"]
    assert [h["action"] for h in done["history"]] == ["CREATED", "PO_UPDATED", "INVOICE_LINKED", "GRN_LINKED",
                                                      "GRN_REFRESHED"]
    assert _stock(part) == before                                            # the register posted nothing
    with pytest.raises(BadRequestError, match="GRN is completed cannot be cancelled"):
        svc.cancel(row["id"], {"reason": "x"})


def test_a_grn_line_is_linked_to_one_challan_only_and_must_match(make):
    store = make.store()
    part, other = make.part(store), make.part(store)
    c1 = svc.create(store, _entry(part, challan_no="DC-1"))
    c2 = svc.create(store, _entry(part, challan_no="DC-2"))
    l1 = svc.link_invoice(c1["id"], {"invoice_no": "I-1", "invoice_date": "2026-10-21", **PO})["links"][0]["id"]
    l2 = svc.link_invoice(c2["id"], {"invoice_no": "I-2", "invoice_date": "2026-10-21", **PO})["links"][0]["id"]
    rr = _grn_line(part, "GRN-1")
    svc.link_grn(c1["id"], l1, {"receiving_register_id": rr.id})
    with pytest.raises(ConflictError, match=f"already linked to {c1['challan_id']}"):
        svc.link_grn(c2["id"], l2, {"receiving_register_id": rr.id})
    assert all(c["id"] != rr.id for c in svc.grn_candidates(c2["id"]))      # not offered either
    wrong = _grn_line(other, "GRN-X")
    with pytest.raises(BadRequestError, match="is for item"):
        svc.link_grn(c2["id"], l2, {"receiving_register_id": wrong.id})
    with pytest.raises(NotFoundError):
        svc.link_grn(c2["id"], l1, {"grn_no": "G"})                         # another line's invoice


def test_vendor_must_match_an_erp_vendor_challan(make):
    store = make.store()
    part = make.part(store)
    _grn_line(part, "OLD", ven_cd="VEN-A", accept=1, grn_date=date(2026, 9, 1))
    row = svc.create(store, _entry(part, supplier_source="ERP_GRN", ven_cd="VEN-A"))
    link = svc.link_invoice(row["id"], {"invoice_no": "I", "invoice_date": "2026-10-21", **PO})["links"][0]["id"]
    rr = _grn_line(part, "GRN-B", ven_cd="VEN-B")
    with pytest.raises(BadRequestError, match="is from vendor VEN-B"):
        svc.link_grn(row["id"], link, {"receiving_register_id": rr.id})


def test_a_typed_grn_waits_for_its_erp_line_and_a_rejection_keeps_it_open(make):
    store = make.store()
    part = make.part(store)
    row = svc.create(store, _entry(part))
    link = svc.link_invoice(row["id"], {"invoice_no": "I-T", "invoice_date": "2026-10-21", **PO})["links"][0]["id"]
    out = svc.link_grn(row["id"], link, {"grn_no": "GRN 55", "grn_date": "2026-10-21"})
    assert (out["status"], out["links"][0]["receiving_register_id"]) == ("GRN Pending", None)
    assert svc.refresh_from_erp(store)["changed"] == 0                      # ERP has no such line yet
    rr = _grn_line(part, "GRN55", reject=5)                                 # it arrives: all rejected
    assert svc.refresh_from_erp(store)["changed"] == 1
    out = svc.detail(row["id"])
    assert (out["status"], out["quality_status"], out["links"][0]["receiving_register_id"]) == (
        "GRN Pending", "Rejected", rr.id)
    tiles = svc.dashboard(store)["tiles"]
    assert (tiles["quality_rejected"], tiles["invoice_received_grn_pending"]) == (1, 1)


# ---- cancel, dashboard, list, export ----------------------------------------------------------

def test_cancel_needs_a_reason_and_keeps_the_record(make):
    store = make.store()
    part = make.part(store)
    row = svc.create(store, _entry(part))
    with pytest.raises(BadRequestError, match="Cancel reason is required"):
        svc.cancel(row["id"], {})
    out = svc.cancel(row["id"], {"reason": "Supplier took it back"}, user="sup")
    assert (out["status"], out["cancel_reason"], out["cancelled_by"]) == ("Cancelled", "Supplier took it back", "sup")
    with pytest.raises(BadRequestError, match="cancelled challan cannot take an invoice"):
        svc.link_invoice(row["id"], {"invoice_no": "I", "invoice_date": "2026-10-21", **PO})
    assert ChallanHistory.query.filter_by(challan_ref=row["id"]).count() == 2


def test_dashboard_ageing_list_filters_and_export(make):
    store = make.store()
    part = make.part(store)
    old = svc.create(store, _entry(part, challan_no="OLD", challan_date="2026-09-01", receipt_date="2026-09-02"))
    svc.create(store, _entry(part, challan_no="NEW", po_status="Not Applicable"))
    gone = svc.create(store, _entry(part, challan_no="GONE"))
    svc.cancel(gone["id"], {"reason": "wrong"})
    d = svc.dashboard(store)
    assert d["tiles"]["total"] == 3 and d["tiles"]["pending_invoice"] == 2 and d["tiles"]["cancelled"] == 1
    assert d["tiles"]["po_approval_pending"] == 1
    assert {a["label"]: a["count"] for a in d["ageing"]} == {
        "0-7 days": 1, "8-15 days": 0, "16-30 days": 0, "Over 30 days": 1}
    assert d["oldest_open"]["challan_id"] == old["challan_id"] and d["oldest_open"]["days"] == 49
    page = svc.list(store, {"status": "Pending Invoice"}, page=1, page_size=1)
    assert (page["total"], page["pages"], len(page["rows"])) == (2, 2, 1)
    assert svc.list(store, {"q": "old"})["rows"][0]["challan_id"] == old["challan_id"]
    assert svc.list(store, {"date_from": "2026-10-01"})["total"] == 2
    with pytest.raises(BadRequestError, match="Status must be one of"):
        svc.list(store, {"status": "Lost"})
    # A tile opens exactly what it counts: a status may be a comma list.
    open_only = {"po_status": "Pending Approval", "status": "Pending Invoice, Invoice Received,GRN Pending"}
    assert svc.list(store, open_only)["total"] == d["tiles"]["po_approval_pending"] == 1
    assert svc.list(store, {"po_status": "Pending Approval"})["total"] == 2          # the cancelled one too
    with pytest.raises(BadRequestError, match="Status must be one of"):
        svc.list(store, {"status": "Pending Invoice,Lost"})
    buf, name = svc.export(store, {})
    rows = list(load_workbook(io.BytesIO(buf.getvalue())).active.iter_rows(values_only=True))
    assert rows[0][0] == "Challan ID" and len(rows) == 4 and name.startswith("Urgent Challan Register")


# ---- the API and its rights --------------------------------------------------------------------

def test_the_right_exists():
    assert "challan_register" in {m["key"] for m in PERMISSION_MODULES}


def test_api_rights_store_access_and_replay(client, make):
    store = make.store()
    part = make.part(store)
    body = dict(_entry(part), store_code=store, request_id="api-1")
    viewer = make.user(stores=(store,), perms={"challan_register": {"view": True}})
    assert _as(client, viewer).post(URL, json=body).status_code == 403          # view only
    assert _as(client, make.user(stores=(store,))).get(f"{URL}?store_code={store}").status_code == 403
    clerk = make.user(stores=(store,), perms={"challan_register": {"view": True, "edit": True}})
    c = _as(client, clerk)
    r = c.post(URL, json=body)
    assert r.status_code == 201, r.get_json()
    pk = r.get_json()["data"]["id"]
    again = c.post(URL, json=body)
    assert again.status_code == 200 and again.get_json()["data"]["id"] == pk
    assert c.get(f"{URL}/{pk}").get_json()["data"]["created_by"] == clerk.username
    r = c.post(f"{URL}/{pk}/invoices", json={"invoice_no": "INV-API", "invoice_date": "2026-10-21", **PO})
    assert r.status_code == 200 and r.get_json()["data"]["status"] == "Invoice Received"
    assert c.get(f"{URL}/dashboard?store_code={store}").get_json()["data"]["tiles"]["total"] == 1
    assert c.get(f"{URL}/export?store_code={store}").status_code == 200
    assert c.get(f"{URL}/parts?store_code={store}&q={part.part_no}").get_json()["data"][0]["id"] == part.id
    assert c.post(URL, json=[1]).status_code == 400
    other_store = make.store()
    outsider = _as(client, make.user(stores=(other_store,), perms={"challan_register": {"view": True,
                                                                                         "edit": True}}))
    assert outsider.get(f"{URL}/{pk}").status_code == 403
    assert outsider.get(f"{URL}?store_code={store}").status_code == 403


def test_an_erp_sync_rewriting_the_grn_register_keeps_every_link(make):
    """The ERP sync deletes and re-inserts the GRN register on every run, so its
    line ids change: a link names its GRN by GRN no. + Part No., the rewrite is
    never blocked by it, and the next read finds the GRN again — two lines of
    the same GRN summed."""
    from app.services.receiving_register_service import ReceivingRegisterService
    store = make.store()
    part = make.part(store)
    feed = f"CRR{_TAG}"                                                 # this test's own ERP snapshot
    row = svc.create(store, _entry(part))
    link = svc.link_invoice(row["id"], {"invoice_no": "I-S", "invoice_date": "2026-10-21", **PO})["links"][0]["id"]
    line = {"grn_no": "GRN-S", "grn_date": TODAY, "ven_cd": "VEN-1", "ven_name": "ERP Vendor",
            "item_cd": part.part_no, "item_desc": part.part_name, "receipt_qty": 3}
    ReceivingRegisterService().replace_from_rows(feed, [line, dict(line, receipt_qty=2)])
    first = ReceivingRegister.query.filter_by(store_code=feed).order_by(ReceivingRegister.id).first()
    out = svc.link_grn(row["id"], link, {"receiving_register_id": first.id})
    assert (out["status"], out["links"][0]["grn_receipt_qty"]) == ("GRN Pending", 5)
    # The next sync: the same GRN, now accepted — every line id new.
    ReceivingRegisterService().replace_from_rows(feed, [dict(line, accept_qty=3), dict(line, receipt_qty=2,
                                                                                       accept_qty=2)])
    assert ReceivingRegister.query.filter_by(store_code=feed).count() == 2     # the rewrite went through
    assert svc.refresh_from_erp(store)["changed"] == 1
    done = svc.detail(row["id"])
    assert (done["status"], done["quality_status"], done["links"][0]["grn_accepted_qty"]) == ("GRN Completed", "OK", 5)
    ReceivingRegisterService().replace_from_rows(feed, [])                # leave no ERP rows behind


def test_a_cancelled_challan_gives_its_grn_up(make):
    store = make.store()
    part = make.part(store)
    wrong = svc.create(store, _entry(part, challan_no="DC-W"))
    right = svc.create(store, _entry(part, challan_no="DC-R"))
    lw = svc.link_invoice(wrong["id"], {"invoice_no": "I-W", "invoice_date": "2026-10-21", **PO})["links"][0]["id"]
    lr = svc.link_invoice(right["id"], {"invoice_no": "I-R", "invoice_date": "2026-10-21", **PO})["links"][0]["id"]
    svc.link_grn(wrong["id"], lw, {"grn_no": "GRN-C"})
    with pytest.raises(ConflictError, match=f"already linked to {wrong['challan_id']}"):
        svc.link_grn(right["id"], lr, {"grn_no": "grn-c"})
    out = svc.cancel(wrong["id"], {"reason": "GRN belongs to DC-R"})
    assert out["history"][-1]["details"]["grns_freed"] == ["GRN-C"] and out["links"][0]["grn_reference"] == "GRN-C"
    assert svc.link_grn(right["id"], lr, {"grn_no": "GRN-C"})["grn_reference"] == "GRN-C"
