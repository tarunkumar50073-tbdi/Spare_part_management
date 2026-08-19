"""The 'VEHICLE WINDOW TIME' monthly sheet builder — a pure function over
master rows + per-day statuses, so no app/DB context is needed here."""
from datetime import time

import openpyxl

from app.utils.window_time_report import build_window_time_report

MASTER = [
    {"vendor_code": "V1", "supplier_name": "SAR Coating", "window_start": time(8, 30),
     "window_end": time(10, 0), "frequency": "Daily"},
    {"vendor_code": "V2", "supplier_name": "NSSI", "window_start": time(11, 30),
     "window_end": time(12, 30), "frequency": "Alternate Day"},
]


def _status(vendor, name, start, end, status, arrival=None, out=None):
    return {"vendor_code": vendor, "supplier_name": name,
            "window_start": start.isoformat(), "window_end": end.isoformat(),
            "actual_arrival": arrival, "actual_out": out, "status": status}


def _load(buf):
    return openpyxl.load_workbook(buf).active


def test_layout_columns_and_title():
    # 30-day month -> 5 fixed + 30*3 day columns
    ws = _load(build_window_time_report(2026, 9, MASTER, {}))
    assert ws.max_column == 5 + 30 * 3
    assert "VEHICLE WINDOW TIME" in ws["A1"].value
    assert "SEPTEMBER 2026" in ws["A1"].value
    # fixed headers + the first day's sub-headers
    assert ws.cell(2, 4).value == "Supplier Name"
    assert [ws.cell(3, c).value for c in (6, 7, 8)] == ["Arrival", "Departure", "Remarks"]
    # plan columns for the first slot
    assert ws.cell(4, 2).value == "08:30"
    assert ws.cell(4, 3).value == "10:00"
    assert ws.cell(4, 4).value == "SAR Coating"


def test_status_words_and_colours():
    day1 = [
        _status("V1", "SAR Coating", time(8, 30), time(10, 0), "On-Time",
                "2026-08-01T08:36:00", "2026-08-01T09:55:00"),
        _status("V2", "NSSI", time(11, 30), time(12, 30), "Late",
                "2026-08-01T12:55:00", "2026-08-01T13:40:00"),
    ]
    day2 = [_status("V1", "SAR Coating", time(8, 30), time(10, 0), "Not Arrived")]
    ws = _load(build_window_time_report(2026, 8, MASTER, {1: day1, 2: day2}))
    # day 1 block starts at column 6 (Arrival/Departure/Remarks)
    assert ws.cell(4, 6).value == "08:36"          # actual arrival
    assert ws.cell(4, 7).value == "09:55"          # actual departure
    assert ws.cell(4, 8).value == "ON TIME"
    assert ws.cell(4, 8).fill.fgColor.rgb.endswith("C6EFCE")   # green
    assert ws.cell(5, 8).value == "LATE"
    assert ws.cell(5, 8).fill.fgColor.rgb.endswith("FFEB9C")   # amber
    # day 2 block starts at column 9; Not Arrived -> red, no actual times
    assert ws.cell(4, 11).value == "NOT COME"
    assert ws.cell(4, 11).fill.fgColor.rgb.endswith("FFC7CE")  # red
    assert ws.cell(4, 9).value is None


def test_pending_and_missing_left_blank():
    # a window still open (Pending Arrival) and a day with no status entry both
    # stay blank — same as the paper sheet before the truck comes
    day1 = [_status("V1", "SAR Coating", time(8, 30), time(10, 0), "Pending Arrival")]
    ws = _load(build_window_time_report(2026, 8, MASTER, {1: day1}))
    assert ws.cell(4, 8).value is None    # V1, pending -> blank remark
    assert ws.cell(5, 8).value is None    # V2, no entry -> blank remark
