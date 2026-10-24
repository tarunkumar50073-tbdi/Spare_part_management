"""The Reports page of a spare store (store complaint of 14 Oct 2026: "some
reports showed old modules and old departments").

GET /reports/catalog?store_code= lists the reports of THAT store: the nine
spare-store reports for MAINTENANCE, the thirteen old ones for CHILD_PART. Every
spare report downloads as Excel / CSV / PDF under "<Label> <DD.MM.YYYY>[ to
<DD.MM.YYYY>]"; a report the store does not offer is a 404. Dated reports cover
plant-clock days, departments and sections come from their masters.
Consumption is not among them: it lives on the Consumption screen
(test_consumption_table.py). The reorder and excess lists' Excel is coloured
like the reorder / excess mail sheets (15 Oct 2026); every other report's Excel
is the plain table.

Everything here is RPT-<tag>-<n>: this file's own stores, sections, departments,
people and users. Afterwards the stores, people and users are switched off and
the departments removed, so nothing leaks into the lists other tests read.
"""
import csv
import io
import uuid
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.lib.units import inch
from werkzeug.http import parse_options_header

from app.extensions import db
from app.models.base import utcnow
from app.models.department import Department
from app.models.employee import Employee
from app.models.part import Part
from app.models.put_away_log import PutAwayLog
from app.models.spare_flow import SpareReturn, SpareTransaction
from app.models.spare_stock import RETURN_ROW_INSPECTED, SpareStock, SpareStockLedger
from app.models.store import Section, Store
from app.models.user import User, UserStoreAccess
from app.services import report_service
from app.services.consumption_analytics_service import _local_offset
from app.services.email_rules import EXCESS_LEGEND
from app.services.reorder_mail_service import REORDER_LEGEND
from app.services.report_service import ReportService
from app.services.spare_report_service import SpareReportService, plant_today
from app.services.spare_stock_service import SpareStockService
from app.utils import pdf_io
from app.utils.excel_io import build_export

_TAG = uuid.uuid4().hex[:6].upper()
_seq = {"n": 0}

# The consumption cuts moved to the Consumption screen ("alag se report me aane
# ki jarurat nhi" — the store, 14 Oct 2026). The Kiosk Part Requests report went
# with the kiosk on 17 Oct 2026.
SPARE_KEYS = ["stock_current", "reorder_list", "excess_list", "issue_register",
              "installation_register", "return_register", "put_away_register",
              "receipt_register", "erp_vs_physical"]
DATED = {"issue_register", "installation_register", "return_register",
         "put_away_register", "receipt_register"}


def _n():
    _seq["n"] += 1
    return _seq["n"]


@pytest.fixture()
def make():
    made = {"stores": [], "people": [], "users": [], "departments": []}

    def store():
        """A spare store with an active and a switched-off section."""
        code = f"RPT{_TAG}{_n()}"
        s = Store(code=code, name=f"Report Store {code}")
        db.session.add(s)
        db.session.flush()
        live = Section(store_id=s.id, name=f"Live Sec {code}")
        old = Section(store_id=s.id, name=f"Old Sec {code}", is_active=False)
        db.session.add_all([live, old])
        db.session.commit()
        made["stores"].append(code)
        return type("S", (), {"code": code, "live": live.id, "old": old.id,
                              "live_name": live.name, "old_name": old.name})

    def department(active=True, name=None, sort_order=100):
        d = Department(name=name or f"RPT Dept {_TAG}-{_n()}", is_active=active, sort_order=sort_order)
        db.session.add(d)
        db.session.commit()
        made["departments"].append(d.id)
        return d

    def person():
        emp = Employee(mobile=f"7{uuid.uuid4().int % 10 ** 9:09d}", name=f"Fitter {_n()}",
                       employee_code=f"E-RPT-{_TAG}-{_n()}")
        db.session.add(emp)
        db.session.commit()
        made["people"].append(emp.id)
        return emp

    def user(stores=("MAINTENANCE",), perms=None, role="STORE_USER"):
        u = User(username=f"rpt{_TAG.lower()}{_n()}", role=role, is_active=True)
        u.set_password("pw")
        if perms:
            u.set_permissions(perms)
        db.session.add(u)
        db.session.flush()
        for code in stores:
            db.session.add(UserStoreAccess(user_id=u.id, store_code=code))
        db.session.commit()
        made["users"].append(u.id)
        return u

    yield type("Make", (), {"store": staticmethod(store), "department": staticmethod(department),
                            "person": staticmethod(person), "user": staticmethod(user)})
    db.session.rollback()
    Store.query.filter(Store.code.in_(made["stores"])).update({"is_active": False}, synchronize_session=False)
    Employee.query.filter(Employee.id.in_(made["people"])).update({"is_active": False}, synchronize_session=False)
    User.query.filter(User.id.in_(made["users"])).update({"is_active": False}, synchronize_session=False)
    Department.query.filter(Department.id.in_(made["departments"])).delete(synchronize_session=False)
    db.session.commit()


@pytest.fixture()
def client(app):
    return app.test_client()


def _as(client, user):
    with client.session_transaction() as sess:
        sess["user_id"] = user.id
    return client


def _part(store, section=None, qty=None, **levels):
    n = _n()
    p = Part(part_no=f"RPT-{_TAG}-{n}", part_name=f"Report Part {n}", uom="NOS", part_type="OTHER",
             store_code=store.code, section_id=section, **levels)
    db.session.add(p)
    db.session.commit()
    if qty is not None:
        SpareStockService().set_quantity(p.id, qty, txn_type="OPENING")
    return p


def _issue(store, department, qty=3, stock=10, section=None, at=None, **kw):
    """A part with `stock` on the shelf and one ISSUE of `qty` -> (part, issue id),
    the issue stamped `at` (a naive UTC instant) when given."""
    p = _part(store, section, qty=stock)
    lid = SpareStockService().issue(p.id, qty, department.name, **kw)["issue"]["id"]
    if at is not None:
        SpareStockLedger.query.filter_by(id=lid).update({"created_at": at})
        db.session.commit()
    return p, lid


def _utc(day, hour=10, minute=0):
    """The UTC instant of a plant-clock time."""
    return datetime.combine(day, time(hour, minute)) - _local_offset()


def _rows(store, key, date_from=None, date_to=None):
    return SpareReportService().build(store.code, key, date_from, date_to)[2]


def _return(store, part, issue_id, at, **split):
    qty = sum(split.values())
    r = SpareReturn(return_no=f"RET-RPT-{_TAG}-{_n()}", part_id=part.id, store_code=store.code,
                    issue_id=issue_id, qty=qty, reason="Did not fit", returned_by="Ravi",
                    received_by="store.user", department="Press Shop", created_at=at, **split)
    r.refresh_status()
    db.session.add(r)
    db.session.commit()
    return r


def _filename(resp):
    return parse_options_header(resp.headers["Content-Disposition"])[1]["filename"]


# ---- the catalogue ------------------------------------------------------------------

def test_catalogue_lists_the_stores_own_reports(client, make):
    c = _as(client, make.user(role="ADMIN", stores=()))
    spare = c.get("/api/v1/reports/catalog?store_code=MAINTENANCE")
    assert spare.status_code == 200
    entries = spare.get_json()["data"]
    assert [e["key"] for e in entries] == SPARE_KEYS
    for e in entries:
        assert set(e) == {"key", "label", "description", "needs_dates", "formats"}
        assert e["label"] and e["description"] and "\n" not in e["description"]
        assert e["needs_dates"] is (e["key"] in DATED)
        assert e["formats"] == ["excel", "csv", "pdf"]
    assert entries[1]["label"] == "Reorder / Critical List"

    child = c.get("/api/v1/reports/catalog?store_code=CHILD_PART").get_json()["data"]
    assert [e["key"] for e in child] == list(ReportService.REPORTS)
    assert [e["label"] for e in child] == [t for t, _, _ in ReportService.REPORTS.values()]
    assert not any(e["needs_dates"] for e in child)

    assert c.get("/api/v1/reports/catalog").status_code == 400
    assert c.get(f"/api/v1/reports/catalog?store_code=NOPE{_TAG}").status_code == 404


def test_catalogue_and_downloads_are_gated(client, make):
    no_right = _as(client, make.user())
    assert no_right.get("/api/v1/reports/catalog?store_code=MAINTENANCE").status_code == 403

    other = make.store()
    elsewhere = _as(client, make.user(stores=(other.code,), perms={"reports": {"view": True}}))
    assert elsewhere.get("/api/v1/reports/catalog?store_code=MAINTENANCE").status_code == 403
    assert elsewhere.get("/api/v1/reports/stock_current/excel?store_code=MAINTENANCE").status_code == 403
    assert elsewhere.get(f"/api/v1/reports/catalog?store_code={other.code}").status_code == 200

    granted = _as(client, make.user(perms={"reports": {"view": True}}))
    assert granted.get("/api/v1/reports/catalog?store_code=MAINTENANCE").status_code == 200
    assert granted.get("/api/v1/reports/excess_list/csv?store_code=MAINTENANCE").status_code == 200


# ---- every report, every format -------------------------------------------------------

def test_every_spare_report_downloads_in_every_format(client, make):
    c = _as(client, make.user(role="ADMIN", stores=()))
    labels = {e.key: e for e in SpareReportService.REPORTS}
    today = f"{plant_today():%d.%m.%Y}"
    for key in SPARE_KEYS:
        report = labels[key]
        dates = "&date_from=2026-10-01&date_to=2026-10-14" if key in DATED else ""
        stem = f"{report.label} 01.10.2026 to 14.10.2026" if key in DATED else f"{report.label} {today}"
        stem = stem.replace(" / ", " ")  # "/" cannot be in a file name
        for fmt, ext in (("excel", "xlsx"), ("csv", "csv"), ("pdf", "pdf")):
            resp = c.get(f"/api/v1/reports/{key}/{fmt}?store_code=MAINTENANCE{dates}")
            assert resp.status_code == 200, (key, fmt, resp.get_json())
            assert _filename(resp) == f"{stem}.{ext}"
            data = resp.get_data()
            if fmt == "excel":
                ws = load_workbook(io.BytesIO(data)).active
                assert [cell.value for cell in ws[1]] == list(report.columns), key
            elif fmt == "csv":
                header = next(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
                assert header == list(report.columns), key
            else:
                assert data[:4] == b"%PDF"


def _drawn_pdf_tables(monkeypatch):
    """[(document, table)] of every PDF built from here on, as handed to reportlab."""
    drawn = []
    real_build = pdf_io.SimpleDocTemplate.build

    def build(doc, story, *args, **kwargs):
        drawn.append((doc, story[-1]))
        return real_build(doc, story, *args, **kwargs)

    monkeypatch.setattr(pdf_io.SimpleDocTemplate, "build", build)
    return drawn


def test_a_wide_pdf_report_is_fitted_to_the_page(client, make, monkeypatch):
    """Current Stock has fifteen columns. Drawn at its natural width it was centred
    and ran off both edges of the page (Section to Description on one side,
    Location to Stock Value on the other). Fitted, the whole table is on the
    page, no line runs past its column and no character is lost; a table that
    fits is drawn as before."""
    s = make.store()
    words = "HYDRAULIC CYLINDER SEAL KIT WITH ROD WIPER, GUIDE RINGS AND O-RINGS FOR THE PRESS LINE"
    p = _part(s, s.live, qty=5, min_stock_level=2, reorder_level=3, max_stock_level=4, unit_price=1234.5,
              location="R1-S1-B1&2&3")
    p.part_name = words
    # One word wider than any column can be.
    p.specification = "SFAU16-192-F29-B15-B8-T29-S15-Q8-SC91-" * 3 + " " + words
    db.session.commit()
    _part(s, s.old, qty=0)
    drawn = _drawn_pdf_tables(monkeypatch)
    c = _as(client, make.user(role="ADMIN", stores=()))

    resp = c.get(f"/api/v1/reports/stock_current/pdf?store_code={s.code}")
    assert resp.status_code == 200 and resp.get_data()[:4] == b"%PDF"
    doc, table = drawn[-1]
    frame = doc.pageTemplates[0].frames[0]
    assert table.wrap(frame._aW, frame._aH)[0] <= frame._aW + 1e-6
    for i, row in enumerate(table._cellvalues):
        for j, cell in enumerate(row):
            style = table._cellStyles[i][j]
            room = table._colWidths[j] - style.leftPadding - style.rightPadding
            assert all(stringWidth(line, style.fontname, style.fontsize) <= room + 1e-6
                       for line in cell.split("\n")), (i, j, cell)
    report, _, rows = SpareReportService().build(s.code, "stock_current")
    expected = [list(report.columns)] + [["" if r[col] is None else str(r[col]) for col in report.columns]
                                         for r in rows]
    def squeeze(cells):  # the characters, line breaks and spaces aside
        return [["".join(cell.split()) for cell in row] for row in cells]

    assert squeeze(table._cellvalues) == squeeze(expected)
    # A long description breaks at its spaces, never inside a word.
    description = table._cellvalues[1][report.columns.index("Description")]
    assert "\n" in description and description.replace("\n", " ") == words

    narrow = c.get(f"/api/v1/reports/put_away_register/pdf?store_code={s.code}")
    assert narrow.status_code == 200
    doc, table = drawn[-1]
    assert doc.leftMargin == doc.rightMargin == inch
    assert not any("\n" in cell for row in table._cellvalues for cell in row)
    assert {style.fontsize for row in table._cellStyles for style in row} == {8}


def test_a_pdf_table_a_little_wider_than_the_page_moves_onto_narrow_margins(monkeypatch):
    """A table wider than the inch margins leave but narrower than the narrow
    margins' room was measured against the narrow room only, so it stayed on
    the inch page and ran off both edges (the Part-wise consumption table,
    short descriptions). Measured against the page it is drawn on, it moves onto the
    narrow margins, type unchanged and nothing broken."""
    from reportlab.platypus import Table

    inch_room = pdf_io._room(pdf_io._document(io.BytesIO(), inch))
    narrow_room = pdf_io._room(pdf_io._document(io.BytesIO(), pdf_io.FIT_MARGIN))
    word = "MMMMMMMM"

    def natural(count):
        table = Table([[word] * count] * 2)
        table.setStyle(pdf_io.STYLE)
        return table.wrap(narrow_room, 1000)[0]

    count = next(n for n in range(1, 40) if natural(n) > inch_room)
    assert natural(count) < narrow_room                     # the case under test
    drawn = _drawn_pdf_tables(monkeypatch)
    columns = [(word, f"c{i}") for i in range(count)]
    buf = pdf_io.build_pdf_table("Between the margins", columns, [{key: word for _, key in columns}])
    assert buf.getvalue()[:4] == b"%PDF"
    doc, table = drawn[-1]
    frame = doc.pageTemplates[0].frames[0]
    assert doc.leftMargin == doc.rightMargin == pdf_io.FIT_MARGIN
    assert table.wrap(frame._aW, frame._aH)[0] <= frame._aW + 1e-6
    assert {style.fontsize for row in table._cellStyles for style in row} == {8}
    assert [list(row) for row in table._cellvalues] == [[word] * count] * 2


def test_dated_report_defaults_to_this_month(client, make):
    c = _as(client, make.user(role="ADMIN", stores=()))
    resp = c.get("/api/v1/reports/issue_register/excel?store_code=MAINTENANCE")
    today = plant_today()
    first = today.replace(day=1)
    expected = (f"Issue Register {today:%d.%m.%Y}.xlsx" if first == today
                else f"Issue Register {first:%d.%m.%Y} to {today:%d.%m.%Y}.xlsx")
    assert resp.status_code == 200 and _filename(resp) == expected


def test_a_report_the_store_does_not_offer_is_404(client, make):
    c = _as(client, make.user(role="ADMIN", stores=()))
    assert c.get("/api/v1/reports/srv/excel?store_code=MAINTENANCE").status_code == 404
    assert c.get("/api/v1/reports/issue_register/excel?store_code=CHILD_PART").status_code == 404
    assert c.get("/api/v1/reports/issue_register/excel").status_code == 404
    assert c.get("/api/v1/reports/stock_current/docx?store_code=MAINTENANCE").status_code == 404
    assert c.get(f"/api/v1/reports/stock_current/excel?store_code=NOPE{_TAG}").status_code == 404
    # The old reports keep working without a store code.
    old = c.get("/api/v1/reports/srv/csv")
    assert old.status_code == 200 and _filename(old) == f"SRV Report {plant_today():%d.%m.%Y}.csv"
    # A bad or reversed range is refused, never guessed.
    bad = "/api/v1/reports/issue_register/excel?store_code=MAINTENANCE&date_from="
    assert c.get(bad + "14-10-2026").status_code == 400
    assert c.get(bad + "2026-10-14&date_to=2026-10-01").status_code == 400


# ---- dates are plant-clock days ---------------------------------------------------------

def test_issue_register_filters_by_plant_day(make):
    s = make.store()
    dept = make.department()
    day = date(2025, 3, 10)
    # 00:30 on the plant clock: still the 9th in UTC wherever the plant is ahead of UTC.
    _, early = _issue(s, dept, at=_utc(day, 0, 30))
    _, before = _issue(s, dept, at=_utc(day - timedelta(days=5)))
    _, after = _issue(s, dept, at=_utc(day + timedelta(days=1), 0, 5))

    rows = _rows(s, "issue_register", "2025-03-10", "2025-03-10")
    assert [r["Issue No"] for r in rows] == [early]
    assert rows[0]["Date & Time"] == "10.03.2025 00:30"
    assert _rows(s, "issue_register", "2025-03-09", "2025-03-09") == []
    span = _rows(s, "issue_register", "2025-03-01", "2025-03-11")
    assert [r["Issue No"] for r in span] == [before, early, after]   # oldest first


# ---- issue register: status words and what is still out --------------------------------

def test_issue_register_status_words(make):
    s = make.store()
    dept = make.department()
    emp = make.person()
    svc = SpareStockService()
    _, waiting = _issue(s, dept, employee_id=emp.id, reference=f"REQ-{_TAG}-1")
    _, none_used = _issue(s, dept)
    # A confirmation with nothing used (refused today, but in the history): PARTIALLY_USED, 0 used.
    SpareStockLedger.query.filter_by(id=none_used).update({
        "status": "PARTIALLY_USED", "actual_qty": 0, "line_name": "Seat Slider", "line_no": "03",
        "machine": "MC-1", "confirmed_by": "fitter", "confirmed_at": utcnow()})
    db.session.commit()
    _, partly = _issue(s, dept)
    svc.confirm_installation(partly, machine="MC-2", actual_qty=1, user="fitter")
    _, full = _issue(s, dept)
    svc.confirm_installation(full, machine="MC-3", die_no="D-7", user="fitter")
    p, back = _issue(s, dept)
    _return(s, p, back, db.session.get(SpareStockLedger, back).created_at, qty_good=1)
    SpareStockLedger.query.filter_by(id=back).update({"status": "PARTIALLY_RETURNED"})
    db.session.commit()

    rows = {r["Issue No"]: r for r in _rows(s, "issue_register")}
    assert rows[waiting]["Status"] == "Waiting for installation"
    assert (rows[waiting]["Emp Code"], rows[waiting]["Mobile"], rows[waiting]["Issued To"]) == (
        emp.employee_code, emp.mobile, emp.name)
    assert rows[waiting]["Reference"] == f"REQ-{_TAG}-1"
    assert rows[waiting]["Used Qty"] is None and rows[waiting]["Still Out"] == 3
    assert rows[none_used]["Status"] == "Not used yet" and rows[none_used]["Used Qty"] == 0
    assert rows[none_used]["Line"] == "Seat Slider 03" and rows[none_used]["Machine"] == "MC-1"
    assert rows[partly]["Status"] == "Partly used" and rows[partly]["Used Qty"] == 1
    assert rows[full]["Status"] == "Installed" and rows[full]["Used Qty"] == 3
    assert rows[full]["Die"] == "D-7"
    assert rows[back]["Status"] == "Partly returned"
    assert (rows[back]["Qty"], rows[back]["Returned Qty"], rows[back]["Still Out"]) == (3, 1, 2)
    # Department is who took it, as typed; Consumer the owner of the part's section.
    assert (rows[back]["Department"], rows[back]["Consumer"]) == (dept.name, "No section")
    columns = next(r.columns for r in SpareReportService.REPORTS if r.key == "issue_register")
    assert columns[columns.index("Department") + 1] == "Consumer"


# ---- installation register: one row per journal entry -----------------------------------

def test_installation_register_reads_the_journal(make):
    s = make.store()
    dept = make.department()
    svc = SpareStockService()
    p, lid = _issue(s, dept, issued_to="Ravi")
    svc.confirm_installation(lid, line_name="Recliner", line_no="02", machine="MC-1", user="fitter.one")
    svc.confirm_installation(lid, machine="MC-2", die_no="D-2", actual_qty=2, user="fitter.two")

    rows = [r for r in _rows(s, "installation_register") if r["Issue No"] == lid]
    assert [(r["Entry"], r["Machine"], r["Used Qty"], r["Recorded By"]) for r in rows] == [
        ("Installation", "MC-1", 3, "fitter.one"), ("Update", "MC-2", 2, "fitter.two")]
    assert all(r["Issued Qty"] == 3 and r["Part No"] == p.part_no and r["Issued To"] == "Ravi"
               and r["Line"] == "Recliner 02" and r["Department"] == dept.name for r in rows)
    assert rows[1]["Die"] == "D-2"
    assert len(rows) == SpareTransaction.query.filter_by(issue_id=lid, txn_type="INSTALLATION").count()


# ---- return register ---------------------------------------------------------------------

def test_return_register_says_where_the_quantity_went(make):
    s = make.store()
    dept = make.department()
    day = date(2025, 5, 6)
    p, lid = _issue(s, dept, at=_utc(day, 9))
    split = _return(s, p, lid, _utc(day, 11), qty_good=1, qty_damaged=1)
    good = _return(s, p, lid, _utc(day, 12), qty_good=1)
    # An older direct return straight to the shelf, and the RETURN row a good
    # inspection writes (its quantity is already on its return).
    db.session.add_all([
        SpareStockLedger(part_id=p.id, qty_delta=1, balance_after=8, txn_type="RETURN",
                         return_of_id=lid, return_no=f"RET-OLD-{_TAG}", return_reason="Wrong part",
                         returned_by="Mohan", department=dept.name, created_by="old.store",
                         created_at=_utc(day, 13)),
        SpareStockLedger(part_id=p.id, qty_delta=1, balance_after=9, txn_type="RETURN",
                         return_of_id=lid, return_no=good.return_no, status=RETURN_ROW_INSPECTED,
                         created_at=_utc(day, 14)),
    ])
    db.session.commit()

    rows = _rows(s, "return_register", "2025-05-06", "2025-05-06")
    assert [r["Return No"] for r in rows] == [split.return_no, good.return_no, f"RET-OLD-{_TAG}"]
    assert [r["Result"] for r in rows] == ["Back in stock 1, Damaged 1", "Back in stock", "Back in stock"]
    assert [r["Reason"] for r in rows] == ["Fitment Issue", "Fitment Issue", "Wrong Part"]
    assert rows[0]["Issue No"] == lid and rows[0]["Qty"] == 2 and rows[0]["Recorded By"] == "store.user"
    assert rows[2]["Returned By"] == "Mohan" and rows[2]["Recorded By"] == "old.store"
    assert _rows(s, "return_register", "2025-05-07", "2025-05-07") == []


# ---- stock lists ---------------------------------------------------------------------------

def test_stock_lists_order_and_quantities(make):
    s = make.store()
    out = _part(s, s.old, qty=0, min_stock_level=5, max_stock_level=10)
    below = _part(s, s.live, qty=2, min_stock_level=5, max_stock_level=20)
    due = _part(s, s.live, qty=6, min_stock_level=5, reorder_level=8, max_stock_level=12)
    cheap = _part(s, s.live, qty=30, max_stock_level=10, unit_price=5)
    dear = _part(s, s.live, qty=15, max_stock_level=10, unit_price=50)
    unrated = _part(s, None, qty=50, max_stock_level=10)
    uncounted = _part(s, s.live)

    reorder = _rows(s, "reorder_list")
    assert [(r["Part No"], r["Status"], r["Order Qty"]) for r in reorder] == [
        (out.part_no, "Out of Stock", 10), (below.part_no, "Below Minimum", 18),
        (due.part_no, "Reorder", 6)]
    assert reorder[0]["Section"] == s.old_name   # a switched-off section, as recorded

    excess = _rows(s, "excess_list")
    assert [(r["Part No"], r["Excess Qty"], r["Excess Value"]) for r in excess] == [
        (dear.part_no, 5, 250), (cheap.part_no, 20, 100), (unrated.part_no, 40, None)]

    stock = {r["Part No"]: r for r in _rows(s, "stock_current")}
    assert len(stock) == 7
    assert stock[uncounted.part_no]["Status"] == "Not Counted" and stock[uncounted.part_no]["In Stock"] is None
    assert stock[unrated.part_no]["Section"] is None
    assert (stock[dear.part_no]["In Stock"], stock[dear.part_no]["Stock Value"]) == (15, 750)


# ---- the reorder and excess lists in the mail sheets' colours ---------------------------------
# The store approved the colours of the reorder / excess mail sheets on 14 Oct
# 2026, then opened the Reports page's reorder download and asked "bhai kahan
# kiya???" (15 Oct 2026): the list must read the same wherever it opens. As
# (fill, text colour, bold); None is no fill / the default text colour.

DARK_RED, RED, ORANGE = ("00DC2626", "00FFFFFF", True), ("00FFE4E6", "009F1239", True), ("00FFEDD5", "009A3412", True)
YELLOW, DARK_AMBER = ("00FEF9C3", "00854D0E", True), ("00D97706", "00FFFFFF", True)
BOLD, PLAIN = (None, None, True), (None, None, False)


def _colour_of(cell):
    fill = cell.fill.fgColor.rgb if cell.fill.fill_type == "solid" else None
    color = cell.font.color
    return fill, color.rgb if color is not None and color.type == "rgb" else None, bool(cell.font.b)


def _layout_of(cell):
    """Everything about a cell but its colours and boldness."""
    side = cell.border
    return (cell.value, cell.number_format, cell.quotePrefix, cell.font.name, cell.font.sz, cell.font.i,
            cell.alignment.horizontal, cell.alignment.vertical, cell.alignment.wrap_text,
            side.left.style, side.right.style, side.top.style, side.bottom.style)


def _excel(client, store, key):
    """(sheet, table) of a report's Excel download: the table is what its filter covers."""
    resp = client.get(f"/api/v1/reports/{key}/excel?store_code={store.code}")
    assert resp.status_code == 200, key
    sheet = load_workbook(io.BytesIO(resp.get_data())).active
    end = range_boundaries(sheet.auto_filter.ref)[3]
    return sheet, [list(row) for row in sheet.iter_rows(max_row=end, values_only=True)]


def _plain_but_for(sheet, table, painted):
    """The download is the plain build_export table of the same rows in all but
    the colours of its `painted` columns: values, formula guard, fonts, borders,
    alignment, header, column widths, frozen header and filter."""
    plain = load_workbook(build_export([(heading, heading) for heading in table[0]],
                                       [dict(zip(table[0], line)) for line in table[1:]])).active
    for number in range(1, len(table) + 1):
        for column, heading in enumerate(table[0], start=1):
            mine, theirs = sheet.cell(number, column), plain.cell(number, column)
            assert _layout_of(mine) == _layout_of(theirs), (number, heading)
            if number == 1 or heading not in painted:
                assert _colour_of(mine) == _colour_of(theirs), (number, heading)
    for column in range(1, len(table[0]) + 1):
        letter = get_column_letter(column)
        assert sheet.column_dimensions[letter].width == plain.column_dimensions[letter].width, letter
    assert (sheet.freeze_panes, sheet.auto_filter.ref) == (plain.freeze_panes, plain.auto_filter.ref)


def _legend_of(sheet, data_rows):
    """The legend under the table: one empty row, then the line in column A —
    the last thing on the sheet, outside the filter, a quiet italic note with
    no border, fill or bold."""
    below = data_rows + 3
    assert sheet.max_row == below
    assert [cell.value for cell in sheet[below - 1]] == [None] * sheet.max_column
    assert [cell.value for cell in sheet[below]][1:] == [None] * (sheet.max_column - 1)
    assert range_boundaries(sheet.auto_filter.ref)[3] == data_rows + 1
    legend = sheet.cell(below, 1)
    assert legend.border.left.style is None and legend.border.bottom.style is None
    assert legend.fill.fill_type is None and not legend.font.b and legend.font.i
    return legend.value


def test_reorder_list_excel_is_coloured_like_the_reorder_mail(client, make, monkeypatch):
    """In Stock and Status by the stock against the min — solid dark red with
    nothing on the shelf (zero or booked below it), red at or below the min (AT
    the min counts, though its status reads Reorder), orange for the rest of the
    list; a min stored as 0 is no min. Order Qty bold, nothing else coloured, and
    the mail's legend under the table. CSV and PDF stay as they were."""
    s = make.store()
    levels = {"min_stock_level": 2, "reorder_level": 4, "max_stock_level": 10, "unit_price": 5}
    zero = _part(s, s.live, qty=0, **levels)
    negative = _part(s, s.live, qty=0, **levels)
    SpareStock.query.filter_by(part_id=negative.id).update({"qty_on_hand": -1})   # booked below zero
    db.session.commit()
    below = _part(s, s.live, qty=1, min_stock_level=4, reorder_level=6, max_stock_level=10)
    at_min = _part(s, s.live, qty=4, min_stock_level=4, reorder_level=6, max_stock_level=10)
    at_reorder = _part(s, s.live, qty=6, min_stock_level=4, reorder_level=6, max_stock_level=10)
    no_min = _part(s, s.live, qty=3, min_stock_level=0, reorder_level=5, max_stock_level=10)
    _part(s, s.live, qty=8, **levels)                                              # healthy: not on the list
    c = _as(client, make.user(role="ADMIN", stores=()))

    sheet, table = _excel(c, s, "reorder_list")
    report = next(r for r in SpareReportService.REPORTS if r.key == "reorder_list")
    assert table[0] == list(report.columns)
    column = {heading: number for number, heading in enumerate(table[0], start=1)}
    parts = [negative, zero, below, at_min, at_reorder, no_min]                    # the mail's order
    assert [(line[column["Part No"] - 1], line[column["In Stock"] - 1], line[column["Status"] - 1])
            for line in table[1:]] == [
        (negative.part_no, -1, "Out of Stock"), (zero.part_no, 0, "Out of Stock"),
        (below.part_no, 1, "Below Minimum"), (at_min.part_no, 4, "Reorder"),
        (at_reorder.part_no, 6, "Reorder"), (no_min.part_no, 3, "Reorder")]
    for number, colour in enumerate([DARK_RED, DARK_RED, RED, RED, ORANGE, ORANGE], start=2):
        assert _colour_of(sheet.cell(number, column["In Stock"])) == colour, number
        assert _colour_of(sheet.cell(number, column["Status"])) == colour, number
        assert _colour_of(sheet.cell(number, column["Order Qty"])) == BOLD, number
        for heading, at in column.items():
            if heading not in ("In Stock", "Status", "Order Qty"):
                assert _colour_of(sheet.cell(number, at)) == PLAIN, (number, heading)
    _plain_but_for(sheet, table, ("In Stock", "Status", "Order Qty"))
    assert sheet.auto_filter.ref == "A1:N7" and sheet.freeze_panes == "A2"
    assert _legend_of(sheet, len(parts)) == REORDER_LEGEND == (
        "In Stock / Status: dark red = no stock · red = at or below min · orange = at or below reorder level")
    assert report.legend is REORDER_LEGEND                                          # the mail's own, not a copy

    # CSV and PDF carry no colours and no legend: the table only, as before.
    csv_rows = list(csv.reader(io.StringIO(
        c.get(f"/api/v1/reports/reorder_list/csv?store_code={s.code}").get_data().decode("utf-8-sig"))))
    assert csv_rows == [[("" if v is None else str(v)) for v in line] for line in table]
    drawn = _drawn_pdf_tables(monkeypatch)
    assert c.get(f"/api/v1/reports/reorder_list/pdf?store_code={s.code}").status_code == 200
    _, pdf_table = drawn[-1]
    assert len(pdf_table._cellvalues) == len(parts) + 1
    assert not any(REORDER_LEGEND in cell for row in pdf_table._cellvalues for cell in row)


def test_excess_list_excel_is_coloured_like_the_excess_mail(client, make):
    """Every In Stock cell yellow; the Excess Value in solid dark amber from
    ₹1,00,000, yellow from ₹25,000 — AT a threshold counts — and no colour
    below that or with no rate. Nothing else coloured, the mail's legend under
    the table."""
    s = make.store()

    def over(qty, rate):
        return _part(s, s.live, qty=qty, max_stock_level=10, unit_price=rate)
    parts = [over(2010, 100), over(1010, 100), over(11, 99999.99), over(260, 100), over(11, 24999),
             over(12, None)]
    _part(s, s.live, qty=8, max_stock_level=10, unit_price=100)                     # within max: not listed
    c = _as(client, make.user(role="ADMIN", stores=()))

    sheet, table = _excel(c, s, "excess_list")
    report = next(r for r in SpareReportService.REPORTS if r.key == "excess_list")
    assert table[0] == list(report.columns)
    column = {heading: number for number, heading in enumerate(table[0], start=1)}
    assert [(line[column["Part No"] - 1], line[column["Excess Value"] - 1]) for line in table[1:]] == [
        (parts[0].part_no, 200000), (parts[1].part_no, 100000), (parts[2].part_no, 99999.99),
        (parts[3].part_no, 25000), (parts[4].part_no, 24999), (parts[5].part_no, None)]
    for number, value_colour in enumerate([DARK_AMBER, DARK_AMBER, YELLOW, YELLOW, PLAIN, PLAIN], start=2):
        assert _colour_of(sheet.cell(number, column["In Stock"])) == YELLOW, number
        assert _colour_of(sheet.cell(number, column["Excess Value"])) == value_colour, number
        for heading, at in column.items():
            if heading not in ("In Stock", "Excess Value"):
                assert _colour_of(sheet.cell(number, at)) == PLAIN, (number, heading)
    _plain_but_for(sheet, table, ("In Stock", "Excess Value"))
    assert sheet.auto_filter.ref == "A1:K7" and sheet.freeze_panes == "A2"
    assert _legend_of(sheet, len(parts)) == EXCESS_LEGEND == (
        "In Stock: yellow = above max · Excess Value: dark amber = ₹1,00,000 or more · "
        "yellow = ₹25,000 to ₹1,00,000")
    assert report.legend is EXCESS_LEGEND


def test_every_other_report_excel_is_the_plain_table(client, make):
    """Only the reorder and excess lists are coloured. Every other report —
    here the Current Stock of the same parts, out of stock and above maximum
    among them — is the plain build_export table: no fill, no bold data cell,
    nothing under it."""
    s = make.store()
    _part(s, s.live, qty=0, min_stock_level=2, reorder_level=4, max_stock_level=10)
    _part(s, s.live, qty=1, min_stock_level=4, max_stock_level=10)
    _part(s, s.live, qty=2010, max_stock_level=10, unit_price=100)
    c = _as(client, make.user(role="ADMIN", stores=()))
    coloured = {"reorder_list", "excess_list"}
    for report in SpareReportService.REPORTS:
        assert (report.paint is None, report.legend is None) == ((report.key not in coloured),) * 2, report.key
    sheet, table = _excel(c, s, "stock_current")
    assert len(table) == 4 and sheet.max_row == 4                                     # nothing under the table
    _plain_but_for(sheet, table, ())
    for key in SPARE_KEYS:
        if key in coloured:
            continue
        sheet, table = _excel(c, s, key)
        assert sheet.max_row == len(table), key
        for row in sheet.iter_rows(min_row=2):
            assert all(_colour_of(cell) == PLAIN for cell in row), key


# ---- put-away, receipts ----------------------------------------------------------------------

def test_put_away_and_receipts(make):
    s = make.store()
    day = date(2025, 7, 1)
    p = _part(s, s.live, qty=5, location="R1-S1-B5(2)", unit_price=12)
    db.session.add(
        PutAwayLog(part_id=p.id, part_no=p.part_no, assigned_location=p.location,
                   scanned_location="R1-S1-B5", gate_entry_no="GE-1", qty=4, store_code=s.code,
                   put_by="putter", confirmed_at=_utc(day, 13), created_at=_utc(day, 13)))
    receipt = SpareStockLedger(part_id=p.id, qty_delta=3, balance_after=8, txn_type="RECEIPT",
                               status="ERP_GRN", reference=f"GRN-{_TAG}", created_at=_utc(day, 14))
    db.session.add(receipt)
    db.session.flush()
    db.session.add(SpareTransaction(part_id=p.id, store_code=s.code, txn_type="GRN", qty=3,
                                    ledger_id=receipt.id, supplier_name="G R TRADERS", grn_no=f"GRN-{_TAG}"))
    db.session.commit()

    put = _rows(s, "put_away_register", "2025-07-01", "2025-07-01")
    assert [(r["Part No"], r["Location"], r["Lot Qty"], r["Gate Entry No"], r["Put By"]) for r in put] == [
        (p.part_no, "R1-S1-B5(2)", 4, "GE-1", "putter")]

    receipts = _rows(s, "receipt_register", "2025-07-01", "2025-07-01")
    assert [(r["Date"], r["GRN No"], r["Qty"], r["Supplier"], r["Rate"], r["Value"], r["Section"])
            for r in receipts] == [("01.07.2025", f"GRN-{_TAG}", 3, "G R TRADERS", 12, 36, s.live_name)]
    assert _rows(s, "receipt_register", "2025-07-02", "2025-07-02") == []


# ---- no "None" / "nan" anywhere --------------------------------------------------------------

def test_no_none_or_nan_cells(client, make, monkeypatch):
    s = make.store()
    dept = make.department()
    _part(s, None)                                  # never counted, no section, no rate
    _, lid = _issue(s, dept)                        # waiting: no line, machine, die, used qty
    c = _as(client, make.user(role="ADMIN", stores=()))
    pdf_rows = []
    real_pdf = report_service.build_pdf_table
    monkeypatch.setattr(report_service, "build_pdf_table",
                        lambda title, columns, rows: pdf_rows.extend(rows) or real_pdf(title, columns, rows))
    bad = {"None", "none", "nan", "NaN", "null"}
    for key in SPARE_KEYS:
        for fmt in ("excel", "csv", "pdf"):
            resp = c.get(f"/api/v1/reports/{key}/{fmt}?store_code={s.code}")
            assert resp.status_code == 200, (key, fmt)
            if fmt == "excel":
                cells = [v for row in load_workbook(io.BytesIO(resp.get_data())).active.iter_rows(values_only=True)
                         for v in row]
                assert not any(isinstance(v, str) and v.strip() in bad for v in cells), key
            elif fmt == "csv":
                cells = [v for row in csv.reader(io.StringIO(resp.get_data().decode("utf-8-sig"))) for v in row]
                assert not any(v.strip() in bad for v in cells), key
    assert pdf_rows and not any(v is None or str(v).strip() in bad for r in pdf_rows for v in r.values())
    issue_rows = SpareReportService().build(s.code, "issue_register")[2]
    assert [r["Issue No"] for r in issue_rows] == [lid]
