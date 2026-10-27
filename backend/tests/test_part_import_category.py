"""Part Master Excel (27 Oct 2026): today's template says "Main Category" (the
store's broad group) and "Category" (the part's kind — sub_category), as the
store's DX sheet does; an older file ("Category (BOP/...)" + "Sub Category")
still reads as before."""
import io

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.part import Part
from app.services.part_service import PartService, TEMPLATE_COLUMNS


def _file(head, *rows):
    wb = Workbook()
    ws = wb.active
    ws.append(head)
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_todays_template_and_an_older_file_both_land_right():
    assert "Category" in TEMPLATE_COLUMNS and "Sub Category" not in TEMPLATE_COLUMNS
    assert any(c.startswith("Main Category") for c in TEMPLATE_COLUMNS)
    new = _file(["Part No", "Part Name", "Main Category (BOP/Import/Maintenance-Tool Room/Consumable/In-Parts)", "Category"],
                ("PIC-NEW-1", "AIR CYLINDER", "Maintenance/Tool Room", "Cylinder"))
    PartService().bulk_import_excel(new, store_code="MAINTENANCE")
    p = Part.query.filter_by(part_no="PIC-NEW-1").one()
    assert (p.sub_category, p.category) == ("Cylinder", "Maintenance/Tool Room")

    old = _file(["Part No", "Part Name", "Category (BOP/Import/Maintenance-Tool Room/Consumable/In-Parts)", "Sub Category"],
                ("PIC-OLD-1", "SEAL", "Maintenance/Tool Room", "Seal"))
    PartService().bulk_import_excel(old, store_code="MAINTENANCE")
    q = Part.query.filter_by(part_no="PIC-OLD-1").one()
    assert (q.sub_category, q.category) == ("Seal", "Maintenance/Tool Room")


def test_the_export_heads_category_as_the_kind():
    db.session.add(Part(part_no="PIC-EXP-1", part_name="BEARING", uom="NOS", part_type="OTHER",
                        store_code="MAINTENANCE", category="Maintenance/Tool Room", sub_category="Bearing"))
    db.session.commit()
    out = PartService().export_all("MAINTENANCE")
    buf = out[0] if isinstance(out, tuple) else out
    ws = load_workbook(buf).active
    head = [c.value for c in next(ws.iter_rows(max_row=1))]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[head.index("Part No")] == "PIC-EXP-1")
    assert row[head.index("Category")] == "Bearing" and row[head.index("Main Category")] == "Maintenance/Tool Room"
    assert "Sub Category" not in head
