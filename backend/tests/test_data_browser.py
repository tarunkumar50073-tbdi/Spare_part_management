"""Data Browser (27 Oct 2026): every table, read only, administrators only,
secrets never shown."""
import io

from openpyxl import load_workbook

from app.extensions import db
from app.models.app_setting import AppSetting
from app.models.part import Part
from app.models.user import User, UserStoreAccess

URL = "/api/v1/data-browser"
_seq = {"n": 0}


def _login(client, role, perms=None):
    _seq["n"] += 1
    u = User(username=f"dbrowse{_seq['n']}", role=role, is_active=True)
    u.set_password("pw-browse-123")
    if perms:
        u.set_permissions(perms)
    db.session.add(u)
    db.session.flush()
    db.session.add(UserStoreAccess(user_id=u.id, store_code="MAINTENANCE"))
    db.session.commit()
    with client.session_transaction() as s:
        s["user_id"] = u.id
    return u


def test_only_an_administrator_may_browse(app):
    c = app.test_client()
    _login(c, "STORE_USER", {"scan_part": {"view": True, "create": True, "update": True}})
    assert c.get(f"{URL}/tables").status_code == 403
    assert c.get(f"{URL}/tables/parts").status_code == 403


def test_tables_rows_search_sort_and_no_secrets(app):
    c = app.test_client()
    me = _login(c, "ADMIN")
    db.session.add(Part(part_no="DBR-ZETA-1", part_name="DATA BROWSER BEARING", uom="NOS", part_type="OTHER",
                        store_code="MAINTENANCE"))
    db.session.commit()
    AppSetting.set_value("dbr_test_mail", {"host": "smtp.local", "smtp_password": "s3cret", "nested": {"api_key": "k"}})

    body = c.get(f"{URL}/tables").get_json()["data"]
    parts = next(t for t in body["tables"] if t["name"] == "parts")
    assert parts["group"] == "Spare store" and parts["rows"] >= 1 and parts["about"]

    got = c.get(f"{URL}/tables/parts", query_string={"q": "dbr-zeta", "sort": "part_no"}).get_json()["data"]
    assert [r["part_no"] for r in got["rows"]] == ["DBR-ZETA-1"] and got["total"] == 1

    users = c.get(f"{URL}/tables/users", query_string={"q": me.username}).get_json()["data"]
    hidden = {col["name"] for col in users["columns"] if col["hidden"]}
    assert "password_hash" in hidden
    row = users["rows"][0]
    assert row["username"] == me.username
    assert all(row[h] in (None, "•••• hidden") for h in hidden)
    assert "pw-browse-123" not in str(users) and me.password_hash not in str(users)

    setting = c.get(f"{URL}/tables/app_settings", query_string={"q": "dbr_test_mail"}).get_json()["data"]["rows"][0]
    assert "smtp.local" in setting["value"] and "s3cret" not in setting["value"] and '"k"' not in setting["value"]

    assert c.get(f"{URL}/tables/no_such_table").status_code == 404
    assert c.get(f"{URL}/tables/parts", query_string={"sort": "no_such_column"}).status_code == 400

    resp = c.get(f"{URL}/tables/users/export", query_string={"q": me.username})
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.data)).active
    head = [c.value for c in ws[1]]
    values = [c.value for c in ws[2]]
    assert values[head.index("username")] == me.username
    assert values[head.index("password_hash")] in (None, "•••• hidden")
