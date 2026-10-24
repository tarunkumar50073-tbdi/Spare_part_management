"""Part Master, 21 Oct 2026: a Model No. column, one photo per part (taken on
the phone or picked from a file, shown beside the part) and rack / category /
sub category filters. Photos are files in their own folder — here a temp folder
per test — never rows, and the nightly backup copies them.

Everything here is PP-<tag>-<n>: this file's own store, parts and users."""
import io
import os
import uuid

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.extensions import db
from app.models.part import Part
from app.models.store import Section, Store
from app.models.user import User, UserStoreAccess
from app.repositories.part_repository import PartRepository
from app.services.backup_service import mirror_part_images
from app.services.part_image_service import PartImageService
from app.services.part_service import PartService
from app.utils.exceptions import BadRequestError, NotFoundError

_TAG = uuid.uuid4().hex[:6].upper()
_seq = {"n": 0}
URL = "/api/v1/parts"


def _n():
    _seq["n"] += 1
    return _seq["n"]


@pytest.fixture()
def photos(app, tmp_path):
    """Each test keeps its photos in its own temp folder."""
    folder = tmp_path / "part_images"
    old = app.config.get("PART_IMAGE_DIR")
    app.config["PART_IMAGE_DIR"] = str(folder)
    yield folder
    app.config["PART_IMAGE_DIR"] = old


@pytest.fixture()
def make():
    made = {"stores": [], "users": []}

    def store():
        code = f"PP{_TAG}{_n()}"
        s = Store(code=code, name=f"Photo Store {code}")
        db.session.add(s)
        db.session.flush()
        sec = Section(store_id=s.id, name=f"Sec {code}")
        db.session.add(sec)
        db.session.commit()
        made["stores"].append(code)
        return code, sec.id

    def part(store_code, section_id=None, **over):
        n = _n()
        p = Part(part_no=f"PP-{_TAG}-{n}", part_name=f"Photo Part {n}", uom="NOS", part_type="OTHER",
                 store_code=store_code, section_id=section_id, **over)
        db.session.add(p)
        db.session.commit()
        return p

    def user(perms, stores=None):
        # A store user reaches the parts of their own stores only (24 Oct
        # 2026): by default, the stores this test made.
        u = User(username=f"pp{_TAG.lower()}{_n()}", role="STORE_USER", is_active=True)
        u.set_password("pw")
        u.set_permissions(perms)
        db.session.add(u)
        db.session.flush()
        for code in (made["stores"] if stores is None else stores):
            db.session.add(UserStoreAccess(user_id=u.id, store_code=code))
        db.session.commit()
        made["users"].append(u.id)
        return u

    yield type("Make", (), {"store": staticmethod(store), "part": staticmethod(part), "user": staticmethod(user)})
    db.session.rollback()
    Part.query.filter(Part.store_code.in_(made["stores"])).delete(synchronize_session=False)
    Store.query.filter(Store.code.in_(made["stores"])).update({"is_active": False}, synchronize_session=False)
    User.query.filter(User.id.in_(made["users"])).update({"is_active": False}, synchronize_session=False)
    db.session.commit()


def _jpeg(size=(2400, 1200), color=(200, 30, 30), fmt="JPEG"):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, fmt)
    buf.seek(0)
    return buf


def _as(client, user):
    with client.session_transaction() as sess:
        sess["user_id"] = user.id
    return client


# ---- the photo -----------------------------------------------------------------------------------

def test_a_photo_is_saved_as_a_file_cut_down_with_a_small_copy(make, photos):
    code, sec = make.store()
    part = make.part(code, sec)
    out = PartImageService().save(part.id, _jpeg())
    name = out.image_file
    assert name.startswith(f"{part.id}-") and name.endswith(".jpg")
    big = Image.open(photos / name)
    small = Image.open(photos / name.replace(".jpg", "_t.jpg"))
    assert max(big.size) == 1600 and max(small.size) == 240 and big.format == "JPEG"
    assert PartImageService().path(part.id).endswith(name)
    assert PartImageService().path(part.id, thumb=True).endswith("_t.jpg")


def test_a_new_photo_replaces_the_old_one_and_remove_clears_it(make, photos):
    code, sec = make.store()
    part = make.part(code, sec)
    first = PartImageService().save(part.id, _jpeg()).image_file
    second = PartImageService().save(part.id, _jpeg(color=(0, 0, 200), fmt="PNG")).image_file
    assert first != second                                              # a new name: never a stale cache
    assert sorted(os.listdir(photos)) == sorted([second, second.replace(".jpg", "_t.jpg")])
    assert PartImageService().delete(part.id).image_file is None
    assert os.listdir(photos) == []
    with pytest.raises(NotFoundError, match="has no photo"):
        PartImageService().path(part.id)


def test_a_file_that_is_not_a_picture_is_refused(make, photos):
    code, sec = make.store()
    part = make.part(code, sec)
    with pytest.raises(BadRequestError, match="not a photo"):
        PartImageService().save(part.id, io.BytesIO(b"MZ this is not a picture"))
    db.session.refresh(part)
    assert part.image_file is None and (not photos.exists() or os.listdir(photos) == [])


def test_deleting_a_part_removes_its_photo(make, photos):
    code, sec = make.store()
    part = make.part(code, sec)
    PartImageService().save(part.id, _jpeg())
    PartService().delete(part.id)
    assert os.listdir(photos) == []


def test_photo_api_needs_the_part_master_edit_right(app, make, photos):
    code, sec = make.store()
    part = make.part(code, sec)
    viewer = make.user({"part_master": {"view": True, "edit": False}})
    editor = make.user({"part_master": {"view": True, "edit": True}})
    stranger = make.user({"part_master": {"view": True, "edit": True}}, stores=[])   # another store's editor
    c = app.test_client()
    assert c.get(f"{URL}/{part.id}/image").status_code == 401                # signed in only
    _as(c, viewer)
    r = c.post(f"{URL}/{part.id}/image", data={"file": (_jpeg(), "a.jpg")}, content_type="multipart/form-data")
    assert r.status_code == 403
    _as(c, stranger)
    r = c.post(f"{URL}/{part.id}/image", data={"file": (_jpeg(), "a.jpg")}, content_type="multipart/form-data")
    assert r.status_code == 403                                              # not a part of their store
    _as(c, editor)
    r = c.post(f"{URL}/{part.id}/image", data={"file": (_jpeg(), "a.jpg")}, content_type="multipart/form-data")
    assert r.status_code == 200 and r.get_json()["data"]["image_file"]
    r = c.post(f"{URL}/{part.id}/image", data={"file": (io.BytesIO(b"not an image"), "x.jpg")},
               content_type="multipart/form-data")
    assert r.status_code == 400
    _as(c, viewer)                                                           # anyone signed in sees it
    for size in ("", "thumb"):
        r = c.get(f"{URL}/{part.id}/image?size={size}")
        assert r.status_code == 200 and r.mimetype == "image/jpeg"
        r.close()
    _as(c, editor)
    assert c.delete(f"{URL}/{part.id}/image").get_json()["data"]["image_file"] is None
    assert c.get(f"{URL}/{part.id}/image").status_code == 404


def test_missing_photo_is_a_chip_like_missing_max(make, photos):
    code, sec = make.store()
    a, b = make.part(code, sec), make.part(code, sec)
    PartImageService().save(a.id, _jpeg())
    assert PartRepository().count_missing(store_code=code)["photo"] == 1
    assert [p.id for p in PartRepository().search(store_code=code, missing="photo")] == [b.id]


def test_photo_mode_tallies_photos_per_rack(make, photos):
    """The Photo Mode works rack by rack: each rack says how many of its parts
    still have no photo, and the scope says how far the store has got."""
    code, sec = make.store()
    a = make.part(code, sec, location="R7-S1-B1")
    make.part(code, sec, location="R7-S1-B2")
    make.part(code, sec, location="R8-S1-B1")
    PartImageService().save(a.id, _jpeg())
    opts = PartRepository().filter_options(store_code=code)
    assert [(r["value"], r["count"], r["no_photo"]) for r in opts["racks"]] == [("R7", 2, 1), ("R8", 1, 1)]
    assert opts["photos"] == {"total": 3, "with_photo": 1}


def test_photo_mode_counts_active_parts_only(app, make, photos):
    """The Photo Mode lists active parts only, so its tally does too — an
    inactive part made it read 4,893 against the 4,892 parts shown (the store,
    23 Oct 2026). Without the flag (Part Master) every part still counts."""
    code, sec = make.store()
    make.part(code, sec, location="R7-S1-B1")
    make.part(code, sec, location="R7-S1-B2", is_active=False)
    repo = PartRepository()
    assert repo.filter_options(store_code=code)["photos"]["total"] == 2
    active = repo.filter_options(store_code=code, is_active=True)
    assert active["photos"] == {"total": 1, "with_photo": 0}
    assert [(r["value"], r["count"]) for r in active["racks"]] == [("R7", 1)]
    viewer = make.user({"part_master": {"view": True}})
    c = _as(app.test_client(), viewer)
    got = c.get("/api/v1/parts/filter-options", query_string={"store_code": code, "is_active": "true"}).get_json()["data"]
    assert got["photos"]["total"] == 1


# ---- model no. -----------------------------------------------------------------------------------

def test_model_no_is_saved_searched_exported_and_imported(app, make):
    code, sec = make.store()
    editor = make.user({"part_master": {"view": True, "edit": True}})
    c = _as(app.test_client(), editor)
    r = c.post(URL, json={"part_no": f"PP-{_TAG}-M", "part_name": "Tower light", "store_code": code,
                          "model_no": "XVC1B3HK"})
    assert r.status_code == 201 and r.get_json()["data"]["model_no"] == "XVC1B3HK"
    pid = r.get_json()["data"]["id"]
    assert c.put(f"{URL}/{pid}", json={"model_no": "XVC1B3HK-2"}).get_json()["data"]["model_no"] == "XVC1B3HK-2"
    rows = c.get(URL, query_string={"store_code": code, "keyword": "xvc1b3"}).get_json()["data"]
    assert [p["id"] for p in rows] == [pid]
    assert c.post(URL, json={"part_no": f"PP-{_TAG}-L", "part_name": "x", "store_code": code,
                             "model_no": "M" * 61}).status_code == 422              # at most 60
    sheet = load_workbook(PartService().export_all(store_code=code)).active
    head = [cell.value for cell in sheet[1]]
    assert head[head.index("Specification") + 1] == "Model No"
    assert sheet.cell(row=2, column=head.index("Model No") + 1).value == "XVC1B3HK-2"
    # Re-import with a "Model" heading fills it in; a blank cell leaves it alone.
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Part No", "Part Name", "Model"])
    wb.active.append([f"PP-{_TAG}-M", "Tower light", "XVC-NEW"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    PartService().bulk_import_excel(buf, store_code=code)
    assert db.session.get(Part, pid).model_no == "XVC-NEW"


# ---- filters -------------------------------------------------------------------------------------

def test_rack_category_and_sub_category_filters(make):
    code, sec = make.store()
    r4a = make.part(code, sec, location="R4-S1-B1 (1)", sub_category="Cable", category="Maintenance/Tool Room")
    r4b = make.part(code, sec, location="r4 -S2-B3", sub_category=" cable ", category="Consumable")
    r14 = make.part(code, sec, location="R14-S1-B9", sub_category="Sensor")
    old = make.part(code, sec, location="OLD STORE", sub_category="Sensor")
    make.part(code, sec, location=None)
    repo = PartRepository()
    ids = lambda **f: sorted(p.id for p in repo.search(store_code=code, **f))   # noqa: E731
    assert ids(rack="R4") == sorted([r4a.id, r4b.id])                            # not R14
    assert ids(rack="R14") == [r14.id] and ids(rack="old store") == [old.id]
    assert ids(sub_category="cable") == sorted([r4a.id, r4b.id])
    assert ids(category="Consumable") == [r4b.id]
    assert ids(rack="R4", category="Consumable") == [r4b.id]
    opts = repo.filter_options(store_code=code)
    # Rack order: R4 before R14 (number order), named places last; one entry
    # per value case aside, no blank one for the part without a location.
    assert [(r["value"].upper(), r["count"]) for r in opts["racks"]] == [("R4", 2), ("R14", 1), ("OLD STORE", 1)]
    assert sorted((s["value"].lower(), s["count"]) for s in opts["sub_categories"]) == [("cable", 2), ("sensor", 2)]
    assert {c["value"]: c["count"] for c in opts["categories"]} == {"Maintenance/Tool Room": 1, "Consumable": 1}


def test_list_api_takes_the_new_filters(app, make):
    code, sec = make.store()
    a = make.part(code, sec, location="R2-S1-B1", sub_category="Bearing")
    make.part(code, sec, location="R3-S1-B1", sub_category="Bearing")
    viewer = make.user({"part_master": {"view": True}})
    c = _as(app.test_client(), viewer)
    rows = c.get(URL, query_string={"store_code": code, "rack": "R2", "sub_category": "bearing"}).get_json()["data"]
    assert [p["id"] for p in rows] == [a.id]
    opts = c.get(f"{URL}/filter-options", query_string={"store_code": code, "section_id": sec}).get_json()["data"]
    assert [r["value"] for r in opts["racks"]] == ["R2", "R3"]


# ---- backup --------------------------------------------------------------------------------------

def test_backup_copies_new_photos_and_never_deletes_there(tmp_path):
    src, backups = tmp_path / "part_images", tmp_path / "backups"
    src.mkdir()
    (src / "1-aa.jpg").write_bytes(b"one")
    (src / "2-bb.jpg.partial").write_bytes(b"half")
    out = mirror_part_images(str(src), str(backups))
    assert out["copied"] == 1 and sorted(os.listdir(backups / "part_images")) == ["1-aa.jpg"]
    os.remove(src / "1-aa.jpg")
    (src / "3-cc.jpg").write_bytes(b"three")
    assert mirror_part_images(str(src), str(backups))["copied"] == 1
    assert sorted(os.listdir(backups / "part_images")) == ["1-aa.jpg", "3-cc.jpg"]
    assert mirror_part_images(str(tmp_path / "none"), str(backups)) == {"copied": 0, "folder": None, "error": None}


def test_photos_live_beside_the_database_named_after_it(app):
    """The TEST copy sits beside the plant's database, so each keeps its own
    folder: dx_store_part_images / dx_store_TEST_part_images."""
    from app.services.backup_service import database_file
    from app.services.part_image_service import image_folder
    old = app.config.get("PART_IMAGE_DIR")
    app.config["PART_IMAGE_DIR"] = None
    try:
        db_path = database_file()
        stem = os.path.splitext(os.path.basename(db_path))[0]
        assert image_folder() == os.path.join(os.path.dirname(db_path), f"{stem}_part_images")
    finally:
        app.config["PART_IMAGE_DIR"] = old


def test_the_supplier_list_can_be_scoped_to_one_stores_parts(app, make):
    """The supplier master is the Child Part store's; the Spare Parts Part Master
    lists only suppliers its own parts are linked to (none, today)."""
    from app.models.supplier import Supplier
    linked, spare = make.store(), make.store()
    sup = Supplier(vendor_code=f"PPV{_TAG}{_n()}"[:20], supplier_name="Linked Vendor", status="Active")
    db.session.add(sup)
    db.session.commit()
    make.part(linked[0], linked[1], supplier_id=sup.id)
    make.part(spare[0], spare[1])
    viewer = make.user({"part_master": {"view": True}})
    c = _as(app.test_client(), viewer)
    ids = lambda code: [s["id"] for s in c.get("/api/v1/suppliers", query_string={"store_code": code, "per_page": 500}).get_json()["data"]]  # noqa: E731
    assert sup.id in ids(linked[0]) and ids(spare[0]) == []
    assert sup.id in [s["id"] for s in c.get("/api/v1/suppliers", query_string={"per_page": 500}).get_json()["data"]]
    Part.query.filter(Part.supplier_id == sup.id).update({"supplier_id": None})
    db.session.delete(sup)
    db.session.commit()


# ---- up to three angles (24 Oct 2026) ------------------------------------------------------------

def test_two_or_three_angles_are_all_kept_the_first_the_main(make, photos):
    """"dono thik lagi to dono rhni chahiye": every angle saved is kept, the
    first in image_file (lists, filters, boards), the rest beside it."""
    code, sec = make.store()
    part = make.part(code, sec)
    svc = PartImageService()
    out = svc.save_set(part.id, [_jpeg(color=(200, 0, 0)), _jpeg(color=(0, 200, 0)), _jpeg(color=(0, 0, 200))])
    assert len(out.photo_files) == 3 and out.image_file == out.photo_files[0]
    assert out.to_dict()["photos"] == out.photo_files
    assert len(os.listdir(photos)) == 6                                  # each with its small copy
    for n, name in enumerate(out.photo_files):
        assert svc.path(part.id, n=n).endswith(name)
        assert svc.path(part.id, thumb=True, n=n).endswith(name.replace(".jpg", "_t.jpg"))
    with pytest.raises(NotFoundError, match="no photo 4"):
        svc.path(part.id, n=3)
    with pytest.raises(BadRequestError, match="At most 3"):
        svc.save_set(part.id, [_jpeg() for _ in range(4)])


def test_one_angle_can_be_taken_out_and_the_next_moves_up(make, photos):
    """"teen khechi and 2 sahi lagi to 1 hta sake"."""
    code, sec = make.store()
    part = make.part(code, sec)
    svc = PartImageService()
    first, second, third = svc.save_set(part.id, [_jpeg(), _jpeg(), _jpeg()]).photo_files
    assert svc.delete(part.id, n=0).photo_files == [second, third]      # the main one gone: the next leads
    assert sorted(os.listdir(photos)) == sorted([second, second.replace(".jpg", "_t.jpg"),
                                                 third, third.replace(".jpg", "_t.jpg")])
    assert svc.delete(part.id, n=1).photo_files == [second]
    assert svc.delete(part.id).photo_files == [] and os.listdir(photos) == []


def test_a_bad_picture_among_the_angles_writes_nothing(make, photos):
    code, sec = make.store()
    part = make.part(code, sec)
    keep = PartImageService().save(part.id, _jpeg()).image_file
    with pytest.raises(BadRequestError, match="not a photo"):
        PartImageService().save_set(part.id, [_jpeg(), io.BytesIO(b"not a picture")])
    db.session.refresh(part)
    assert part.photo_files == [keep]
    assert sorted(os.listdir(photos)) == sorted([keep, keep.replace(".jpg", "_t.jpg")])


def test_deleting_a_part_removes_every_angle(make, photos):
    code, sec = make.store()
    part = make.part(code, sec)
    PartImageService().save_set(part.id, [_jpeg(), _jpeg()])
    PartService().delete(part.id)
    assert os.listdir(photos) == []


def test_the_photo_api_takes_and_serves_several_angles(app, make, photos):
    code, sec = make.store()
    part = make.part(code, sec)
    editor = make.user({"part_master": {"view": True, "update": True}})
    c = _as(app.test_client(), editor)
    r = c.post(f"{URL}/{part.id}/image", content_type="multipart/form-data",
               data={"file": [(_jpeg(), "a.jpg"), (_jpeg(color=(0, 90, 0)), "b.jpg")]})
    assert r.status_code == 200 and len(r.get_json()["data"]["photos"]) == 2
    for n in (0, 1):
        got = c.get(f"{URL}/{part.id}/image?n={n}&size=thumb")
        assert got.status_code == 200 and got.mimetype == "image/jpeg"
        got.close()
    assert c.get(f"{URL}/{part.id}/image?n=2").status_code == 404
    r = c.delete(f"{URL}/{part.id}/image?n=1")
    assert r.status_code == 200 and len(r.get_json()["data"]["photos"]) == 1
    r = c.post(f"{URL}/{part.id}/image", content_type="multipart/form-data",
               data={"file": [(_jpeg(), f"{i}.jpg") for i in range(4)]})
    assert r.status_code == 400


def test_a_live_database_without_the_photos_column_gets_it_on_boot(app):
    """The live server's database may not be SQLite, where _patch_schema never
    runs: the column for the 2nd / 3rd photo is added by the portable step."""
    from sqlalchemy import inspect, text
    from app import _add_late_columns
    names = lambda: {c["name"] for c in inspect(db.engine).get_columns("parts")}
    db.session.execute(text("ALTER TABLE parts DROP COLUMN extra_images"))
    db.session.commit()
    assert "extra_images" not in names()
    _add_late_columns()
    assert "extra_images" in names()
    _add_late_columns()                                          # a second boot changes nothing
    assert "extra_images" in names()
