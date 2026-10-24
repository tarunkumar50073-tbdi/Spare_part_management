"""The Daily Stock Report workbook (services/stock_mis_report) and its download
route (GET /api/v1/spare-stock/daily-report).

The report is whole-store, so it gets a store of its own — every figure on it
is then known from the parts below, whatever else the shared session DB holds:

    section MIS Alpha
      MISA01  0,  min 2, max 10, rate 100       zero stock, order 10 (to max) = 1,000
      MISA02  1,  min 5, no max, rate 10        below min,  order 4 (to min)  = 40
      MISA03  6,  min 2, reorder 6, max 12, r 5 reorder,    order 6           = 30
      MISA04  30, max 10, rate 50               excess 20 = 1,000
      MISA05  12, max 10, no rate               excess 2, not valued
      MISA06  5,  min 2, max 10, rate 20.5      healthy, 102.50; has a gate entry (supplier)
      MISA07  never counted, rate 99            not counted, no value
    section MIS Beta
      MISB01  0, levels stored as 0, no rate    zero stock, no order quantity, no value
      MISB02  8, max 5, rate 300                excess 3 = 900
      MISB03  2.5 KGS, named "=SUM(1,1)",       healthy, 10.00
              specification "-710619"
    no section
      MISU01  3, levels stored as 0, rate 10    healthy, 30.00 -> "Unassigned"
    MISX01 is inactive and must not appear anywhere.

MISA01, MISA05 and MISA07 each have a category to themselves: one worth a real
0.00, two whose value is not known (no rate; never counted).
"""
import io
import re
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.gate_entry import GateEntry
from app.models.part import Part
from app.models.spare_stock import SpareStock, SpareStockLedger
from app.models.store import Section, Store
from app.models.user import User, UserStoreAccess
from app.services import stock_mis_report
from app.services.stock_mis_report import MONEY, WHOLE, XLSX_MIME, daily_stock_report, stock_only_report

STORE = "MISRPT"
# Noon UTC: the same calendar day on the plant clock wherever the server runs.
NOW = datetime(2026, 10, 13, 12, 0)
SHEETS = ["Summary", "Stock Detail", "Critical & Reorder", "Excess", "Category Summary"]
DETAIL_HEAD = ["Sr No", "Section", "Category", "Part No", "Description", "Specification", "Supplier",
               "Location", "UOM", "In Stock", "Min", "Reorder Level", "Max", "Status", "Rate (₹)",
               "Stock Value (₹)"]
URL = "/api/v1/spare-stock/daily-report"
STOCK_ONLY_URL = "/api/v1/spare-stock/daily-stock-only"
STOCK_ONLY_HEAD = ["Sr No", "Section", "Part No", "Description", "Location", "UOM", "In Stock", "Min",
                   "Reorder Level", "Max"]


def _part(section, part_no, qty, name=None, category=None, mn=0, mx=None, ro=None, rate=None,
          uom="NOS", active=True, spec=None):
    part = Part(part_no=part_no, part_name=name or f"Part {part_no}", uom=uom, part_type="OTHER",
                store_code=STORE, section_id=section.id if section else None,
                sub_category=category, specification=spec or f"Spec {part_no}", location="R1-S1-B1",
                min_stock_level=mn, max_stock_level=mx, reorder_level=ro, unit_price=rate,
                is_active=active)
    db.session.add(part)
    db.session.flush()
    if qty is not None:
        db.session.add(SpareStock(part_id=part.id, qty_on_hand=qty))
    return part


def _seed():
    store = Store(code=STORE, name="MIS Test Store")
    db.session.add(store)
    db.session.flush()
    alpha, beta = Section(store_id=store.id, name="MIS Alpha"), Section(store_id=store.id, name="MIS Beta")
    db.session.add_all([alpha, beta])
    db.session.flush()
    _part(alpha, "MISA01", 0, category="Coupling", mn=2, mx=10, rate=100)
    _part(alpha, "MISA02", 1, category="Bearing", mn=5, rate=10)
    _part(alpha, "MISA03", 6, category="Bearing", mn=2, ro=6, mx=12, rate=5)
    _part(alpha, "MISA04", 30, category="Bearing", mn=2, mx=10, rate=50)
    _part(alpha, "MISA05", 12, category="Adapter", mn=2, mx=10)
    moved = _part(alpha, "MISA06", 5, category="Sensor", mn=2, mx=10, rate=20.5)
    _part(alpha, "MISA07", None, category="Gauge", mn=2, mx=10, rate=99)
    _part(beta, "MISB01", 0, category="Belt", ro=0, mx=0)
    _part(beta, "MISB02", 8, category="Belt", mx=5, rate=300)
    _part(beta, "MISB03", 2.5, name="=SUM(1,1)", spec="-710619", mn=1, mx=5, rate=4, uom="KGS")
    _part(None, "MISU01", 3, category="Bearing", ro=0, mx=0, rate=10)
    _part(alpha, "MISX01", 4, category="Bearing", rate=1000, active=False)
    db.session.add(GateEntry(store_code=STORE, item_code="MISA06", ven_cd="V-ACME", ven_name="ACME TOOLS",
                             gate_entry_dt=date(2026, 9, 1), qty=1, gate_entry_no="GE-MIS-1"))
    # Movement inside the report's week: 5 received the day before, 2 issued on the day.
    db.session.add_all([
        SpareStockLedger(part_id=moved.id, qty_delta=5, balance_after=7, txn_type="RECEIPT",
                         created_at=datetime(2026, 10, 12, 12, 0)),
        SpareStockLedger(part_id=moved.id, qty_delta=-2, balance_after=5, txn_type="ISSUE",
                         created_at=datetime(2026, 10, 13, 12, 0)),
    ])
    db.session.commit()


def _cleanup():
    part_ids = [p.id for p in Part.query.filter_by(store_code=STORE)]
    for model in (SpareStockLedger, SpareStock):
        model.query.filter(model.part_id.in_(part_ids)).delete(synchronize_session=False)
    GateEntry.query.filter_by(store_code=STORE).delete()
    Part.query.filter_by(store_code=STORE).delete()
    for user in User.query.filter(User.username.like("misuser%")):
        db.session.delete(user)
    store = Store.query.filter_by(code=STORE).first()
    if store:
        db.session.delete(store)   # takes its sections with it
    db.session.commit()


@pytest.fixture(scope="module")
def report(app):
    """The seeded store's report, built once: the attachment and the workbook
    read back from its bytes."""
    with app.app_context():
        _seed()
        try:
            built = daily_stock_report(STORE, now=NOW)
            yield built, load_workbook(io.BytesIO(built["content"]))
        finally:
            db.session.rollback()
            _cleanup()


@pytest.fixture()
def wb(report):
    return report[1]


@pytest.fixture()
def client(app):
    return app.test_client()


# ---- reading the workbook back ----

def _table(ws, heading):
    """A Summary table as {row label: {column header: value}}: the rows under
    `heading`, down to the first empty row."""
    top = next(c.row for c in ws["A"] if c.value == heading) + 1
    heads = {c.column: c.value for c in ws[top] if c.value is not None}
    first = min(heads)
    rows = {}
    for r in range(top + 1, ws.max_row + 1):
        label = ws.cell(r, first).value
        if label is None:
            break
        rows[label] = {heads[c]: ws.cell(r, c).value for c in heads if c != first}
    return rows


def _tile(ws, label):
    """(figure, small print) of the key-figure tile with this label."""
    cell = next(c for row in ws.iter_rows() for c in row if c.value == label)
    return ws.cell(cell.row + 1, cell.column).value, ws.cell(cell.row + 2, cell.column).value


def _list(ws):
    """A list sheet as (headers, [row dict]): header on row 4, data from row 5."""
    heads = [c.value for c in ws[4]]
    return heads, [dict(zip(heads, row)) for row in ws.iter_rows(min_row=5, values_only=True)]


def _values(ws, row):
    return [c.value for c in ws[row] if c.value is not None]


# ---- the workbook ----

def test_attachment_and_workbook_shape(report):
    built, wb = report
    assert built["filename"] == "Daily Stock Report 13.10.2026.xlsx"
    assert built["filename"].isascii()
    assert built["mime"] == XLSX_MIME
    assert built["content"][:2] == b"PK"

    assert wb.sheetnames == SHEETS
    assert wb.active.title == "Summary"
    assert wb.properties.title == "Daily Stock Report"
    assert wb.properties.creator == "TBDI Digital SMP"
    for ws in wb.worksheets:
        assert ws.sheet_view.showGridLines is False
        assert ws.sheet_properties.tabColor is not None
        assert ws.sheet_view.selection[-1].activeCell == "A1"
    assert len(wb["Summary"]._images) == 1      # the logo


def test_summary_header_and_key_figures(wb):
    ws = wb["Summary"]
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "DAILY STOCK REPORT" in texts
    assert "Toyota Boshoku Device India · MIS Test Store" in texts
    assert "Report date: 13.10.2026" in texts
    assert any(t.startswith("Generated: 13.10.2026 ") for t in texts)
    assert "System-generated report · TBDI Digital SMP" in texts
    for box in ("Prepared by (Store)", "Checked by", "Approved by"):
        assert box in texts

    assert _tile(ws, "Total Parts") == (11, "1 not counted yet")
    assert _tile(ws, "Stock Value (₹)") == (4082.5, "2 with no rate")
    # Critical is zero stock and below minimum together, with the split beneath.
    assert _tile(ws, "Critical") == (3, "zero stock 2\nbelow minimum 1")
    assert _tile(ws, "Reorder") == (1, "at/below reorder level")
    assert _tile(ws, "Excess") == (3, "₹ 1,900 tied up")
    # MISU01 is Healthy with no level to fall foul of: the tile says how many are.
    assert _tile(ws, "Healthy") == (3, "1 with no levels set")


def test_tile_small_print_fits_its_tile(wb):
    """A tile is two grid columns wide — about 24 characters of its small print.
    A longer line breaks wherever it happens to ("zero stock 353 · below" over
    "minimum 142"), so every line is short and Critical breaks between its halves."""
    ws = wb["Summary"]
    for label in ("Total Parts", "Stock Value (₹)", "Critical", "Reorder", "Excess", "Healthy"):
        tile = next(c for row in ws.iter_rows() for c in row if c.value == label)
        small = ws.cell(tile.row + 2, tile.column)
        assert max(len(line) for line in small.value.split("\n")) <= 24, (label, small.value)
        assert small.alignment.wrap_text, label


def test_caption_figures_round_like_the_cells():
    """A caption's rupees are rounded as Excel rounds the same amount in a cell:
    a half goes up, where round() would print 12,344 beside a cell showing 12,345."""
    figures = (12344.5, 2.5, 0.5, 9342159.53, 1234567, 999, 0)
    assert [stock_mis_report._grouped(v) for v in figures] == [
        "12,345", "3", "1", "93,42,160", "12,34,567", "999", "0"]


def test_summary_status_table_reconciles(wb):
    rows = _table(wb["Summary"], "Stock by status")
    assert list(rows) == ["Critical", "of which zero stock", "of which below minimum", "Reorder",
                          "Excess (above maximum)", "Healthy", "Discontinued (not required)", "Not counted",
                          "Total"]
    parts = {label: row["Parts"] for label, row in rows.items()}
    value = {label: row["Stock Value (₹)"] for label, row in rows.items()}
    assert parts == {"Critical": 3, "of which zero stock": 2, "of which below minimum": 1, "Reorder": 1,
                     "Excess (above maximum)": 3, "Healthy": 3, "Discontinued (not required)": 0,
                     "Not counted": 1, "Total": 11}
    # MISA05 (no rate) and MISA07 (never counted) add parts but no value.
    assert value == {"Critical": 10, "of which zero stock": 0, "of which below minimum": 10, "Reorder": 30,
                     "Excess (above maximum)": 3900, "Healthy": 142.5, "Discontinued (not required)": 0,
                     "Not counted": 0, "Total": 4082.5}

    # The total is the whole store; the two indented lines are inside Critical,
    # not on top of it.
    lines = ("Critical", "Reorder", "Excess (above maximum)", "Healthy", "Not counted")
    assert sum(parts[k] for k in lines) == parts["Total"] == 11
    assert sum(value[k] for k in lines) == value["Total"]
    assert parts["of which zero stock"] + parts["of which below minimum"] == parts["Critical"]

    assert rows["Critical"]["% of parts"] == pytest.approx(3 / 11)
    assert rows["Excess (above maximum)"]["% of value"] == pytest.approx(3900 / 4082.5)
    assert rows["Total"]["% of parts"] == rows["Total"]["% of value"] == 1
    ws = wb["Summary"]
    pct = next(c for row in ws.iter_rows() for c in row if c.value == rows["Critical"]["% of parts"])
    assert pct.number_format == "0.0%"


def test_summary_section_table(wb):
    rows = _table(wb["Summary"], "Stock by section")
    figures = {label: tuple(row.values()) for label, row in rows.items() if label in
               ("MIS Alpha", "MIS Beta", "Unassigned", "Total")}
    #                       parts crit reorder excess healthy  value  excess value
    assert figures == {
        "MIS Alpha":  (7, 2, 1, 2, 1, 1642.5, 1000),
        "MIS Beta":   (3, 1, 0, 1, 1, 2410, 900),
        "Unassigned": (1, 0, 0, 0, 1, 30, 0),       # the part with no section
        "Total":      (11, 3, 1, 3, 3, 4082.5, 1900),
    }
    assert list(rows)[:4] == ["MIS Alpha", "MIS Beta", "Unassigned", "Total"]


def test_summary_movement_table(wb):
    rows = _table(wb["Summary"], "Material movement - last 7 days")
    days = [d for d in rows if d != "Total"]
    assert days == [datetime(2026, 10, d) for d in range(7, 14)]      # newest last, real dates
    assert rows[datetime(2026, 10, 12)] == {"Received Qty": 5, "Issued Qty": 0}
    assert rows[datetime(2026, 10, 13)] == {"Received Qty": 0, "Issued Qty": 2}
    assert rows["Total"] == {"Received Qty": 5, "Issued Qty": 2}
    ws = wb["Summary"]
    day = next(c for c in ws["A"] if c.value == datetime(2026, 10, 12))
    assert day.number_format == "DD.MM.YYYY"


def test_detail_sheet(wb):
    ws = wb["Stock Detail"]
    heads, rows = _list(ws)
    assert heads == DETAIL_HEAD
    # By section, then part no; the part with no section last; the inactive part nowhere.
    assert [r["Part No"] for r in rows] == ["MISA01", "MISA02", "MISA03", "MISA04", "MISA05", "MISA06",
                                            "MISA07", "MISB01", "MISB02", "MISB03", "MISU01"]
    assert [r["Sr No"] for r in rows] == list(range(1, 12))
    by = {r["Part No"]: r for r in rows}

    assert by["MISA06"] == {
        "Sr No": 6, "Section": "MIS Alpha", "Category": "Sensor", "Part No": "MISA06",
        "Description": "Part MISA06", "Specification": "Spec MISA06", "Supplier": "ACME TOOLS",
        "Location": "R1-S1-B1", "UOM": "NOS", "In Stock": 5, "Min": 2, "Reorder Level": None, "Max": 10,
        "Status": "Healthy", "Rate (₹)": 20.5, "Stock Value (₹)": 102.5,
    }
    assert by["MISA01"]["Supplier"] is None                 # never seen at the gate: an empty cell
    assert by["MISU01"]["Section"] == "Unassigned"
    assert by["MISB03"]["Category"] is None
    assert [by[p]["Status"] for p in ("MISA01", "MISA02", "MISA03", "MISA04", "MISA06", "MISA07")] == [
        "Out of Stock", "Below Minimum", "Reorder", "Above Maximum", "Healthy", "Not Counted"]
    # Not counted / not rated: empty, never a zero.
    assert (by["MISA07"]["In Stock"], by["MISA07"]["Stock Value (₹)"]) == (None, None)
    assert (by["MISA05"]["Rate (₹)"], by["MISA05"]["Stock Value (₹)"]) == (None, None)
    assert by["MISA01"]["Stock Value (₹)"] == 0             # counted, rated, none on the shelf
    # A level stored as 0 is a level not set: an empty cell, never "Max 0" beside 3 in stock.
    for part_no in ("MISB01", "MISU01"):
        assert [by[part_no][k] for k in ("Min", "Reorder Level", "Max")] == [None, None, None]
    assert (by["MISU01"]["In Stock"], by["MISU01"]["Status"]) == (3, "Healthy")
    assert by["MISB01"]["In Stock"] == 0                    # a count of zero is still a figure

    # Totals sit above the header, outside the filter range.
    assert _values(ws, 3) == ["Total", 11, "Stock value (₹)", 4082.5]
    assert ws.auto_filter.ref == "A4:P15"
    assert ws.freeze_panes == "E5"                           # under the header, after Part No

    cell = {head: ws.cell(5 + 9, col) for col, head in enumerate(heads, start=1)}     # MISB03
    # Text starting with = + - @ stays text, never a formula or a number — by
    # Excel's quote prefix, not an apostrophe in the text, which would print.
    for head, text in (("Description", "=SUM(1,1)"), ("Specification", "-710619")):
        assert (cell[head].value, cell[head].data_type, cell[head].quotePrefix) == (text, "s", True)
    assert not cell["Part No"].quotePrefix
    # Whole quantities are integers; a fraction keeps 2 places.
    assert (cell["In Stock"].value, cell["In Stock"].number_format) == (2.5, MONEY)
    whole = ws.cell(5, heads.index("Max") + 1)
    assert (whole.value, type(whole.value), whole.number_format) == (10, int, WHOLE)
    assert cell["Rate (₹)"].number_format == MONEY and cell["Rate (₹)"].alignment.horizontal == "right"
    assert cell["Description"].alignment.wrap_text and cell["Specification"].alignment.wrap_text
    # The Status cell wears the status' tint.
    status = heads.index("Status") + 1
    assert ws.cell(5, status).fill.fgColor.rgb.endswith("FFE4E6")     # Out of Stock
    assert ws.cell(10, status).fill.fgColor.rgb.endswith("D1FAE5")    # Healthy
    assert ws.cell(4, 1).fill.fgColor.rgb.endswith("1F3A5F") and ws.cell(4, 1).font.bold


def test_critical_and_reorder_sheet(wb):
    ws = wb["Critical & Reorder"]
    heads, rows = _list(ws)
    assert heads == ["Sr No", "Section", "Part No", "Description", "Supplier", "Location", "UOM", "In Stock",
                     "Min", "Reorder Level", "Max", "Status", "Order Qty", "Rate (₹)", "Order Value (₹)"]
    # Worst first: zero stock (largest gap first), below minimum, then reorder.
    assert [(r["Part No"], r["Status"]) for r in rows] == [
        ("MISA01", "Out of Stock"), ("MISB01", "Out of Stock"), ("MISA02", "Below Minimum"),
        ("MISA03", "Reorder")]
    # Up to the maximum; with no maximum, up to the minimum; nothing planned, nothing to order.
    assert [r["Order Qty"] for r in rows] == [10, None, 4, 6]
    assert [r["Order Value (₹)"] for r in rows] == [1000, None, 40, 30]
    assert [rows[1][k] for k in ("Min", "Reorder Level", "Max")] == [None, None, None]     # stored as 0
    assert [rows[2][k] for k in ("Min", "Reorder Level", "Max")] == [5, None, None]
    assert _values(ws, 3) == ["Total", 4, "Order value (₹)", 1070]
    assert "1 part has no rate or no order quantity and is not valued" in _values(ws, 2)[-1]
    assert (ws.auto_filter.ref, ws.freeze_panes) == ("A4:O8", "D5")


def test_excess_sheet(wb):
    ws = wb["Excess"]
    heads, rows = _list(ws)
    assert heads == ["Sr No", "Section", "Part No", "Description", "Supplier", "Location", "UOM", "In Stock",
                     "Max", "Excess Qty", "Rate (₹)", "Excess Value (₹)"]
    # Dearest excess first; the part with no rate after the rated ones.
    assert [r["Part No"] for r in rows] == ["MISA04", "MISB02", "MISA05"]
    assert [r["Excess Qty"] for r in rows] == [20, 3, 2]
    assert [r["Excess Value (₹)"] for r in rows] == [1000, 900, None]
    assert _values(ws, 3) == ["Total", 3, "Excess value (₹)", 1900]
    assert "1 part has no rate and is not valued" in _values(ws, 2)[-1]
    assert (ws.auto_filter.ref, ws.freeze_panes) == ("A4:L7", "D5")


def test_category_summary_sheet(wb):
    ws = wb["Category Summary"]
    assert [c.value for c in ws[4]] == ["Section", "Category", "Parts", "Critical", "Reorder", "Excess",
                                        "Healthy", "Stock Value (₹)"]
    rows = [row for row in ws.iter_rows(min_row=5, values_only=True)]
    # By section, the category holding the most value first, a subtotal under each.
    # A category none of whose parts is both counted and rated has no value to
    # show: an empty cell, after the ones that have — a real 0.00 included.
    assert rows == [
        ("MIS Alpha", "Bearing", 3, 1, 1, 1, 0, 1540),
        ("MIS Alpha", "Sensor", 1, 0, 0, 0, 1, 102.5),
        ("MIS Alpha", "Coupling", 1, 1, 0, 0, 0, 0),        # rated, none on the shelf
        ("MIS Alpha", "Adapter", 1, 0, 0, 1, 0, None),      # 12 on the shelf, no rate
        ("MIS Alpha", "Gauge", 1, 0, 0, 0, 0, None),        # rated, never counted
        ("MIS Alpha", "Subtotal", 7, 2, 1, 2, 1, 1642.5),
        ("MIS Beta", "Belt", 2, 1, 0, 1, 0, 2400),
        ("MIS Beta", "Uncategorised", 1, 0, 0, 0, 1, 10),
        ("MIS Beta", "Subtotal", 3, 1, 0, 1, 1, 2410),
        ("Unassigned", "Bearing", 1, 0, 0, 0, 1, 30),
        ("Unassigned", "Subtotal", 1, 0, 0, 0, 1, 30),
    ]
    assert _values(ws, 3) == ["Grand total", 11, 3, 1, 3, 3, 4082.5]
    assert "3 parts have no rate or no count and are not valued" in _values(ws, 2)[-1]
    assert ws.cell(10, 3).font.bold                          # a subtotal row
    assert ws.auto_filter.ref is None                        # subtotals stay where they are
    assert ws.freeze_panes == "C5"


def test_no_placeholder_text_and_real_numbers(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    text = cell.value.strip()
                    assert text.lower() not in ("", "none", "nan", "null", "-", "—"), (ws.title, cell.coordinate)
                    assert "None" not in text and "—" not in text, (ws.title, cell.coordinate)
    # Figures are numbers in the cells, not text that looks like one.
    for title in SHEETS[1:4]:
        heads, rows = _list(wb[title])
        for row in rows:
            for head in heads:
                if head in ("Sr No", "In Stock", "Min", "Reorder Level", "Max", "Order Qty", "Excess Qty") \
                        or "₹" in head:
                    assert row[head] is None or isinstance(row[head], (int, float)), (title, head, row[head])


def test_print_setup(wb):
    summary = wb["Summary"]
    assert summary.page_setup.orientation == "portrait"
    assert (summary.page_setup.fitToWidth, summary.page_setup.fitToHeight) == (1, 1)     # one page
    assert summary.sheet_properties.pageSetUpPr.fitToPage is True
    assert re.fullmatch(r"('?Summary'?!)?\$A\$1:\$L\$\d+", summary.print_area)

    for title in SHEETS[1:]:
        ws = wb[title]
        assert ws.page_setup.orientation == "landscape", title
        assert (ws.page_setup.fitToWidth, ws.page_setup.fitToHeight) == (1, 0), title    # one page wide
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True, title
        assert ws.print_title_rows == "$4:$4", title                                    # header on every page

    for ws in wb.worksheets:
        assert int(ws.page_setup.paperSize) == 9, ws.title                               # A4
        assert ws.print_options.horizontalCentered, ws.title
        assert ws.oddFooter.left.text == "Daily Stock Report · 13.10.2026", ws.title
        assert ws.oddFooter.right.text == "Page &P of &N", ws.title


def test_default_clock_and_missing_logo(report, monkeypatch):
    monkeypatch.setattr(stock_mis_report, "LOGO", stock_mis_report.LOGO + ".missing")
    built = daily_stock_report(STORE)
    assert built["filename"] == f"Daily Stock Report {datetime.now():%d.%m.%Y}.xlsx"     # the plant clock
    wb = load_workbook(io.BytesIO(built["content"]))
    assert wb.sheetnames == SHEETS
    assert wb["Summary"]._images == []                        # built without the picture, no crash


def test_empty_store_builds():
    built = daily_stock_report("MIS-NO-SUCH-STORE", now=NOW)
    wb = load_workbook(io.BytesIO(built["content"]))
    assert wb.sheetnames == SHEETS
    ws = wb["Summary"]
    assert _tile(ws, "Total Parts")[0] == 0
    total = _table(ws, "Stock by status")["Total"]
    assert (total["Parts"], total["% of parts"], total["Stock Value (₹)"]) == (0, None, 0)
    assert "No material was received or issued in the last 7 days." in [c.value for c in ws["A"]]
    for title, ref in (("Stock Detail", "A4:P4"), ("Critical & Reorder", "A4:O4"), ("Excess", "A4:L4")):
        assert wb[title].auto_filter.ref == ref
        assert wb[title].max_row == 4


# ---- the download route ----

_seq = {"n": 0}


def _login(client, role, stores=()):
    _seq["n"] += 1
    user = User(username=f"misuser{_seq['n']}", role=role, is_active=True)
    user.set_password("pw")
    db.session.add(user)
    db.session.flush()
    for code in stores:
        db.session.add(UserStoreAccess(user_id=user.id, store_code=code))
    db.session.commit()
    with client.session_transaction() as sess:
        sess["user_id"] = user.id


def test_route_sends_the_workbook(report, client):
    _login(client, "ADMIN")
    resp = client.get(URL, query_string={"store_code": STORE})
    assert resp.status_code == 200
    assert resp.mimetype == XLSX_MIME
    disposition = resp.headers["Content-Disposition"]
    assert disposition.startswith("attachment")
    assert re.search(r'filename="?Daily Stock Report \d{2}\.\d{2}\.\d{4}\.xlsx"?', disposition)
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == SHEETS
    assert _tile(wb["Summary"], "Total Parts")[0] == 11


# ---- the stock-only workbook ----

def test_stock_only_workbook(report):
    """One sheet with only the stock; each part as the MIS Stock Detail has it."""
    built = stock_only_report(STORE, now=NOW)
    assert built["filename"] == "Daily Stock Report (Stock Only) 13.10.2026.xlsx"
    assert built["mime"] == XLSX_MIME
    wb = load_workbook(io.BytesIO(built["content"]))
    assert wb.sheetnames == ["Stock List"]
    ws = wb["Stock List"]
    heads, rows = _list(ws)
    assert heads == STOCK_ONLY_HEAD
    detail = _list(report[1]["Stock Detail"])[1]
    assert [r["Part No"] for r in rows] == [r["Part No"] for r in detail]     # same parts, same order
    for mine, mis in zip(rows, detail):
        # A part never counted says so here; the MIS leaves the cell empty.
        want = {k: mis[k] for k in STOCK_ONLY_HEAD}
        if mis["In Stock"] is None:
            want["In Stock"] = "Not counted"
        assert mine == want, mine["Part No"]
    by = {r["Part No"]: r for r in rows}
    assert by["MISA07"]["In Stock"] == "Not counted"        # never counted: said, never a 0
    assert by["MISB01"]["In Stock"] == 0
    assert by["MISB03"]["In Stock"] == 2.5

    assert ws.cell(1, 1).value == "Daily Stock Report · Stock List"
    assert ws.cell(1, 7).value == "Report date: 13.10.2026"          # over the last four columns
    assert "G1:J1" in {str(r) for r in ws.merged_cells.ranges}
    # 11 parts, 8 of them with stock (MISA01 and MISB01 at 0, MISA07 not counted).
    assert _values(ws, 3) == ["Total", 11, "With stock: 8"]
    assert ws.auto_filter.ref == "A4:J15"
    assert ws.freeze_panes == "D5"
    # A4 landscape, one page wide, as many long as it needs, header on every page.
    assert ws.page_setup.orientation == "landscape"
    assert (ws.page_setup.fitToWidth, ws.page_setup.fitToHeight) == (1, 0)
    assert int(ws.page_setup.paperSize) == 9
    assert ws.print_title_rows == "$4:$4"
    assert ws.oddFooter.left.text == "Daily Stock Report (Stock Only) · 13.10.2026"


def test_stock_only_colours_in_stock_by_the_levels():
    """Solid dark red at zero, red at or below min / reorder (AT the level
    too), green in between, yellow above max, grey "Not counted"; no colour
    without levels (the store's rule, 14 Oct 2026)."""
    ws = load_workbook(io.BytesIO(stock_only_report(STORE, now=NOW)["content"]))["Stock List"]
    heads = [c.value for c in ws[4]]
    col = heads.index("In Stock") + 1
    cells = {ws.cell(r, 3).value: ws.cell(r, col) for r in range(5, ws.max_row + 1)}
    red, green, yellow, zero, grey = "FFE4E6", "D1FAE5", "FEF9C3", "DC2626", "E5E7EB"
    expected = {
        "MISA01": zero,      # 0, min 2
        "MISA02": red,       # 1, min 5
        "MISA03": red,       # 6 = its reorder level 6: at the level is red
        "MISA04": yellow,    # 30, max 10
        "MISA05": yellow,    # 12, max 10
        "MISA06": green,     # 5, between min 2 and max 10
        "MISB01": zero,      # 0, no levels: nothing on the shelf is dark red all the same
        "MISA07": grey,      # never counted
        "MISB02": yellow,    # 8, max 5
        "MISB03": green,     # 2.5, between min 1 and max 5
    }
    for part_no, fill in expected.items():
        cell = cells[part_no]
        assert cell.fill.fgColor.rgb.endswith(fill), part_no
        assert cell.font.bold, part_no
    assert cells["MISB03"].number_format == MONEY                    # a fraction keeps 2 places
    assert cells["MISA01"].font.color.rgb.endswith("FFFFFF")          # white on the dark red
    assert cells["MISA07"].value == "Not counted"
    assert cells["MISU01"].fill.fill_type is None                     # 3 in stock, no levels: no colour
    # The row banding leaves the coloured column alone, or even rows would hide it.
    assert [str(cf.sqref) for cf in ws.conditional_formatting] == ["A5:F15 H5:J15"]


def test_stock_only_at_the_minimum_is_red_without_a_reorder_level():
    band = stock_mis_report._level_band
    assert band({"qty": 2, "min": 2, "reorder": None, "max": 10}) == "low"
    assert band({"qty": 3, "min": 2, "reorder": None, "max": 10}) == "ok"
    assert band({"qty": 10, "min": 2, "reorder": None, "max": 10}) == "ok"         # AT max is not excess
    assert band({"qty": 11, "min": None, "reorder": None, "max": 10}) == "high"
    assert band({"qty": 4, "min": None, "reorder": None, "max": None}) is None
    assert band({"qty": 0, "min": None, "reorder": None, "max": None}) == "zero"
    assert band({"qty": None, "min": 2, "reorder": 3, "max": 10}) == "uncounted"


def test_stock_only_empty_store_builds():
    wb = load_workbook(io.BytesIO(stock_only_report("MIS-NO-SUCH-STORE", now=NOW)["content"]))
    ws = wb["Stock List"]
    assert (ws.max_row, ws.auto_filter.ref) == (4, "A4:J4")
    assert _values(ws, 3) == ["Total", 0, "With stock: 0"]


def test_stock_only_route(report, client):
    _login(client, "ADMIN")
    resp = client.get(STOCK_ONLY_URL, query_string={"store_code": STORE})
    assert resp.status_code == 200 and resp.mimetype == XLSX_MIME
    assert re.search(r'filename="?Daily Stock Report \(Stock Only\) \d{2}\.\d{2}\.\d{4}\.xlsx"?',
                     resp.headers["Content-Disposition"])
    assert load_workbook(io.BytesIO(resp.data)).sheetnames == ["Stock List"]


def test_route_is_gated_like_the_stock_export(report, client):
    export = "/api/v1/spare-stock/export"
    for url in (export, URL, STOCK_ONLY_URL):
        # Not signed in.
        assert client.get(url, query_string={"store_code": STORE}).status_code == 401
    # A store user of ANOTHER store is refused both; one of this store gets both.
    _login(client, "STORE_USER", stores=["MAINTENANCE"])
    for url in (export, URL, STOCK_ONLY_URL):
        assert client.get(url, query_string={"store_code": STORE}).status_code == 403
        assert client.get(url).status_code == 400            # store_code is required
    _login(client, "STORE_USER", stores=[STORE])
    for url in (export, URL, STOCK_ONLY_URL):
        assert client.get(url, query_string={"store_code": STORE}).status_code == 200
